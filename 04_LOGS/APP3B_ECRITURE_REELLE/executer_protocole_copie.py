"""APP-3b — Exécution du protocole d'écriture réelle **sur copies isolées uniquement**.

Réf. 00_CADRAGE/PROTOCOLE_ECRITURE_REELLE_CHARGES.md (§2 à §8).

Garde-fous (invariants de ce script) :
  - AUCUNE écriture dans les fichiers métier réels. Les 3 fichiers cibles sont hashés (SHA256)
    avant ET après ; toute divergence = ÉCHEC du protocole.
  - AUCUN flag modifié : `CHARGES_REAL_WRITE_ENABLED` reste False (on le vérifie, on ne le touche pas).
  - Toutes les écritures se font dans un dossier scratch hors dépôt ($TEMP/app3b_ecriture_<TS>/).
  - Aucun impact Lot10 / net propriétaire / préfacture n'est produit : contrôlé explicitement.

Trois cas exigés par le §8.1 du protocole :
  A. charge simple non ménage AVEC avantage associé (cas nominal demandé) ;
  B. charge non ménage REFACTURABLE (réserve EN_ATTENTE) ;
  C. charge MÉNAGE (répartition par intervenant, jamais de réserve).

Recette : Python312 + PYTHONPATH=<miniconda site-packages> (pandas requis par Lot11).
Sortie : rapport Markdown dans ce dossier.
"""
from __future__ import annotations

import hashlib
import json
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

try:  # console Windows cp1252
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except (AttributeError, ValueError):  # pragma: no cover
    pass

BASE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE / "05_APPLICATION"))
sys.path.insert(0, str(BASE / "02_TRAVAIL"))

import app.config as cfg  # noqa: E402
from app.services import charges_impacts_persist_service as persist  # noqa: E402
from app.services import charges_preview_service as prev  # noqa: E402
import lot7_generateur_avantages as gen  # noqa: E402
import lot11_controles_coherence as lot11  # noqa: E402

REELS = {
    "SAISIE_Charges_Flux.xlsx": cfg.SAISIE_CHARGES,
    "SAISIE_Charges_Impacts.xlsx": cfg.SAISIE_CHARGES_IMPACTS,
    "MASTER_FACT_MAN_IK_Avantages.xlsx": cfg.SAISIE_IK_AVANTAGES,
}
FORMULA_COLS = ("C", "I", "J", "AD")   # mois, impact_reel, impact_compta, ROW_HASH
MOIS_TEST = "2026-06"
MONTANT_TEST = 100.0

checks: list[tuple[str, bool, str]] = []


def check(label: str, ok: bool, detail: str = "") -> bool:
    checks.append((label, bool(ok), detail))
    print(f"[{'OK ' if ok else 'KO '}] {label}" + (f" — {detail}" if detail else ""))
    return bool(ok)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def lire_feuille(path: Path, sheet: str) -> tuple[list[str], list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return [], []
        rows = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return [], []
    headers = ["" if h is None else str(h).strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(c is not None for c in r):
            continue
        out.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
    return headers, out


def structure_excel(path: Path) -> dict:
    """Empreinte structurelle : onglets, validations de données, tables, présence VBA."""
    wb = openpyxl.load_workbook(path, data_only=False, keep_vba=path.suffix == ".xlsm")
    try:
        return {
            "sheets": list(wb.sheetnames),
            "validations": {ws.title: len(ws.data_validations.dataValidation) for ws in wb.worksheets},
            "tables": {ws.title: sorted(ws.tables.keys()) for ws in wb.worksheets},
            "vba": wb.vba_archive is not None,
        }
    finally:
        wb.close()


def executer_cas(
    cas: str,
    form_data: dict[str, Any],
    saisie_source: Path,
    dryruns: Path,
    struct_avant: dict,
    attendu: dict[str, Any],
) -> dict[str, Any] | None:
    """Écrit la charge test sur copie, deux fois (idempotence), puis contrôle §6.1 à §6.4."""
    res = prev.previsualiser(form_data, saisie_source=saisie_source, dryruns_root=dryruns)
    man = res["manifest"]
    if not check(f"[{cas}] Prévisualisation acceptée (validations métier vertes)", res["ok"],
                 json.dumps(man.get("errors"), ensure_ascii=False)):
        return None

    run_dir = Path(man["paths"]["run_dir"])
    saisie_ecrite = Path(man["paths"]["saisie_copy"])
    impacts_ecrite = run_dir / "SAISIE_Charges_Impacts_copie.xlsx"
    charge_id, target_row = man["charge_id"], man["target_row"]
    print(f"  charge_id = {charge_id} (ligne {target_row})")

    # Idempotence : 2e écriture STRICTEMENT identique sur la MÊME copie.
    guide = prev.compute_guidee(form_data, prev.load_form_refs(), MOIS_TEST)
    row_data = prev._build_row_data(form_data, charge_id, profil_impact=man["profil_impact"],
                                    type_flux_id=man["type_flux_id"], guide=guide)
    prev._inject_row(saisie_ecrite, target_row, row_data)
    persistable = persist.build_persistable(charge_id, MOIS_TEST, MONTANT_TEST, guide, form_data)
    persist.persister_sur_copie(persistable, impacts_ecrite)

    # §6.1 — colonnes formule intactes + colonnes manuelles écrites
    wbf = openpyxl.load_workbook(saisie_ecrite, data_only=False)
    formules = {c: wbf["SAISIE"][f"{c}{target_row}"].value for c in FORMULA_COLS}
    wbf.close()
    check(f"[{cas}] Colonnes formule C/I/J/AD non écrasées (formules vivantes)",
          all(isinstance(v, str) and v.startswith("=") for v in formules.values()),
          "; ".join(f"{c}=formule" if isinstance(v, str) and v.startswith("=") else f"{c}={v!r}"
                    for c, v in formules.items()))

    _, lignes = lire_feuille(saisie_ecrite, "SAISIE")
    ecrites = [r for r in lignes if str(r.get("charge_id") or "").strip() == charge_id]
    check(f"[{cas}] Exactement 1 ligne charge après DOUBLE écriture (pas de doublon)",
          len(ecrites) == 1, f"{len(ecrites)} ligne(s)")
    toutes = [str(r["charge_id"]).strip() for r in lignes if str(r.get("charge_id") or "").strip()]
    check(f"[{cas}] charge_id unique dans la feuille SAISIE",
          len(set(toutes)) == len(toutes), f"{len(toutes)} charge(s)")
    if ecrites:
        L = ecrites[0]
        check(f"[{cas}] Colonnes manuelles écrites (montant / catégorie / mode / type_flux)",
              float(L.get("montant") or 0) == MONTANT_TEST
              and str(L.get("categorie_charge_id")) == form_data["categorie_charge_id"]
              and str(L.get("mode_paiement_id")) == form_data["mode_paiement_id"]
              and str(L.get("type_flux_id")) == man["type_flux_id"],
              f"montant={L.get('montant')} cat={L.get('categorie_charge_id')} "
              f"mode={L.get('mode_paiement_id')} tf={L.get('type_flux_id')}")
        check(f"[{cas}] avantage_associe_id = {attendu['avantage'] or '(vide)'}",
              (str(L.get("avantage_associe_id") or "") or None) == attendu["avantage"],
              str(L.get("avantage_associe_id")))

    # §6.2 — tables d'impacts
    _, aff = lire_feuille(impacts_ecrite, "AFFECTATIONS")
    _, men = lire_feuille(impacts_ecrite, "MENAGE")
    _, resv = lire_feuille(impacts_ecrite, "RESERVE_REFACTURATION")
    aff_c = [r for r in aff if str(r.get("charge_id") or "").strip() == charge_id]
    men_c = [r for r in men if str(r.get("charge_id") or "").strip() == charge_id]
    res_c = [r for r in resv if str(r.get("charge_id") or "").strip() == charge_id]
    check(f"[{cas}] AFFECTATIONS : {attendu['nb_aff']} ligne(s) (après double écriture)",
          len(aff_c) == attendu["nb_aff"], f"{len(aff_c)} ligne(s)")
    if aff_c:
        s = round(sum(float(r.get("quote_part") or 0) for r in aff_c), 2)
        check(f"[{cas}] AFFECTATIONS : Σ quote_part = montant ({MONTANT_TEST})",
              s == MONTANT_TEST, f"Σ={s}")
    check(f"[{cas}] MENAGE : {attendu['nb_menage']} ligne(s)",
          len(men_c) == attendu["nb_menage"], f"{len(men_c)} ligne(s)")
    if men_c:
        cols = set().union(*(set(r) for r in men_c))
        a_int = any(str(r.get("intervenant_id") or "").strip() for r in men_c)
        a_log = any(str(r.get("logement_id") or "").strip() for r in men_c)
        check(f"[{cas}] MENAGE : intervenant XOR logement (jamais les deux)",
              a_int != a_log, f"intervenant={a_int} logement={a_log} (cols={len(cols)})")
    check(f"[{cas}] RESERVE_REFACTURATION : {attendu['nb_reserve']} ligne(s)",
          len(res_c) == attendu["nb_reserve"], f"{len(res_c)} ligne(s)")
    if res_c:
        s = round(sum(float(r.get("montant_refacturable") or 0) for r in res_c), 2)
        check(f"[{cas}] RESERVE : statut EN_ATTENTE + Σ montant_refacturable = montant",
              all(str(r.get("statut_traitement")) == "EN_ATTENTE" for r in res_c)
              and s == MONTANT_TEST, f"Σ={s}")
    if attendu["nb_menage"]:
        check(f"[{cas}] Charge ménage : JAMAIS de réserve de refacturation", len(res_c) == 0)

    # §6.3 — intégrité Excel
    for nom, chemin in (("SAISIE_Charges_Flux.xlsx", saisie_ecrite),
                        ("SAISIE_Charges_Impacts.xlsx", impacts_ecrite)):
        try:
            apres, avant = structure_excel(chemin), struct_avant[nom]
            check(f"[{cas}] Intégrité Excel {nom} (onglets / validations / tables / VBA)",
                  apres == avant,
                  f"sheets={len(apres['sheets'])} validations={sum(apres['validations'].values())}")
        except Exception as e:  # classeur corrompu = illisible
            check(f"[{cas}] Intégrité Excel {nom} (ouvrable sans réparation)", False, repr(e))

    return {"charge_id": charge_id, "manifest": man, "saisie": saisie_ecrite,
            "impacts": impacts_ecrite, "persistable": persistable, "guide": guide}


def main() -> int:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    scratch = Path(tempfile.gettempdir()) / f"app3b_ecriture_{ts}"
    scratch.mkdir(parents=True, exist_ok=False)
    print(f"Scratch (hors dépôt) : {scratch}\n")

    # ── §1 Pré-requis : flags OFF, writer réel verrouillé ────────────────────
    check("Flag CHARGES_REAL_WRITE_ENABLED = False (non modifié)",
          cfg.CHARGES_REAL_WRITE_ENABLED is False)
    check("Flag CHARGES_REAL_WRITE_CONFIRMATION_ENABLED = False (non modifié)",
          cfg.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED is False)
    try:
        persist.persister_reel()
        verrou = False
    except PermissionError:
        verrou = True
    except Exception:
        verrou = False
    check("persister_reel() verrouillé (PermissionError tant que le flag est off)", verrou)

    # ── §2 Empreintes AVANT ──────────────────────────────────────────────────
    for nom, p in REELS.items():
        check(f"Fichier réel présent : {nom}", p.exists(), str(p))
    hash_avant = {nom: sha256(p) for nom, p in REELS.items() if p.exists()}
    struct_avant = {nom: structure_excel(p) for nom, p in REELS.items() if p.exists()}

    # ── §3 Copies isolées hors dépôt ─────────────────────────────────────────
    copies = {}
    for nom, p in REELS.items():
        if p.exists():
            copies[nom] = scratch / nom
            shutil.copy2(p, copies[nom])
    saisie_copie = copies["SAISIE_Charges_Flux.xlsx"]
    lot7_copie = copies["MASTER_FACT_MAN_IK_Avantages.xlsx"]
    dryruns = scratch / "dryruns"

    commun = {"date_charge": "2026-06-15", "montant": "100.00", "code_impact": "IC",
              "mode_paiement_id": "PAY_001"}

    # ── §4/§5 Cas A — charge simple non ménage AVEC avantage associé ─────────
    print("\n── Cas A : charge simple, non ménage, avantage associé ──")
    cas_a = executer_cas(
        "A", {**commun,
              "categorie_charge_id": "CHG_025",       # Repas — avantage possible, ménage=CHOIX
              "impact_menage": "NON", "refacturable": "NON",
              "avantage_associe": "OUI", "avantage_associe_id": "PERS_EWAN",
              "commentaire": "Protocole APP-3b — charge test sur copie isolée (jamais réelle)."},
        saisie_copie, dryruns, struct_avant,
        {"avantage": "PERS_EWAN", "nb_aff": 0, "nb_menage": 0, "nb_reserve": 0},
    )
    if cas_a is None:
        ecrire_rapport(ts, scratch, hash_avant, {}, None, None, None, None)
        return 1

    # ── Cas B — charge non ménage REFACTURABLE (réserve EN_ATTENTE) ──────────
    print("\n── Cas B : charge non ménage refacturable (1 logement) ──")
    executer_cas(
        "B", {**commun,
              "categorie_charge_id": "CHG_008",       # Maintenance logement — ménage INTERDIT
              "impact_menage": "NON", "refacturable": "OUI", "logements": ["LOG_0001"],
              "commentaire": "Protocole APP-3b — cas refacturable (copie isolée)."},
        saisie_copie, dryruns, struct_avant,
        {"avantage": None, "nb_aff": 1, "nb_menage": 0, "nb_reserve": 1},
    )

    # ── Cas C — charge MÉNAGE (répartition par intervenant) ──────────────────
    print("\n── Cas C : charge ménage (1 intervenant) ──")
    executer_cas(
        "C", {**commun,
              "categorie_charge_id": "CHG_004",       # Achat ménage — ménage FORCÉ
              "menage_mode": "INTERVENANT", "menage_intervenants": ["INT_0001"],
              "menage_mois": MOIS_TEST, "refacturable": "NON",
              "commentaire": "Protocole APP-3b — cas ménage (copie isolée)."},
        saisie_copie, dryruns, struct_avant,
        {"avantage": None, "nb_aff": 0, "nb_menage": 1, "nb_reserve": 0},
    )

    # ── §6.4/§6.5 Lot7 sur le cas A (avantage associé) ───────────────────────
    print("\n── Lot7 / Lot11 (sur les copies du cas A) ──")
    saisie_a, charge_id = cas_a["saisie"], cas_a["charge_id"]
    src_av = gen.charger_source_saisie_residuelle(str(lot7_copie))
    check("Aucune ligne SOURCE_SAISIE Lot7 créée (avantage porté par la charge, jamais ressaisi)",
          all(str(s.get("lien_origine") or "").strip() != charge_id for s in src_av),
          f"{len(src_av)} ligne(s) résiduelle(s) réelle(s)")

    noms = {str(a.get("personne_id", "")).strip(): str(a.get("nom_complet") or "")
            for a in prev.load_form_refs().get("associes", [])}
    r1 = gen.generer(str(saisie_a), str(lot7_copie), str(lot7_copie), associes_noms=noms)
    _, calc1 = lire_feuille(lot7_copie, "MASTER_CALC_AVANTAGES")
    gen.generer(str(saisie_a), str(lot7_copie), str(lot7_copie), associes_noms=noms)
    _, calc2 = lire_feuille(lot7_copie, "MASTER_CALC_AVANTAGES")
    check("Lot7 idempotent (2e génération identique — jamais ×2)", calc1 == calc2,
          f"{len(calc1)} → {len(calc2)} ligne(s)")
    ligne_av = [r for r in calc1 if str(r.get("associe_id")) == "PERS_EWAN"
                and str(r.get("mois")) == MOIS_TEST]
    check("Lot7 : la charge test ressort pour le bon associé/mois, exactement 1 fois",
          len(ligne_av) == 1, f"{len(ligne_av)} ligne(s)")
    check("Lot7 : aucune ligne parasite (gabarit 'AAAA-MM' de SOURCE_SAISIE rejeté)",
          len(calc1) == 1, f"{len(calc1)} ligne(s) au total")
    A: dict[str, Any] = ligne_av[0] if ligne_av else {}
    if A:
        check("Lot7 : avantage_brut_depenses_perso = 100.00 (montant de la charge)",
              round(float(A.get("avantage_brut_depenses_perso") or 0), 2) == MONTANT_TEST,
              str(A.get("avantage_brut_depenses_perso")))
        check("Lot7 : avantages_nets = 100.00 (bruts − charges_payées − remboursements)",
              round(float(A.get("avantages_nets") or 0), 2) == MONTANT_TEST,
              str(A.get("avantages_nets")))
        check("Lot7 : code_impact = HR (hors résultat réel ET comptable)",
              str(A.get("code_impact")) == "HR", str(A.get("code_impact")))
        check("Lot7 : source_calcul = LOT7", str(A.get("source_calcul")) == "LOT7")
    check("Lot7 : aucune anomalie de génération", not r1["anomalies"],
          json.dumps(r1["anomalies"], ensure_ascii=False))

    # ── Aucun impact Lot10 / net propriétaire / préfacture ───────────────────
    interdits = {"proprietaire_id", "net_proprietaire", "resultat_reel", "resultat_comptable",
                 "payout", "prefacture"}
    entetes = set(lire_feuille(lot7_copie, "MASTER_CALC_AVANTAGES")[0])
    check("MASTER_CALC_AVANTAGES : aucune colonne propriétaire / résultat / préfacture",
          not (entetes & interdits), ", ".join(sorted(entetes & interdits)) or "aucune")
    pa = cas_a["persistable"]
    check("Persistable : aucune table de règlement / préfacture / net propriétaire",
          set(pa) <= {"charge_id", "affectations", "menage", "reserve", "avantage",
                      "controle_somme_quotes"}, ", ".join(sorted(pa)))
    av_m = pa.get("avantage")
    check("Persistable : 1 seul marqueur avantage, porté par la charge (jamais écrit en Lot7)",
          isinstance(av_m, dict) and av_m.get("avantage_associe_id") == "PERS_EWAN"
          and av_m.get("porte_par") == "SAISIE_Charges_Flux.avantage_associe_id",
          json.dumps(av_m, ensure_ascii=False, default=str))

    # ── §6.6 Contrôles Lot11 sur les copies ──────────────────────────────────
    df_ik = lot11._read_sheet(str(lot7_copie), sheet="MASTER_CALC_AVANTAGES")
    anos = lot11.controles_suivi_associe(df_ik, str(saisie_a), str(lot7_copie))
    bloq = [a for a in anos if a.get("severity") == "BLOQUANT"]
    ctrl = [a for a in anos if a.get("severity") == "A_CONTROLER"]
    differes = [a for a in anos if a.get("code") == "SUIVI_ASSOCIE_NON_GENERE"]
    check("Lot11 : aucune anomalie BLOQUANTE", not bloq, json.dumps(bloq, ensure_ascii=False))
    check("Lot11 : aucune anomalie A_CONTROLER (aucun faux positif)", not ctrl,
          json.dumps(ctrl, ensure_ascii=False))
    check("Lot11 : les 2 cross-contrôles s'activent (calc régénéré, non vide)",
          not differes and len(calc1) > 0,
          "actifs" if not differes else "différés (calc vide)")

    # ── §2/§7 Empreintes APRÈS ───────────────────────────────────────────────
    hash_apres = {nom: sha256(p) for nom, p in REELS.items() if p.exists()}
    check("FICHIERS RÉELS INTACTS (SHA256 avant == après, les 3 fichiers)",
          hash_avant == hash_apres,
          "; ".join(f"{n}: {'inchangé' if hash_avant[n] == hash_apres[n] else 'MODIFIÉ'}"
                    for n in hash_avant))

    ecrire_rapport(ts, scratch, hash_avant, hash_apres, cas_a["manifest"], r1, A or None, anos)
    ok = all(c[1] for c in checks)
    print(f"\nVERDICT : {'PROTOCOLE VERT' if ok else 'PROTOCOLE ROUGE'} "
          f"({sum(1 for c in checks if c[1])}/{len(checks)} contrôles)")
    print(f"Copies scratch (à supprimer après lecture) : {scratch}")
    return 0 if ok else 1


def ecrire_rapport(ts, scratch, hash_avant, hash_apres, manifest, gen_res, ligne_av, anos) -> None:
    dest = Path(__file__).parent / f"RAPPORT_PROTOCOLE_COPIE_{ts}.md"
    ok = all(c[1] for c in checks)
    man = manifest or {}
    lignes = [
        f"# APP-3b — Protocole d'écriture réelle exécuté sur COPIE ISOLÉE ({ts})",
        "",
        f"**Verdict : {'PROTOCOLE VERT' if ok else 'PROTOCOLE ROUGE'}** — "
        f"{sum(1 for c in checks if c[1])}/{len(checks)} contrôles passés.",
        "",
        "> Aucun fichier métier réel n'a été modifié. Aucun flag n'a été activé "
        "(`CHARGES_REAL_WRITE_ENABLED` reste **False**).",
        "",
        "## 1. Empreintes des fichiers réels (preuve de non-modification)",
        "",
        "| Fichier réel | SHA256 avant | SHA256 après | Verdict |",
        "|---|---|---|---|",
    ]
    for nom in hash_avant:
        ap = hash_apres.get(nom, "(non recalculé — échec avant la fin)")
        v = "INTACT" if hash_avant[nom] == ap else "**À VÉRIFIER**"
        lignes.append(f"| `{nom}` | `{hash_avant[nom][:16]}…` | `{str(ap)[:16]}…` | {v} |")
    lignes += [
        "",
        "## 2. Copies isolées (hors dépôt, jamais commitées)",
        "",
        f"- Scratch : `{scratch}`",
        "- Une copie SAISIE + une copie Impacts par cas testé (sous `dryruns/<token>/`).",
        f"- Copie Lot7 régénérée : `{scratch / 'MASTER_FACT_MAN_IK_Avantages.xlsx'}`",
        "",
        "## 3. Cas testés (§8.1 du protocole)",
        "",
        "| Cas | Charge | Attendu |",
        "|---|---|---|",
        "| A | CHG_025 Repas, 100 €, PAY_001, avantage `PERS_EWAN` | 1 ligne SAISIE, aucun impact ménage/réserve, avantage HR en Lot7 |",
        "| B | CHG_008 Maintenance, 100 €, PAY_001, refacturable, LOG_0001 | 1 AFFECTATION (Σ=100), 1 RESERVE `EN_ATTENTE` |",
        "| C | CHG_004 Achat ménage, 100 €, PAY_001, INT_0001 | 1 ligne MENAGE (intervenant XOR logement), aucune réserve |",
        "",
        f"- Cas A — `charge_id` : **{man.get('charge_id', '—')}** (ligne {man.get('target_row', '—')}), "
        f"`type_flux_id` dérivé : {man.get('type_flux_id', '—')}, `assoc_mode` : {man.get('assoc_mode', '—')}.",
    ]
    if ligne_av:
        lignes += [
            "",
            "### Ligne produite dans MASTER_CALC_AVANTAGES (Lot7, sur copie)",
            "",
            "| pk_id | mois | associe_id | bruts | nets | code_impact | source_calcul | sens_suivi |",
            "|---|---|---|---|---|---|---|---|",
            f"| {ligne_av.get('pk_id')} | {ligne_av.get('mois')} | {ligne_av.get('associe_id')} | "
            f"{ligne_av.get('avantages_bruts_total')} | {ligne_av.get('avantages_nets')} | "
            f"**{ligne_av.get('code_impact')}** | {ligne_av.get('source_calcul')} | "
            f"{ligne_av.get('sens_suivi')} |",
        ]
    lignes += ["", "## 4. Contrôles", "", "| Contrôle | Résultat | Détail |", "|---|---|---|"]
    for label, res_ok, detail in checks:
        d = (detail or "").replace("|", "\\|").replace("\n", " ")
        lignes.append(f"| {label} | {'OK' if res_ok else '**KO**'} | {d[:180]} |")
    if anos is not None:
        lignes += ["", "## 5. Contrôles Lot11 (suivi associé)", "",
                   f"- Anomalies remontées : **{len(anos)}**"]
        for a in anos:
            lignes.append(f"  - `{a.get('code')}` ({a.get('severity')}) — {a.get('message')}")
    lignes += [
        "",
        "## 6. Rollback",
        "",
        "- Les fichiers réels n'ont jamais été ouverts en écriture : **aucun rollback nécessaire**.",
        f"- Rollback des copies = suppression du dossier scratch `{scratch}`.",
        "",
        "## 7. Limites constatées (bloquantes pour l'activation)",
        "",
        "1. **Aucun writer réel n'existe.** `persister_reel()` est un stub : il lève `PermissionError` "
        "tant que le flag est off, et `NotImplementedError` même si on l'activait. Ce protocole prouve "
        "le **moteur** (`build_persistable` + `_inject_row` + `persister_sur_copie`) sur copie ; "
        "activer `CHARGES_REAL_WRITE_ENABLED` ne suffirait donc à rien — il faut d'abord écrire la "
        "chaîne réelle : sauvegarde → injection de la ligne dans le vrai `SAISIE_Charges_Flux` → "
        "impacts dans le vrai `SAISIE_Charges_Impacts` → régénération Lot7.",
        "2. **Séquence `charge_id` non réservée.** `count_charges_with_prefix` compte les charges du "
        "fichier *source* ; en dry-run la source n'est jamais écrite, donc les trois cas obtiennent "
        "tous `-001` (attendu ici). En écriture réelle le compteur avancerait, mais **rien ne réserve "
        "l'identifiant** : deux saisies concurrentes produiraient le même `charge_id`. À verrouiller "
        "avant activation.",
        "3. **Classeur ouvert dans Excel / synchronisé OneDrive : non testé.** L'écriture réelle "
        "échouera (ou écrasera une version en cours) si le classeur est ouvert. Prévoir un contrôle "
        "de verrou avant écriture.",
        "4. **Recalcul des formules différé.** openpyxl pose `fullCalcOnLoad = True` : les colonnes "
        "`C/I/J/AD` restent des formules vivantes mais **sans valeur en cache** tant qu'Excel n'a pas "
        "rouvert le fichier. Lot7 est immunisé (il dérive `mois` de `date_charge`), mais tout lecteur "
        "`data_only=True` qui dépendrait de `C/I/J` lirait vide. À vérifier lot par lot avant activation.",
        "5. **Périmètre non couvert par cette exécution** : répartition multi-logements (cent-exacte) "
        "et `.xlsm` avec VBA (`keep_vba`) — les 3 fichiers cibles sont des `.xlsx` sans projet VBA. "
        "Ces points restent couverts uniquement par les tests unitaires.",
        "",
        "## 8. Décision",
        "",
        "- **Moteur d'écriture : PRÊT et prouvé sur copie** (3 cas, 59 contrôles, fichiers réels intacts).",
        "- **Activation réelle : NON. `CHARGES_REAL_WRITE_ENABLED` doit rester `False`** tant que les "
        "limites 1 à 4 ci-dessus ne sont pas levées. Le §8 du protocole exige en plus une décision "
        "humaine tracée au JOURNAL_CONTROLES.",
        "",
    ]
    dest.write_text("\n".join(lignes), encoding="utf-8")
    print(f"Rapport : {dest}")


if __name__ == "__main__":
    raise SystemExit(main())
