#!/usr/bin/env python3
"""Ajoute à data_recette/ les sources amont FICTIVES requises par la chaîne aval.

Principe : on seede des **entrées**, jamais des sorties de lot. `lot4quater` doit réellement
tourner et produire lui-même `MASTER_CALC_Reservations_Resolues.xlsx` ; seeder ce fichier
reviendrait à court-circuiter le moteur et à fabriquer un faux succès.

Produit donc :
- `02_TRAVAIL/Lot4bis_TableCommune/MASTER_CALC_Reservations.xlsx` (onglet MASTER) : table live
  consommée par lot4quater — 4 réservations « propres » 2026-06 (1000 € chacune, LOG_A1/A2/B1/C1)
  + remplissage pour dépasser le seuil de volume CTR-9-003 (>= 1000 lignes) ;
- `02_TRAVAIL/Lot1_Hostaway/MASTER_CALC_HA_Payout.xlsx` (onglet data) : payout correspondant ;
- `02_TRAVAIL/Lot6c_MenagesExternes/MASTER_FACT_MEN_MenagesExternes.xlsx` : en-tête seul
  (0 flux ménage externe — lot6c n'est pas dans la chaîne aval).

Le fichier bancaire `Lot8_Banque/BANQUE_LOT8_IMPORT.xlsx` n'est PAS écrit ici : `build_banque()`
de `build_data_recette.py` en produit une version complète (schéma lot8a, 26 colonnes) que lot9
sait lire. Ne pas l'écraser par un en-tête vide.

Aucune donnée réelle. Montants ronds pour des résultats vérifiables :
  LOG_A1 15% → commission 150, net 850 ; LOG_A2 19% → 190 / 810 ;
  LOG_B1 15% → 150 / 850 ; LOG_C1 19% → 190 / 810.
"""
from __future__ import annotations

import hashlib
import os
from datetime import date, timedelta
from pathlib import Path

import openpyxl

WT = Path(os.environ.get("WT", Path(__file__).resolve().parent.parent)).resolve()
REC = WT / "data_recette"

# Colonnes de la table live lot4bis — doivent couvrir BASE_COLS de lot4quater, qui recopie
# `{c: r.get(c) for c in BASE_COLS}`. Une colonne absente arriverait à None dans la résolue.
LIVE_HDR = ['reservation_calc_id', 'ROW_HASH', 'source', 'reservation_id_hostaway',
            'reservation_hh_id', 'mois', 'logement_id', 'proprietaire_id',
            'date_arrivee', 'date_depart', 'nuits', 'guestCount', 'source_guestCount',
            'montant_retenu', 'source_montant', 'code_impact', 'impact_resultat_reel',
            'impact_resultat_comptable', 'statut_controle', 'niveau_anomalie', 'code_anomalie',
            'commentaire', 'source_module', 'source_table', 'source_pk', 'date_integration']

PAYOUT_HDR = ['reservation_id', 'listingMapId', 'source', 'channel_type', 'statut_calcul_payout',
              'payout_calcule', 'source_payout', 'menage_retenu', 'assiette_commission',
              'inclure_resultat_auto', 'extrait_le', 'ROW_HASH', 'menage_retenu_source',
              'cout_standard_id', 'cout_standard_menage_snapshot', 'cout_standard_date_debut_validite',
              'cout_standard_date_fin_validite', 'logement_id_snapshot', 'type_logement_id_snapshot',
              'date_reference_cout_menage']

MEN_HDR = ['menage_externe_id', 'ROW_HASH', 'facture_id', 'date_facture', 'date_menage', 'mois',
           'logement_id', 'proprietaire_id', 'intervenant_id', 'montant', 'source_pk',
           'statut_controle', 'niveau_anomalie', 'code_anomalie', 'commentaire']

PARC = [("LOG_A1", "PROP_A", "TYPE_001"), ("LOG_A2", "PROP_A", "TYPE_002"),
        ("LOG_B1", "PROP_B", "TYPE_001"), ("LOG_C1", "PROP_C", "TYPE_003")]


def _hash(*parts):
    return hashlib.sha256("|".join(str(p) for p in parts).encode()).hexdigest()[:16]


def _live_row(rid_calc, rid_ha, mois, lg, pr, arr, montant):
    dep = arr + timedelta(days=2)
    return {
        'reservation_calc_id': rid_calc, 'ROW_HASH': _hash(rid_calc), 'source': 'HOSTAWAY_AIRBNB',
        'reservation_id_hostaway': rid_ha, 'reservation_hh_id': None, 'mois': mois,
        'logement_id': lg, 'proprietaire_id': pr, 'date_arrivee': arr.isoformat(),
        'date_depart': dep.isoformat(), 'nuits': 2, 'guestCount': 2,
        'source_guestCount': 'HOSTAWAY', 'montant_retenu': montant,
        'source_montant': 'HOSTAWAY_PAYOUT', 'code_impact': 'IC', 'impact_resultat_reel': 'OUI',
        'impact_resultat_comptable': 'OUI', 'statut_controle': 'VALIDE', 'niveau_anomalie': 'INFO',
        'code_anomalie': None, 'commentaire': 'FICTIF recette', 'source_module': 'lot4bis',
        'source_table': 'MASTER_FACT_HA_Reservations', 'source_pk': rid_calc,
        'date_integration': '2026-06-15T00:00:00',
    }


def _payout_row(rid_ha, lg, typ, montant):
    return {
        'reservation_id': rid_ha, 'listingMapId': 900000, 'source': 'airbnbOfficial',
        'channel_type': 'AIRBNB', 'statut_calcul_payout': 'NORMAL', 'payout_calcule': montant,
        'source_payout': 'airbnbExpectedPayoutAmount', 'menage_retenu': 0,
        'assiette_commission': montant, 'inclure_resultat_auto': 'OUI',
        'extrait_le': '2026-06-15T00:00:00', 'ROW_HASH': _hash(rid_ha), 'menage_retenu_source': 'REF',
        'cout_standard_id': 'COUT_MEN_001', 'cout_standard_menage_snapshot': 0,
        'cout_standard_date_debut_validite': '2026-01-01', 'cout_standard_date_fin_validite': None,
        'logement_id_snapshot': lg, 'type_logement_id_snapshot': typ,
        'date_reference_cout_menage': '2026-06-06',
    }


def build():
    live_rows, payout_rows = [], []
    n = 0
    # 4 réservations propres 2026-06 (1000 € chacune)
    for lg, pr, typ in PARC:
        n += 1
        rid_ha = 700000 + n
        live_rows.append(_live_row(f"RES-2026-06-{lg}", rid_ha, "2026-06", lg, pr,
                                   date(2026, 6, 6), 1000))
        payout_rows.append(_payout_row(rid_ha, lg, typ, 1000))
    # Remplissage volume (>= 1000 lignes) sur mois antérieurs, montant 10 €, hors 2026-06.
    # Uniquement sur des mois couverts par un taux de commission historisé (>= 2026-01).
    months = [f"2026-{m:02d}" for m in range(1, 6)]
    filler_target = 1010
    i = 0
    while len(live_rows) < filler_target:
        lg, pr, typ = PARC[i % len(PARC)]
        mo = months[i % len(months)]
        yr, mm = int(mo[:4]), int(mo[5:7])
        n += 1
        rid_ha = 700000 + n
        live_rows.append(_live_row(f"RES-{mo}-F{n}", rid_ha, mo, lg, pr, date(yr, mm, 10), 10))
        payout_rows.append(_payout_row(rid_ha, lg, typ, 10))
        i += 1

    # Table live lot4bis (entrée de lot4quater)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(LIVE_HDR)
    for r in live_rows:
        ws.append([r.get(h) for h in LIVE_HDR])
    dst = REC / "02_TRAVAIL" / "Lot4bis_TableCommune" / "MASTER_CALC_Reservations.xlsx"
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)

    # payout
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(PAYOUT_HDR)
    for r in payout_rows:
        ws.append([r.get(h) for h in PAYOUT_HDR])
    dst = REC / "02_TRAVAIL" / "Lot1_Hostaway" / "MASTER_CALC_HA_Payout.xlsx"
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)

    # ménages externes (en-tête seul)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(MEN_HDR)
    dst = REC / "02_TRAVAIL" / "Lot6c_MenagesExternes" / "MASTER_FACT_MEN_MenagesExternes.xlsx"
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)

    print(f"   réservations live : {len(live_rows)} lignes (dont 4 propres 2026-06) "
          f"| payout : {len(payout_rows)}")


if __name__ == "__main__":
    build()
