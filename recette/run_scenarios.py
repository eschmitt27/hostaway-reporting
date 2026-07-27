#!/usr/bin/env python3
"""Exécute les scénarios de charges en RECETTE : reset → prévisualisation → confirmation réelle →
lecture des chiffres écrits (charge, affectations, réserve). Valeurs attendues codées explicitement.

Doit être lancé avec l'environnement recette :
  PYTHONPATH=<...>/05_APPLICATION PROJECT_ROOT=<data_recette> APP_DATA_DIR=<data_recette>/data
  RECETTE_MODE=1 RECETTE_ROOT=<data_recette> CHARGES_REAL_WRITE_ENABLED=1
  CHARGES_REAL_WRITE_CONFIRMATION_ENABLED=1 LOT4A_ENGINE_PYTHON=<python-pandas>  WT=<worktree>

Sortie : une ligne par scénario + « ALL SCENARIOS OK » et sortie 0 si tout est conforme.
"""
import contextlib
import importlib.util
import io
import os
import sys
from pathlib import Path

import openpyxl

WT = Path(os.environ["WT"]).resolve()
REC = WT / "data_recette"
ENGINE = Path(os.environ.get("LOT4A_ENGINE_PYTHON", r"C:\Program Files\Python312\python.exe"))


def reset():
    spec = importlib.util.spec_from_file_location("bdr", WT / "recette" / "build_data_recette.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    with contextlib.redirect_stdout(io.StringIO()):
        mod.main()
    from app.db.connection import apply_migrations
    apply_migrations()


def read_charge():
    """Charge écrite par le scénario courant = la DERNIÈRE ligne de la SAISIE.

    Lisait `rows[0]` : juste tant que la SAISIE de recette était vide, faux dès qu'elle porte des
    charges de départ — le writer ajoute à la suite. Les vérifications portaient alors sur la
    première charge du jeu de recette et passaient par coïncidence quand elle avait les mêmes
    valeurs attendues. `reset()` précède chaque scénario : la dernière ligne est bien la sienne.
    """
    wb = openpyxl.load_workbook(REC / "01_SOURCES_BRUTES" / "Charges" / "SAISIE_Charges_Flux.xlsx",
                               read_only=True, data_only=True)
    ws = wb["SAISIE"]; hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    wb.close(); return rows[-1] if rows else {}


def read_impacts():
    wb = openpyxl.load_workbook(REC / "01_SOURCES_BRUTES" / "Charges" / "SAISIE_Charges_Impacts.xlsx",
                               read_only=True, data_only=True)
    out = {}
    for sh in ("AFFECTATIONS", "RESERVE_REFACTURATION"):
        if sh in wb.sheetnames:
            ws = wb[sh]; hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
            out[sh] = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)
                       if any(v is not None for v in r)]
    wb.close(); return out


SC = {
 "A": {"date_charge": "2026-06-15", "montant": "100", "categorie_charge_id": "CHG_008", "code_impact": "IC", "logements": ["LOG_A1"], "mode_paiement_id": "PAY_001", "refacturable": "NON"},
 "B": {"date_charge": "2026-06-15", "montant": "100", "categorie_charge_id": "CHG_008", "code_impact": "IC", "logements": ["LOG_A1"], "mode_paiement_id": "PAY_001", "refacturable": "OUI"},
 "C": {"date_charge": "2026-06-15", "montant": "100", "categorie_charge_id": "CHG_008", "code_impact": "IC", "logements": ["LOG_A1"], "mode_paiement_id": "PAY_004", "associe_id": "PERS_X", "refacturable": "NON", "commentaire": "FICTIF compte perso"},
 "D": {"date_charge": "2026-06-15", "montant": "100", "categorie_charge_id": "CHG_008", "code_impact": "HC", "logements": ["LOG_A1"], "mode_paiement_id": "PAY_001", "refacturable": "NON", "commentaire": "FICTIF hors compta"},
 "H": {"date_charge": "2026-06-15", "montant": "100", "categorie_charge_id": "CHG_008", "code_impact": "IC", "logements": ["LOG_A1", "LOG_A2"], "mode_paiement_id": "PAY_001", "refacturable": "NON"},
 "I": {"date_charge": "2026-06-15", "montant": "100", "categorie_charge_id": "CHG_008", "code_impact": "IC", "logements": ["LOG_A1", "LOG_B1"], "mode_paiement_id": "PAY_001", "refacturable": "NON"},
}

EXPECT = {
 "A": {"prise_en_compta": "OUI", "nb_affect": 1, "quotes": {"LOG_A1": 100}, "reserve_total": 0},
 "B": {"prise_en_compta": "OUI", "nb_affect": 1, "quotes": {"LOG_A1": 100}, "reserve_total": 100},
 "C": {"prise_en_compta": "OUI", "nb_affect": 1, "quotes": {"LOG_A1": 100}, "reserve_total": 0},
 "D": {"prise_en_compta": "NON", "nb_affect": 1, "quotes": {"LOG_A1": 100}, "reserve_total": 0},
 "H": {"prise_en_compta": "OUI", "nb_affect": 2, "quotes": {"LOG_A1": 50, "LOG_A2": 50}, "reserve_total": 0},
 "I": {"prise_en_compta": "OUI", "nb_affect": 2, "quotes": {"LOG_A1": 50, "LOG_B1": 50}, "reserve_total": 0},
}


def _num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def main():
    from app.services.charges_preview_service import previsualiser
    from app.services import charges_confirmation_service as conf

    all_ok = True
    for name, form in SC.items():
        reset()
        r = previsualiser(dict(form))
        assert r["ok"], f"{name}: preview refusé {[e['code'] for e in r['manifest']['errors']]}"
        res = conf.confirmer(r["token"], python_moteur=ENGINE)
        d = res.as_dict()
        assert d.get("statut") == "SUCCES", f"{name}: confirm {d.get('statut')} {d.get('code')}"
        post = d.get("post_ecriture") or {}
        assert (post.get("lot3") or {}).get("statut") == "OK", f"{name}: lot3 KO"
        assert (post.get("lot11") or {}).get("statut") == "OK", f"{name}: lot11 KO"

        charge = read_charge()
        imp = read_impacts()
        aff = imp.get("AFFECTATIONS", [])
        quotes = {a.get("logement_id"): _num(a.get("quote_part") or a.get("montant")) for a in aff}
        reserve_total = round(sum(_num(x.get("montant_refacturable") or x.get("montant")) or 0
                                  for x in imp.get("RESERVE_REFACTURATION", [])), 2)

        exp = EXPECT[name]
        checks = {
            "prise_en_compta": charge.get("prise_en_compta") == exp["prise_en_compta"],
            "nb_affect": len(aff) == exp["nb_affect"],
            "quotes": quotes == {k: float(v) for k, v in exp["quotes"].items()},
            "somme=montant": round(sum(v or 0 for v in quotes.values()), 2) == 100.0,
            "reserve_total": reserve_total == float(exp["reserve_total"]),
        }
        ok = all(checks.values())
        all_ok = all_ok and ok
        print(f"[{'PASS' if ok else 'FAIL'}] {name}: impact={charge.get('code_impact')} "
              f"compta={charge.get('prise_en_compta')} quotes={quotes} reserve={reserve_total}"
              f"{'' if ok else ' ECHECS=' + str([k for k, v in checks.items() if not v])}")

    print("ALL SCENARIOS OK" if all_ok else "SOME SCENARIOS FAILED")
    return 0 if all_ok else 1


if __name__ == "__main__":
    sys.exit(main())
