"""LOT4A — Transformateur SAISIE ReservationsHH -> MASTER de test. Mode UNIQUE : --dry-run.

Genere un MASTER de test (feuilles MASTER + VUE_ACTIVE, 34 colonnes canoniques) UNIQUEMENT sous
04_LOGS/LOT4A_DRY_RUN/<horodatage-UTC>/. Jamais de --write-master, jamais Power Query, jamais Excel COM.
Ne modifie aucun fichier metier (SAISIE, MASTER reel, REF_Setup, sources, exports, aval, application).
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import sys
from pathlib import Path

import openpyxl

from lib_lot4a_reservations_hh import (
    PROJECT_ROOT, SAISIE_PATH, MASTER_PATH, REF_SETUP_PATH,
    MASTER_COLUMNS, SHEET_MASTER, SHEET_VUE_ACTIVE,
    read_saisie_values, read_taux_rows, build_master,
    file_fingerprint, assert_output_under, _norm,
)

LOGS_ROOT = PROJECT_ROOT / "04_LOGS" / "LOT4A_DRY_RUN"
MASTER_FILENAME = "MASTER_FACT_MAN_ReservationsHorsHostaway.xlsx"
TMP_FILENAME = ".tmp_master_test.xlsx"


class AsOfError(ValueError):
    pass


def parse_as_of(raw: str | None) -> str:
    """Exige un ISO-8601 UTC explicite. Retourne la forme canonique 'YYYY-MM-DDTHH:MM:SSZ'."""
    if not raw or not raw.strip():
        raise AsOfError("--as-of obligatoire (ISO-8601 UTC, ex. 2026-07-02T00:00:00Z)")
    s = raw.strip()
    try:
        dt = _dt.datetime.fromisoformat(s)
    except ValueError:
        raise AsOfError(f"--as-of format invalide/ambigu: {raw!r} (attendu ISO-8601 UTC)")
    if dt.tzinfo is None:
        raise AsOfError(f"--as-of sans fuseau (date naive interdite): {raw!r}")
    if dt.utcoffset() != _dt.timedelta(0):
        raise AsOfError(f"--as-of doit etre en UTC (offset non nul interdit): {raw!r}")
    return dt.astimezone(_dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _cell(value):
    """Vide -> None (cellule vide) ; sinon valeur telle quelle (nombres float, textes)."""
    return None if (value is None or value == "") else value


def _write_sheet(ws, rows: list[dict]):
    ws.append(MASTER_COLUMNS)
    for rec in rows:
        ws.append([_cell(rec.get(col)) for col in MASTER_COLUMNS])


def _atomic_write_master(out_dir: Path, master_rows: list[dict], vue_rows: list[dict]) -> Path:
    """Ecrit temp -> valide -> remplacement atomique. Garde de chemin AVANT chaque ecriture."""
    tmp_path = assert_output_under(out_dir / TMP_FILENAME, LOGS_ROOT)
    final_path = assert_output_under(out_dir / MASTER_FILENAME, LOGS_ROOT)

    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        _write_sheet(wb.create_sheet(SHEET_MASTER), master_rows)
        _write_sheet(wb.create_sheet(SHEET_VUE_ACTIVE), vue_rows)
        wb.save(str(tmp_path))
        wb.close()

        _validate_workbook(tmp_path, master_rows, vue_rows)

        import os
        os.replace(str(tmp_path), str(final_path))  # atomique (meme dossier)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()
    return final_path


def _validate_workbook(path: Path, master_rows, vue_rows):
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        assert wb.sheetnames == [SHEET_MASTER, SHEET_VUE_ACTIVE], f"feuilles inattendues: {wb.sheetnames}"
        for sheet, expected_rows in ((SHEET_MASTER, master_rows), (SHEET_VUE_ACTIVE, vue_rows)):
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            header = [_norm(h) for h in rows[0]]
            assert header == MASTER_COLUMNS, f"{sheet}: schema/ordre invalide"
            assert len(rows) - 1 == len(expected_rows), f"{sheet}: nb lignes incoherent"
        # oracle
        ws = wb[SHEET_MASTER]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [_norm(h) for h in rows[0]]
        for r in rows[1:]:
            rec = dict(zip(hdr, r))
            if rec.get("reservation_hh_id") == "RESHH-2026-05-001":
                assert float(rec["taux_commission"]) == 0.15, "oracle taux"
                assert float(rec["commission"]) == 343.27, "oracle commission"
                assert float(rec["acompte_facture"]) == 2343.48, "oracle acompte"
    finally:
        wb.close()


def run_dry_run(as_of_raw: str | None) -> int:
    horodatage = _dt.datetime.now(_dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    try:
        as_of = parse_as_of(as_of_raw)
    except AsOfError as exc:
        print(f"ERREUR_TECHNIQUE : {exc}", file=sys.stderr)
        return 2

    for p in (SAISIE_PATH, REF_SETUP_PATH):
        if not p.exists():
            print(f"ERREUR_TECHNIQUE : fichier introuvable {p}", file=sys.stderr)
            return 2

    sources = {
        "SAISIE_ReservationsHorsHostaway.xlsx": SAISIE_PATH,
        "MASTER_FACT_MAN_ReservationsHorsHostaway.xlsx": MASTER_PATH,
        "REF_Setup.xlsm": REF_SETUP_PATH,
    }
    fp_avant = {n: file_fingerprint(p) for n, p in sources.items()}

    try:
        saisie_list = read_saisie_values()
        taux_rows = read_taux_rows()
        result = build_master(saisie_list, taux_rows, as_of)
    except Exception as exc:  # noqa: BLE001
        print(f"ERREUR_TECHNIQUE : {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2

    out_dir = assert_output_under(LOGS_ROOT / horodatage, LOGS_ROOT)
    out_dir.mkdir(parents=True, exist_ok=True)

    statut = result["statut"]
    master_test_genere = False
    final_path = None

    if statut in ("ANALYSE_BLOQUEE_TAUX", "ANALYSE_BLOQUEE_DONNEES"):
        # Aucun MASTER de test, aucune sortie partielle
        _write_anomalies(out_dir, result)
    else:
        final_path = _atomic_write_master(out_dir, result["master_rows"], result["vue_active_rows"])
        master_test_genere = True

    fp_apres = {n: file_fingerprint(p) for n, p in sources.items()}
    sources_inchangees = fp_avant == fp_apres

    manifest = {
        "horodatage_utc": horodatage,
        "mode": "--dry-run",
        "as_of_normalise": as_of,
        "python_executable": sys.executable,
        "python_version": sys.version.split()[0],
        "openpyxl": openpyxl.__version__,
        "statut": statut,
        "master_test_genere": master_test_genere,
        "motif_blocage": result.get("motif_blocage"),
        "nb_saisie": len(saisie_list),
        "nb_master_rows": len(result["master_rows"]),
        "nb_vue_active": len(result["vue_active_rows"]),
        "anomalies_taux": result["anomalies_taux"],
        "anomalies_donnees": result["anomalies_donnees"],
        "fingerprints_avant": fp_avant,
        "fingerprints_apres": fp_apres,
        "sources_inchangees": sources_inchangees,
        "master_test_path": str(final_path) if final_path else None,
    }
    manifest_path = assert_output_under(out_dir / "manifest.json", LOGS_ROOT)
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    rapport_path = assert_output_under(out_dir / "rapport.md", LOGS_ROOT)
    rapport_path.write_text(_render_md(manifest, result), encoding="utf-8")

    if not sources_inchangees:
        print("ERREUR_TECHNIQUE : une source metier a change pendant l'execution", file=sys.stderr)
        return 2

    print(f"Statut            : {statut}")
    print(f"MASTER de test    : {final_path if final_path else '(non genere — blocage)'}")
    print(f"Dossier de run    : {out_dir}")
    print(f"Manifest          : {manifest_path}")
    print(f"master_test_genere: {master_test_genere}")
    return 0


def _write_anomalies(out_dir: Path, result: dict):
    path = assert_output_under(out_dir / "rapport_anomalies.md", LOGS_ROOT)
    L = ["# LOT4A dry-run — BLOCAGE\n", f"**Statut** : `{result['statut']}`  ",
         f"**Motif** : `{result.get('motif_blocage')}`\n"]
    if result["anomalies_taux"]:
        L.append("## Taux bloquants\n")
        for a in result["anomalies_taux"]:
            L.append(f"- {a['reservation_hh_id']} : {a['statut']} — {a['message']}")
    if result["anomalies_donnees"]:
        L.append("\n## Donnees manuelles invalides\n")
        for a in result["anomalies_donnees"]:
            L.append(f"- {a['reservation_hh_id']} : {', '.join(a['codes'])}")
    L.append("\n**Aucun MASTER de test genere. Aucune sortie .xlsx partielle.**\n")
    path.write_text("\n".join(L), encoding="utf-8")


def _render_md(manifest: dict, result: dict) -> str:
    L = ["# LOT4A dry-run — Transformateur SAISIE ReservationsHH -> MASTER de test\n"]
    L.append(f"**Statut** : `{manifest['statut']}`  ")
    L.append(f"**as-of (date_integration)** : `{manifest['as_of_normalise']}`  ")
    L.append(f"**MASTER de test genere** : {manifest['master_test_genere']}  ")
    L.append(f"**Interpreteur** : `{manifest['python_executable']}` (Python {manifest['python_version']})  ")
    L.append(f"**openpyxl** {manifest['openpyxl']}\n")
    L.append(f"- lignes SAISIE utiles : {manifest['nb_saisie']}")
    L.append(f"- lignes MASTER generees : {manifest['nb_master_rows']}")
    L.append(f"- lignes VUE_ACTIVE (statut_controle=VALIDE) : {manifest['nb_vue_active']}\n")
    # oracle
    oracle = next((r for r in result["master_rows"] if r.get("reservation_hh_id") == "RESHH-2026-05-001"), None)
    if oracle:
        L.append("## Oracle RESHH-2026-05-001\n")
        for f in ("taux_commission", "commission", "acompte_facture", "date_integration"):
            L.append(f"- {f} = `{oracle[f]}`")
    L.append("\n## Invariance des sources reelles (lecture seule)\n")
    for name, fp in manifest["fingerprints_avant"].items():
        inchange = fp == manifest["fingerprints_apres"].get(name, {})
        L.append(f"- `{name}` : {'INCHANGE' if inchange else 'MODIFIE !!!'} (sha256 {fp.get('sha256','?')[:12]}…)")
    return "\n".join(L) + "\n"


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="LOT4A transformateur dry-run (MASTER de test, lecture seule des sources).")
    parser.add_argument("--dry-run", action="store_true", help="Mode unique disponible.")
    parser.add_argument("--as-of", default=None, help="ISO-8601 UTC obligatoire (ex. 2026-07-02T00:00:00Z).")
    args = parser.parse_args(argv)
    if not args.dry_run:
        parser.error("Seul --dry-run est disponible a ce stade (aucun --write-master).")
    return run_dry_run(args.as_of)


if __name__ == "__main__":
    raise SystemExit(main())
