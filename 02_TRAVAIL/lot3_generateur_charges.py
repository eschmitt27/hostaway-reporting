"""Lot3 — Générateur réel du MASTER charges (Option A : Python/openpyxl déterministe).

MOTEUR RÉEL de Lot3. Le classeur MASTER_FACT_MAN_Charges.xlsx ne contient AUCUN Power Query
vivant (pas de connections.xml / DataMashup / queryTables) : l'onglet POWER_QUERY_CODE est
DOCUMENTAIRE. Rien n'alimentait le MASTER — une charge saisie restait invisible pour Lot9, Lot10,
Lot11, Lot12 et la page Charges de l'app, qui lisent tous le MASTER. C'est ce module qui produit
réellement MASTER_FACT_MAN_Charges à partir de la source durable.

  SAISIE_Charges_Flux.xlsx (vérité métier)  ──►  MASTER_FACT_MAN_Charges.xlsx (sortie calculée)

Le contrat de sortie reproduit à l'identique celui décrit par le M-code documentaire (37 colonnes =
31 colonnes SAISIE + 6 colonnes dérivées : sens, filtre_vue_menage, source_module, source_table,
source_pk, date_integration), plus l'onglet VUE_MENAGE (filtre_vue_menage=OUI ET statut_controle=VALIDE).

INDÉPENDANCE AU CACHE EXCEL — les colonnes portées par des formules dans la SAISIE sont recalculées
en Python, jamais lues depuis leur cache (openpyxl préserve les formules mais ne les recalcule pas) :
  - `mois`                      ← dérivé de `date_charge` (jamais la colonne formule C) ;
  - `impact_resultat_reel`      ← dérivé de `code_impact` (D012) ;
  - `impact_resultat_comptable` ← dérivé de `code_impact` (D012) ;
  - `ROW_HASH`                  ← recalculé (jamais la colonne formule AD).

Règles métier préservées : une charge = UNE charge économique unique (aucune duplication par
affectation analytique, ménage ou réserve) ; les avantages associés restent Lot7 (HR) ; les réserves
refacturables n'appliquent rien automatiquement en préfacture.

Écriture : uniquement sur le classeur cible passé en paramètre (`sortie_path`). Les tests n'écrivent
que sur copies/fixtures. Ce module ne modifie JAMAIS la SAISIE ni REF_Setup (lecture seule).
"""
from __future__ import annotations

import datetime as dt
from typing import Any, Iterable

import openpyxl

SAISIE_SHEET = "SAISIE"
MASTER_SHEET = "MASTER"
VUE_MENAGE_SHEET = "VUE_MENAGE"
REF_CATEGORIES_SHEET = "REF_Categories_Charges"

# Colonnes reprises telles quelles de la SAISIE (A..AE), dans l'ordre du contrat.
SAISIE_COLS: list[str] = [
    "charge_id", "date_charge", "mois", "montant", "sens_flux",
    "categorie_charge_id", "type_flux_id", "code_impact",
    "impact_resultat_reel", "impact_resultat_comptable", "prise_en_compta",
    "associe_id", "mode_paiement_id", "carte_id",
    "affectation_type", "logement_id", "proprietaire_id", "reservation_id",
    "refacturable", "source_flux", "methode_traitement",
    "paye_avec_montant_recupere", "lien_virement_banque",
    "statut_controle", "niveau_anomalie", "code_anomalie", "statut_rapprochement",
    "justificatif", "commentaire", "ROW_HASH", "date_saisie",
]

# Contrat de sortie : 37 colonnes (31 SAISIE + 6 dérivées). Ordre figé — des lecteurs
# (Lot10, Lot11, app) lisent par nom d'en-tête ; l'ordre reste celui du M-code documentaire.
MASTER_HEADERS: list[str] = [
    "charge_id", "date_charge", "mois",
    "montant", "sens_flux", "sens",
    "categorie_charge_id", "filtre_vue_menage", "type_flux_id",
    "code_impact", "impact_resultat_reel", "impact_resultat_comptable",
    "prise_en_compta",
    "associe_id", "mode_paiement_id", "carte_id",
    "affectation_type", "logement_id", "proprietaire_id",
    "reservation_id", "refacturable",
    "source_flux", "methode_traitement",
    "paye_avec_montant_recupere", "lien_virement_banque",
    "statut_controle", "niveau_anomalie", "code_anomalie", "statut_rapprochement",
    "justificatif", "commentaire",
    "ROW_HASH", "date_saisie",
    "source_module", "source_table", "source_pk", "date_integration",
]

# D012 — codes d'impact. IC : résultat réel ET comptable. HC : réel seul. HR : ni l'un ni l'autre.
# Même mapping que lib_lot4a_reservations_hh (IMPACT_REEL / IMPACT_COMPTA), appliqué aux charges.
IMPACT_REEL: dict[str, str] = {"IC": "OUI", "HC": "OUI", "HR": "NON"}
IMPACT_COMPTA: dict[str, str] = {"IC": "OUI", "HC": "NON", "HR": "NON"}
IMPACT_INCONNU = "A_CONTROLER"

# `sens` dérivé de `sens_flux` (M-code documentaire, requête 3).
SENS_BY_SENS_FLUX: dict[str, str] = {
    "DEPENSE": "CHARGE",
    "RECUPERATION": "PRODUIT",
    "REMBOURSEMENT": "NEUTRALISATION",
    "REFACTURATION": "CHARGE",   # charge avancée puis récupérée sur le propriétaire
    "NEUTRE": "NEUTRALISATION",
}

SOURCE_MODULE = "LOT3_CHARGES"
SOURCE_TABLE = "SAISIE_Charges_Flux"

STATUT_VALIDE = "VALIDE"
FILTRE_OUI = "OUI"


# ── Utilitaires ──────────────────────────────────────────────────────────────

def _txt(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _montant(v: Any) -> float:
    try:
        return round(float(v or 0.0), 2)
    except (TypeError, ValueError):
        return 0.0


def _est_placeholder(charge_id: str) -> bool:
    """Ligne gabarit / instruction (jamais une donnée). Ex. « [Charge par Power Query — …] »."""
    return charge_id.startswith("[") or charge_id.startswith("#")


def _date(v: Any) -> dt.date | None:
    """Normalise date_charge en `date`. Retourne None si non dérivable (jamais un faux mois)."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = _txt(v)
    if not s:
        return None
    try:
        return dt.date.fromisoformat(s[:10])
    except ValueError:
        return None


def mois_de(date_charge: Any) -> str:
    """mois = YYYY-MM dérivé de date_charge. JAMAIS le cache de la colonne formule C."""
    d = _date(date_charge)
    return d.strftime("%Y-%m") if d else ""


def row_hash(date_charge: Any, montant: Any, categorie_id: Any, mode_paiement_id: Any) -> str:
    """ROW_HASH recalculé en Python (équivalent de la formule AD, sans dépendre de son cache)."""
    d = _date(date_charge)
    if not d:
        return ""
    return (
        f"{d.strftime('%Y%m%d')}|{_montant(montant):.2f}|"
        f"{_txt(categorie_id)}|{_txt(mode_paiement_id)}"
    )


# ── Lecture des sources (lecture seule) ──────────────────────────────────────

def charger_charges_saisie(saisie_path: str) -> list[dict[str, Any]]:
    """Lit SAISIE_Charges_Flux.xlsx (onglet SAISIE). Ne garde que les vraies charges.

    Ignore : lignes vides, lignes sans charge_id, lignes gabarit/instruction.
    La source n'est jamais modifiée.
    """
    wb = openpyxl.load_workbook(saisie_path, read_only=True, data_only=True)
    try:
        ws = wb[SAISIE_SHEET]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [_txt(h) for h in rows[0]]
    idx = {name: headers.index(name) for name in SAISIE_COLS if name in headers}

    charges: list[dict[str, Any]] = []
    for r in rows[1:]:
        if not any(c is not None for c in r):
            continue
        get = lambda name: r[idx[name]] if name in idx and idx[name] < len(r) else None
        cid = _txt(get("charge_id"))
        if not cid or _est_placeholder(cid):
            continue
        charges.append({name: get(name) for name in SAISIE_COLS})
    return charges


def charger_filtre_vue_menage(ref_path: str) -> dict[str, str]:
    """Lit REF_Categories_Charges (REF_Setup) → {categorie_charge_id: filtre_vue_menage}."""
    wb = openpyxl.load_workbook(ref_path, read_only=True, data_only=True, keep_vba=True)
    try:
        if REF_CATEGORIES_SHEET not in wb.sheetnames:
            return {}
        rows = list(wb[REF_CATEGORIES_SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return {}
    headers = [_txt(h) for h in rows[0]]
    if "categorie_charge_id" not in headers or "filtre_vue_menage" not in headers:
        return {}
    i_id, i_f = headers.index("categorie_charge_id"), headers.index("filtre_vue_menage")
    out: dict[str, str] = {}
    for r in rows[1:]:
        cid = _txt(r[i_id]) if i_id < len(r) else ""
        if cid:
            out[cid] = _txt(r[i_f]) if i_f < len(r) else ""
    return out


# ── Construction déterministe des lignes MASTER ──────────────────────────────

def construire_lignes(
    charges: Iterable[dict[str, Any]],
    filtres_vue_menage: dict[str, str] | None = None,
    date_integration: Any = None,
) -> dict[str, Any]:
    """Construit les lignes MASTER (37 colonnes) + les anomalies.

    Une charge d'entrée = UNE ligne de sortie. Aucune duplication, jamais.
    `date_integration` : passer une valeur fixe rend la sortie strictement reproductible.
    """
    filtres = filtres_vue_menage or {}
    horodatage = date_integration if date_integration is not None else dt.datetime.now()
    anomalies: list[dict[str, Any]] = []
    rows: list[list[Any]] = []
    vus: set[str] = set()

    for c in charges:
        cid = _txt(c.get("charge_id"))
        if cid in vus:
            anomalies.append({
                "code": "CHARGE_ID_DOUBLON", "niveau": "BLOQUANT",
                "detail": f"charge_id={cid} présent en double dans la SAISIE — première occurrence gardée",
            })
            continue
        vus.add(cid)

        mois = mois_de(c.get("date_charge"))
        if not mois:
            anomalies.append({
                "code": "DATE_CHARGE_INVALIDE", "niveau": "BLOQUANT",
                "detail": f"charge_id={cid} : date_charge={c.get('date_charge')!r} non exploitable — "
                          f"mois non dérivable (la charge ne sera rattachée à aucun mois)",
            })

        code_impact = _txt(c.get("code_impact")).upper()
        if code_impact and code_impact not in IMPACT_REEL:
            anomalies.append({
                "code": "CODE_IMPACT_INCONNU", "niveau": "A_CONTROLER",
                "detail": f"charge_id={cid} : code_impact={code_impact!r} hors IC/HC/HR",
            })

        sens_flux = _txt(c.get("sens_flux")).upper()
        sens = SENS_BY_SENS_FLUX.get(sens_flux)
        if sens is None and sens_flux:
            anomalies.append({
                "code": "SENS_FLUX_INCONNU", "niveau": "A_CONTROLER",
                "detail": f"charge_id={cid} : sens_flux={sens_flux!r} non mappé — sens laissé vide",
            })

        categorie = _txt(c.get("categorie_charge_id"))
        valeurs: dict[str, Any] = dict(c)
        valeurs["mois"] = mois
        valeurs["impact_resultat_reel"] = IMPACT_REEL.get(code_impact, IMPACT_INCONNU)
        valeurs["impact_resultat_comptable"] = IMPACT_COMPTA.get(code_impact, IMPACT_INCONNU)
        valeurs["ROW_HASH"] = row_hash(
            c.get("date_charge"), c.get("montant"), categorie, c.get("mode_paiement_id")
        )
        valeurs["montant"] = _montant(c.get("montant"))
        valeurs["sens"] = sens
        valeurs["filtre_vue_menage"] = filtres.get(categorie, "")
        valeurs["source_module"] = SOURCE_MODULE
        valeurs["source_table"] = SOURCE_TABLE
        valeurs["source_pk"] = cid
        valeurs["date_integration"] = horodatage

        rows.append([valeurs.get(h) for h in MASTER_HEADERS])

    return {"rows": rows, "anomalies": anomalies}


def lignes_vue_menage(rows: list[list[Any]]) -> list[list[Any]]:
    """VUE_MENAGE = MASTER filtré sur filtre_vue_menage=OUI ET statut_controle=VALIDE (D028)."""
    i_f = MASTER_HEADERS.index("filtre_vue_menage")
    i_s = MASTER_HEADERS.index("statut_controle")
    return [
        r for r in rows
        if _txt(r[i_f]).upper() == FILTRE_OUI and _txt(r[i_s]).upper() == STATUT_VALIDE
    ]


# ── Écriture idempotente du classeur cible ───────────────────────────────────

def _ecrire_onglet(wb: Any, sheet: str, rows: list[list[Any]]) -> None:
    """Recrée l'onglet à sa position d'origine, avec l'en-tête canonique puis les lignes."""
    idx = wb.sheetnames.index(sheet) if sheet in wb.sheetnames else None
    if idx is not None:
        del wb[sheet]
        ws = wb.create_sheet(sheet, idx)
    else:
        ws = wb.create_sheet(sheet)
    ws.append(MASTER_HEADERS)
    for row in rows:
        ws.append(row)


def ecrire_master(sortie_path: str, rows: list[list[Any]]) -> None:
    """Écrit MASTER + VUE_MENAGE dans le classeur cible. Idempotent (jamais d'accumulation).

    Les autres onglets (dont POWER_QUERY_CODE, documentaire) sont préservés.
    """
    wb = openpyxl.load_workbook(sortie_path)
    try:
        _ecrire_onglet(wb, MASTER_SHEET, rows)
        _ecrire_onglet(wb, VUE_MENAGE_SHEET, lignes_vue_menage(rows))
        wb.save(sortie_path)
    finally:
        wb.close()


def generer(
    saisie_path: str,
    ref_path: str,
    sortie_path: str,
    date_integration: Any = None,
) -> dict[str, Any]:
    """Orchestrateur : lit la SAISIE + REF_Setup, construit les lignes, écrit le classeur cible.

    `sortie_path` est le classeur MASTER à produire (en test : toujours une copie).
    Retourne { nb_lignes, nb_vue_menage, anomalies, cible }.
    """
    charges = charger_charges_saisie(saisie_path)
    filtres = charger_filtre_vue_menage(ref_path)
    res = construire_lignes(charges, filtres, date_integration=date_integration)
    ecrire_master(sortie_path, res["rows"])
    return {
        "nb_lignes": len(res["rows"]),
        "nb_vue_menage": len(lignes_vue_menage(res["rows"])),
        "anomalies": res["anomalies"],
        "cible": sortie_path,
    }
