import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def write_book(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
    wb.save(path)
    wb.close()


def read_rows(path: Path, sheet: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    headers = list(next(it))
    rows = [dict(zip(headers, row)) for row in it]
    wb.close()
    return rows


class GuestCountCanapeIntegrationTest(unittest.TestCase):
    def test_guestcount_canape_end_to_end_in_temp_project(self):
        with tempfile.TemporaryDirectory() as tmp:
            project = Path(tmp) / "Pilotage_Conciergerie"
            work = project / "02_TRAVAIL"
            work.mkdir(parents=True)
            (project / "02_TRAVAIL/Lot11_Controles").mkdir(parents=True)
            for path in (ROOT / "02_TRAVAIL").glob("*.py"):
                shutil.copy2(path, work / path.name)

            self._create_sources(project)

            scripts = [
                "lot4bis_charger_reservations.py",
                "lot4quater_resoudre_source_reservations.py",
            ]
            for script in scripts:
                subprocess.run([sys.executable, str(work / script)], cwd=project, check=True)

            self._create_flux_from_resolved(project)

            for script in [
                "lot10_calculer_resultats.py",
                "lot11_controles_coherence.py",
                "lot12_generer_factures.py",
            ]:
                subprocess.run([sys.executable, str(work / script)], cwd=project, check=True)

            comm = read_rows(project / "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_Commissions.xlsx", "COMMISSIONS")
            ac = read_rows(project / "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_Commissions.xlsx", "A_CONTROLER")
            reg = read_rows(project / "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_NetProprietaire.xlsx", "REGLEMENT")
            lines = read_rows(project / "02_TRAVAIL/Lot12_Factures/MASTER_FACT_Proprietaires.xlsx", "FACT_FACTURE_LIGNES")
            ctrl = read_rows(project / "02_TRAVAIL/Lot11_Controles/MASTER_CTRL_Coherence.xlsx", "MASTER")

            calc_rows = read_rows(project / "02_TRAVAIL/Lot4bis_TableCommune/MASTER_CALC_Reservations.xlsx", "MASTER")
            resolved_rows = read_rows(project / "02_TRAVAIL/Lot4quater_SourceResolue/MASTER_CALC_Reservations_Resolues.xlsx", "MASTER")
            calc_a = next(r for r in calc_rows if str(r["reservation_id_hostaway"]).split(".")[0] == "1001")
            resolved_a = next(r for r in resolved_rows if str(r["reservation_id_hostaway"]).split(".")[0] == "1001")
            self.assertEqual(float(calc_a["guestCount"]), 3.0)
            self.assertEqual(calc_a["source_guestCount"], "API_LIST")
            self.assertEqual(float(resolved_a["guestCount"]), 3.0)
            self.assertEqual(resolved_a["source_guestCount"], "API_LIST")

            row_a = next(r for r in comm if str(r["reservation_id_hostaway"]).split(".")[0] == "1001")
            self.assertEqual(float(row_a["guestCount"]), 3.0)
            self.assertEqual(float(row_a["preparation_canape_voyageurs"]), 10.0)
            self.assertEqual(row_a["controle_preparation_canape"], "OK")
            self.assertEqual(float(row_a["assiette_commission"]), 80.0)
            self.assertEqual(float(row_a["commission_conciergerie"]), 20.0)
            self.assertEqual(float(row_a["net_proprietaire"]), 50.0)

            row_b = next(r for r in comm if str(r["reservation_id_hostaway"]).split(".")[0] == "1002")
            self.assertIsNone(row_b["guestCount"])
            self.assertEqual(float(row_b["preparation_canape_voyageurs"]), 0.0)
            self.assertEqual(row_b["controle_preparation_canape"], "A_CONTROLER")
            self.assertEqual(float(row_b["net_proprietaire"]), 60.0)

            ac_codes = {r.get("code_anomalie_lot10") for r in ac}
            self.assertIn("GUEST_COUNT_MANQUANT_PREPARATION_CANAPE", ac_codes)
            ctrl_codes = {r.get("code_controle") for r in ctrl}
            self.assertIn("GUEST_COUNT_MANQUANT_PREPARATION_CANAPE", ctrl_codes)

            reg_june = next(r for r in reg if r["mois"] == "2026-06" and r["logement_id"] == "LOG_0006")
            reg_july = next(r for r in reg if r["mois"] == "2026-07" and r["logement_id"] == "LOG_0006")
            self.assertEqual(float(reg_june["total_preparation_canape_mois"]), 10.0)
            self.assertEqual(float(reg_june["montant_du_conciergerie"]), 50.0)
            self.assertEqual(float(reg_june["net_proprietaire_apres_charge_mois"]), 50.0)
            self.assertEqual(float(reg_july["total_preparation_canape_mois"]), 0.0)
            self.assertEqual(float(reg_july["montant_du_conciergerie"]), 40.0)

            canape_lines = [r for r in lines if r["type_ligne"] == "PREPARATION_CANAPE"]
            self.assertEqual(len(canape_lines), 1)
            self.assertEqual(float(canape_lines[0]["montant"]), 10.0)
            self.assertEqual(canape_lines[0]["libelle"], "Préparation du canapé payée par les voyageurs")

            self.assertFalse(any(str(r.get("reservation_id_hostaway")).split(".")[0] == "1003" for r in comm))
            self.assertFalse(any(str(r.get("reservation_id_hostaway")).split(".")[0] == "1004" for r in comm))
            self.assertFalse(any(r.get("logement_id") == "LOGEMENT_DIVERS" for r in reg))

    def _create_sources(self, project: Path) -> None:
        ref = project / "01_SOURCES_BRUTES/REF_Setup/REF_Setup.xlsm"
        write_book(ref, {
            "REF_Logements": (
                ["logement_id", "nom_logement_officiel", "nom_court", "ville", "type_logement_id",
                 "sur_hostaway", "actif", "statut_parc", "forfait_logiciel_consommables_mensuel",
                 "seuil_voyageurs_preparation_canape", "montant_preparation_canape"],
                [
                    {"logement_id": "LOG_0006", "nom_logement_officiel": "Canape 3", "nom_court": "C3",
                     "ville": "X", "type_logement_id": "T2", "sur_hostaway": "OUI", "actif": "OUI",
                     "statut_parc": "GERE", "forfait_logiciel_consommables_mensuel": 0,
                     "seuil_voyageurs_preparation_canape": 3, "montant_preparation_canape": 10},
                    {"logement_id": "LOGEMENT_DIVERS", "nom_logement_officiel": "Divers", "nom_court": "DIV",
                     "ville": "X", "type_logement_id": "TECH", "sur_hostaway": "NON", "actif": "OUI",
                     "statut_parc": "HORS_PARC_TECHNIQUE", "forfait_logiciel_consommables_mensuel": 0},
                ],
            ),
            "REF_Proprietaires": (
                ["proprietaire_id", "prenom_proprietaire", "nom_proprietaire", "adresse_facturation",
                 "mode_facturation", "actif"],
                [{"proprietaire_id": "PROP_1", "prenom_proprietaire": "P", "nom_proprietaire": "One",
                  "adresse_facturation": "1 rue", "mode_facturation": "EMAIL", "actif": "OUI"}],
            ),
            "REF_Mapping_Logements": (
                ["mapping_logement_id", "source", "champ_source", "valeur_source", "logement_id", "actif"],
                [
                    {"mapping_logement_id": "MAP1", "source": "Hostaway", "champ_source": "listingMapId", "valeur_source": 101, "logement_id": "LOG_0006", "actif": "OUI"},
                    {"mapping_logement_id": "MAP2", "source": "Hostaway", "champ_source": "listingMapId", "valeur_source": 104, "logement_id": "LOGEMENT_DIVERS", "actif": "OUI"},
                ],
            ),
            "REF_Taux_Commission": (
                ["taux_commission_id", "proprietaire_id", "logement_id", "taux_commission",
                 "date_debut", "date_fin", "actif", "justification", "commentaire"],
                [{"taux_commission_id": "TX1", "proprietaire_id": "PROP_1", "logement_id": "LOG_0006",
                  "taux_commission": 0.25, "date_debut": "2020-01-01", "date_fin": None,
                  "actif": "OUI", "justification": "test", "commentaire": ""}],
            ),
            "REF_Gestion_Logements_Hist": (
                ["gestion_id", "logement_id", "proprietaire_id", "date_debut", "date_fin",
                 "statut_gestion", "source", "commentaire", "actif"],
                [{"gestion_id": "G1", "logement_id": "LOG_0006", "proprietaire_id": "PROP_1",
                  "date_debut": "2020-01-01", "date_fin": None, "statut_gestion": "ACTIVE",
                  "source": "test", "commentaire": "", "actif": "OUI"}],
            ),
            "REF_Cloture_Mensuelle": (["mois", "statut_mois"], []),
        })

        ha_headers = ["reservation_id", "listingMapId", "source", "channel_type", "source_financiere",
                      "status", "paymentStatus", "checkInDate", "checkOutDate", "nights", "numberOfGuests", "guestCount", "source_guestCount",
                      "guestName", "totalPrice", "cleaningFee_res", "channelCommission",
                      "airbnbExpectedPayout", "is_ownerStay", "inclure_resultat", "updatedOn", "createdOn",
                      "extrait_le", "ROW_HASH"]
        write_book(project / "02_TRAVAIL/Lot1_Hostaway/MASTER_FACT_HA_Reservations.xlsx", {
            "data": (ha_headers, [
                {"reservation_id": 1001, "listingMapId": 101, "source": "airbnbOfficial", "channel_type": "AIRBNB",
                 "source_financiere": "AIRBNB", "status": "new", "checkInDate": "2026-06-10",
                 "checkOutDate": "2026-06-12", "nights": 2, "numberOfGuests": 3, "guestCount": 3,
                 "source_guestCount": "API_LIST", "totalPrice": 100, "inclure_resultat": "OUI"},
                {"reservation_id": 1002, "listingMapId": 101, "source": "airbnbOfficial", "channel_type": "AIRBNB",
                 "source_financiere": "AIRBNB", "status": "new", "checkInDate": "2026-07-10",
                 "checkOutDate": "2026-07-12", "nights": 2, "totalPrice": 100, "inclure_resultat": "OUI"},
                {"reservation_id": 1003, "listingMapId": 101, "source": "airbnbOfficial", "channel_type": "AIRBNB",
                 "source_financiere": "AIRBNB", "status": "ownerStay", "checkInDate": "2026-08-10",
                 "checkOutDate": "2026-08-12", "nights": 2, "totalPrice": 100, "inclure_resultat": "NON"},
                {"reservation_id": 1004, "listingMapId": 104, "source": "airbnbOfficial", "channel_type": "AIRBNB",
                 "source_financiere": "AIRBNB", "status": "new", "checkInDate": "2026-09-10",
                 "checkOutDate": "2026-09-12", "nights": 2, "totalPrice": 100, "inclure_resultat": "OUI"},
            ]),
        })
        write_book(project / "02_TRAVAIL/Lot1_Hostaway/MASTER_FACT_HA_ReservationDetails.xlsx", {
            "data": (["reservation_id", "json_snapshot", "extrait_le", "ROW_HASH"], [
                {"reservation_id": 1001, "json_snapshot": '{"id":1001,"hostawayReservationId":"1001","numberOfGuests":3'},
            ]),
        })
        pay_headers = ["reservation_id", "listingMapId", "source", "channel_type", "statut_calcul_payout",
                       "payout_calcule", "source_payout", "menage_retenu", "assiette_commission",
                       "inclure_resultat_auto", "menage_retenu_source", "cout_standard_id",
                       "cout_standard_menage_snapshot", "cout_standard_date_debut_validite",
                       "cout_standard_date_fin_validite", "logement_id_snapshot", "type_logement_id_snapshot",
                       "date_reference_cout_menage"]
        write_book(project / "02_TRAVAIL/Lot1_Hostaway/MASTER_CALC_HA_Payout.xlsx", {
            "data": (pay_headers, [
                {"reservation_id": 1001, "listingMapId": 101, "channel_type": "AIRBNB", "statut_calcul_payout": "NORMAL",
                 "payout_calcule": 100, "source_payout": "test", "menage_retenu": 20, "assiette_commission": 80,
                 "inclure_resultat_auto": "OUI", "menage_retenu_source": "TEST", "logement_id_snapshot": "LOG_0006",
                 "type_logement_id_snapshot": "T2", "date_reference_cout_menage": "2026-06-10"},
                {"reservation_id": 1002, "listingMapId": 101, "channel_type": "AIRBNB", "statut_calcul_payout": "NORMAL",
                 "payout_calcule": 100, "source_payout": "test", "menage_retenu": 20, "assiette_commission": 80,
                 "inclure_resultat_auto": "OUI", "menage_retenu_source": "TEST", "logement_id_snapshot": "LOG_0006",
                 "type_logement_id_snapshot": "T2", "date_reference_cout_menage": "2026-07-10"},
            ]),
        })
        write_book(project / "02_TRAVAIL/Lot1_Hostaway/MASTER_CTRL_HA_Anomalies.xlsx", {
            "data": (["reservation_id", "code_anomalie", "severite", "description", "statut", "date_detection", "ROW_HASH"], []),
        })
        write_book(project / "02_TRAVAIL/Lot4_ReservationsHH/MASTER_FACT_MAN_ReservationsHorsHostaway.xlsx", {
            "MASTER": (["reservation_hh_id", "statut_controle", "reservation_id_hostaway", "guestCount"], []),
        })
        write_book(project / "02_TRAVAIL/Lot4bis_TableCommune/MASTER_CALC_Reservations.xlsx", {
            "MASTER": (["placeholder"], []),
            "VUE_FLUX": (["placeholder"], []),
            "POWER_QUERY_CODE": (["placeholder"], []),
        })
        for rel, sheet, headers in [
            ("02_TRAVAIL/Lot3_Charges/MASTER_FACT_MAN_Charges.xlsx", "MASTER", ["charge_id", "statut_controle"]),
            ("02_TRAVAIL/Lot6c_MenagesExternes/MASTER_FACT_MEN_MenagesExternes.xlsx", "MASTER", ["menage_id", "statut_controle"]),
            ("02_DONNEES_NORMALISEES/menages/M04_MENAGES_PowerQuery.xlsx", "MASTER", ["menage_calc_id", "statut_controle"]),
            ("02_TRAVAIL/Lot7_IK_Avantages/MASTER_FACT_MAN_IK_Avantages.xlsx", "MASTER_CALC_AVANTAGES", ["pk_id", "statut_controle"]),
        ]:
            write_book(project / rel, {sheet: (headers, [])})
        write_book(project / "02_TRAVAIL/Lot5_AcomptesProprietaires/MASTER_FACT_MAN_AcomptesProprietaires.xlsx", {
            "MASTER": (
                ["acompte_id", "mois", "logement_id", "proprietaire_id", "montant_acompte", "statut_controle"],
                [{"acompte_id": "ACC1", "mois": "2026-06", "logement_id": "LOG_0006",
                  "proprietaire_id": "PROP_1", "montant_acompte": 1, "statut_controle": "VALIDE"}],
            ),
        })

    def _create_flux_from_resolved(self, project: Path) -> None:
        vue = read_rows(project / "02_TRAVAIL/Lot4quater_SourceResolue/MASTER_CALC_Reservations_Resolues.xlsx", "VUE_FLUX")
        rows = []
        for i, r in enumerate(vue, 1):
            rows.append({
                "flux_id": f"FLUX-{i}",
                "ROW_HASH": f"H{i}",
                "source_module": "test",
                "source_table": "MASTER_CALC_Reservations_Resolues",
                "source_pk": r["reservation_calc_id"],
                "date_flux": r["date_arrivee"],
                "mois": r["mois"],
                "logement_id": r["logement_id"],
                "proprietaire_id": r["proprietaire_id"],
                "associe_id": None,
                "type_flux_id": "TYPE_FLUX_017",
                "sens": "PRODUIT",
                "montant": r["montant_retenu"],
                "code_impact": r["code_impact"],
                "inclure_resultat_reel": "OUI",
                "inclure_resultat_comptable": "OUI",
                "inclure_resultat_hors_compta": "NON",
                "statut_controle": "VALIDE",
                "niveau_anomalie": "INFO",
                "code_anomalie": None,
                "commentaire": "test",
                "date_integration": "2026-06-30",
            })
        write_book(project / "02_TRAVAIL/Lot9_FluxUnifie/MASTER_CALC_Flux.xlsx", {
            "MASTER": ([
                "flux_id", "ROW_HASH", "source_module", "source_table", "source_pk",
                "date_flux", "mois", "logement_id", "proprietaire_id", "associe_id",
                "type_flux_id", "sens", "montant", "code_impact",
                "inclure_resultat_reel", "inclure_resultat_comptable", "inclure_resultat_hors_compta",
                "statut_controle", "niveau_anomalie", "code_anomalie", "commentaire", "date_integration",
            ], rows),
        })


if __name__ == "__main__":
    unittest.main()
