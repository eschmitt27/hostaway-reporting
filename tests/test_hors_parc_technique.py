import sys
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "02_TRAVAIL"))

from lib_parc import (
    A_CONTROLER,
    GERE,
    HORS_PARC_TECHNIQUE,
    STATUT_PARC_INVALIDE,
    is_gere,
    is_hors_parc_technique,
    is_statut_parc_a_controler,
    normalise_statut_parc,
    statut_parc_traitement,
)
from lot10_calculer_resultats import build_charge_fixe, build_commissions


PAYOUT_COLS = [
    "reservation_id", "channel_type", "statut_calcul_payout", "payout_calcule", "source_payout",
    "menage_retenu", "assiette_commission", "inclure_resultat_auto", "menage_retenu_source",
    "cout_standard_id", "cout_standard_menage_snapshot", "cout_standard_date_debut_validite",
    "cout_standard_date_fin_validite", "logement_id_snapshot", "type_logement_id_snapshot",
    "date_reference_cout_menage",
]


def _reservation_frame(logement_id):
    return pd.DataFrame([{
        "reservation_calc_id": "RES-2026-05-HA-001",
        "reservation_id_hostaway": 123,
        "reservation_hh_id": None,
        "source": "HOSTAWAY_AIRBNB",
        "source_montant": "HOSTAWAY_PAYOUT",
        "montant_retenu": 100.0,
        "logement_id": logement_id,
        "proprietaire_id": "PROP_1",
        "date_arrivee": "2026-05-10",
        "date_depart": "2026-05-12",
        "nuits": 2,
        "guestCount": 2,
        "canal": "AIRBNB",
        "payout_calcule": 100.0,
        "menage_retenu": 20.0,
        "assiette_commission": 80.0,
    }])


def _flux_frame(logement_id, proprietaire_id="PROP_1"):
    return pd.DataFrame([{
        "type_flux_id": "TYPE_FLUX_017",
        "source_pk": "RES-2026-05-HA-001",
        "flux_id": "FLUX_1",
        "mois": "2026-05",
        "logement_id": logement_id,
        "proprietaire_id": proprietaire_id,
    }])


class HorsParcTechniqueTests(unittest.TestCase):
    def test_ref_setup_marks_only_technical_codes_hors_parc(self):
        wb = load_workbook(ROOT / "01_SOURCES_BRUTES" / "REF_Setup" / "REF_Setup.xlsm", read_only=True, data_only=True, keep_vba=True)
        ws = wb["REF_Logements"]
        rows = list(ws.iter_rows(values_only=True))
        headers = list(rows[0])
        self.assertIn("statut_parc", headers)
        idx_id = headers.index("logement_id")
        idx_statut = headers.index("statut_parc")
        values = {row[idx_id]: row[idx_statut] for row in rows[1:] if row[idx_id]}
        self.assertEqual(values["APPARTEMENT_DIVERS"], HORS_PARC_TECHNIQUE)
        self.assertEqual(values["LOGEMENT_DIVERS"], HORS_PARC_TECHNIQUE)
        for logement_id, statut in values.items():
            if logement_id not in {"APPARTEMENT_DIVERS", "LOGEMENT_DIVERS"}:
                self.assertEqual(statut, GERE)
        wb.close()

    def test_helper_distinguishes_gere_hors_parc_and_invalid(self):
        self.assertEqual(normalise_statut_parc(" gere "), GERE)
        self.assertEqual(statut_parc_traitement({"statut_parc": GERE}), GERE)
        self.assertTrue(is_gere({"statut_parc": GERE}))
        self.assertEqual(statut_parc_traitement({"statut_parc": HORS_PARC_TECHNIQUE}), HORS_PARC_TECHNIQUE)
        self.assertTrue(is_hors_parc_technique({"statut_parc": HORS_PARC_TECHNIQUE}))
        for row in ({"statut_parc": None}, {"statut_parc": ""}, {"statut_parc": "INCONNU"}, {"logement_id": "LOG_1"}, None):
            self.assertEqual(statut_parc_traitement(row), A_CONTROLER)
            self.assertTrue(is_statut_parc_a_controler(row))
            self.assertFalse(is_gere(row))
            self.assertFalse(is_hors_parc_technique(row))

    def test_gere_reservation_keeps_normal_commission_path_when_rate_exists(self):
        df_log = pd.DataFrame([{"logement_id": "LOG_GERE", "statut_parc": GERE}])
        df_taux = pd.DataFrame([{
            "taux_commission_id": "TC_1",
            "proprietaire_id": "PROP_1",
            "logement_id": "LOG_GERE",
            "taux_commission": 0.2,
            "date_debut": "2026-01-01",
            "date_fin": None,
            "actif": "OUI",
        }])
        df_comm, df_ac, _ = build_commissions(
            _flux_frame("LOG_GERE"),
            _reservation_frame("LOG_GERE"),
            pd.DataFrame([{
                "reservation_id": 123,
                "channel_type": "AIRBNB",
                "statut_calcul_payout": "NORMAL",
                "payout_calcule": 100.0,
                "source_payout": "HOSTAWAY_PAYOUT",
                "menage_retenu": 20.0,
                "assiette_commission": 80.0,
                "inclure_resultat_auto": "OUI",
                "menage_retenu_source": "TEST",
                "cout_standard_id": None,
                "cout_standard_menage_snapshot": None,
                "cout_standard_date_debut_validite": None,
                "cout_standard_date_fin_validite": None,
                "logement_id_snapshot": "LOG_GERE",
                "type_logement_id_snapshot": None,
                "date_reference_cout_menage": "2026-05-10",
            }]),
            pd.DataFrame(columns=["reservation_hh_id", "total_percu", "menage", "commission", "taux_commission"]),
            df_log,
            pd.DataFrame(),
            df_taux,
        )
        self.assertEqual(len(df_comm), 1)
        self.assertNotIn(STATUT_PARC_INVALIDE, set(df_ac.get("code_anomalie_lot10", pd.Series(dtype=object)).dropna()))

    def test_hors_parc_reservation_does_not_generate_commission_or_net(self):
        df_log = pd.DataFrame([{"logement_id": "APPARTEMENT_DIVERS", "statut_parc": HORS_PARC_TECHNIQUE}])
        df_comm, df_ac, _ = build_commissions(
            _flux_frame("APPARTEMENT_DIVERS", None),
            _reservation_frame("APPARTEMENT_DIVERS"),
            pd.DataFrame(columns=PAYOUT_COLS),
            pd.DataFrame(columns=["reservation_hh_id", "total_percu", "menage", "commission", "taux_commission"]),
            df_log,
            pd.DataFrame(),
            pd.DataFrame(),
        )
        self.assertEqual(len(df_comm), 0)
        self.assertIn(HORS_PARC_TECHNIQUE, set(df_ac["code_anomalie_lot10"].dropna()))

    def test_empty_statut_parc_is_a_controler_without_economic_calculation(self):
        df_log = pd.DataFrame([{"logement_id": "LOG_EMPTY", "statut_parc": ""}])
        df_comm, df_ac, _ = build_commissions(
            _flux_frame("LOG_EMPTY"),
            _reservation_frame("LOG_EMPTY"),
            pd.DataFrame(columns=PAYOUT_COLS),
            pd.DataFrame(columns=["reservation_hh_id", "total_percu", "menage", "commission", "taux_commission"]),
            df_log,
            pd.DataFrame(),
            pd.DataFrame(),
        )
        self.assertEqual(len(df_comm), 0)
        self.assertIn(STATUT_PARC_INVALIDE, set(df_ac["code_anomalie_lot10"].dropna()))
        self.assertIn(A_CONTROLER, set(df_ac["niveau"].dropna()))

    def test_unknown_statut_parc_is_a_controler_without_economic_calculation(self):
        df_log = pd.DataFrame([{"logement_id": "LOG_UNKNOWN", "statut_parc": "AUTRE"}])
        df_comm, df_ac, _ = build_commissions(
            _flux_frame("LOG_UNKNOWN"),
            _reservation_frame("LOG_UNKNOWN"),
            pd.DataFrame(columns=PAYOUT_COLS),
            pd.DataFrame(columns=["reservation_hh_id", "total_percu", "menage", "commission", "taux_commission"]),
            df_log,
            pd.DataFrame(),
            pd.DataFrame(),
        )
        self.assertEqual(len(df_comm), 0)
        self.assertIn(STATUT_PARC_INVALIDE, set(df_ac["code_anomalie_lot10"].dropna()))

    def test_invalid_statut_parc_blocks_charge_fixe(self):
        df_cfix, controls = build_charge_fixe(_flux_frame("LOG_BAD"), pd.DataFrame([{
            "logement_id": "LOG_BAD",
            "statut_parc": "BAD",
            "forfait_logiciel_consommables_mensuel": 35,
            "actif": "OUI",
        }]))
        self.assertEqual(len(df_cfix), 0)
        self.assertIn(STATUT_PARC_INVALIDE, {c["code_anomalie"] for c in controls})
        self.assertIn(A_CONTROLER, {c["niveau"] for c in controls})

    def test_gere_logement_without_owner_remains_blocking_for_charge_fixe(self):
        df_flux = pd.DataFrame([{
            "type_flux_id": "TYPE_FLUX_017",
            "mois": "2026-05",
            "logement_id": "LOG_GERE",
            "proprietaire_id": None,
        }])
        df_log = pd.DataFrame([{
            "logement_id": "LOG_GERE",
            "statut_parc": GERE,
            "forfait_logiciel_consommables_mensuel": 35,
            "actif": "OUI",
        }])
        df_gest = pd.DataFrame([{
            "gestion_id": "GST_NO_OWNER",
            "logement_id": "LOG_GERE",
            "proprietaire_id": None,
            "date_debut": "2026-01-01",
            "date_fin": None,
            "statut_gestion": "ACTIF",
        }])
        df_cfix, controls = build_charge_fixe(df_flux, df_log, df_gest)
        self.assertEqual(len(df_cfix), 0)
        self.assertIn("GESTION_LOGEMENT_MISSING_OWNER", {c["code_anomalie"] for c in controls})
        self.assertIn("BLOQUANT", {c["niveau"] for c in controls})

    def test_lot12_excludes_hors_parc_and_invalid_before_invoice_creation(self):
        lot12 = (ROOT / "02_TRAVAIL" / "lot12_generer_factures.py").read_text(encoding="utf-8")
        self.assertIn("is_hors_parc_technique(log_row)", lot12)
        self.assertIn("is_statut_parc_a_controler(log_row)", lot12)
        self.assertIn("STATUT_PARC_INVALIDE", lot12)
        self.assertIn("continue", lot12)


if __name__ == "__main__":
    unittest.main()


class SuppressionDatesGestionTests(unittest.TestCase):
    def test_ref_logements_no_longer_contains_duplicate_management_dates(self):
        wb = load_workbook(ROOT / "01_SOURCES_BRUTES" / "REF_Setup" / "REF_Setup.xlsm", read_only=True, data_only=True, keep_vba=True)
        ws = wb["REF_Logements"]
        headers = list(next(ws.iter_rows(values_only=True)))
        wb.close()
        self.assertNotIn("date_entree_gestion", headers)
        self.assertNotIn("date_sortie_gestion", headers)
        self.assertIn("statut_parc", headers)

    def test_active_scripts_do_not_fallback_to_old_management_date_columns(self):
        for rel in [
            "lot4bis_charger_reservations.py",
            "lot6c_menages_externes.py",
            "lot10_calculer_resultats.py",
            "lot11_controles_coherence.py",
            "lot13_export_powerbi.py",
        ]:
            text = (ROOT / "02_TRAVAIL" / rel).read_text(encoding="utf-8")
            self.assertNotIn("date_entree_gestion", text, rel)
            self.assertNotIn("date_sortie_gestion", text, rel)
