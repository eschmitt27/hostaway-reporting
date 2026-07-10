"""Lot7 — preuve du GÉNÉRATEUR RÉEL (Option A, Python/openpyxl) sur COPIES contrôlées.

Aucune donnée métier réelle : chaque test construit ses fixtures en tmp, exécute le générateur
et lit MASTER_CALC_AVANTAGES produit. Couvre les 5 cas exigés + idempotence.
"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "02_TRAVAIL"))

import lot7_generateur_avantages as gen  # noqa: E402

SAISIE_COLS = ["charge_id", "date_charge", "mois", "montant", "type_flux_id",
               "associe_id", "mode_paiement_id", "avantage_associe_id"]
SOURCE_COLS = ["mois", "associe_id", "type_flux_id", "type_remboursement",
               "nature", "montant", "mode_paiement_id", "lien_origine", "commentaire"]


def _make_saisie(path, charges):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    ws.append(SAISIE_COLS)
    for c in charges:
        ws.append([c.get(k) for k in SAISIE_COLS])
    wb.save(path)
    wb.close()


def _make_lot7(path, source_rows=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SOURCE_SAISIE"
    ws.append(SOURCE_COLS)
    # Ligne d'instruction (doit être ignorée : mois non YYYY-MM).
    ws.append(["AAAA-MM", "PERS_EWAN ou PERS_WAFA", "", "", "", None, "", "", ""])
    for s in source_rows:
        ws.append([s.get(k) for k in SOURCE_COLS])
    wb.save(path)
    wb.close()


def _read_calc(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["MASTER_CALC_AVANTAGES"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    headers = list(rows[0])
    out = {}
    for r in rows[1:]:
        d = dict(zip(headers, r))
        out[(d["associe_id"], d["mois"])] = d
    return out


class Lot7GenerateurTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.saisie = os.path.join(self.tmp, "SAISIE_Charges_Flux.xlsx")
        self.lot7 = os.path.join(self.tmp, "MASTER_FACT_MAN_IK_Avantages.xlsx")

    def _run(self):
        return gen.generer(self.saisie, self.lot7, self.lot7)

    # ── Cas 1 — Banque pro (PAY_001) + avantage explicite ────────────────────
    def test_cas1_banque_pro_avantage_explicite_idempotent(self):
        _make_saisie(self.saisie, [
            {"charge_id": "C1", "date_charge": "2026-06-10", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_EWAN"},
        ])
        _make_lot7(self.lot7)
        self._run()
        calc = _read_calc(self.lot7)
        self.assertEqual(calc[("PERS_EWAN", "2026-06")]["avantage_brut_depenses_perso"], 100.0)
        self.assertEqual(calc[("PERS_EWAN", "2026-06")]["avantages_bruts_total"], 100.0)
        # exactement une ligne pour cet associé/mois
        self.assertEqual(sum(1 for (a, m) in calc if a == "PERS_EWAN"), 1)
        # 2e exécution → toujours 100, jamais 200
        self._run()
        calc2 = _read_calc(self.lot7)
        self.assertEqual(calc2[("PERS_EWAN", "2026-06")]["avantage_brut_depenses_perso"], 100.0)
        self.assertEqual(len(calc2), len(calc))

    # ── Cas 2 — Carte associée (PAY_003/004) sans avantage explicite ─────────
    def test_cas2_carte_sans_avantage_aucun_avantage(self):
        _make_saisie(self.saisie, [
            {"charge_id": "C2", "date_charge": "2026-06-11", "montant": 50,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "PERS_WAFA",
             "mode_paiement_id": "PAY_003", "avantage_associe_id": ""},
        ])
        _make_lot7(self.lot7)
        res = self._run()
        calc = _read_calc(self.lot7)
        # aucun avantage automatique créé (ni pour le payeur, ni pour personne)
        self.assertEqual(calc, {})
        self.assertEqual(res["nb_lignes"], 0)

    # ── Cas 3 — Carte associée + avantage explicite (bénéficiaire ≠ payeur) ──
    def test_cas3_carte_avantage_explicite_beneficiaire_prime(self):
        _make_saisie(self.saisie, [
            {"charge_id": "C3", "date_charge": "2026-06-12", "montant": 80,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "PERS_WAFA",
             "mode_paiement_id": "PAY_004", "avantage_associe_id": "PERS_EWAN"},
        ])
        _make_lot7(self.lot7)
        self._run()
        calc = _read_calc(self.lot7)
        # avantage attribué au bénéficiaire EWAN, pas au payeur WAFA
        self.assertEqual(calc[("PERS_EWAN", "2026-06")]["avantage_brut_depenses_perso"], 80.0)
        self.assertNotIn(("PERS_WAFA", "2026-06"), calc)

    # ── Cas 4 — SOURCE_SAISIE résiduelle (autonome) + anti-double ────────────
    def test_cas4_source_saisie_residuelle_prise_en_compte(self):
        # charge C1 en Lot3 ; une ligne résiduelle autonome (virement) sans lien_origine.
        _make_saisie(self.saisie, [
            {"charge_id": "C1", "date_charge": "2026-06-10", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_EWAN"},
        ])
        _make_lot7(self.lot7, source_rows=[
            {"mois": "2026-06", "associe_id": "PERS_WAFA", "type_flux_id": "TYPE_FLUX_001",
             "montant": 200, "lien_origine": ""},
        ])
        res = self._run()
        calc = _read_calc(self.lot7)
        # résiduel WAFA (virement autonome) pris en compte, sans interférer avec l'avantage charge EWAN
        self.assertEqual(calc[("PERS_WAFA", "2026-06")]["avantage_brut_virements"], 200.0)
        self.assertEqual(calc[("PERS_EWAN", "2026-06")]["avantage_brut_depenses_perso"], 100.0)
        self.assertEqual(res["anomalies"], [])

    def test_cas4b_source_saisie_lien_origine_existant_rejetee(self):
        # ligne résiduelle référant une charge Lot3 déjà existante → rejet anti-double-comptage.
        _make_saisie(self.saisie, [
            {"charge_id": "C1", "date_charge": "2026-06-10", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_EWAN"},
        ])
        _make_lot7(self.lot7, source_rows=[
            {"mois": "2026-06", "associe_id": "PERS_EWAN", "type_flux_id": "TYPE_FLUX_001",
             "montant": 100, "lien_origine": "C1"},
        ])
        res = self._run()
        calc = _read_calc(self.lot7)
        # le virement résiduel est rejeté : EWAN garde uniquement l'avantage porté par la charge (100)
        self.assertEqual(calc[("PERS_EWAN", "2026-06")]["avantage_brut_virements"], 0.0)
        self.assertEqual(calc[("PERS_EWAN", "2026-06")]["avantage_brut_depenses_perso"], 100.0)
        codes = [a["code"] for a in res["anomalies"]]
        self.assertIn("SAISIE_LOT7_SOURCE_DEJA_EXISTANTE", codes)

    # ── Cas 5 — Double charge_id ─────────────────────────────────────────────
    def test_cas5_double_charge_id_compte_une_fois(self):
        _make_saisie(self.saisie, [
            {"charge_id": "C9", "date_charge": "2026-06-13", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_WAFA"},
            {"charge_id": "C9", "date_charge": "2026-06-13", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_WAFA"},
        ])
        _make_lot7(self.lot7)
        res = self._run()
        calc = _read_calc(self.lot7)
        # 100 une seule fois, jamais 200
        self.assertEqual(calc[("PERS_WAFA", "2026-06")]["avantage_brut_depenses_perso"], 100.0)
        self.assertEqual(res["doublons_charge_id"], 1)
        codes = [a["code"] for a in res["anomalies"]]
        self.assertIn("CHARGE_ID_DOUBLON", codes)

    # ── Règle historique TYPE_FLUX_002 préservée (sans flag) ─────────────────
    def test_type_flux_002_historique_par_associe_paiement(self):
        _make_saisie(self.saisie, [
            {"charge_id": "C7", "date_charge": "2026-06-14", "montant": 30,
             "type_flux_id": "TYPE_FLUX_002", "associe_id": "PERS_WAFA",
             "mode_paiement_id": "PAY_004", "avantage_associe_id": ""},
        ])
        _make_lot7(self.lot7)
        self._run()
        calc = _read_calc(self.lot7)
        self.assertEqual(calc[("PERS_WAFA", "2026-06")]["avantage_brut_depenses_perso"], 30.0)

    # ── Données réelles vides → sortie vide (pas de faux avantage) ────────────
    def test_saisie_sans_avantage_produit_master_vide(self):
        _make_saisie(self.saisie, [
            {"charge_id": "C0", "date_charge": "2026-06-01", "montant": 10,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_002",
             "avantage_associe_id": ""},
        ])
        _make_lot7(self.lot7)
        self._run()
        self.assertEqual(_read_calc(self.lot7), {})


if __name__ == "__main__":
    unittest.main()
