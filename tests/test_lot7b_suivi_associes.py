"""Lot7B — suivi associé (compte courant associé) : preuves métier sur COPIES / fixtures.

HR STRICT : aucune ligne n'impacte le résultat conciergerie ni le net propriétaire. Ce n'est PAS un
règlement (aucun virement, aucun solde de trésorerie). Aucune donnée métier réelle : fixtures en tmp.
"""
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "02_TRAVAIL"))
import lot7_generateur_avantages as gen  # noqa: E402

SAISIE_COLS = ["charge_id", "date_charge", "mois", "montant", "type_flux_id",
               "associe_id", "mode_paiement_id", "avantage_associe_id"]
SOURCE_COLS = ["mois", "associe_id", "type_flux_id", "type_remboursement",
               "nature", "montant", "mode_paiement_id", "lien_origine", "commentaire"]


def _make_saisie(path, charges):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SAISIE"
    ws.append(SAISIE_COLS)
    for c in charges:
        ws.append([c.get(k) for k in SAISIE_COLS])
    wb.save(path); wb.close()


def _make_lot7(path, source_rows=()):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SOURCE_SAISIE"
    ws.append(SOURCE_COLS)
    ws.append(["AAAA-MM", "instruction", "", "", "", None, "", "", ""])  # ligne instruction ignorée
    for s in source_rows:
        ws.append([s.get(k) for k in SOURCE_COLS])
    wb.save(path); wb.close()


def _read_calc(path):
    wb = openpyxl.load_workbook(path); ws = wb["MASTER_CALC_AVANTAGES"]
    rows = list(ws.iter_rows(values_only=True)); wb.close()
    headers = list(rows[0])
    return headers, {(_d["associe_id"], _d["mois"]): _d
                     for _d in (dict(zip(headers, r)) for r in rows[1:])}


class Lot7BSuiviAssociesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.saisie = os.path.join(self.tmp, "SAISIE_Charges_Flux.xlsx")
        self.lot7 = os.path.join(self.tmp, "MASTER_FACT_MAN_IK_Avantages.xlsx")

    def _run(self, noms=None):
        return gen.generer(self.saisie, self.lot7, self.lot7, associes_noms=noms)

    # ── Garantie transverse : HR strict + aucune colonne propriétaire/résultat ──
    def _assert_hr_et_sans_proprietaire(self, headers, calc):
        # Aucune colonne d'impact propriétaire / résultat dans la sortie suivi.
        interdits = {"proprietaire_id", "net_proprietaire", "resultat_reel",
                     "resultat_comptable", "revenu_net_exploitation_proprietaire"}
        self.assertEqual(interdits & set(headers), set())
        for d in calc.values():
            self.assertEqual(d["code_impact"], "HR")       # hors résultat réel ET comptable
            self.assertEqual(d["source_calcul"], "LOT7")

    # ── Cas A — Avantage banque pro (PAY_001) + avantage explicite ───────────
    def test_casA_banque_pro_avantage_hr_idempotent(self):
        _make_saisie(self.saisie, [
            {"charge_id": "A1", "date_charge": "2026-06-10", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_EWAN"}])
        _make_lot7(self.lot7)
        self._run(noms={"PERS_EWAN": "Ewan"})
        headers, calc = _read_calc(self.lot7)
        d = calc[("PERS_EWAN", "2026-06")]
        self.assertEqual(d["avantages_bruts_total"], 100.0)
        self.assertEqual(d["avantages_nets"], 100.0)
        self.assertEqual(d["code_impact"], "HR")
        self.assertEqual(d["sens_suivi"], "A_CONTROLER_POSITIF")
        self.assertEqual(d["associe_nom"], "Ewan")
        self._assert_hr_et_sans_proprietaire(headers, calc)
        # idempotence : 2e génération → toujours 100, jamais 200
        self._run(noms={"PERS_EWAN": "Ewan"})
        _, calc2 = _read_calc(self.lot7)
        self.assertEqual(calc2[("PERS_EWAN", "2026-06")]["avantages_nets"], 100.0)
        self.assertEqual(len(calc2), len(calc))

    # ── Cas B — Carte associée (PAY_003/004) sans avantage ───────────────────
    def test_casB_carte_sans_avantage_aucun_net(self):
        _make_saisie(self.saisie, [
            {"charge_id": "B1", "date_charge": "2026-06-11", "montant": 50,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "PERS_WAFA",
             "mode_paiement_id": "PAY_003", "avantage_associe_id": ""}])
        _make_lot7(self.lot7)
        res = self._run()
        _, calc = _read_calc(self.lot7)
        self.assertEqual(calc, {})
        self.assertEqual(res["nb_lignes"], 0)

    # ── Cas C — Charge payée pour la société (TF004/008) déduite, pas doublée ─
    def test_casC_charge_payee_societe_deduite(self):
        _make_saisie(self.saisie, [
            {"charge_id": "C1", "date_charge": "2026-06-12", "montant": 40,
             "type_flux_id": "TYPE_FLUX_004", "associe_id": "PERS_WAFA",
             "mode_paiement_id": "PAY_003", "avantage_associe_id": ""}])
        _make_lot7(self.lot7)
        self._run()
        _, calc = _read_calc(self.lot7)
        d = calc[("PERS_WAFA", "2026-06")]
        self.assertEqual(d["charges_payees_pour_societe"], 40.0)
        self.assertEqual(d["avantage_brut_depenses_perso"], 0.0)   # pas doublée en avantage brut
        self.assertEqual(d["avantages_nets"], -40.0)               # net = 0 - 40
        self.assertEqual(d["sens_suivi"], "A_CONTROLER_NEGATIF")
        self.assertEqual(d["code_impact"], "HR")

    # ── Cas D — IK (circuit Lot7 SOURCE_SAISIE), jamais dans Lot3 ─────────────
    def test_casD_ik_dans_suivi_jamais_dans_charges(self):
        _make_saisie(self.saisie, [])  # aucune charge Lot3 : l'IK ne vient PAS de SAISIE_Charges_Flux
        _make_lot7(self.lot7, source_rows=[
            {"mois": "2026-06", "associe_id": "PERS_EWAN", "type_flux_id": "TYPE_FLUX_015",
             "montant": 75, "lien_origine": ""}])
        self._run()
        _, calc = _read_calc(self.lot7)
        d = calc[("PERS_EWAN", "2026-06")]
        self.assertEqual(d["avantage_brut_ik"], 75.0)
        self.assertEqual(d["avantages_bruts_total"], 75.0)
        self.assertEqual(d["avantage_brut_depenses_perso"], 0.0)   # aucune charge Lot3 impliquée
        self.assertEqual(d["code_impact"], "HR")

    # ── Cas E — SOURCE_SAISIE résiduelle autonome, sans doublon Lot3 ─────────
    def test_casE_source_saisie_residuelle(self):
        _make_saisie(self.saisie, [
            {"charge_id": "E1", "date_charge": "2026-06-10", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_EWAN"}])
        _make_lot7(self.lot7, source_rows=[
            {"mois": "2026-06", "associe_id": "PERS_WAFA", "type_flux_id": "TYPE_FLUX_001",
             "montant": 200, "lien_origine": ""}])
        res = self._run()
        _, calc = _read_calc(self.lot7)
        self.assertEqual(calc[("PERS_WAFA", "2026-06")]["avantage_brut_virements"], 200.0)
        self.assertEqual(calc[("PERS_EWAN", "2026-06")]["avantage_brut_depenses_perso"], 100.0)
        self.assertEqual(res["anomalies"], [])

    # ── Cas F — Doublons (charge_id ET lien_origine) ─────────────────────────
    def test_casF_double_charge_id_compte_une_fois(self):
        _make_saisie(self.saisie, [
            {"charge_id": "F1", "date_charge": "2026-06-13", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_WAFA"},
            {"charge_id": "F1", "date_charge": "2026-06-13", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_WAFA"}])
        _make_lot7(self.lot7)
        res = self._run()
        _, calc = _read_calc(self.lot7)
        self.assertEqual(calc[("PERS_WAFA", "2026-06")]["avantage_brut_depenses_perso"], 100.0)
        self.assertEqual(res["doublons_charge_id"], 1)

    def test_casF_lien_origine_existant_rejete(self):
        _make_saisie(self.saisie, [
            {"charge_id": "G1", "date_charge": "2026-06-10", "montant": 100,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_EWAN"}])
        _make_lot7(self.lot7, source_rows=[
            {"mois": "2026-06", "associe_id": "PERS_EWAN", "type_flux_id": "TYPE_FLUX_001",
             "montant": 100, "lien_origine": "G1"}])
        res = self._run()
        _, calc = _read_calc(self.lot7)
        self.assertEqual(calc[("PERS_EWAN", "2026-06")]["avantage_brut_virements"], 0.0)
        self.assertIn("SAISIE_LOT7_SOURCE_DEJA_EXISTANTE", [a["code"] for a in res["anomalies"]])

    # ── Solde nul → SOLDE_NUL (jamais « à payer/rembourser ») ────────────────
    def test_solde_nul(self):
        # avantage 60 (flag) puis charge payée société 60 → net 0.
        _make_saisie(self.saisie, [
            {"charge_id": "H1", "date_charge": "2026-06-10", "montant": 60,
             "type_flux_id": "TYPE_FLUX_020", "associe_id": "", "mode_paiement_id": "PAY_001",
             "avantage_associe_id": "PERS_WAFA"},
            {"charge_id": "H2", "date_charge": "2026-06-11", "montant": 60,
             "type_flux_id": "TYPE_FLUX_008", "associe_id": "PERS_WAFA",
             "mode_paiement_id": "PAY_003", "avantage_associe_id": ""}])
        _make_lot7(self.lot7)
        self._run()
        _, calc = _read_calc(self.lot7)
        d = calc[("PERS_WAFA", "2026-06")]
        self.assertEqual(d["avantages_nets"], 0.0)
        self.assertEqual(d["sens_suivi"], "SOLDE_NUL")


if __name__ == "__main__":
    unittest.main()
