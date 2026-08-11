import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def write_book(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
    wb.save(path)
    wb.close()


def read_rows(path: Path, sheet: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    headers = list(next(it))
    rows = [dict(zip(headers, row)) for row in it]
    wb.close()
    return rows


class Lot10ReservationExclueDedupTest(unittest.TestCase):
    """Reproduit puis verifie le correctif : une reservation deja resolue via HH
    (branche HH de Lot10, presente dans COMMISSIONS) ne doit pas AUSSI apparaitre
    dans A_CONTROLER/RESERVATION_EXCLUE_A_CONTROLER a partir du statut brut Lot1.
    Une reservation genuinement non resolue (aucune saisie HH) doit, elle,
    rester A_CONTROLER."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.project = Path(self._tmp.name) / "Pilotage_Conciergerie"
        self.work = self.project / "02_TRAVAIL"
        self.work.mkdir(parents=True)
        for path in (ROOT / "02_TRAVAIL").glob("*.py"):
            shutil.copy2(path, self.work / path.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _create_sources(self):
        write_book(
            self.project / "01_SOURCES_BRUTES/REF_Setup/REF_Setup.xlsm",
            {
                "REF_Logements": (
                    ["logement_id", "type_logement_id", "seuil_voyageurs_preparation_canape",
                     "montant_preparation_canape", "forfait_logiciel_consommables_mensuel",
                     "statut_parc"],
                    [{"logement_id": "LOG_0001", "type_logement_id": "T1",
                      "forfait_logiciel_consommables_mensuel": 0, "statut_parc": "GERE"}],
                ),
                "REF_Proprietaires": (
                    ["proprietaire_id", "actif"],
                    [{"proprietaire_id": "PROP_0001", "actif": "OUI"}],
                ),
                "REF_Taux_Commission": (
                    ["taux_commission_id", "proprietaire_id", "logement_id", "taux_commission",
                     "date_debut", "date_fin", "actif"],
                    [{"taux_commission_id": "TX1", "proprietaire_id": "PROP_0001",
                      "logement_id": None, "taux_commission": 0.15,
                      "date_debut": "2020-01-01", "date_fin": None, "actif": "OUI"}],
                ),
                "REF_Gestion_Logements_Hist": (
                    ["gestion_id", "logement_id", "proprietaire_id", "date_debut", "date_fin",
                     "statut_gestion"],
                    [{"gestion_id": "G1", "logement_id": "LOG_0001", "proprietaire_id": "PROP_0001",
                      "date_debut": "2020-01-01", "date_fin": None, "statut_gestion": "ACTIVE"}],
                ),
                "REF_Couts_Standards_Menage": (
                    ["type_logement_id", "actif", "cout_standard_menage",
                     "date_debut_validite", "date_fin_validite"],
                    [{"type_logement_id": "T1", "actif": "OUI", "cout_standard_menage": 30.0}],
                ),
                "REF_Cloture_Mensuelle": (["mois", "statut_mois"], []),
            },
        )
        (self.project / "02_TRAVAIL/Lot3_Charges").mkdir(parents=True, exist_ok=True)
        (self.project / "02_TRAVAIL/Lot6c_MenagesExternes").mkdir(parents=True, exist_ok=True)
        (self.project / "02_DONNEES_NORMALISEES/menages").mkdir(parents=True, exist_ok=True)
        (self.project / "02_TRAVAIL/Lot7_IK_Avantages").mkdir(parents=True, exist_ok=True)
        (self.project / "02_TRAVAIL/Lot5_AcomptesProprietaires").mkdir(parents=True, exist_ok=True)
        for rel, sheet, headers in [
            ("02_TRAVAIL/Lot3_Charges/MASTER_FACT_MAN_Charges.xlsx", "MASTER", ["charge_id", "statut_controle"]),
            ("02_TRAVAIL/Lot6c_MenagesExternes/MASTER_FACT_MEN_MenagesExternes.xlsx", "MASTER", ["menage_id", "statut_controle"]),
            ("02_DONNEES_NORMALISEES/menages/M04_MENAGES_PowerQuery.xlsx", "MASTER", ["menage_calc_id", "statut_controle"]),
            ("02_TRAVAIL/Lot7_IK_Avantages/MASTER_FACT_MAN_IK_Avantages.xlsx", "MASTER_CALC_AVANTAGES", ["pk_id", "statut_controle"]),
            ("02_TRAVAIL/Lot5_AcomptesProprietaires/MASTER_FACT_MAN_AcomptesProprietaires.xlsx", "MASTER",
             ["acompte_id", "mois", "logement_id", "proprietaire_id", "montant_acompte", "statut_controle"]),
        ]:
            write_book(self.project / rel, {sheet: (headers, [])})

    def _write_payout(self, rows):
        headers = [
            "reservation_id", "listingMapId", "source", "channel_type", "statut_calcul_payout",
            "payout_calcule", "source_payout", "menage_retenu", "assiette_commission",
            "inclure_resultat_auto", "menage_retenu_source", "cout_standard_id",
            "cout_standard_menage_snapshot", "cout_standard_date_debut_validite",
            "cout_standard_date_fin_validite", "logement_id_snapshot", "type_logement_id_snapshot",
            "date_reference_cout_menage",
        ]
        write_book(self.project / "02_TRAVAIL/Lot1_Hostaway/MASTER_CALC_HA_Payout.xlsx",
                   {"data": (headers, rows)})

    def _write_flux(self, rows):
        headers = [
            "flux_id", "ROW_HASH", "source_module", "source_table", "source_pk",
            "date_flux", "mois", "logement_id", "proprietaire_id", "associe_id",
            "type_flux_id", "sens", "montant", "code_impact",
            "inclure_resultat_reel", "inclure_resultat_comptable", "inclure_resultat_hors_compta",
            "statut_controle", "niveau_anomalie", "code_anomalie", "commentaire", "date_integration",
        ]
        write_book(self.project / "02_TRAVAIL/Lot9_FluxUnifie/MASTER_CALC_Flux.xlsx",
                   {"MASTER": (headers, rows)})

    def _write_resolved(self, hh_rows):
        headers = [
            "reservation_calc_id", "ROW_HASH", "source", "reservation_id_hostaway",
            "reservation_hh_id", "mois", "logement_id", "proprietaire_id", "date_arrivee",
            "date_depart", "nuits", "guestCount", "source_guestCount", "montant_retenu",
            "source_montant", "code_impact", "impact_resultat_reel", "impact_resultat_comptable",
            "statut_controle", "niveau_anomalie", "code_anomalie", "commentaire", "source_module",
            "source_table", "source_pk", "date_integration", "canal", "etat_mois",
            "origine_initiale", "source_ligne", "methode", "payout_calcule", "menage_retenu",
            "assiette_commission",
        ]
        write_book(self.project / "02_TRAVAIL/Lot4quater_SourceResolue/MASTER_CALC_Reservations_Resolues.xlsx",
                   {"MASTER": (headers, hh_rows), "VUE_FLUX": (headers, [r for r in hh_rows if r.get("statut_controle") == "VALIDE"])})

    def _write_hh(self, rows):
        headers = ["reservation_hh_id", "reservation_id_hostaway", "total_percu", "menage",
                   "commission", "taux_commission", "statut_controle"]
        write_book(self.project / "02_TRAVAIL/Lot4_ReservationsHH/MASTER_FACT_MAN_ReservationsHorsHostaway.xlsx",
                   {"MASTER": (headers, rows)})

    def _run(self, script):
        r = subprocess.run(
            [sys.executable, str(self.work / script)],
            cwd=self.project, capture_output=True, text=True,
        )
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)

    def test_reservation_resolue_via_hh_absente_de_a_controler(self):
        """La resa 9001 (DIRECT) a un statut_calcul_payout=A_CONTROLER cote Lot1 (aucun
        payout plateforme, normal pour Direct), MAIS elle est resolue via une saisie HH
        (montant connu). Elle doit apparaitre dans COMMISSIONS, jamais dans A_CONTROLER."""
        self._create_sources()
        self._write_payout([{
            "reservation_id": 9001, "listingMapId": 1, "source": "direct", "channel_type": "DIRECT",
            "statut_calcul_payout": "A_CONTROLER", "payout_calcule": None,
            "source_payout": "DIRECT_HORS_HOSTAWAY", "menage_retenu": 0, "assiette_commission": None,
            "inclure_resultat_auto": "NON", "menage_retenu_source": "NON_APPLICABLE",
        }])
        self._write_hh([{
            "reservation_hh_id": "RESHH-2026-06-001", "reservation_id_hostaway": 9001,
            "total_percu": 500.0, "menage": 30.0, "commission": None, "taux_commission": None,
            "statut_controle": "VALIDE",
        }])
        self._write_resolved([{
            "reservation_calc_id": "RES-2026-06-HH-001", "reservation_id_hostaway": 9001,
            "reservation_hh_id": "RESHH-2026-06-001", "mois": "2026-06", "logement_id": "LOG_0001",
            "proprietaire_id": "PROP_0001", "date_arrivee": "2026-06-10", "date_depart": "2026-06-12",
            "nuits": 2, "montant_retenu": 500.0, "code_impact": "HC", "impact_resultat_reel": "OUI",
            "impact_resultat_comptable": "NON", "statut_controle": "VALIDE", "canal": "DIRECT",
            "etat_mois": "OUVERT",
        }])
        self._write_flux([{
            "flux_id": "FLUX-1", "ROW_HASH": "H1", "source_module": "test",
            "source_table": "MASTER_CALC_Reservations_Resolues", "source_pk": "RES-2026-06-HH-001",
            "date_flux": "2026-06-10", "mois": "2026-06", "logement_id": "LOG_0001",
            "proprietaire_id": "PROP_0001", "type_flux_id": "TYPE_FLUX_017", "sens": "PRODUIT",
            "montant": 500.0, "code_impact": "HC", "inclure_resultat_reel": "OUI",
            "inclure_resultat_comptable": "NON", "inclure_resultat_hors_compta": "OUI",
            "statut_controle": "VALIDE", "niveau_anomalie": "INFO", "date_integration": "2026-06-30",
        }])

        self._run("lot10_calculer_resultats.py")

        out = self.project / "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_Commissions.xlsx"
        comm = read_rows(out, "COMMISSIONS")
        ac = read_rows(out, "A_CONTROLER")

        self.assertTrue(any(r.get("reservation_id_hostaway") == 9001 for r in comm),
                         "reservation resolue via HH absente de COMMISSIONS")
        self.assertFalse(any(r.get("reservation_id") == 9001 for r in ac),
                          "reservation deja resolue via HH toujours listee en A_CONTROLER -> doublon")

    def test_reservation_non_resolue_reste_a_controler(self):
        """La resa 9002 (VRBO) n'a AUCUNE saisie HH : elle doit rester A_CONTROLER,
        le correctif ne doit pas la faire disparaitre a tort."""
        self._create_sources()
        self._write_payout([{
            "reservation_id": 9002, "listingMapId": 1, "source": "vrboical", "channel_type": "VRBO",
            "statut_calcul_payout": "A_CONTROLER", "payout_calcule": None,
            "source_payout": "VRBO_UNKNOWN", "menage_retenu": 0, "assiette_commission": None,
            "inclure_resultat_auto": "NON", "menage_retenu_source": "NON_APPLICABLE",
        }])
        self._write_hh([])
        self._write_resolved([{
            "reservation_calc_id": "RES-2026-06-HA-002", "reservation_id_hostaway": 9002,
            "mois": "2026-06", "logement_id": "LOG_0001", "proprietaire_id": "PROP_0001",
            "date_arrivee": "2026-06-15", "date_depart": "2026-06-17", "nuits": 2,
            "montant_retenu": None, "code_impact": "HC", "impact_resultat_reel": "NON",
            "impact_resultat_comptable": "NON", "statut_controle": "A_CONTROLER", "canal": "VRBO",
            "etat_mois": "OUVERT", "code_anomalie": "VRBO_MONTANT_NON_RENSEIGNE",
        }])
        self._write_flux([{
            "flux_id": "FLUX-DUMMY", "ROW_HASH": "HD", "source_module": "test",
            "source_table": "dummy", "source_pk": "DUMMY", "date_flux": "2026-06-01",
            "mois": "2026-06", "logement_id": "LOG_0001", "proprietaire_id": "PROP_0001",
            "type_flux_id": "TYPE_FLUX_099", "sens": "PRODUIT", "montant": 0.0,
            "code_impact": "IC", "inclure_resultat_reel": "NON", "inclure_resultat_comptable": "NON",
            "inclure_resultat_hors_compta": "NON", "statut_controle": "VALIDE",
            "niveau_anomalie": "INFO", "date_integration": "2026-06-30",
        }])

        self._run("lot10_calculer_resultats.py")

        out = self.project / "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_Commissions.xlsx"
        comm = read_rows(out, "COMMISSIONS")
        ac = read_rows(out, "A_CONTROLER")

        self.assertFalse(any(r.get("reservation_id_hostaway") == 9002 for r in comm))
        self.assertTrue(any(r.get("reservation_id") == 9002 and r.get("code_anomalie_lot10") == "RESERVATION_EXCLUE_A_CONTROLER" for r in ac),
                         "reservation genuinement non resolue doit rester A_CONTROLER")


if __name__ == "__main__":
    unittest.main()
