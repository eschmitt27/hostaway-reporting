import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def write_book(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
    wb.save(path)
    wb.close()


def read_rows(path: Path, sheet: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    headers = list(next(it))
    rows = [dict(zip(headers, row)) for row in it]
    wb.close()
    return rows


class Lot4terGuestcountPersistenceTest(unittest.TestCase):
    """Reproduit puis verifie le correctif : HIST_Reservations_Cloturees doit
    conserver guestCount a travers un run normal de lot4ter (upsert), sans
    jamais backfiller automatiquement les anciennes lignes depuis le live."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.project = Path(self._tmp.name) / "Pilotage_Conciergerie"
        self.work = self.project / "02_TRAVAIL"
        self.work.mkdir(parents=True)
        for path in (ROOT / "02_TRAVAIL").glob("*.py"):
            shutil.copy2(path, self.work / path.name)
        self._create_sources()

    def tearDown(self):
        self._tmp.cleanup()

    def _create_sources(self):
        write_book(
            self.project / "01_SOURCES_BRUTES/REF_Setup/REF_Setup.xlsm",
            {
                "REF_Cloture_Mensuelle": (
                    ["mois", "statut_mois"],
                    [{"mois": "2025-06", "statut_mois": "CLOTURE"}],
                ),
                "REF_Logements": (
                    ["logement_id", "type_logement_id"],
                    [{"logement_id": "LOG_0001", "type_logement_id": "T1"}],
                ),
                "REF_Couts_Standards_Menage": (
                    ["type_logement_id", "actif", "cout_standard_menage",
                     "date_debut_validite", "date_fin_validite"],
                    [{"type_logement_id": "T1", "actif": "OUI",
                      "cout_standard_menage": 30.0,
                      "date_debut_validite": None, "date_fin_validite": None}],
                ),
            },
        )
        write_book(
            self.project / "02_TRAVAIL/Lot1_Hostaway/MASTER_CALC_HA_Payout.xlsx",
            {"data": (["reservation_id", "payout_calcule", "menage_retenu",
                       "assiette_commission"], [])},
        )
        (self.project / "02_TRAVAIL/Lot4bis_TableCommune").mkdir(parents=True, exist_ok=True)
        (self.project / "02_DONNEES_NORMALISEES/historique_reservations").mkdir(parents=True, exist_ok=True)

    def _write_live_master(self, rows):
        headers = [
            "reservation_calc_id", "source", "reservation_id_hostaway", "reservation_hh_id",
            "mois", "logement_id", "proprietaire_id", "date_arrivee", "date_depart", "nuits",
            "guestCount", "montant_retenu", "code_impact", "impact_resultat_reel",
            "impact_resultat_comptable", "statut_controle", "niveau_anomalie", "code_anomalie",
        ]
        write_book(
            self.project / "02_TRAVAIL/Lot4bis_TableCommune/MASTER_CALC_Reservations.xlsx",
            {"MASTER": (headers, rows)},
        )

    def _write_hist(self, rows):
        headers = [
            "cle_historisation", "reservation_calc_id", "reservation_id_hostaway",
            "reservation_hh_id", "canal", "logement_id", "proprietaire_id", "mois",
            "date_arrivee", "date_depart", "nuits", "montant_retenu", "payout_calcule",
            "menage_retenu", "assiette_commission", "code_impact", "impact_resultat_reel",
            "impact_resultat_comptable", "statut_controle", "niveau_anomalie", "code_anomalie",
            "origine_initiale", "source_ligne", "source_montant", "methode", "mois_cloture",
            "fige_le", "ROW_HASH", "guestCount",
        ]
        write_book(
            self.project / "02_DONNEES_NORMALISEES/historique_reservations/HIST_Reservations_Cloturees.xlsx",
            {"HIST_Reservations_Cloturees": (headers, rows)},
        )

    def _run_lot4ter(self):
        subprocess.run(
            [sys.executable, str(self.work / "lot4ter_historiser_reservations_cloturees.py")],
            cwd=self.project, check=True, capture_output=True, text=True,
        )

    def _hist_rows(self):
        return read_rows(
            self.project / "02_DONNEES_NORMALISEES/historique_reservations/HIST_Reservations_Cloturees.xlsx",
            "HIST_Reservations_Cloturees",
        )

    def test_ancienne_ligne_avec_guestcount_corrige_est_conservee(self):
        """Reproduit exactement le bug : une ligne deja historisee avec guestCount
        corrige ne doit PAS perdre cette valeur lors d'un run normal (upsert)."""
        self._write_hist([{
            "cle_historisation": "1001", "reservation_calc_id": "RES-2025-06-HA-001",
            "reservation_id_hostaway": 1001, "canal": "AIRBNB", "logement_id": "LOG_0001",
            "proprietaire_id": "PROP_0001", "mois": "2025-06", "date_arrivee": "2025-06-01",
            "date_depart": "2025-06-05", "nuits": 4, "montant_retenu": 200.0,
            "statut_controle": "VALIDE", "origine_initiale": "API_HOSTAWAY",
            "guestCount": 4,
        }])
        self._write_live_master([])  # rien a historiser de nouveau

        self._run_lot4ter()

        rows = self._hist_rows()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["reservation_id_hostaway"], 1001)
        self.assertEqual(rows[0]["guestCount"], 4,
                          "guestCount deja corrige efface par un run normal (upsert) -> regression")

    def test_nouvelle_historisation_ecrit_guestcount_depuis_le_live(self):
        """Une reservation pas encore historisee doit archiver son guestCount live
        des la premiere historisation."""
        self._write_hist([])
        self._write_live_master([{
            "reservation_calc_id": "RES-2025-06-HA-002", "source": "HOSTAWAY_AIRBNB",
            "reservation_id_hostaway": 1002, "mois": "2025-06", "logement_id": "LOG_0001",
            "proprietaire_id": "PROP_0001", "date_arrivee": "2025-06-10",
            "date_depart": "2025-06-12", "nuits": 2, "guestCount": 5,
            "montant_retenu": 150.0, "code_impact": "IC", "impact_resultat_reel": "OUI",
            "impact_resultat_comptable": "OUI", "statut_controle": "VALIDE",
        }])

        self._run_lot4ter()

        rows = self._hist_rows()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["reservation_id_hostaway"], 1002)
        self.assertEqual(rows[0]["guestCount"], 5)

    def test_ancienne_ligne_sans_guestcount_reste_vide_sans_backfill_live(self):
        """Interdit : aller chercher automatiquement la valeur live pour une ligne
        deja historisee sans guestCount. Elle doit rester vide."""
        self._write_hist([{
            "cle_historisation": "1003", "reservation_calc_id": "RES-2025-06-HA-003",
            "reservation_id_hostaway": 1003, "canal": "AIRBNB", "logement_id": "LOG_0001",
            "proprietaire_id": "PROP_0001", "mois": "2025-06", "date_arrivee": "2025-06-15",
            "date_depart": "2025-06-17", "nuits": 2, "montant_retenu": 100.0,
            "statut_controle": "VALIDE", "origine_initiale": "API_HOSTAWAY",
            "guestCount": None,
        }])
        # Le live porte desormais un guestCount (re-extraction posterieure) mais la
        # ligne est deja historisee -> ne doit PAS etre utilise.
        self._write_live_master([{
            "reservation_calc_id": "RES-2025-06-HA-003", "source": "HOSTAWAY_AIRBNB",
            "reservation_id_hostaway": 1003, "mois": "2025-06", "logement_id": "LOG_0001",
            "proprietaire_id": "PROP_0001", "date_arrivee": "2025-06-15",
            "date_depart": "2025-06-17", "nuits": 2, "guestCount": 6,
            "montant_retenu": 100.0, "code_impact": "IC", "impact_resultat_reel": "OUI",
            "impact_resultat_comptable": "OUI", "statut_controle": "VALIDE",
        }])

        self._run_lot4ter()

        rows = self._hist_rows()
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0]["guestCount"],
                           "backfill automatique depuis le live detecte sur une ligne deja historisee -> interdit")

    def test_schema_et_autres_champs_inchanges(self):
        """L'ajout de guestCount ne doit changer ni les cles, ni le nombre de
        lignes, ni les autres colonnes d'une ligne deja historisee."""
        self._write_hist([{
            "cle_historisation": "1004", "reservation_calc_id": "RES-2025-06-HA-004",
            "reservation_id_hostaway": 1004, "canal": "BOOKING", "logement_id": "LOG_0001",
            "proprietaire_id": "PROP_0001", "mois": "2025-06", "date_arrivee": "2025-06-20",
            "date_depart": "2025-06-22", "nuits": 2, "montant_retenu": 300.0,
            "payout_calcule": 300.0, "menage_retenu": 30.0, "assiette_commission": 270.0,
            "statut_controle": "VALIDE", "origine_initiale": "API_HOSTAWAY",
            "code_impact": "IC", "impact_resultat_reel": "OUI", "impact_resultat_comptable": "OUI",
            "guestCount": 3,
        }])
        self._write_live_master([])

        self._run_lot4ter()

        rows = self._hist_rows()
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual(r["canal"], "BOOKING")
        self.assertEqual(r["montant_retenu"], 300.0)
        self.assertEqual(r["payout_calcule"], 300.0)
        self.assertEqual(r["menage_retenu"], 30.0)
        self.assertEqual(r["assiette_commission"], 270.0)
        self.assertEqual(r["statut_controle"], "VALIDE")
        self.assertEqual(r["guestCount"], 3)


if __name__ == "__main__":
    unittest.main()
