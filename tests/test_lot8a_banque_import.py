"""Lot8a — import bancaire : détection de format et adaptateurs (D-8a-FORMAT).

Couvre le format historique Crédit Mutuel natif (régression, jamais cassé) et le nouveau format
« relevé bancaire consolidé » (fusion outillée de plusieurs relevés). Aucune donnée bancaire
réelle : chaque test construit ses fixtures fictives en tmp_path et exécute le script réel via
`runpy.run_path` (pas de duplication de moteur, pas de subprocess).
"""
from __future__ import annotations

import runpy
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

SCRIPT = Path(__file__).resolve().parents[1] / "02_TRAVAIL" / "lot8a_banque_import.py"

NATIF_HDR = ["Date", "Valeur", "Libellé", "Débit", "Crédit", "Solde", "Devise"]
CONSOLIDE_HDR = ["N°", "Date opération", "Date de valeur", "Libellé", "Débit", "Crédit",
                 "Montant net", "Solde consolidé", "Devise", "Source du relevé", "Mois",
                 "Ligne source"]


def _make_natif(path: Path, rows: list[list]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Cpt 02211 00021321603")
    for _ in range(4):
        ws.append([None])
    ws.append(NATIF_HDR)
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


def _make_consolide(path: Path, mouvements: list[list], *, avec_sources=True,
                     avec_controles=True) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws_s = wb.create_sheet("Synthese")
    ws_s.append(["Relevé bancaire consolidé"])
    ws_m = wb.create_sheet("Mouvements")
    for _ in range(4):
        ws_m.append([None])
    ws_m.append(CONSOLIDE_HDR)
    for r in mouvements:
        ws_m.append(r)
    wb.create_sheet("Mensuel").append(["Mois", "Crédits", "Débits"])
    if avec_controles:
        wb.create_sheet("Controles").append(["Contrôle", "Résultat", "Détail"])
    if avec_sources:
        wb.create_sheet("Sources").append(["Fichier source", "Situation"])
    wb.save(path)
    wb.close()


def _mvt(n, date_op, libelle, debit, credit, solde=0.0, devise="EUR", source="Relevé test"):
    return [n, date_op, date_op, libelle, debit, credit, credit - debit, solde, devise,
            source, date_op.replace(day=1), n]


def _run(monkeypatch, tmp_path: Path, brut: Path):
    out = tmp_path / "Lot8_Banque" / "BANQUE_LOT8_IMPORT.xlsx"
    monkeypatch.setenv("LOT8A_BRUT_FILE_OVERRIDE", str(brut))
    monkeypatch.setenv("LOT8A_OUT_FILE_OVERRIDE", str(out))
    g = runpy.run_path(str(SCRIPT), run_name="lot8a_test_run")
    return g, out


def _read_norm(out_path: Path):
    wb = openpyxl.load_workbook(out_path, data_only=True)
    ws = wb["NORM_Banque"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return rows


def _read_ctrl_codes(out_path: Path):
    wb = openpyxl.load_workbook(out_path, data_only=True)
    ws = wb["CTRL_A_CONTROLER"]
    hdr = [c.value for c in ws[1]]
    idx = hdr.index("code_controle")
    codes = [r[idx] for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return codes


def _read_log_format(out_path: Path):
    wb = openpyxl.load_workbook(out_path, data_only=True)
    ws = wb["LOG_Traitement"]
    hdr = [c.value for c in ws[1]]
    idx = hdr.index("format_source")
    val = ws.cell(2, idx + 1).value
    wb.close()
    return val


# ── FORMAT_CREDIT_MUTUEL_NATIF — régression historique ───────────────────────

def test_format_natif_historique_toujours_vert(tmp_path, monkeypatch):
    brut = tmp_path / "releve_natif.xlsx"
    _make_natif(brut, [
        [datetime(2026, 3, 3), datetime(2026, 3, 3), "PAIEMENT CB", 8.64, None, 100.0, "EUR"],
        [datetime(2026, 3, 4), datetime(2026, 3, 4), "VIR AIRBNB", None, 361.92, 461.92, "EUR"],
    ])
    g, out = _run(monkeypatch, tmp_path, brut)
    assert out.exists()
    assert _read_log_format(out) == g["FORMAT_CREDIT_MUTUEL_NATIF"]
    rows = _read_norm(out)
    assert len(rows) == 2
    assert {r["sens"] for r in rows} == {"DEBIT", "CREDIT"}
    assert round(sum(r["montant"] for r in rows if r["sens"] == "DEBIT"), 2) == 8.64
    assert round(sum(r["montant"] for r in rows if r["sens"] == "CREDIT"), 2) == 361.92


# ── FORMAT_RELEVE_CONSOLIDE — nouveau ────────────────────────────────────────

def test_format_consolide_detecte_et_produit_sortie_canonique(tmp_path, monkeypatch):
    brut = tmp_path / "releve_consolide.xlsx"
    _make_consolide(brut, [
        _mvt(1, datetime(2026, 3, 3), "PAIEMENT CB 0211", 8.64, 0, solde=100.0),
        _mvt(2, datetime(2026, 3, 4), "VIR AIRBNB PAYMENTS", 0, 361.92, solde=461.92),
    ])
    g, out = _run(monkeypatch, tmp_path, brut)
    assert out.exists()
    assert _read_log_format(out) == g["FORMAT_RELEVE_CONSOLIDE"]
    rows = _read_norm(out)
    assert len(rows) == 2
    debit_rows = [r for r in rows if r["sens"] == "DEBIT"]
    credit_rows = [r for r in rows if r["sens"] == "CREDIT"]
    assert len(debit_rows) == 1 and round(debit_rows[0]["montant"], 2) == 8.64
    assert len(credit_rows) == 1 and round(credit_rows[0]["montant"], 2) == 361.92
    # provenance tracée dans le champ commentaire existant, jamais une colonne inventée
    assert all(r["commentaire"] == "Relevé test" for r in rows)


def test_format_consolide_debit_credit_zero_ne_declenche_pas_double(tmp_path, monkeypatch):
    """Le format consolidé renseigne toujours les deux colonnes (0 côté inactif) — sans la
    conversion 0->None, chaque ligne déclencherait à tort BANQUE_DEBIT_CREDIT_DOUBLES."""
    brut = tmp_path / "releve_consolide.xlsx"
    _make_consolide(brut, [_mvt(1, datetime(2026, 3, 3), "PAIEMENT CB", 8.64, 0)])
    _, out = _run(monkeypatch, tmp_path, brut)
    codes = _read_ctrl_codes(out)
    assert "BANQUE_DEBIT_CREDIT_DOUBLES" not in codes


def test_format_consolide_montants_totaux_coherents_avec_synthese(tmp_path, monkeypatch):
    brut = tmp_path / "releve_consolide.xlsx"
    _make_consolide(brut, [
        _mvt(1, datetime(2026, 3, 3), "A", 10.0, 0),
        _mvt(2, datetime(2026, 3, 4), "B", 0, 25.5),
        _mvt(3, datetime(2026, 3, 5), "C", 5.0, 0),
    ])
    _, out = _run(monkeypatch, tmp_path, brut)
    rows = _read_norm(out)
    total_debit = round(sum(r["montant"] for r in rows if r["sens"] == "DEBIT"), 2)
    total_credit = round(sum(r["montant"] for r in rows if r["sens"] == "CREDIT"), 2)
    assert total_debit == 15.0
    assert total_credit == 25.5


def test_format_consolide_multi_mois_signale_a_controler_pas_bloquant(tmp_path, monkeypatch):
    brut = tmp_path / "releve_consolide.xlsx"
    _make_consolide(brut, [
        _mvt(1, datetime(2025, 11, 3), "A", 10.0, 0),
        _mvt(2, datetime(2026, 8, 1), "B", 0, 25.5),
    ])
    _, out = _run(monkeypatch, tmp_path, brut)
    codes = _read_ctrl_codes(out)
    assert "BANQUE_FICHIER_PERIODE_INCOHERENTE" in codes
    rows = _read_norm(out)
    # signale, mais jamais bloquant : les deux lignes restent traitees
    assert len(rows) == 2


def test_format_consolide_doublon_exact_detecte_pas_masque(tmp_path, monkeypatch):
    m = _mvt(1, datetime(2026, 3, 3), "PAIEMENT CB IDENTIQUE", 12.0, 0)
    m2 = _mvt(2, datetime(2026, 3, 3), "PAIEMENT CB IDENTIQUE", 12.0, 0)
    brut = tmp_path / "releve.xlsx"
    _make_consolide(brut, [m, m2])
    _, out = _run(monkeypatch, tmp_path, brut)
    rows = _read_norm(out)
    assert len(rows) == 2
    doublons = [r for r in rows if "DOUBLON" in (r["codes_anomalie"] or "")]
    assert len(doublons) == 1
    total_debit = round(sum(r["montant"] for r in rows if r["sens"] == "DEBIT"), 2)
    assert total_debit == 24.0  # jamais masque : le montant du doublon reste visible, pas soustrait


def test_format_consolide_idempotent_deux_executions(tmp_path, monkeypatch):
    brut = tmp_path / "releve_consolide.xlsx"
    _make_consolide(brut, [
        _mvt(1, datetime(2026, 3, 3), "A", 10.0, 0),
        _mvt(2, datetime(2026, 3, 4), "B", 0, 25.5),
    ])
    _, out = _run(monkeypatch, tmp_path, brut)
    rows1 = _read_norm(out)
    _, out2 = _run(monkeypatch, tmp_path, brut)
    rows2 = _read_norm(out2)
    assert len(rows1) == len(rows2) == 2
    ids1 = sorted(r["mouvement_id"] for r in rows1)
    ids2 = sorted(r["mouvement_id"] for r in rows2)
    assert ids1 == ids2
    total1 = round(sum(r["montant"] for r in rows1), 2)
    total2 = round(sum(r["montant"] for r in rows2), 2)
    assert total1 == total2


def test_format_consolide_incomplet_sans_sources_est_inconnu(tmp_path, monkeypatch):
    """Sources manquante => les 3 feuilles requises ne sont pas toutes presentes => INCONNU."""
    brut = tmp_path / "releve_incomplet.xlsx"
    _make_consolide(brut, [_mvt(1, datetime(2026, 3, 3), "A", 10.0, 0)], avec_sources=False)
    out = tmp_path / "Lot8_Banque" / "BANQUE_LOT8_IMPORT.xlsx"
    monkeypatch.setenv("LOT8A_BRUT_FILE_OVERRIDE", str(brut))
    monkeypatch.setenv("LOT8A_OUT_FILE_OVERRIDE", str(out))
    with pytest.raises(SystemExit) as exc:
        runpy.run_path(str(SCRIPT), run_name="lot8a_test_run")
    assert exc.value.code == 1
    assert not out.exists()


# ── FORMAT_INCONNU ────────────────────────────────────────────────────────────

def test_format_inconnu_refuse_explicitement(tmp_path, monkeypatch):
    brut = tmp_path / "releve_inconnu.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "FeuilleQuelconque"
    wb.save(brut)
    wb.close()
    out = tmp_path / "Lot8_Banque" / "BANQUE_LOT8_IMPORT.xlsx"
    monkeypatch.setenv("LOT8A_BRUT_FILE_OVERRIDE", str(brut))
    monkeypatch.setenv("LOT8A_OUT_FILE_OVERRIDE", str(out))
    with pytest.raises(SystemExit) as exc:
        runpy.run_path(str(SCRIPT), run_name="lot8a_test_run")
    assert exc.value.code == 1
    assert not out.exists()


def test_fichier_absent_refuse_explicitement(tmp_path, monkeypatch):
    brut = tmp_path / "n_existe_pas.xlsx"
    out = tmp_path / "Lot8_Banque" / "BANQUE_LOT8_IMPORT.xlsx"
    monkeypatch.setenv("LOT8A_BRUT_FILE_OVERRIDE", str(brut))
    monkeypatch.setenv("LOT8A_OUT_FILE_OVERRIDE", str(out))
    with pytest.raises(SystemExit) as exc:
        runpy.run_path(str(SCRIPT), run_name="lot8a_test_run")
    assert exc.value.code == 1


# ── Confidentialité ───────────────────────────────────────────────────────────

def test_stdout_ne_contient_pas_de_chemin_absolu_du_compte(tmp_path, monkeypatch, capsys):
    brut = tmp_path / "releve_consolide.xlsx"
    _make_consolide(brut, [_mvt(1, datetime(2026, 3, 3), "A", 10.0, 0)])
    _run(monkeypatch, tmp_path, brut)
    captured = capsys.readouterr()
    # le compte reste un identifiant opaque déjà construit (CM_...), jamais un IBAN/RIB brut
    assert "10278" not in captured.out
