"""Lot3 — preuve du GÉNÉRATEUR RÉEL du MASTER charges (Option A, Python/openpyxl).

Aucune donnée métier réelle : chaque test construit ses fixtures en tmp, exécute le générateur et
relit le MASTER produit. Les fixtures n'ont AUCUNE valeur en cache pour les colonnes formule
(mois, impact_resultat_reel/comptable, ROW_HASH) — c'est précisément l'état d'un classeur écrit par
openpyxl sans réouverture Excel. Les tests prouvent que la sortie est correcte malgré cela.
"""
import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "02_TRAVAIL"))

import lot3_generateur_charges as gen  # noqa: E402

# Colonnes SAISIE (A..AE) + les colonnes applicatives (AF..AJ), comme le classeur réel.
SAISIE_COLS = gen.SAISIE_COLS + [
    "affectable_menage", "intervenant_concerne", "profil_impact_charge",
    "libelle_categorie_personnalise", "avantage_associe_id",
]

FIXE = dt.datetime(2026, 7, 13, 12, 0, 0)   # date_integration figée → sortie reproductible


def _charge(**kw):
    """Charge type : les colonnes formule sont laissées VIDES (aucun cache Excel)."""
    base = {c: None for c in SAISIE_COLS}
    base.update({
        "charge_id": "C1", "date_charge": dt.date(2026, 6, 15), "montant": 100.0,
        "sens_flux": "DEPENSE", "categorie_charge_id": "CHG_025", "type_flux_id": "TYPE_FLUX_020",
        "code_impact": "IC", "mode_paiement_id": "PAY_001", "statut_controle": "VALIDE",
        # mois / impact_resultat_reel / impact_resultat_comptable / ROW_HASH : VIDES exprès.
    })
    base.update(kw)
    return base


def _make_saisie(path, charges):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SAISIE"
    ws.append(SAISIE_COLS)
    for c in charges:
        ws.append([c.get(k) for k in SAISIE_COLS])
    wb.save(path); wb.close()


def _make_ref(path, filtres=None):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "REF_Categories_Charges"
    ws.append(["categorie_charge_id", "filtre_vue_menage"])
    for cid, f in (filtres or {"CHG_025": "NON", "CHG_004": "OUI", "CHG_003": "OUI"}).items():
        ws.append([cid, f])
    wb.save(path); wb.close()


def _make_master_cible(path):
    """Classeur cible : onglets MASTER / VUE_MENAGE vides + POWER_QUERY_CODE (documentaire)."""
    wb = openpyxl.Workbook()
    wb.active.title = "MASTER"
    wb.create_sheet("VUE_MENAGE")
    pq = wb.create_sheet("POWER_QUERY_CODE")
    pq["A1"] = "POWER QUERY M-CODE — documentaire"
    wb.save(path); wb.close()


def _lire(path, sheet="MASTER"):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(c is not None for c in r)]


class Lot3GenerateurCharges(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.saisie = self.tmp / "SAISIE_Charges_Flux.xlsx"
        self.ref = self.tmp / "REF_Setup.xlsx"
        self.master = self.tmp / "MASTER_FACT_MAN_Charges.xlsx"
        _make_ref(self.ref)
        _make_master_cible(self.master)

    def _run(self, date_integration=FIXE):
        return gen.generer(str(self.saisie), str(self.ref), str(self.master),
                           date_integration=date_integration)

    # 1. mois dérivé de date_charge, SANS cache Excel de la colonne C.
    def test_1_mois_derive_sans_cache_excel(self):
        _make_saisie(self.saisie, [_charge(date_charge=dt.date(2026, 6, 15))])
        self._run()
        rows = _lire(self.master)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["mois"], "2026-06")
        self.assertEqual(rows[0]["ROW_HASH"], "20260615|100.00|CHG_025|PAY_001")

    # 2. Lignes gabarit / instruction / vides → ignorées.
    def test_2_gabarit_ignore(self):
        _make_saisie(self.saisie, [
            _charge(charge_id="[Charge par Power Query — requete MASTER_FACT_MAN_Charges]"),
            _charge(charge_id=""),
            _charge(charge_id="C1"),
        ])
        res = self._run()
        self.assertEqual(res["nb_lignes"], 1)
        self.assertEqual([r["charge_id"] for r in _lire(self.master)], ["C1"])

    # 3. Charge IC → impact réel OUI, comptable OUI.
    def test_3_impact_ic(self):
        _make_saisie(self.saisie, [_charge(code_impact="IC")])
        self._run()
        r = _lire(self.master)[0]
        self.assertEqual((r["impact_resultat_reel"], r["impact_resultat_comptable"]), ("OUI", "OUI"))

    # 4. Charge HC → réel OUI, comptable NON.
    def test_4_impact_hc(self):
        _make_saisie(self.saisie, [_charge(code_impact="HC")])
        self._run()
        r = _lire(self.master)[0]
        self.assertEqual((r["impact_resultat_reel"], r["impact_resultat_comptable"]), ("OUI", "NON"))

    # 5. Charge HR → hors résultat réel ET comptable.
    def test_5_impact_hr(self):
        _make_saisie(self.saisie, [_charge(code_impact="HR")])
        self._run()
        r = _lire(self.master)[0]
        self.assertEqual((r["impact_resultat_reel"], r["impact_resultat_comptable"]), ("NON", "NON"))

    # 6. Charge ménage : le mois vu par Lot6f vient de date_charge, jamais de la formule C.
    def test_6_charge_menage_mois_derive_pour_lot6f(self):
        c = _charge(charge_id="CM1", categorie_charge_id="CHG_004", affectable_menage="OUI",
                    date_charge=dt.date(2026, 5, 20), statut_controle="VALIDE")
        _make_saisie(self.saisie, [c])
        # La colonne `mois` de la SAISIE est vide (pas de cache) — Lot6f ne doit pas s'y fier.
        self.assertIsNone(c["mois"])
        self.assertEqual(gen.mois_de(c["date_charge"]), "2026-05")
        res = self._run()
        r = _lire(self.master)[0]
        self.assertEqual(r["mois"], "2026-05")
        # filtre_vue_menage=OUI + statut VALIDE → la charge alimente VUE_MENAGE.
        self.assertEqual(r["filtre_vue_menage"], "OUI")
        self.assertEqual(res["nb_vue_menage"], 1)
        self.assertEqual([v["charge_id"] for v in _lire(self.master, "VUE_MENAGE")], ["CM1"])

    # 7. Idempotence : deux générations produisent exactement le même contenu.
    def test_7_idempotence(self):
        _make_saisie(self.saisie, [_charge(charge_id="C1"), _charge(charge_id="C2", montant=50.0)])
        self._run()
        avant = _lire(self.master)
        self._run()
        apres = _lire(self.master)
        self.assertEqual(avant, apres)
        self.assertEqual(len(apres), 2)          # jamais d'accumulation

    # 8. Sources jamais modifiées (le générateur n'écrit que la cible).
    def test_8_sources_intactes(self):
        import hashlib

        def h(p):
            return hashlib.sha256(Path(p).read_bytes()).hexdigest()

        _make_saisie(self.saisie, [_charge()])
        h_saisie, h_ref = h(self.saisie), h(self.ref)
        self._run()
        self.assertEqual(h(self.saisie), h_saisie)
        self.assertEqual(h(self.ref), h_ref)

    # ── Contrat de sortie et règles métier ────────────────────────────────────

    def test_contrat_37_colonnes_et_onglet_documentaire_preserve(self):
        _make_saisie(self.saisie, [_charge()])
        self._run()
        wb = openpyxl.load_workbook(self.master, read_only=True)
        try:
            self.assertEqual(wb.sheetnames, ["MASTER", "VUE_MENAGE", "POWER_QUERY_CODE"])
            entetes = [c.value for c in next(wb["MASTER"].iter_rows(max_row=1))]
        finally:
            wb.close()
        self.assertEqual(entetes, gen.MASTER_HEADERS)
        self.assertEqual(len(entetes), 37)

    def test_une_charge_saisie_donne_une_seule_ligne_master(self):
        """Une charge = UNE charge économique. Aucune duplication (affectation/ménage/réserve)."""
        _make_saisie(self.saisie, [_charge(charge_id="C1", affectable_menage="OUI")])
        res = self._run()
        self.assertEqual(res["nb_lignes"], 1)

    def test_sens_derive_du_sens_flux(self):
        _make_saisie(self.saisie, [
            _charge(charge_id="C1", sens_flux="DEPENSE"),
            _charge(charge_id="C2", sens_flux="RECUPERATION"),
            _charge(charge_id="C3", sens_flux="REMBOURSEMENT"),
            _charge(charge_id="C4", sens_flux="REFACTURATION"),
        ])
        self._run()
        sens = {r["charge_id"]: r["sens"] for r in _lire(self.master)}
        self.assertEqual(sens, {"C1": "CHARGE", "C2": "PRODUIT",
                                "C3": "NEUTRALISATION", "C4": "CHARGE"})

    def test_date_charge_invalide_donne_controle_clair(self):
        """Jamais un faux « aucune charge saisie » : la date inexploitable est signalée."""
        _make_saisie(self.saisie, [_charge(charge_id="C1", date_charge="pas-une-date")])
        res = self._run()
        codes = [a["code"] for a in res["anomalies"]]
        self.assertIn("DATE_CHARGE_INVALIDE", codes)
        # Cellule vide (openpyxl relit "" comme None) — jamais un mois inventé.
        self.assertIn(_lire(self.master)[0]["mois"], (None, ""))

    def test_charge_id_doublon_compte_une_seule_fois(self):
        _make_saisie(self.saisie, [_charge(charge_id="C1"), _charge(charge_id="C1", montant=999.0)])
        res = self._run()
        self.assertEqual(res["nb_lignes"], 1)
        self.assertIn("CHARGE_ID_DOUBLON", [a["code"] for a in res["anomalies"]])
        self.assertEqual(_lire(self.master)[0]["montant"], 100.0)   # première occurrence gardée

    def test_vue_menage_exclut_non_valide(self):
        """VUE_MENAGE = filtre_vue_menage=OUI ET statut_controle=VALIDE (D028)."""
        _make_saisie(self.saisie, [
            _charge(charge_id="C1", categorie_charge_id="CHG_004", statut_controle="A_CONTROLER"),
            _charge(charge_id="C2", categorie_charge_id="CHG_004", statut_controle="VALIDE"),
            _charge(charge_id="C3", categorie_charge_id="CHG_025", statut_controle="VALIDE"),
        ])
        res = self._run()
        self.assertEqual(res["nb_vue_menage"], 1)
        self.assertEqual([v["charge_id"] for v in _lire(self.master, "VUE_MENAGE")], ["C2"])


if __name__ == "__main__":
    unittest.main()
