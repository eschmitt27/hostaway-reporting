"""Lot11 — tests des contrôles du suivi des avantages associés (fonctions pures + intégration)."""
import os
import sys
import tempfile
import unittest

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "02_TRAVAIL"))
import lib_controles_avantages as ctl  # noqa: E402
import lot7_generateur_avantages as gen  # noqa: E402


def _charge(cid, montant, tf="TYPE_FLUX_020", associe="", avantage="", mois="2026-06"):
    return {"charge_id": cid, "date_charge": f"{mois}-10", "montant": montant,
            "type_flux_id": tf, "associe_id": associe, "avantage_associe_id": avantage}


def _calc(associe, mois="2026-06", bruts=0.0, ik=0.0, dp=0.0, ch=0.0, ravs=0.0, rsva=0.0,
          nets=None, code_impact="HR"):
    nets = (bruts - ch - ravs + rsva) if nets is None else nets
    return {"associe_id": associe, "mois": mois, "avantage_brut_ik": ik,
            "avantage_brut_depenses_perso": dp, "avantages_bruts_total": bruts,
            "charges_payees_pour_societe": ch, "remboursements_associe_vers_societe": ravs,
            "remboursements_societe_vers_associe": rsva, "avantages_nets": nets,
            "code_impact": code_impact}


class ControlesAvantagesUnitaires(unittest.TestCase):
    def test_avantage_absent(self):
        charges = [_charge("C1", 100, avantage="PERS_EWAN")]
        self.assertEqual(ctl.ctrl_avantage_absent_du_suivi(charges, [_calc("PERS_EWAN", bruts=100, dp=100)]), [])
        anos = ctl.ctrl_avantage_absent_du_suivi(charges, [])
        self.assertEqual(anos[0]["code"], "AVANTAGE_ABSENT_DU_SUIVI")

    def test_charge_id_double(self):
        self.assertEqual(ctl.ctrl_charge_id_double([_charge("C1", 10)]), [])
        anos = ctl.ctrl_charge_id_double([_charge("C1", 10), _charge("C1", 10)])
        self.assertEqual(anos[0]["code"], "CHARGE_ID_AVANTAGE_DOUBLE")

    def test_source_saisie_lien_deja_lot3(self):
        res = [{"lien_origine": "C1"}]
        self.assertEqual(ctl.ctrl_source_saisie_lien_deja_lot3(res, []), [])
        anos = ctl.ctrl_source_saisie_lien_deja_lot3(res, ["C1"])
        self.assertEqual(anos[0]["code"], "SOURCE_SAISIE_LIEN_DEJA_LOT3")

    def test_code_impact_hr(self):
        self.assertEqual(ctl.ctrl_code_impact_hr([_calc("PERS_EWAN", code_impact="HR")]), [])
        anos = ctl.ctrl_code_impact_hr([_calc("PERS_EWAN", code_impact="IC")])
        self.assertEqual(anos[0]["code"], "AVANTAGE_CODE_IMPACT_NON_HR")

    def test_pas_impact_proprietaire(self):
        self.assertEqual(ctl.ctrl_pas_impact_proprietaire(["mois", "associe_id", "avantages_nets"]), [])
        anos = ctl.ctrl_pas_impact_proprietaire(["mois", "proprietaire_id"])
        self.assertEqual(anos[0]["code"], "AVANTAGE_IMPACT_PROPRIETAIRE_INTERDIT")

    def test_cle_suivi_presente(self):
        self.assertEqual(ctl.ctrl_cle_suivi_presente([_calc("PERS_EWAN")]), [])
        anos = ctl.ctrl_cle_suivi_presente([_calc("", mois="")])
        self.assertEqual(anos[0]["code"], "SUIVI_CLE_MANQUANTE")

    def test_ik_hors_charges_lot3(self):
        self.assertEqual(ctl.ctrl_ik_hors_charges_lot3([_charge("C1", 10, tf="TYPE_FLUX_020")]), [])
        anos = ctl.ctrl_ik_hors_charges_lot3([_charge("C1", 10, tf="TYPE_FLUX_015")])
        self.assertEqual(anos[0]["code"], "IK_DANS_CHARGES_LOT3")

    def test_charges_payees_reprises(self):
        charges = [_charge("C1", 40, tf="TYPE_FLUX_004", associe="PERS_WAFA")]
        ok = [_calc("PERS_WAFA", ch=40, nets=-40)]
        self.assertEqual(ctl.ctrl_charges_payees_reprises(charges, ok), [])
        anos = ctl.ctrl_charges_payees_reprises(charges, [])
        self.assertEqual(anos[0]["code"], "CHARGE_PAYEE_NON_REPRISE")

    def test_avantage_net_coherent(self):
        self.assertEqual(ctl.ctrl_avantage_net_coherent([_calc("PERS_EWAN", bruts=100, ch=40, nets=60)]), [])
        anos = ctl.ctrl_avantage_net_coherent([_calc("PERS_EWAN", bruts=100, ch=40, nets=999)])
        self.assertEqual(anos[0]["code"], "AVANTAGE_NET_INCOHERENT")


class ControlesAvantagesIntegration(unittest.TestCase):
    """La sortie réelle du générateur passe TOUS les contrôles (aucune anomalie)."""
    def test_sortie_generateur_sans_anomalie(self):
        tmp = tempfile.mkdtemp()
        saisie = os.path.join(tmp, "S.xlsx")
        lot7 = os.path.join(tmp, "L.xlsx")
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SAISIE"
        cols = ["charge_id", "date_charge", "mois", "montant", "type_flux_id",
                "associe_id", "mode_paiement_id", "avantage_associe_id"]
        ws.append(cols)
        ws.append(["C1", "2026-06-10", None, 100, "TYPE_FLUX_020", "", "PAY_001", "PERS_EWAN"])
        ws.append(["C2", "2026-06-11", None, 40, "TYPE_FLUX_004", "PERS_WAFA", "PAY_003", ""])
        wb.save(saisie); wb.close()
        wbl = openpyxl.Workbook(); wsl = wbl.active; wsl.title = "SOURCE_SAISIE"
        wsl.append(["mois", "associe_id", "type_flux_id", "type_remboursement",
                    "nature", "montant", "mode_paiement_id", "lien_origine", "commentaire"])
        wbl.save(lot7); wbl.close()

        gen.generer(saisie, lot7, lot7)
        wb2 = openpyxl.load_workbook(lot7); ws2 = wb2["MASTER_CALC_AVANTAGES"]
        rows = list(ws2.iter_rows(values_only=True)); wb2.close()
        headers = list(rows[0])
        calc = [dict(zip(headers, r)) for r in rows[1:]]
        charges = gen.charger_charges_saisie(saisie)
        residuels = gen.charger_source_saisie_residuelle(lot7)
        anomalies = ctl.controler_suivi_associes(charges, residuels, calc, headers)
        self.assertEqual(anomalies, [])


if __name__ == "__main__":
    unittest.main()
