"""Lot 8c doit lire l'etat reel du master Lot 5, pas l'affirmer sans le verifier.

Contexte : `lot8c_rapprochement_banque.py` ecrivait en dur, pour chaque virement proprietaire,
le prerequis "MASTER_FACT_MAN_AcomptesProprietaires vide - attendre saisie Lot 5" et le controle
`LOT5_PREREQUIS_MANQUANT`, sans jamais ouvrir ce fichier. Consequence : alimenter Lot 5 puis
relancer Lot 8c ne changeait rien et ne le signalait pas -- l'action demandee a l'utilisateur
n'aboutissait a aucun effet observable.

Ces tests ne definissent AUCUNE regle de rapprochement Lot5 <-> Banque (montant, tolerance,
fenetre de date, groupement) : cet arbitrage est metier et reste ouvert. Ils verifient seulement
que le message et le controle refletent l'etat reel du fichier.
"""
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

PROJET = Path(__file__).resolve().parent.parent
LOT8C = PROJET / "02_TRAVAIL" / "lot8c_rapprochement_banque.py"

BANQUE_REL = Path("02_TRAVAIL") / "Lot8_Banque" / "BANQUE_LOT8_IMPORT.xlsx"
MASTER5_REL = (Path("02_TRAVAIL") / "Lot5_AcomptesProprietaires"
               / "MASTER_FACT_MAN_AcomptesProprietaires.xlsx")

NORM_COLS = [
    "mouvement_id", "ROW_HASH", "import_id", "ligne_source", "date_operation", "date_valeur",
    "libelle", "libelle_brut", "montant", "sens", "devise", "compte_id", "tiers_detecte",
    "categorie", "type_flux_id", "code_impact", "source_classification", "source_economique",
    "statut_controle", "niveau_risque", "codes_anomalie", "date_integration", "commentaire",
    "statut_classification", "niveau_anomalie", "regle_id_appliquee",
]

MASTER5_COLS = [
    "acompte_id", "ROW_HASH", "mois", "proprietaire_id", "logement_id", "facture_ref",
    "source_acompte", "source_hh_id", "montant_acompte", "report_mois_precedent",
    "mode_paiement_id", "code_impact", "impact_resultat_reel", "impact_resultat_comptable",
    "statut_controle", "niveau_anomalie", "code_anomalie", "commentaire",
    "source_module", "source_table", "source_pk", "date_integration",
]


def _mouvement_proprietaire(i):
    row = {c: None for c in NORM_COLS}
    row.update({
        "mouvement_id": f"MVT-FIXTURE-{i:03d}",
        "date_operation": "2026-03-1%d" % (i % 10),
        "libelle": "VIR RECU FIXTURE",
        "montant": 100.0 * i,
        "sens": "CREDIT",
        "tiers_detecte": "PROP_0002",
        "categorie": "VIREMENT_PROPRIETAIRE_A_RAPPROCHER",
        "statut_classification": "RAPPROCHEMENT_REQUIS",
    })
    return [row[c] for c in NORM_COLS]


def _construire_projet(racine, lignes_lot5):
    """Arborescence minimale : Banque avec 2 virements proprietaires + master Lot 5 donne."""
    banque = racine / BANQUE_REL
    banque.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NORM_Banque"
    ws.append(NORM_COLS)
    for i in (1, 2):
        ws.append(_mouvement_proprietaire(i))
    wb.create_sheet("BRUT_Banque").append(["import_id"])
    ws_log = wb.create_sheet("LOG_Traitement")
    ws_log.append(["run_id", "import_id", "etape", "timestamp", "nb_lignes_lues",
                   "nb_lignes_exclues_pied", "nb_lignes_brut", "nb_lignes_norm", "nb_bloquants",
                   "nb_a_controler", "nb_doublons", "periode_incoherente", "commentaire",
                   "format_source"])
    wb.save(banque)
    wb.close()

    master5 = racine / MASTER5_REL
    master5.parent.mkdir(parents=True, exist_ok=True)
    wb5 = openpyxl.Workbook()
    ws5 = wb5.active
    ws5.title = "MASTER"
    ws5.append(MASTER5_COLS)
    for ligne in lignes_lot5:
        ws5.append(ligne)
    wb5.save(master5)
    wb5.close()


def _acompte(i):
    d = {c: None for c in MASTER5_COLS}
    d.update({
        "acompte_id": f"ACPT-FIXTURE-{i:03d}",
        "mois": "2026-03",
        "proprietaire_id": "PROP_0002",
        "facture_ref": f"FACT-{i:03d}",
        "montant_acompte": 100.0 * i,
        "statut_controle": "VALIDE",
    })
    return [d[c] for c in MASTER5_COLS]


def _lancer(racine):
    # --project-root est obligatoire : lot8c derive sinon sa racine de __file__, donc du projet reel.
    (racine / "99_ARCHIVES").mkdir(parents=True, exist_ok=True)
    res = subprocess.run([sys.executable, str(LOT8C), "--project-root", str(racine)],
                         cwd=str(racine), capture_output=True, text=True)
    if res.returncode != 0:
        raise AssertionError(f"lot8c a echoue ({res.returncode})\n{res.stdout}\n{res.stderr}")
    return res


def _lire(racine, onglet):
    wb = openpyxl.load_workbook(racine / BANQUE_REL, read_only=True, data_only=True)
    ws = wb[onglet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = rows[0]
    return [dict(zip(hdr, r)) for r in rows[1:]]


class TestLot8cLitLot5(unittest.TestCase):

    def test_lot5_vide_message_inchange(self):
        """Non-regression : Lot 5 vide -> comportement actuel strictement conserve."""
        with tempfile.TemporaryDirectory() as tmp:
            racine = Path(tmp)
            _construire_projet(racine, lignes_lot5=[])
            _lancer(racine)

            lignes = _lire(racine, "RAPPROCH_PROPRIETAIRES_ATTENTE")
            self.assertEqual(len(lignes), 2)
            for l in lignes:
                self.assertIn("vide", l["prerequis_rapprochement"].lower())
                self.assertEqual(l["statut_rapprochement"], "EN_ATTENTE_SAISIE_ACOMPTE")

            ctrl = {c["code_controle"] for c in _lire(racine, "CTRL_RAPPROCHEMENT_8C")}
            self.assertIn("LOT5_PREREQUIS_MANQUANT", ctrl)

    def test_lot5_alimente_ne_doit_pas_affirmer_vide(self):
        """Lot 5 alimente -> lot8c ne doit plus affirmer que le master est vide."""
        with tempfile.TemporaryDirectory() as tmp:
            racine = Path(tmp)
            _construire_projet(racine, lignes_lot5=[_acompte(1), _acompte(2), _acompte(3)])
            _lancer(racine)

            lignes = _lire(racine, "RAPPROCH_PROPRIETAIRES_ATTENTE")
            self.assertEqual(len(lignes), 2)
            for l in lignes:
                self.assertNotIn(
                    "vide", l["prerequis_rapprochement"].lower(),
                    "lot8c affirme que le master Lot 5 est vide alors qu'il contient 3 objets",
                )

    def test_lot5_alimente_controle_distinct(self):
        """Lot 5 alimente -> le controle ne doit plus demander d'alimenter Lot 5, et doit
        signaler que les regles de rapprochement restent a arbitrer (aucune regle inventee)."""
        with tempfile.TemporaryDirectory() as tmp:
            racine = Path(tmp)
            _construire_projet(racine, lignes_lot5=[_acompte(1)])
            _lancer(racine)

            ctrl = _lire(racine, "CTRL_RAPPROCHEMENT_8C")
            codes = {c["code_controle"] for c in ctrl}
            self.assertNotIn(
                "LOT5_PREREQUIS_MANQUANT", codes,
                "lot8c demande encore d'alimenter Lot 5 alors que Lot 5 est alimente",
            )
            self.assertIn("LOT5_REGLES_RAPPROCHEMENT_A_ARBITRER", codes)

    def test_aucun_rapprochement_automatique_invente(self):
        """Garde-fou : meme avec des montants identiques, lot8c ne rapproche rien
        automatiquement — aucune regle metier n'a ete definie."""
        with tempfile.TemporaryDirectory() as tmp:
            racine = Path(tmp)
            # montants 100 et 200, exactement ceux des 2 mouvements bancaires
            _construire_projet(racine, lignes_lot5=[_acompte(1), _acompte(2)])
            _lancer(racine)

            for l in _lire(racine, "RAPPROCH_PROPRIETAIRES_ATTENTE"):
                self.assertNotEqual(l["statut_rapprochement"], "RAPPROCHE")
                self.assertIsNone(l["date_rapprochement"] if l["date_rapprochement"] == "" else None)


if __name__ == "__main__":
    unittest.main()
