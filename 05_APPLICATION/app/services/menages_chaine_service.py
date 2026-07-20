"""Chaîne ménages COMPLÈTE sur copies (recette) — sources → MASTER → contrôles.

Ce service orchestre la relance de TOUTE la chaîne ménages dans un workspace isolé :

    stub Hostaway (extraction simulée, AUCUNE requête API)
    → lot6b (déclarations internes, source Google Sheet COPIÉE — stub lib_sheet_source)
    → lot6c (ménages externes depuis les factures PDF copiées)
    → lot6d (rapprochement) → lot6e (gain/perte) → lot6f (coût complet) → lot11 (contrôles).

Ordre AUDITÉ : lot6c lit la VUE_COMPTAGE du MASTER Hostaway (Lot6a) pour son onglet
VUE_ECART_HOSTAWAY — l'étape Hostaway précède donc obligatoirement lot6c. lot6b est
indépendant de 6a/6c et passe avant 6c par convention. Les trois sources sont TOUJOURS
régénérées avant lot6d.

Deux modes :
  MODE_COPIES — seule exécution possible. Tout se passe sous un workspace de data/ ;
                aucun fichier réel n'est modifié (vérifié par sha256 avant/après).
  MODE_REEL   — réservé au futur bouton « Actualiser les sources et recalculer ».
                BLOQUÉ tant que MENAGES_REAL_RECALC_ENABLED est False : refus tracé,
                aucune exécution, aucune requête réseau.

Ce module n'importe JAMAIS le moteur (test_no_import_of_travail_modules) : les scripts
sont copiés puis exécutés en sous-processus par le runner générique
`runners/menages_recalcul_runner.py` (interpréteur moteur), qui refuse tout chemin hors
workspace, capture stdout/stderr et vérifie chaque sortie.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import app.config as cfg
from app.services import menages_recalcul_service as recalc
from app.services import snapshot_service
from app.services import saisie_charges_lock_service as verrou_lib
from app.services.audit_service import log_event
from app.services.menages_recalcul_service import (
    STATUT_BLOQUE, STATUT_ECHEC, STATUT_PARTIEL, STATUT_SUCCES, STATUT_VERROUILLE,
    _enregistrer_run, _git_head, _lire_comparaison, _sha256, detecter_excel_ouvert,
)

MODE_COPIES = "COPIES"
MODE_REEL = "REEL"
MODE_RUN_DB = "CHAINE_COPIES"          # valeur `mode` enregistrée dans menages_recalcul_runs
LOCK_NAME = ".menages_chaine.lock"
OPERATION = "menages_chaine"

# ── Sources RÉELLES copiées dans le workspace ────────────────────────────────
# Cœur ménages : indispensables aux étapes Lot6a-Lot6e ; préflight bloquant si absentes.
SOURCES_COEUR = [
    "01_SOURCES_BRUTES/REF_Setup/REF_Setup.xlsm",
    "02_TRAVAIL/Lot1_Hostaway/MASTER_FACT_HA_CleaningTasks_Discovery.xlsx",
    "02_DONNEES_NORMALISEES/menages/M04_MENAGES_PowerQuery.xlsx",
]
# Sources FACULTATIVES : utilisées seulement par Lot6f (pools courses/conso — vides aujourd'hui)
# et/ou les contrôles transverses. Leur absence n'empêche PAS Lot6a-Lot6e : elle ne rend donc
# JAMAIS la recette Ménages « rouge » — simple avertissement.
SOURCES_FACULTATIVES = [
    "01_SOURCES_BRUTES/Charges/SAISIE_Charges_Flux.xlsx",
]
# Libellés utilisateur des étapes (les noms techniques restent dans un détail repliable de l'UI).
LIBELLES_ETAPES = {
    "hostaway_stub": "Récupération des tâches Hostaway",
    "lot6b_declarations_internes": "Import des déclarations internes",
    "lot6c_menages_externes": "Extraction des factures externes PDF",
    "lot6d_rapprochement": "Rapprochement des ménages",
    "lot6e_gainperte": "Calcul des écarts de coût",
    "lot6f_cout_complet": "Calcul du coût complet",
    "lot11_controles": "Contrôles de cohérence",
}
# Dossier des factures PDF (copié fichier par fichier — jamais déplacé ni renommé).
PDF_DIR_REL = "01_SOURCES_BRUTES/MenagesExternes/Factures_PDF"
# Entrées lot11 (contrôles transverses) : copiées si présentes ; une absence est
# remontée en avertissement (lot11 échouerait alors explicitement, jamais en silence).
SOURCES_LOT11 = [
    "02_TRAVAIL/Lot9_FluxUnifie/MASTER_CALC_Flux.xlsx",
    "02_TRAVAIL/Lot4quater_SourceResolue/MASTER_CALC_Reservations_Resolues.xlsx",
    "02_TRAVAIL/Lot1_Hostaway/MASTER_CALC_HA_Payout.xlsx",
    "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_Commissions.xlsx",
    "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_NetProprietaire.xlsx",
    "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_Resultats.xlsx",
    "02_TRAVAIL/Lot1_Hostaway/MASTER_CTRL_HA_Anomalies.xlsx",
    "02_TRAVAIL/Lot3_Charges/MASTER_FACT_MAN_Charges.xlsx",
    "01_SOURCES_BRUTES/Charges/SAISIE_Charges_Impacts.xlsx",
    "02_TRAVAIL/Lot4_ReservationsHH/MASTER_FACT_MAN_ReservationsHorsHostaway.xlsx",
    "02_TRAVAIL/Lot5_AcomptesProprietaires/MASTER_FACT_MAN_AcomptesProprietaires.xlsx",
    "02_TRAVAIL/Lot7_IK_Avantages/MASTER_FACT_MAN_IK_Avantages.xlsx",
    "02_TRAVAIL/Lot8_Banque/BANQUE_LOT8_IMPORT.xlsx",
    "01_SOURCES_BRUTES/AirCover/SAISIE_AirCover.xlsx",
    "01_SOURCES_BRUTES/ImputationsAirbnb/SAISIE_ImputationsAirbnb.xlsx",
    "01_SOURCES_BRUTES/AjustementsPostCloture/SAISIE_Ajustements_PostCloture.xlsx",
]

# Scripts moteur copiés (fermeture d'imports auditée : lib_* + lot3/lot7 importés).
SCRIPTS_MOTEUR = [
    "02_TRAVAIL/lot6b_m04_menages_internes.py",
    "02_TRAVAIL/lot6c_menages_externes.py",
    "02_TRAVAIL/lib_menages_externes_pdf.py",   # extraction PDF (mode PDF_AUTOMATIQUE de lot6c)
    "02_TRAVAIL/lot6d_rapprochement_menages.py",
    "02_TRAVAIL/lot6e_gainperte_menages.py",
    "02_TRAVAIL/lot6f_cout_complet_menages.py",
    "02_TRAVAIL/lot11_controles_coherence.py",
    "02_TRAVAIL/lib_parc.py",
    "02_TRAVAIL/lib_ref_history.py",
    "02_TRAVAIL/lib_menage_costs.py",
    "02_TRAVAIL/lib_controls.py",
    "02_TRAVAIL/lib_cloture.py",
    "02_TRAVAIL/lib_settlements.py",
    "02_TRAVAIL/lib_controles_avantages.py",
    "02_TRAVAIL/lib_avantages.py",
    "02_TRAVAIL/lot7_generateur_avantages.py",
    "02_TRAVAIL/lot3_generateur_charges.py",
]

# Stubs contrôlés (copiés depuis runners/stubs_menages/ — remplacent réseau/API dans le
# workspace UNIQUEMENT ; les fichiers réels du dépôt ne sont jamais touchés).
STUB_SHEET_SOURCE = ("stub_lib_sheet_source.py", "02_TRAVAIL/lib_sheet_source.py")
STUB_HOSTAWAY = ("stub_lot6a_hostaway.py", "02_TRAVAIL/stub_lot6a_hostaway.py")

# Fichier « source copiée » des déclarations internes lu par le stub lib_sheet_source.
DECLARATIONS_COPIEES_REL = "02_DONNEES_NORMALISEES/menages/source_sheet_copiee.csv"
INJECTION_HOSTAWAY_REL = "_stub_hostaway_injection.json"

# Sorties régénérées dans le workspace (fichier -> [(onglet, nb_lignes_minimum)]).
SORTIES_CHAINE: dict[str, list[tuple[str, int]]] = {
    "02_TRAVAIL/Lot1_Hostaway/MASTER_FACT_HA_CleaningTasks_Discovery.xlsx":
        [("data", 1), ("MASTER_ENRICHI", 1), ("VUE_COMPTAGE", 1)],
    "02_TRAVAIL/Lot6b_DeclarationsInternes/MASTER_NORM_Declarations_Internes.xlsx":
        [("MASTER_NORMALISE", 1)],
    "02_DONNEES_NORMALISEES/menages/M04_MENAGES_PowerQuery.xlsx":
        [("SOURCE_RAW", 1), ("MASTER", 1)],
    "02_TRAVAIL/Lot6c_MenagesExternes/MASTER_FACT_MEN_MenagesExternes.xlsx":
        [("SOURCE_RAW", 1), ("MASTER", 1), ("VUE_ACTIVE", 1), ("VUE_ECART_HOSTAWAY", 1)],
    "02_TRAVAIL/Lot6d_Rapprochement_Menages/MASTER_CTRL_Rapprochement_Menages.xlsx":
        [("TABLEAU_COMPARAISON", 1)],
    "02_TRAVAIL/Lot6e_GainPerte_Menages/MASTER_CALC_GainPerte_Menages.xlsx":
        [("DETAIL_ECART_COUT", 1)],
    "02_TRAVAIL/Lot6f_CoutComplet_Menages/MASTER_CALC_CoutComplet_Menages.xlsx":
        [("DETAIL_COUT_COMPLET", 1)],
    "02_TRAVAIL/Lot11_Controles/MASTER_CTRL_Coherence.xlsx":
        [("MASTER", 1)],
}

# Étapes exécutées par le runner générique, DANS CET ORDRE (voir audit en tête de module).
STEPS_CHAINE: list[dict[str, Any]] = [
    {"name": "hostaway_stub",
     "script": "02_TRAVAIL/stub_lot6a_hostaway.py",
     "produces": ["02_TRAVAIL/Lot1_Hostaway/MASTER_FACT_HA_CleaningTasks_Discovery.xlsx"]},
    {"name": "lot6b_declarations_internes",
     "script": "02_TRAVAIL/lot6b_m04_menages_internes.py",
     "produces": ["02_TRAVAIL/Lot6b_DeclarationsInternes/MASTER_NORM_Declarations_Internes.xlsx",
                  "02_DONNEES_NORMALISEES/menages/M04_MENAGES_PowerQuery.xlsx"]},
    {"name": "lot6c_menages_externes",
     "script": "02_TRAVAIL/lot6c_menages_externes.py",
     "produces": ["02_TRAVAIL/Lot6c_MenagesExternes/MASTER_FACT_MEN_MenagesExternes.xlsx"]},
    {"name": "lot6d_rapprochement",
     "script": "02_TRAVAIL/lot6d_rapprochement_menages.py",
     "produces": ["02_TRAVAIL/Lot6d_Rapprochement_Menages/MASTER_CTRL_Rapprochement_Menages.xlsx"]},
    {"name": "lot6e_gainperte",
     "script": "02_TRAVAIL/lot6e_gainperte_menages.py",
     "produces": ["02_TRAVAIL/Lot6e_GainPerte_Menages/MASTER_CALC_GainPerte_Menages.xlsx"]},
    {"name": "lot6f_cout_complet",
     "script": "02_TRAVAIL/lot6f_cout_complet_menages.py",
     "produces": ["02_TRAVAIL/Lot6f_CoutComplet_Menages/MASTER_CALC_CoutComplet_Menages.xlsx"]},
    {"name": "lot11_controles",
     "script": "02_TRAVAIL/lot11_controles_coherence.py",
     "produces": ["02_TRAVAIL/Lot11_Controles/MASTER_CTRL_Coherence.xlsx"]},
]


# ── Chemins ──────────────────────────────────────────────────────────────────

def _project_root() -> Path:
    return Path(cfg.PROJECT_ROOT)


def _scripts_root() -> Path:
    """Racine du CODE moteur exécuté par la recette = worktree applicatif COURANT.

    Distinct de PROJECT_ROOT (les DONNÉES) : la recette valide le code du worktree
    (correctifs inclus, ex. lot6c) en lisant les données réelles du projet ciblé. Quand
    PROJECT_ROOT n'est pas surchargé, les deux coïncident (comportement par défaut).
    """
    return Path(cfg.APP_ROOT).parent


def _stubs_dir() -> Path:
    return Path(cfg.MENAGES_STUBS_DIR)


def _source_recette_defaut() -> str | None:
    """CSV de recette par défaut (déclarations simulées) — utilisé si aucune source fournie.

    Jamais le réseau, jamais un cache réel : une source COPIÉE, versionnée avec les stubs.
    """
    p = _stubs_dir() / "source_sheet_recette.csv"
    return p.read_text(encoding="utf-8") if p.exists() else None


def _lock_path() -> Path:
    return Path(cfg.DATA_DIR) / LOCK_NAME


def _pdf_reels() -> list[Path]:
    d = _project_root() / PDF_DIR_REL
    return sorted(d.glob("*.pdf")) if d.exists() else []


def _sources_reelles_presentes() -> list[Path]:
    root = _project_root()
    fichiers = [root / rel for rel in SOURCES_COEUR + SOURCES_LOT11]
    return [p for p in fichiers if p.exists()] + _pdf_reels()


def _sorties_reelles() -> list[Path]:
    root = _project_root()
    return [root / rel for rel in SORTIES_CHAINE]


# ── Préflight ────────────────────────────────────────────────────────────────

def preparer_chaine(mode: str = MODE_COPIES) -> dict[str, Any]:
    """Diagnostic préalable de la chaîne complète. Ne lance rien, n'écrit rien."""
    root = _project_root()
    coeur = [{"chemin_relatif": rel, "nom": Path(rel).name, "present": (root / rel).exists()}
             for rel in SOURCES_COEUR]
    facultatives = [{"chemin_relatif": rel, "nom": Path(rel).name, "present": (root / rel).exists()}
                    for rel in SOURCES_FACULTATIVES]
    lot11 = [{"chemin_relatif": rel, "nom": Path(rel).name, "present": (root / rel).exists()}
             for rel in SOURCES_LOT11]
    pdf = [p.name for p in _pdf_reels()]

    coeur_absents = [s["nom"] for s in coeur if not s["present"]]
    facultatives_absentes = [s["nom"] for s in facultatives if not s["present"]]
    lot11_absents = [s["nom"] for s in lot11 if not s["present"]]
    excel_ouverts = detecter_excel_ouvert(
        [root / rel for rel in SOURCES_COEUR + SOURCES_LOT11] + _sorties_reelles())
    etat_verrou = verrou_lib.inspecter_verrou(_lock_path())

    bloque = mode == MODE_REEL and not cfg.MENAGES_REAL_RECALC_ENABLED
    warnings: list[str] = []
    if coeur_absents:
        warnings.append("Source(s) cœur absente(s) : " + ", ".join(coeur_absents)
                        + " — indispensables à Lot6a-Lot6e.")
    if facultatives_absentes:
        warnings.append("Source(s) facultative(s) absente(s) : " + ", ".join(facultatives_absentes)
                        + " — utilisées seulement par Lot6f/contrôles (non bloquant).")
    if lot11_absents:
        warnings.append("Entrée(s) lot11 absente(s) : " + ", ".join(lot11_absents)
                        + " — seule l'étape « Contrôles de cohérence » serait affectée (non bloquant pour 6a-6f).")
    if not pdf:
        warnings.append("Aucun PDF dans " + cfg.MENAGES_PDF_DIR_REL + ".")
    if excel_ouverts:
        warnings.append("Fichier(s) probablement ouverts dans Excel : " + ", ".join(excel_ouverts) + ".")
    if etat_verrou["etat"] != verrou_lib.ETAT_ABSENT:
        warnings.append("Un verrou de chaîne est présent (" + etat_verrou["etat"] + ").")

    # PRÊTE ne dépend QUE des sources cœur (6a-6e) + verrou + Excel : les facultatives et lot11
    # n'empêchent jamais la recette 6a-6f.
    prete = (not bloque) and (not coeur_absents) and (not excel_ouverts) \
        and etat_verrou["etat"] == verrou_lib.ETAT_ABSENT
    return {
        "mode": mode,
        "etapes": [{"name": s["name"], "libelle": LIBELLES_ETAPES.get(s["name"], s["name"])}
                   for s in STEPS_CHAINE],
        "sources_coeur": coeur, "sources_facultatives": facultatives, "sources_lot11": lot11,
        "facultatives_absentes": facultatives_absentes,
        "pdf_dossier_relatif": cfg.MENAGES_PDF_DIR_REL, "pdf_presents": pdf, "nb_pdf": len(pdf),
        "coeur_absents": coeur_absents, "lot11_absents": lot11_absents,
        "excel_ouverts": excel_ouverts, "verrou": etat_verrou,
        "reel_active": bool(cfg.MENAGES_REAL_RECALC_ENABLED),
        "bloque": bloque,
        "raison": ("Le mode RÉEL est désactivé (MENAGES_REAL_RECALC_ENABLED = False). "
                   "Seule la recette sur copies est exécutable.") if bloque else None,
        "prete": prete, "warnings": warnings,
    }


# ── Workspace ────────────────────────────────────────────────────────────────

def _construire_workspace(run_ts: str, declarations_csv: str,
                          injection_hostaway: dict | None) -> tuple[Path, list[str]]:
    """Miroir minimal du projet + stubs + source déclarations copiée + injection Hostaway.

    Retourne (workspace, avertissements). Copies uniquement — aucun fichier réel déplacé,
    renommé ou modifié.
    """
    base = Path(cfg.MENAGES_CHAINE_WORKSPACE)
    workspace = base / run_ts
    if workspace.exists():
        shutil.rmtree(workspace)
    workspace.mkdir(parents=True, exist_ok=True)
    root = _project_root()          # DONNÉES (projet ciblé)
    code_root = _scripts_root()     # CODE moteur (worktree courant, correctifs inclus)
    warnings: list[str] = []

    # Données réelles (sources cœur) : copiées depuis le projet ciblé.
    for rel in SOURCES_COEUR:
        src = root / rel
        dst = workspace / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
    # Sources facultatives (Lot6f pools) : copiées SI présentes, jamais bloquantes.
    for rel in SOURCES_FACULTATIVES:
        src = root / rel
        if not src.exists():
            warnings.append(f"Source facultative absente, non copiée : {rel}")
            continue
        dst = workspace / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
    # Scripts moteur : copiés depuis le worktree courant (lot6c corrigé, etc.).
    for rel in SCRIPTS_MOTEUR:
        src = code_root / rel
        dst = workspace / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
    for rel in SOURCES_LOT11:
        src = root / rel
        if not src.exists():
            warnings.append(f"Entrée lot11 absente, non copiée : {rel}")
            continue
        dst = workspace / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)

    # Copie des PDF (étape 3 de l'ordre imposé) — fichier par fichier, octets identiques.
    pdf_dst = workspace / PDF_DIR_REL
    pdf_dst.mkdir(parents=True, exist_ok=True)
    for p in _pdf_reels():
        shutil.copy2(p, pdf_dst / p.name)

    # Stubs contrôlés : lib_sheet_source (remplace le réseau) + extraction Hostaway simulée.
    for nom_stub, rel_dst in (STUB_SHEET_SOURCE, STUB_HOSTAWAY):
        shutil.copy2(_stubs_dir() / nom_stub, workspace / rel_dst)

    # Source déclarations COPIÉE (jamais le réseau en mode copies).
    decl = workspace / DECLARATIONS_COPIEES_REL
    decl.parent.mkdir(parents=True, exist_ok=True)
    decl.write_text(declarations_csv, encoding="utf-8")

    if injection_hostaway:
        (workspace / INJECTION_HOSTAWAY_REL).write_text(
            json.dumps(injection_hostaway, ensure_ascii=False, indent=2), encoding="utf-8")

    for rel in SORTIES_CHAINE:
        (workspace / Path(rel).parent).mkdir(parents=True, exist_ok=True)
    return workspace, warnings


# ── Vérification des sorties (onglets + lignes + lisibilité) ─────────────────

def _verifier_sortie_excel(path: Path, attentes: list[tuple[str, int]]) -> dict[str, Any]:
    if not path.exists():
        return {"present": False, "lisible": False, "onglets": {}, "ok": False}
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return {"present": True, "lisible": False, "onglets": {},
                "ok": False, "erreur": f"{type(exc).__name__}: {exc}"}
    onglets: dict[str, Any] = {}
    ok = True
    try:
        for onglet, mini in attentes:
            if onglet not in wb.sheetnames:
                onglets[onglet] = {"present": False, "nb_lignes": 0, "ok": False}
                ok = False
                continue
            ws = wb[onglet]
            nb = sum(1 for r in ws.iter_rows(min_row=2, values_only=True)
                     if any(c is not None for c in r))
            onglets[onglet] = {"present": True, "nb_lignes": nb, "ok": nb >= mini}
            ok = ok and nb >= mini
    finally:
        wb.close()
    return {"present": True, "lisible": True, "onglets": onglets, "ok": ok}


# ── Exécution ────────────────────────────────────────────────────────────────

def executer_chaine(mode: str = MODE_COPIES, declarations_csv: str | None = None,
                    injection_hostaway: dict | None = None,
                    db_path: Path | None = None) -> dict[str, Any]:
    """Relance la chaîne complète sur copies. MODE_REEL refusé tant que le flag est False."""
    db_path = db_path or cfg.DB_PATH   # lu À CHAUD
    if declarations_csv is None:
        declarations_csv = _source_recette_defaut()

    if mode == MODE_REEL and not cfg.MENAGES_REAL_RECALC_ENABLED:
        run_id = _enregistrer_run(
            db_path, mode=MODE_RUN_DB, periode=None, statut=STATUT_BLOQUE,
            git_head=_git_head(), erreur_code="E_MODE_REEL_DESACTIVE",
            erreur_resume="MENAGES_REAL_RECALC_ENABLED=False — aucune exécution, aucune requête réseau.",
        )
        log_event("MENAGE_CHAINE_BLOQUE", {"mode": MODE_REEL}, db_path=db_path)
        return {"ok": False, "run_id": run_id, "statut": STATUT_BLOQUE,
                "erreur_code": "E_MODE_REEL_DESACTIVE",
                "message": "Chaîne réelle désactivée. Aucun fichier réel touché, aucune requête envoyée."}
    if mode not in (MODE_COPIES, MODE_REEL):
        raise ValueError(f"Mode inconnu : {mode}")

    if not declarations_csv:
        run_id = _enregistrer_run(
            db_path, mode=MODE_RUN_DB, periode=None, statut=STATUT_ECHEC, git_head=_git_head(),
            erreur_code="E_SOURCE_DECLARATIONS_ABSENTE",
            erreur_resume="Mode copies : la source déclarations copiée/simulée est obligatoire "
                          "(jamais de reprise silencieuse du réseau ou d'un cache).",
        )
        return {"ok": False, "run_id": run_id, "statut": STATUT_ECHEC,
                "erreur_code": "E_SOURCE_DECLARATIONS_ABSENTE",
                "message": "Fournir la source déclarations copiée (CSV) pour la recette sur copies."}

    plan = preparer_chaine(mode)
    if plan["coeur_absents"]:
        run_id = _enregistrer_run(
            db_path, mode=MODE_RUN_DB, periode=None, statut=STATUT_ECHEC, git_head=_git_head(),
            erreur_code="E_SOURCE_ABSENTE", erreur_resume="Absentes : " + ", ".join(plan["coeur_absents"]),
        )
        return {"ok": False, "run_id": run_id, "statut": STATUT_ECHEC,
                "erreur_code": "E_SOURCE_ABSENTE",
                "message": "Source(s) cœur absente(s) : " + ", ".join(plan["coeur_absents"])}
    if plan["excel_ouverts"]:
        run_id = _enregistrer_run(
            db_path, mode=MODE_RUN_DB, periode=None, statut=STATUT_ECHEC, git_head=_git_head(),
            erreur_code="E_EXCEL_OUVERT", erreur_resume="Ouverts : " + ", ".join(plan["excel_ouverts"]),
        )
        return {"ok": False, "run_id": run_id, "statut": STATUT_ECHEC,
                "erreur_code": "E_EXCEL_OUVERT",
                "message": "Fichier(s) ouverts dans Excel : " + ", ".join(plan["excel_ouverts"])}

    try:
        verrou = verrou_lib.acquerir_verrou(operation=OPERATION, lock_path=_lock_path())
    except verrou_lib.VerrouSaisieChargesDejaPrisError as exc:
        run_id = _enregistrer_run(
            db_path, mode=MODE_RUN_DB, periode=None, statut=STATUT_VERROUILLE,
            git_head=_git_head(), erreur_code="E_VERROU", erreur_resume=str(exc),
        )
        return {"ok": False, "run_id": run_id, "statut": STATUT_VERROUILLE,
                "erreur_code": "E_VERROU", "message": "Une chaîne est déjà en cours."}

    date_debut = datetime.now(timezone.utc).isoformat()
    t0 = time.monotonic()
    try:
        sources_reelles = _sources_reelles_presentes()
        sha_sources_avant = {str(p.relative_to(_project_root())): _sha256(p) for p in sources_reelles}
        sha_sorties_reelles_avant = {rel: (_sha256(_project_root() / rel)
                                           if (_project_root() / rel).exists() else None)
                                     for rel in SORTIES_CHAINE}
        snap = snapshot_service.create_snapshot(
            "MENAGES_CHAINE_AVANT",
            [_project_root() / rel for rel in SOURCES_COEUR]
            + [p for p in _sorties_reelles() if p.exists()],
            db_path=db_path)

        run_ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        workspace, ws_warnings = _construire_workspace(run_ts, declarations_csv, injection_hostaway)
        requete = {"allowed_root": str(workspace.resolve()), "workspace": str(workspace.resolve()),
                   "steps": STEPS_CHAINE, "timeout": cfg.MENAGES_CHAINE_TIMEOUT_SECONDS}
        requete_path = workspace / "_requete.json"
        reponse_path = workspace / "_reponse.json"
        requete_path.write_text(json.dumps(requete, ensure_ascii=False, indent=2), encoding="utf-8")

        runner = Path(cfg.MENAGES_RECALC_RUNNER)
        engine_py = Path(cfg.MENAGES_ENGINE_PYTHON)
        erreur_code = erreur_resume = None
        etapes: list[dict[str, Any]] = []
        if not runner.exists():
            statut, erreur_code, erreur_resume = STATUT_ECHEC, "E_RUNNER", f"Runner absent : {runner}"
        elif not engine_py.exists():
            statut, erreur_code, erreur_resume = STATUT_ECHEC, "E_RUNNER", f"Interpréteur moteur absent : {engine_py}"
        else:
            try:
                proc = subprocess.run(
                    [str(engine_py), str(runner), str(requete_path), str(reponse_path)],
                    capture_output=True, text=True,
                    timeout=cfg.MENAGES_CHAINE_TIMEOUT_SECONDS + 60,
                )
            except subprocess.TimeoutExpired:
                statut, erreur_code, erreur_resume = STATUT_ECHEC, "E_TIMEOUT", "Le runner a dépassé le délai."
            else:
                if not reponse_path.exists():
                    statut = STATUT_ECHEC
                    erreur_code = "E_REPONSE_INVALIDE"
                    erreur_resume = f"Pas de réponse runner (rc={proc.returncode}). {proc.stderr[-500:]}"
                else:
                    reponse = json.loads(reponse_path.read_text(encoding="utf-8"))
                    etapes = reponse.get("steps", [])
                    nb_ok = sum(1 for e in etapes if e.get("statut") == "OK")
                    if reponse.get("ok") and nb_ok == len(STEPS_CHAINE):
                        statut = STATUT_SUCCES
                    elif nb_ok == 0:
                        statut = STATUT_ECHEC
                    else:
                        statut = STATUT_PARTIEL
                    if statut != STATUT_SUCCES:
                        premiere = next((e for e in etapes if e.get("statut") != "OK"), {})
                        erreur_code = premiere.get("erreur_code") or "E_CONTROLE_ECHEC"
                        erreur_resume = premiere.get("detail")

        # ── Vérification des sorties DU WORKSPACE (onglets + lignes + lisibilité) ─
        verif_sorties = {}
        sha_sorties_apres = {}
        for rel, attentes in SORTIES_CHAINE.items():
            produit = workspace / rel
            verif_sorties[rel] = _verifier_sortie_excel(produit, attentes)
            sha_sorties_apres[rel] = _sha256(produit) if produit.exists() else None
        if statut == STATUT_SUCCES and not all(v["ok"] for v in verif_sorties.values()):
            statut = STATUT_ECHEC
            erreur_code = "E_SORTIE_INVALIDE"
            erreur_resume = "Sortie produite mais onglet/lignes attendus manquants (voir verif_sorties)."

        comparaison = {
            "avant": _lire_comparaison(
                _project_root() / "02_TRAVAIL/Lot6d_Rapprochement_Menages/MASTER_CTRL_Rapprochement_Menages.xlsx"),
            "apres": _lire_comparaison(
                workspace / "02_TRAVAIL/Lot6d_Rapprochement_Menages/MASTER_CTRL_Rapprochement_Menages.xlsx"),
        }

        # ── Garde : AUCUN fichier réel modifié (sources + sorties + PDF) ──────────
        sha_sources_apres = {str(p.relative_to(_project_root())): (_sha256(p) if p.exists() else None)
                             for p in sources_reelles}
        sha_sorties_reelles_apres = {rel: (_sha256(_project_root() / rel)
                                           if (_project_root() / rel).exists() else None)
                                     for rel in SORTIES_CHAINE}
        reel_intact = (sha_sources_apres == sha_sources_avant
                       and sha_sorties_reelles_apres == sha_sorties_reelles_avant)
        if not reel_intact:
            statut = STATUT_ECHEC
            erreur_code = "E_REEL_MODIFIE"
            erreur_resume = "ANOMALIE GRAVE : un fichier réel a changé pendant la chaîne sur copies."

        duree = round(time.monotonic() - t0, 2)
        date_fin = datetime.now(timezone.utc).isoformat()
        run_id = _enregistrer_run(
            db_path, mode=MODE_RUN_DB, periode=None, statut=statut,
            date_debut=date_debut, date_fin=date_fin, duree_secondes=duree,
            snapshot_id=snap["id"], git_head=_git_head(), workspace_path=str(workspace),
            sha256_sources_avant=json.dumps(sha_sources_avant, ensure_ascii=False),
            sha256_sorties_avant=json.dumps(sha_sorties_reelles_avant, ensure_ascii=False),
            sha256_sorties_apres=json.dumps(sha_sorties_apres, ensure_ascii=False),
            comparaison_json=json.dumps({"rapprochement": comparaison, "reel_intact": reel_intact,
                                         "verif_sorties": verif_sorties}, ensure_ascii=False),
            etapes_json=json.dumps(etapes, ensure_ascii=False, default=str),
            erreur_code=erreur_code, erreur_resume=erreur_resume,
        )
        log_event("MENAGE_CHAINE", {"mode": mode, "statut": statut, "run_id": run_id,
                                    "reel_intact": reel_intact}, db_path=db_path)
        return {
            "ok": statut == STATUT_SUCCES, "run_id": run_id, "statut": statut, "mode": mode,
            "duree_secondes": duree, "snapshot_id": snap["id"], "workspace": str(workspace),
            "etapes": etapes, "verif_sorties": verif_sorties, "comparaison": comparaison,
            "sha256_sorties_apres": sha_sorties_apres, "reel_intact": reel_intact,
            "warnings": ws_warnings, "erreur_code": erreur_code, "erreur_resume": erreur_resume,
        }
    except Exception as exc:
        # Défense en profondeur : jamais de 500 ni de traceback navigateur.
        run_id = _enregistrer_run(
            db_path, mode=MODE_RUN_DB, periode=None, statut=STATUT_ECHEC, date_debut=date_debut,
            git_head=_git_head(), erreur_code="E_INATTENDU",
            erreur_resume=f"{type(exc).__name__}: {exc}",
        )
        log_event("MENAGE_CHAINE_ERREUR", {"erreur": type(exc).__name__}, db_path=db_path)
        return {"ok": False, "run_id": run_id, "statut": STATUT_ECHEC, "erreur_code": "E_INATTENDU",
                "message": "Une erreur inattendue a interrompu la recette "
                           f"({type(exc).__name__}). Aucun fichier réel n'a été touché."}
    finally:
        try:
            verrou_lib.liberer_verrou(verrou)
        except verrou_lib.LiberationVerrouSaisieChargesError:
            log_event("MENAGE_CHAINE_VERROU_NON_LIBERE", {"lock": str(_lock_path())}, db_path=db_path)
