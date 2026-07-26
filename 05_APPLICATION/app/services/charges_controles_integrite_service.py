"""Contrôles d'intégrité des charges avant recalcul aval (esprit Lot11, côté application).

Motif : Lot10 n'infère JAMAIS le propriétaire. Une charge `refacturable=OUI` + `VALIDE` sans
propriétaire déterminé ne peut donc produire ni préfacture ni montant dû : elle serait silencieusement
perdue. Ces contrôles la rendent visible et BLOQUANTE avant tout recalcul.

Lecture seule. Aucun contrôle ne modifie une donnée.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

import app.config as cfg

BLOQUANT = "BLOQUANT"
INFO = "INFO"

C_REFAC_ORPHELINE = "CTRL_CHG_REFAC_SANS_PROPRIETAIRE"
C_PROP_INCONNU = "CTRL_CHG_PROPRIETAIRE_INCONNU"
C_PROP_INCOHERENT = "CTRL_CHG_PROPRIETAIRE_INCOHERENT_LOGEMENT"
C_STATUT_INVALIDE = "CTRL_CHG_STATUT_NON_INGERABLE"

MESSAGES = {
    C_REFAC_ORPHELINE: ("Charge refacturable validée sans propriétaire déterminé : impossible de "
                        "calculer la préfacture et le montant dû."),
    C_PROP_INCONNU: "Propriétaire de la charge inconnu dans le référentiel.",
    C_PROP_INCOHERENT: ("Propriétaire de la charge incohérent avec le propriétaire du logement à "
                        "la période."),
    C_STATUT_INVALIDE: "Statut de charge non ingérable par le calcul aval.",
}


def _lire(path: Path, sheet: str) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[sheet]
    data = list(ws.iter_rows(values_only=True))
    wb.close()
    if not data:
        return []
    hdr = list(data[0])
    out = []
    for r in data[1:]:
        d = dict(zip(hdr, r))
        cid = str(d.get("charge_id") or d.get("logement_id") or "").strip()
        if cid and cid[0] in "#[<←-*":
            continue
        out.append(d)
    return out


def _txt(v) -> str:
    return str(v or "").strip()


def controler(master_charges: Path | None = None, ref_path: Path | None = None) -> dict[str, Any]:
    """Retourne les anomalies détectées sur les charges du MASTER (Lot3).

    Un contrôle BLOQUANT signifie : le recalcul aval ne doit pas être considéré comme fiable tant
    que l'anomalie n'est pas corrigée.
    """
    charges = _lire(Path(master_charges or cfg.MASTER_CHARGES), "MASTER")
    ref = Path(ref_path or cfg.REF_SETUP)
    proprios = {_txt(r.get("proprietaire_id")) for r in _lire(ref, "REF_Proprietaires")}
    proprios.discard("")
    gestion = _lire(ref, "REF_Gestion_Logements_Hist")

    def _proprio_du_logement(logement_id: str, mois: str) -> set[str]:
        out = set()
        for g in gestion:
            if _txt(g.get("logement_id")) != logement_id:
                continue
            if _txt(g.get("statut_gestion")).upper() != "ACTIF":
                continue
            debut = _txt(g.get("date_debut"))[:7]
            fin = _txt(g.get("date_fin"))[:7]
            if debut and mois and debut > mois:
                continue
            if fin and mois and fin < mois:
                continue
            p = _txt(g.get("proprietaire_id"))
            if p:
                out.add(p)
        return out

    anomalies: list[dict[str, Any]] = []

    def _ajout(code, severite, charge, detail=""):
        anomalies.append({
            "code": code, "severite": severite, "message": MESSAGES.get(code, code),
            "charge_id": _txt(charge.get("charge_id")), "logement_id": _txt(charge.get("logement_id")),
            "proprietaire_id": _txt(charge.get("proprietaire_id")), "mois": _txt(charge.get("mois")),
            "montant": charge.get("montant"), "detail": detail,
        })

    for c in charges:
        statut = _txt(c.get("statut_controle")).upper()
        refac = _txt(c.get("refacturable")).upper() == "OUI"
        prop = _txt(c.get("proprietaire_id"))
        log = _txt(c.get("logement_id"))
        mois = _txt(c.get("mois"))

        if statut == "VALIDE" and refac and not prop:
            _ajout(C_REFAC_ORPHELINE, BLOQUANT, c,
                   "refacturable=OUI, statut=VALIDE, proprietaire_id vide")

        if prop and proprios and prop not in proprios:
            _ajout(C_PROP_INCONNU, BLOQUANT, c, f"proprietaire_id={prop} absent du référentiel")

        if prop and log:
            attendus = _proprio_du_logement(log, mois)
            if attendus and prop not in attendus:
                _ajout(C_PROP_INCOHERENT, BLOQUANT, c,
                       f"attendu(s) {sorted(attendus)} pour {log} en {mois}")

        if statut not in ("VALIDE", "A_CONTROLER", "REJETE", ""):
            _ajout(C_STATUT_INVALIDE, INFO, c, f"statut={statut}")

    bloquants = [a for a in anomalies if a["severite"] == BLOQUANT]
    return {
        "nb_charges": len(charges),
        "anomalies": anomalies,
        "nb_bloquants": len(bloquants),
        "recalcul_fiable": not bloquants,
        "statut": BLOQUANT if bloquants else "OK",
    }
