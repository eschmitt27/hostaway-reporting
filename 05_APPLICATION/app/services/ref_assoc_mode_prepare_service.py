"""APP-3b-0 — Préparation contrôlée du référentiel REF_Assoc_Mode.

Modes :
    diagnostiquer(ref_path)
        Lecture seule — hash + état feuille REF_Assoc_Mode dans REF_Setup.xlsm.

    preparer_sur_copie(ref_source, output_dir)
        Copie ref_source dans output_dir, crée REF_Assoc_Mode + table structurée sur
        la copie de travail, vérifie cohérence et préservation VBA/package sensible.
        Ne touche jamais au fichier source.

    executer_migration_reelle(confirmation, ...)
        Migration réelle future. Gated par cfg.REF_ASSOC_MODE_REAL_WRITE_ENABLED.
        Si False → MIGRATION_REELLE_NON_ACTIVE, aucune écriture.

Gardes :
    - output_dir != dossier du fichier source
    - Feuille incoherente existante → refus (pas d'écrasement)
    - Source hash inchangé vérifié avant et après chaque opération
"""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
from typing import Any
import zipfile

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

import app.config as cfg


# ── Référentiel cible ──────────────────────────────────────────────────────────

SHEET_NAME = "REF_Assoc_Mode"
TABLE_NAME = "tblRefAssocMode"
CONFIRMATION_EXECUTION = "MIGRER_REF_ASSOC_MODE"
REF_PREPARE_NAME = "REF_Setup_assoc_mode_prepare.xlsm"
REF_REF_NAME = "REF_Setup_assoc_mode_ref.xlsm"
MANIFEST_NAME = "manifest_ref_assoc_mode.json"

EXPECTED_HEADERS = [
    "assoc_mode_id",
    "mode_paiement_id",
    "associe_id",
    "assoc_mode",
    "actif",
    "commentaire",
]

INITIAL_ROWS: list[tuple[str, str, str, str, str, str]] = [
    ("AM_001", "PAY_001", "",          "BANQUE",     "OUI", "Banque professionnelle"),
    ("AM_002", "PAY_002", "",          "LIQ",        "OUI", "Espèces caisse"),
    ("AM_003", "PAY_003", "PERS_EWAN", "EWAN-CB",    "OUI", "Carte associée Ewan"),
    ("AM_004", "PAY_003", "PERS_WAFA", "WAFA-CB",    "OUI", "Carte associée Wafa"),
    ("AM_005", "PAY_004", "PERS_EWAN", "EWAN-PERSO", "OUI", "Compte personnel Ewan"),
    ("AM_006", "PAY_004", "PERS_WAFA", "WAFA-PERSO", "OUI", "Compte personnel Wafa"),
    ("AM_007", "PAY_005", "",          "ADEF",       "OUI", "Mode à définir — contrôle obligatoire"),
]

# ── Parties ZIP sensibles à préserver ─────────────────────────────────────────

_VBA_PROJECT_PATH = "xl/vbaProject.bin"
_VBA_SIG_PATH = "xl/vbaProjectSignature.bin"
_CONTENT_TYPES_PATH = "[Content_Types].xml"
_WORKBOOK_RELS_PATH = "xl/_rels/workbook.xml.rels"
_VBA_MACRO_ENABLED_MARKER = "macroEnabled"
_VBA_RELATIONSHIP_TYPE = "relationships/vbaProject"

_SENSITIVE_PREFIXES = (
    "xl/activeX/",
    "xl/ctrlProps/",
    "xl/embeddings/",
    "xl/externalLinks/",
    "xl/connections.xml",
    "customUI/",
    "docProps/custom.xml",
    "xl/printerSettings/",
)


# ── Utilitaires ZIP/VBA ────────────────────────────────────────────────────────

def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(str(path), "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _zip_entry_sha256(zf: zipfile.ZipFile, name: str) -> str | None:
    try:
        return hashlib.sha256(zf.read(name)).hexdigest()
    except KeyError:
        return None


def _vba_snapshot(path: Path) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(str(path), "r") as zf:
            names = set(zf.namelist())
            vba_present = _VBA_PROJECT_PATH in names
            vba_sha256 = _zip_entry_sha256(zf, _VBA_PROJECT_PATH) if vba_present else None
            sig_present = _VBA_SIG_PATH in names
            sig_sha256 = _zip_entry_sha256(zf, _VBA_SIG_PATH) if sig_present else None
            macro_enabled = False
            vba_rel_valid = False
            if _CONTENT_TYPES_PATH in names:
                ct = zf.read(_CONTENT_TYPES_PATH).decode("utf-8", errors="replace")
                macro_enabled = _VBA_MACRO_ENABLED_MARKER in ct
            if _WORKBOOK_RELS_PATH in names:
                wr = zf.read(_WORKBOOK_RELS_PATH).decode("utf-8", errors="replace")
                vba_rel_valid = _VBA_RELATIONSHIP_TYPE in wr
            return {
                "vba_project_present": vba_present,
                "vba_project_sha256": vba_sha256,
                "vba_signature_present": sig_present,
                "vba_signature_sha256": sig_sha256,
                "macro_enabled_declared": macro_enabled,
                "vba_relationship_valid": vba_rel_valid,
            }
    except Exception as exc:
        return {"error": str(exc)}


def _check_vba_integrity(ref_path: Path, work_path: Path) -> list[str]:
    violations: list[str] = []
    try:
        snap_r = _vba_snapshot(ref_path)
        snap_w = _vba_snapshot(work_path)
        if "error" in snap_r or "error" in snap_w:
            return [f"VBA_PRESERVATION_ECHEC: lecture ZIP impossible "
                    f"(ref={snap_r.get('error')}, work={snap_w.get('error')})"]
        if snap_r["vba_project_present"]:
            if not snap_w["vba_project_present"]:
                violations.append("VBA_PRESERVATION_ECHEC: xl/vbaProject.bin disparu apres preparation")
            elif snap_r["vba_project_sha256"] != snap_w["vba_project_sha256"]:
                violations.append(
                    f"VBA_PRESERVATION_ECHEC: xl/vbaProject.bin modifie "
                    f"(ref={snap_r['vba_project_sha256'][:16]}... "
                    f"work={snap_w['vba_project_sha256'][:16]}...)"
                )
        if snap_r["vba_signature_present"]:
            if not snap_w["vba_signature_present"]:
                violations.append("VBA_PRESERVATION_ECHEC: xl/vbaProjectSignature.bin disparu")
            elif snap_r["vba_signature_sha256"] != snap_w["vba_signature_sha256"]:
                violations.append("VBA_PRESERVATION_ECHEC: xl/vbaProjectSignature.bin modifie")
        if snap_r["macro_enabled_declared"] and not snap_w["macro_enabled_declared"]:
            violations.append("VBA_PRESERVATION_ECHEC: macro-enabled disparu de Content_Types.xml")
        if snap_r["vba_relationship_valid"] and not snap_w["vba_relationship_valid"]:
            violations.append("VBA_PRESERVATION_ECHEC: relation VBA disparue de workbook.xml.rels")
    except Exception as exc:
        violations.append(f"VBA_PRESERVATION_ECHEC: verification impossible : {exc}")
    return violations


def _check_sensitive_parts(ref_path: Path, work_path: Path) -> list[str]:
    violations: list[str] = []
    try:
        with (
            zipfile.ZipFile(str(ref_path), "r") as zf_r,
            zipfile.ZipFile(str(work_path), "r") as zf_w,
        ):
            ref_names = set(zf_r.namelist())
            work_names = set(zf_w.namelist())
            for entry in sorted(ref_names):
                for prefix in _SENSITIVE_PREFIXES:
                    if entry == prefix or entry.startswith(prefix):
                        if entry not in work_names:
                            violations.append(
                                f"PACKAGE_SENSIBLE_ECHEC: {entry} disparu apres preparation"
                            )
                        elif _zip_entry_sha256(zf_r, entry) != _zip_entry_sha256(zf_w, entry):
                            violations.append(
                                f"PACKAGE_SENSIBLE_ECHEC: {entry} modifie apres preparation"
                            )
                        break
    except Exception as exc:
        violations.append(f"PACKAGE_SENSIBLE_ECHEC: verification impossible : {exc}")
    return violations


# ── Lecture feuille ────────────────────────────────────────────────────────────

def _read_sheet_rows(path: Path, sheet_name: str) -> list[list[Any]]:
    """Toutes les lignes (header inclus) de la feuille, data_only=True."""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True, keep_vba=True)
    try:
        if sheet_name not in wb.sheetnames:
            return []
        ws = wb[sheet_name]
        return [list(row) for row in ws.iter_rows(values_only=True)
                if any(c is not None for c in row)]
    finally:
        wb.close()


def _sheetnames(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True, keep_vba=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


# ── Vérification cohérence ────────────────────────────────────────────────────

def _verifier_coherence_feuille(path: Path) -> list[str]:
    """Vérifie la feuille REF_Assoc_Mode dans le fichier path (après écriture)."""
    violations: list[str] = []
    rows = _read_sheet_rows(path, SHEET_NAME)
    if not rows:
        return [f"FEUILLE_VIDE_OU_ABSENTE: {SHEET_NAME}"]

    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    if headers != EXPECTED_HEADERS:
        violations.append(f"HEADERS_INCORRECTS: attendu {EXPECTED_HEADERS}, lu {headers}")
        return violations

    data = rows[1:]
    if len(data) != len(INITIAL_ROWS):
        violations.append(f"NOMBRE_LIGNES_INCORRECT: attendu {len(INITIAL_ROWS)}, lu {len(data)}")

    ids: list[str] = []
    mode_assoc_pairs: list[tuple[str, str]] = []
    abbrevs: list[str] = []

    for i, row in enumerate(data):
        row_vals = [str(c).strip() if c is not None else "" for c in row]
        row_dict = dict(zip(EXPECTED_HEADERS, row_vals))

        am_id = row_dict.get("assoc_mode_id", "")
        mode = row_dict.get("mode_paiement_id", "")
        assoc = row_dict.get("associe_id", "")
        abbrev = row_dict.get("assoc_mode", "")

        ids.append(am_id)
        mode_assoc_pairs.append((mode, assoc))
        abbrevs.append(abbrev)

    # Unicité assoc_mode_id
    seen_ids: set[str] = set()
    for v in ids:
        if v in seen_ids:
            violations.append(f"DOUBLON_ASSOC_MODE_ID: {v!r}")
        seen_ids.add(v)

    # Unicité (mode_paiement_id, associe_id)
    seen_pairs: set[tuple[str, str]] = set()
    for p in mode_assoc_pairs:
        if p in seen_pairs:
            violations.append(f"DOUBLON_MODE_ASSOCIE: {p}")
        seen_pairs.add(p)

    # Unicité assoc_mode
    seen_abbrev: set[str] = set()
    for v in abbrevs:
        if v in seen_abbrev:
            violations.append(f"DOUBLON_ASSOC_MODE: {v!r}")
        seen_abbrev.add(v)

    return violations


def _feuille_coherente_vs_attendue(path: Path) -> bool:
    """True si la feuille existante contient exactement les données attendues."""
    rows = _read_sheet_rows(path, SHEET_NAME)
    if not rows:
        return False
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    if headers != EXPECTED_HEADERS:
        return False
    data = rows[1:]
    if len(data) != len(INITIAL_ROWS):
        return False
    for row, expected in zip(data, INITIAL_ROWS):
        actual = tuple(str(c).strip() if c is not None else "" for c in row)
        if actual != expected:
            return False
    return True


# ── Création feuille ───────────────────────────────────────────────────────────

def _ajouter_feuille(work_path: Path) -> None:
    """Ouvre work_path (xlsm), crée REF_Assoc_Mode avec table structurée, sauvegarde."""
    wb = openpyxl.load_workbook(str(work_path), data_only=False, keep_vba=True)
    try:
        ws = wb.create_sheet(SHEET_NAME)
        ws.append(EXPECTED_HEADERS)
        for row in INITIAL_ROWS:
            ws.append(list(row))

        n_rows = 1 + len(INITIAL_ROWS)
        n_cols = len(EXPECTED_HEADERS)
        last_col_letter = openpyxl.utils.get_column_letter(n_cols)
        table_ref = f"A1:{last_col_letter}{n_rows}"

        tbl = Table(displayName=TABLE_NAME, ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleLight2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tbl.tableStyleInfo = style
        ws.add_table(tbl)
    finally:
        wb.save(str(work_path))
        wb.close()


# ── Diagnostic ────────────────────────────────────────────────────────────────

def diagnostiquer(ref_path: Path | None = None) -> dict[str, Any]:
    """Lecture seule. Hash + état feuille REF_Assoc_Mode. Ne modifie jamais la source."""
    p = Path(ref_path or cfg.REF_SETUP)
    errors: list[str] = []
    feuille_presente = False
    coherente = False
    violations: list[str] = []

    sha256 = None
    try:
        sha256 = _sha256_file(p)
    except Exception as exc:
        errors.append(f"HASH_ILLISIBLE: {exc}")

    try:
        sheets = _sheetnames(p)
        feuille_presente = SHEET_NAME in sheets
    except Exception as exc:
        errors.append(f"SHEETNAMES_ILLISIBLES: {exc}")
        sheets = []

    if feuille_presente:
        try:
            violations = _verifier_coherence_feuille(p)
            coherente = not bool(violations)
        except Exception as exc:
            errors.append(f"COHERENCE_ILLISIBLE: {exc}")

    vba = _vba_snapshot(p) if p.exists() else {}

    return {
        "status": "OK" if not errors else "ERREUR",
        "path": str(p),
        "sha256": sha256,
        "feuille_presente": feuille_presente,
        "feuille_coherente": coherente if feuille_presente else None,
        "coherence_violations": violations,
        "vba": vba,
        "errors": errors,
        "migration_necessaire": not feuille_presente or not coherente,
    }


# ── Préparation sur copie ─────────────────────────────────────────────────────

def _assert_output_dir_safe(ref_source: Path, output_dir: Path) -> None:
    """output_dir ne doit pas être le dossier du fichier source ni le fichier lui-même."""
    src_dir = Path(ref_source).resolve().parent
    out = Path(output_dir).resolve()
    if out == src_dir or out == Path(ref_source).resolve():
        raise RuntimeError(
            f"output_dir interdit : {out} correspond au dossier ou au fichier source {src_dir}"
        )
    # Jamais sous cfg.REF_SETUP parent (pour les tests aussi)
    real_ref_dir = cfg.REF_SETUP.resolve().parent
    if out == real_ref_dir or str(out).startswith(str(real_ref_dir) + os.sep):
        raise RuntimeError(
            f"output_dir interdit : {out} est sous le dossier du fichier REF_Setup réel"
        )


def preparer_sur_copie(
    ref_source: Path | None = None,
    output_dir: Path | None = None,
) -> dict[str, Any]:
    """Prépare REF_Assoc_Mode sur une copie uniquement.

    Ne touche jamais au fichier source.
    Retourne un manifest avec status OK / FEUILLE_INCOHERENTE_REFUS_ECRASEMENT / ERREUR.
    """
    p_src = Path(ref_source or cfg.REF_SETUP)
    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    p_out = Path(output_dir or (Path(tempfile.gettempdir()) / f"ref_assoc_mode_{ts}"))
    p_out.mkdir(parents=True, exist_ok=True)

    errors: list[str] = []

    _assert_output_dir_safe(p_src, p_out)

    # Hash source avant
    try:
        sha256_avant = _sha256_file(p_src)
    except Exception as exc:
        return {"status": "ERREUR", "error": f"SHA256 source illisible : {exc}"}

    # Diagnostic initial
    diag_avant = diagnostiquer(p_src)

    # Vérifier feuille déjà présente (idempotence)
    if diag_avant["feuille_presente"]:
        if diag_avant["feuille_coherente"]:
            return {
                "status": "OK",
                "feuille_deja_presente": True,
                "path": str(p_src),
                "sha256_avant": sha256_avant,
                "message": "REF_Assoc_Mode déjà présente et cohérente — aucune action.",
            }
        else:
            return {
                "status": "FEUILLE_INCOHERENTE_REFUS_ECRASEMENT",
                "path": str(p_src),
                "coherence_violations": diag_avant["coherence_violations"],
                "message": "REF_Assoc_Mode présente mais incohérente — refus d'écrasement.",
            }

    # Copies : référence immuable + copie de travail
    p_ref_copy = p_out / REF_REF_NAME
    p_work_copy = p_out / REF_PREPARE_NAME
    shutil.copy2(str(p_src), str(p_ref_copy))
    shutil.copy2(str(p_src), str(p_work_copy))

    original_sheets = _sheetnames(p_ref_copy)
    vba_snap_avant = _vba_snapshot(p_ref_copy)

    # Ajout de la feuille sur la copie de travail
    try:
        _ajouter_feuille(p_work_copy)
    except Exception as exc:
        errors.append(f"CREATION_FEUILLE_ECHEC: {exc}")

    # Vérification cohérence dans la copie de travail
    coherence_violations: list[str] = []
    if not errors:
        try:
            coherence_violations = _verifier_coherence_feuille(p_work_copy)
            if coherence_violations:
                errors.extend(coherence_violations)
        except Exception as exc:
            errors.append(f"VERIFICATION_COHERENCE_ECHEC: {exc}")

    # Feuilles originales préservées
    preserved_sheets_errors: list[str] = []
    if not errors:
        try:
            work_sheets = _sheetnames(p_work_copy)
            for s in original_sheets:
                if s not in work_sheets:
                    preserved_sheets_errors.append(f"FEUILLE_SUPPRIMEE: {s}")
            if SHEET_NAME not in work_sheets:
                preserved_sheets_errors.append(f"FEUILLE_ABSENTE_APRES_PREPARATION: {SHEET_NAME}")
            errors.extend(preserved_sheets_errors)
        except Exception as exc:
            errors.append(f"SHEETS_VERIFICATION_ECHEC: {exc}")

    # Intégrité VBA
    vba_violations = _check_vba_integrity(p_ref_copy, p_work_copy)
    if vba_violations:
        errors.extend(vba_violations)

    # Parties sensibles
    sensitive_violations = _check_sensitive_parts(p_ref_copy, p_work_copy)
    if sensitive_violations:
        errors.extend(sensitive_violations)

    vba_snap_apres = _vba_snapshot(p_work_copy)

    # Invariant : source inchangée
    try:
        sha256_apres = _sha256_file(p_src)
    except Exception as exc:
        sha256_apres = None
        errors.append(f"SHA256_SOURCE_APRES_ILLISIBLE: {exc}")

    if sha256_apres is not None and sha256_apres != sha256_avant:
        errors.append(
            f"SOURCE_MODIFIEE_PENDANT_PREPARATION: "
            f"avant={sha256_avant[:16]}... apres={sha256_apres[:16]}..."
        )

    status = "OK" if not errors else "ERREUR"

    manifest = {
        "operation": "APP-3b-0 REF_Assoc_Mode sur copie",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": status,
        "source_hash_avant": sha256_avant,
        "source_hash_apres": sha256_apres,
        "source_inchangee": sha256_apres == sha256_avant,
        "feuille_deja_presente": False,
        "paths": {
            "source": str(p_src),
            "output_dir": str(p_out),
            "ref_copy": str(p_ref_copy),
            "ref_setup_prepare": str(p_work_copy),
        },
        "original_sheets": original_sheets,
        "vba_snapshots": {"avant": vba_snap_avant, "apres": vba_snap_apres},
        "coherence_violations": coherence_violations,
        "vba_violations": vba_violations,
        "sensitive_violations": sensitive_violations,
        "preserved_sheets_errors": preserved_sheets_errors,
        "errors": errors,
    }

    manifest_path = p_out / MANIFEST_NAME
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )

    return manifest


# ── Exécution réelle (gated) ───────────────────────────────────────────────────

def executer_migration_reelle(
    *,
    confirmation: str,
    ref_path: Path | None = None,
    work_dir: Path | None = None,
) -> dict[str, Any]:
    """Migration réelle future.

    Gated par cfg.REF_ASSOC_MODE_REAL_WRITE_ENABLED (False en APP-3b-0).
    Ne jamais appeler sans activation explicite du flag et confirmation humaine.
    """
    if not cfg.REF_ASSOC_MODE_REAL_WRITE_ENABLED:
        return {
            "status": "MIGRATION_REELLE_NON_ACTIVE",
            "reason": "cfg.REF_ASSOC_MODE_REAL_WRITE_ENABLED = False",
            "message": (
                "Migration réelle désactivée. "
                "Activer cfg.REF_ASSOC_MODE_REAL_WRITE_ENABLED avant d'exécuter."
            ),
        }

    if confirmation != CONFIRMATION_EXECUTION:
        return {
            "status": "REFUSE",
            "reason": "CONFIRMATION_INCORRECTE",
            "message": f"Attendu : {CONFIRMATION_EXECUTION!r}",
        }

    p_ref = Path(ref_path or cfg.REF_SETUP)
    if p_ref.resolve() != cfg.REF_SETUP.resolve():
        return {
            "status": "REFUSE",
            "reason": "CIBLE_NON_AUTORISEE",
            "message": f"Cible {p_ref} != {cfg.REF_SETUP}",
        }

    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    root = Path(work_dir or (Path(tempfile.gettempdir()) / f"ref_assoc_mode_real_{ts}"))
    root.mkdir(parents=True, exist_ok=True)

    before_hash = _sha256_file(p_ref)
    backup_path = root / f"{p_ref.name}.backup"
    shutil.copy2(str(p_ref), str(backup_path))

    # Préparation sur copie d'abord
    manifest = preparer_sur_copie(ref_source=p_ref, output_dir=root / "tmp_validation")
    if manifest.get("status") != "OK" or manifest.get("feuille_deja_presente"):
        manifest["real_status"] = (
            "DEJA_OK" if manifest.get("feuille_deja_presente")
            else "REFUSE_VALIDATION_TEMPORAIRE"
        )
        return manifest

    tmp_work = Path(manifest["paths"]["ref_setup_prepare"])
    try:
        os.replace(str(tmp_work), str(p_ref))
        after_hash = _sha256_file(p_ref)
        manifest.update({
            "real_status": "OK",
            "backup_path": str(backup_path),
            "real_hash_before": before_hash,
            "real_hash_after": after_hash,
        })
    except Exception as exc:
        os.replace(str(backup_path), str(p_ref))
        rollback_hash = _sha256_file(p_ref)
        manifest.update({
            "real_status": "ROLLBACK",
            "error": str(exc),
            "rollback_hash": rollback_hash,
            "rollback_hash_verified": rollback_hash == before_hash,
        })

    return manifest
