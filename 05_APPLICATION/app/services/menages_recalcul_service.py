"""APP-2b — Service de recalcul ménages (sécurisé, sur copies).

Ce service orchestre un cycle de recalcul CONTRÔLÉ du rapprochement ménages, sans jamais importer
le moteur (interdiction `test_no_import_of_travail_modules`) et sans jamais toucher un fichier
métier réel.

Deux modes (cf. `config.MENAGES_REAL_RECALC_ENABLED`) :

  MODE_COPIES  — recette isolée. On copie le sous-arbre nécessaire dans un workspace sous `data/`,
                 on y exécute lot6d puis lot6e via le runner hors paquet, et on compare l'état
                 produit à l'état réel. **Aucun fichier réel n'est modifié.** Toujours autorisé.

  MODE_REEL    — régénérerait les MASTER ménages EN PLACE (via `run_menages_pipeline`, qui commence
                 par lot6b → Google Sheet + réécriture M04/MASTER_NORM). **Bloqué** tant que le flag
                 est False : `confirmer(MODE_REEL)` refuse et trace un run BLOQUE, sans rien exécuter.

Garde-fous repris de la chaîne charges (APP-3b) :
  * verrou interprocessus atomique (fichier de verrou dédié), libéré dans un `finally` quoi qu'il arrive ;
  * snapshot horodaté + manifeste sha256 des sources réelles AVANT tout ;
  * un run n'est SUCCES que si chaque étape rend rc==0 ET produit ses sorties (jamais sur le seul code 0) ;
  * vérification que les fichiers réels sont bits-pour-bits identiques avant/après.

lot6f (coût complet) est **exclu** du recalcul sur copies : il lit la Google Sheet (réseau) et n'est
donc pas déterministe hors ligne. Le recalcul sur copies rafraîchit le rapprochement (lot6d) et la
vue gain/perte (lot6e), tous deux hors ligne.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import app.config as cfg
from app.db.connection import get_db
from app.services import snapshot_service
from app.services import saisie_charges_lock_service as verrou_lib  # verrou interprocessus générique
from app.services.audit_service import log_event

MODE_COPIES = "COPIES"
MODE_REEL = "REEL"
LOCK_NAME = ".menages_recalcul.lock"
OPERATION = "menages_recalcul"

# Le moteur fige le mois du rapprochement (lot6d/6e : MONTH = "2026-05").
# Le recalcul ne peut donc porter que ce mois tant que le moteur n'est pas paramétré.
MOIS_MOTEUR = "2026-05"

# Étapes exécutées sur copies — hors ligne, déterministes. lot6f exclu (réseau, cf. docstring).
STEPS_COPIES: list[dict[str, Any]] = [
    {"name": "lot6d_rapprochement",
     "script": "02_TRAVAIL/lot6d_rapprochement_menages.py",
     "produces": ["02_TRAVAIL/Lot6d_Rapprochement_Menages/MASTER_CTRL_Rapprochement_Menages.xlsx"]},
    {"name": "lot6e_gainperte",
     "script": "02_TRAVAIL/lot6e_gainperte_menages.py",
     "produces": ["02_TRAVAIL/Lot6e_GainPerte_Menages/MASTER_CALC_GainPerte_Menages.xlsx"]},
]

# Fichiers RÉELS lus par lot6d/6e — copiés dans le workspace (chemins relatifs au projet).
SOURCES_A_COPIER = [
    "01_SOURCES_BRUTES/REF_Setup/REF_Setup.xlsm",
    "02_TRAVAIL/Lot1_Hostaway/MASTER_FACT_HA_CleaningTasks_Discovery.xlsx",
    "02_TRAVAIL/Lot6b_DeclarationsInternes/MASTER_NORM_Declarations_Internes.xlsx",
    "02_TRAVAIL/Lot6c_MenagesExternes/MASTER_FACT_MEN_MenagesExternes.xlsx",
]
# Scripts moteur copiés (exécutés dans le workspace ; ROOT y résout).
SCRIPTS_A_COPIER = [
    "02_TRAVAIL/lot6d_rapprochement_menages.py",
    "02_TRAVAIL/lot6e_gainperte_menages.py",
]
# Sorties réelles régénérées par le recalcul (pour comparaison avant/après).
SORTIES_REELLES = {
    "rapprochement": "02_TRAVAIL/Lot6d_Rapprochement_Menages/MASTER_CTRL_Rapprochement_Menages.xlsx",
    "gainperte": "02_TRAVAIL/Lot6e_GainPerte_Menages/MASTER_CALC_GainPerte_Menages.xlsx",
}

STATUT_SUCCES = "SUCCES"
STATUT_ECHEC = "ECHEC"
STATUT_PARTIEL = "PARTIEL"
STATUT_VERROUILLE = "VERROUILLE"
STATUT_BLOQUE = "BLOQUE"


# ── Chemins ──────────────────────────────────────────────────────────────────

def _project_root() -> Path:
    return Path(cfg.PROJECT_ROOT)


def _lock_path() -> Path:
    return Path(cfg.DATA_DIR) / LOCK_NAME


def _sha256(path: Path) -> str:
    return snapshot_service._sha256(path)  # même implémentation, une seule source de vérité


def _git_head() -> str | None:
    try:
        r = subprocess.run(
            ["git", "-C", str(_project_root()), "rev-parse", "--short", "HEAD"],
            capture_output=True, text=True, timeout=10,
        )
        return r.stdout.strip() or None if r.returncode == 0 else None
    except Exception:
        return None


# ── Détection Excel ouvert (fichiers de verrou Office ~$) ─────────────────────

def _fichier_office_verrou(path: Path) -> Path:
    return path.parent / ("~$" + path.name)


def detecter_excel_ouvert(paths: list[Path]) -> list[str]:
    """Retourne les noms des fichiers dont un verrou Office (~$) existe — donc probablement ouverts."""
    ouverts = []
    for p in paths:
        p = Path(p)
        if _fichier_office_verrou(p).exists():
            ouverts.append(p.name)
    return ouverts


def _sources_reelles() -> list[Path]:
    root = _project_root()
    return [root / rel for rel in SOURCES_A_COPIER]


def _sorties_reelles() -> list[Path]:
    root = _project_root()
    return [root / rel for rel in SORTIES_REELLES.values()]


# ── Préparation (préflight, aucune mutation) ─────────────────────────────────

def preparer(mode: str = MODE_COPIES, mois: str | None = None) -> dict[str, Any]:
    """Diagnostic préalable. Ne lance rien, n'écrit rien. Backe l'écran de préparation."""
    mois = mois or MOIS_MOTEUR
    root = _project_root()
    sources = [{"nom": Path(rel).name, "chemin_relatif": rel, "present": (root / rel).exists()}
               for rel in SOURCES_A_COPIER]
    sorties = [{"cle": cle, "nom": Path(rel).name, "chemin_relatif": rel, "present": (root / rel).exists()}
               for cle, rel in SORTIES_REELLES.items()]

    excel_ouverts = detecter_excel_ouvert(_sources_reelles() + _sorties_reelles())
    etat_verrou = verrou_lib.inspecter_verrou(_lock_path())
    sources_absentes = [s["nom"] for s in sources if not s["present"]]

    bloque = False
    raison = None
    if mode == MODE_REEL and not cfg.MENAGES_REAL_RECALC_ENABLED:
        bloque = True
        raison = ("Le mode RÉEL est désactivé (MENAGES_REAL_RECALC_ENABLED = False). "
                  "Seule la recette sur copies est exécutable : elle ne modifie aucun fichier réel.")

    warnings: list[str] = []
    if excel_ouverts:
        warnings.append("Fichier(s) probablement ouverts dans Excel : " + ", ".join(excel_ouverts)
                        + ". Fermez-les avant de recalculer.")
    if sources_absentes:
        warnings.append("Source(s) absente(s) : " + ", ".join(sources_absentes) + ".")
    if etat_verrou["etat"] != verrou_lib.ETAT_ABSENT:
        warnings.append("Un verrou de recalcul est présent (" + etat_verrou["etat"] + ") : "
                        + etat_verrou["raison"])

    prete = (not bloque) and (not excel_ouverts) and (not sources_absentes) \
        and etat_verrou["etat"] == verrou_lib.ETAT_ABSENT

    return {
        "mode": mode,
        "mois": mois,
        "mois_moteur_fige": MOIS_MOTEUR,
        "sources": sources,
        "sorties": sorties,
        "etapes": [s["name"] for s in STEPS_COPIES],
        "excel_ouverts": excel_ouverts,
        "verrou": etat_verrou,
        "sources_absentes": sources_absentes,
        "reel_active": bool(cfg.MENAGES_REAL_RECALC_ENABLED),
        "bloque": bloque,
        "raison": raison,
        "prete": prete,
        "warnings": warnings,
    }


# ── Fraîcheur : sources plus récentes que le dernier rapprochement ? ─────────

def etat_fraicheur() -> dict[str, Any]:
    """Indicateur (non comptable) : une source a-t-elle été modifiée après le dernier rapprochement ?

    Comparaison de mtime uniquement — un simple signal de fraîcheur, pas une preuve comptable.
    """
    root = _project_root()
    master = root / SORTIES_REELLES["rapprochement"]
    if not master.exists():
        return {"etat": "JAMAIS_CALCULE", "sources_recentes": [],
                "master_maj": None, "recommander_relance": False,
                "libelle": "Jamais calculé"}
    master_mtime = master.stat().st_mtime
    recentes = [Path(rel).name for rel in SOURCES_A_COPIER
                if (root / rel).exists() and (root / rel).stat().st_mtime > master_mtime]
    return {
        "etat": "SOURCES_PLUS_RECENTES" if recentes else "A_JOUR",
        "sources_recentes": recentes,
        "master_maj": datetime.fromtimestamp(master_mtime).strftime("%Y-%m-%d %H:%M"),
        "recommander_relance": bool(recentes),
        "libelle": "Sources modifiées depuis le dernier rapprochement" if recentes else "À jour",
    }


# ── Construction du workspace (copies) ───────────────────────────────────────

def _construire_workspace(run_ts: str) -> Path:
    """Miroir minimal du sous-arbre projet. ROOT des scripts copiés y résout → lit/écrit ici."""
    base = Path(cfg.MENAGES_RECALC_WORKSPACE)
    workspace = base / run_ts
    if workspace.exists():
        shutil.rmtree(workspace)
    workspace.mkdir(parents=True, exist_ok=True)

    root = _project_root()
    for rel in SOURCES_A_COPIER + SCRIPTS_A_COPIER:
        src = root / rel
        dst = workspace / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
    # Crée les dossiers de sortie attendus par les scripts (ils font os.makedirs de toute façon).
    for rel in SORTIES_REELLES.values():
        (workspace / Path(rel).parent).mkdir(parents=True, exist_ok=True)
    return workspace


# ── Lecture légère d'un TABLEAU_COMPARAISON (comparaison avant/après) ─────────

def _lire_comparaison(path: Path) -> dict[str, Any] | None:
    """Compte lignes / statuts / écarts d'un MASTER rapprochement, sans dépendre du reader métier."""
    if not Path(path).exists():
        return None
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "TABLEAU_COMPARAISON" not in wb.sheetnames:
            wb.close()
            return None
        ws = wb["TABLEAU_COMPARAISON"]
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        wb.close()
    except Exception:
        return None
    if not rows:
        return {"nb_lignes": 0, "statuts": {}, "nb_ecarts": 0}
    hdr = [str(c) for c in rows[0]]
    data = [dict(zip(hdr, r)) for r in rows[1:]]
    statuts: dict[str, int] = {}
    nb_ecarts = 0
    for d in data:
        st = str(d.get("statut_controle") or "")
        statuts[st] = statuts.get(st, 0) + 1
        try:
            if int(d.get("ecart") or 0) != 0:
                nb_ecarts += 1
        except (TypeError, ValueError):
            pass
    return {"nb_lignes": len(data), "statuts": statuts, "nb_ecarts": nb_ecarts}


# ── Enregistrement d'un run ──────────────────────────────────────────────────

def _enregistrer_run(db_path: Path, **champs: Any) -> int:
    conn = get_db(db_path)
    try:
        cols = ", ".join(champs.keys())
        placeholders = ", ".join("?" for _ in champs)
        cur = conn.execute(
            f"INSERT INTO menages_recalcul_runs ({cols}) VALUES ({placeholders})",
            tuple(champs.values()),
        )
        conn.commit()
        return cur.lastrowid
    finally:
        conn.close()


# ── Confirmation / exécution ─────────────────────────────────────────────────

def confirmer(mode: str = MODE_COPIES, mois: str | None = None,
              db_path: Path | None = None) -> dict[str, Any]:
    """Exécute le recalcul. MODE_COPIES seulement ; MODE_REEL refusé tant que le flag est False."""
    db_path = db_path or cfg.DB_PATH          # lu À CHAUD — jamais le défaut figé à l'import (défaut get_db)
    mois = mois or MOIS_MOTEUR

    # ── Garde mode réel ──────────────────────────────────────────────────────
    if mode == MODE_REEL and not cfg.MENAGES_REAL_RECALC_ENABLED:
        run_id = _enregistrer_run(
            db_path, mode=MODE_REEL, periode=mois, statut=STATUT_BLOQUE,
            git_head=_git_head(), erreur_code="E_MODE_REEL_DESACTIVE",
            erreur_resume="MENAGES_REAL_RECALC_ENABLED=False — aucune exécution.",
        )
        log_event("MENAGE_RECALC_BLOQUE", {"mode": MODE_REEL, "mois": mois}, db_path=db_path)
        return {"ok": False, "run_id": run_id, "statut": STATUT_BLOQUE,
                "erreur_code": "E_MODE_REEL_DESACTIVE",
                "message": "Recalcul réel désactivé. Aucun fichier réel n'a été touché."}

    if mode not in (MODE_COPIES, MODE_REEL):
        raise ValueError(f"Mode inconnu : {mode}")

    # ── Préflight bloquant (Excel ouvert / source absente) ───────────────────
    excel_ouverts = detecter_excel_ouvert(_sources_reelles() + _sorties_reelles())
    if excel_ouverts:
        run_id = _enregistrer_run(
            db_path, mode=mode, periode=mois, statut=STATUT_ECHEC, git_head=_git_head(),
            erreur_code="E_EXCEL_OUVERT", erreur_resume="Ouverts : " + ", ".join(excel_ouverts),
        )
        return {"ok": False, "run_id": run_id, "statut": STATUT_ECHEC,
                "erreur_code": "E_EXCEL_OUVERT",
                "message": "Fichier(s) ouverts dans Excel : " + ", ".join(excel_ouverts)}

    sources_reelles = _sources_reelles()
    absentes = [p.name for p in sources_reelles if not p.exists()]
    if absentes:
        run_id = _enregistrer_run(
            db_path, mode=mode, periode=mois, statut=STATUT_ECHEC, git_head=_git_head(),
            erreur_code="E_SOURCE_ABSENTE", erreur_resume="Absentes : " + ", ".join(absentes),
        )
        return {"ok": False, "run_id": run_id, "statut": STATUT_ECHEC,
                "erreur_code": "E_SOURCE_ABSENTE",
                "message": "Source(s) absente(s) : " + ", ".join(absentes)}

    # Scripts moteur (lot6d/lot6e) : vérifiés AVANT le workspace (sinon shutil.copy2 lève
    # FileNotFoundError non capturée -> 500). Cas typique d'un PROJECT_ROOT de démonstration.
    scripts_absents = [Path(rel).name for rel in SCRIPTS_A_COPIER
                       if not (_project_root() / rel).exists()]
    if scripts_absents:
        run_id = _enregistrer_run(
            db_path, mode=mode, periode=mois, statut=STATUT_ECHEC, git_head=_git_head(),
            erreur_code="E_SCRIPT_MOTEUR_ABSENT", erreur_resume="Absents : " + ", ".join(scripts_absents),
        )
        return {"ok": False, "run_id": run_id, "statut": STATUT_ECHEC,
                "erreur_code": "E_SCRIPT_MOTEUR_ABSENT",
                "message": "Script(s) moteur absent(s) du projet : " + ", ".join(scripts_absents)
                           + ". La simulation n'a pas été lancée ; aucun fichier réel touché."}

    # ── Verrou interprocessus (dédié) — libéré dans finally quoi qu'il arrive ─
    try:
        verrou = verrou_lib.acquerir_verrou(operation=OPERATION, lock_path=_lock_path())
    except verrou_lib.VerrouSaisieChargesDejaPrisError as exc:
        run_id = _enregistrer_run(
            db_path, mode=mode, periode=mois, statut=STATUT_VERROUILLE, git_head=_git_head(),
            erreur_code="E_VERROU", erreur_resume=str(exc),
        )
        return {"ok": False, "run_id": run_id, "statut": STATUT_VERROUILLE,
                "erreur_code": "E_VERROU", "message": "Un recalcul est déjà en cours."}

    date_debut = datetime.now(timezone.utc).isoformat()
    t0 = time.monotonic()
    try:
        # ── Snapshot des sources réelles AVANT (traçabilité, lecture seule) ──
        sha_sources_avant = {p.name: _sha256(p) for p in sources_reelles}
        sha_sorties_avant = {cle: (_sha256(_project_root() / rel)
                                   if (_project_root() / rel).exists() else None)
                             for cle, rel in SORTIES_REELLES.items()}
        snap = snapshot_service.create_snapshot(
            "MENAGES_RECALC_AVANT", sources_reelles + _sorties_reelles(), db_path=db_path)

        # ── Workspace copies + requête runner ───────────────────────────────
        run_ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        workspace = _construire_workspace(run_ts)
        requete = {"allowed_root": str(workspace.resolve()), "workspace": str(workspace.resolve()),
                   "steps": STEPS_COPIES, "timeout": cfg.MENAGES_RECALC_TIMEOUT_SECONDS}
        requete_path = workspace / "_requete.json"
        reponse_path = workspace / "_reponse.json"
        requete_path.write_text(json.dumps(requete, ensure_ascii=False, indent=2), encoding="utf-8")

        # ── Invocation du runner HORS paquet app/, interpréteur moteur ──────
        runner = Path(cfg.MENAGES_RECALC_RUNNER)
        engine_py = Path(cfg.MENAGES_ENGINE_PYTHON)
        erreur_code = None
        erreur_resume = None
        etapes: list[dict[str, Any]] = []
        if not runner.exists():
            statut = STATUT_ECHEC
            erreur_code, erreur_resume = "E_RUNNER", f"Runner absent : {runner}"
        elif not engine_py.exists():
            statut = STATUT_ECHEC
            erreur_code, erreur_resume = "E_RUNNER", f"Interpréteur moteur absent : {engine_py}"
        else:
            try:
                proc = subprocess.run(
                    [str(engine_py), str(runner), str(requete_path), str(reponse_path)],
                    capture_output=True, text=True,
                    timeout=cfg.MENAGES_RECALC_TIMEOUT_SECONDS + 30,
                )
            except subprocess.TimeoutExpired:
                statut = STATUT_ECHEC
                erreur_code, erreur_resume = "E_TIMEOUT", "Le runner a dépassé le délai."
            else:
                if not reponse_path.exists():
                    statut = STATUT_ECHEC
                    erreur_code = "E_REPONSE_INVALIDE"
                    erreur_resume = f"Pas de réponse runner (rc={proc.returncode}). {proc.stderr[-500:]}"
                else:
                    reponse = json.loads(reponse_path.read_text(encoding="utf-8"))
                    etapes = reponse.get("steps", [])
                    nb_ok = sum(1 for e in etapes if e.get("statut") == "OK")
                    if reponse.get("ok") and nb_ok == len(STEPS_COPIES):
                        statut = STATUT_SUCCES
                    elif nb_ok == 0:
                        statut = STATUT_ECHEC
                    else:
                        statut = STATUT_PARTIEL
                    if statut != STATUT_SUCCES:
                        premiere = next((e for e in etapes if e.get("statut") != "OK"), {})
                        erreur_code = premiere.get("erreur_code") or "E_CONTROLE_ECHEC"
                        erreur_resume = premiere.get("detail")

        # ── sha256 des sorties APRÈS — depuis le WORKSPACE (jamais le réel) ──
        sha_sorties_apres = {}
        comparaison = {}
        for cle, rel in SORTIES_REELLES.items():
            produit = workspace / rel
            sha_sorties_apres[cle] = _sha256(produit) if produit.exists() else None
            if cle == "rapprochement":
                comparaison = {
                    "avant": _lire_comparaison(_project_root() / rel),
                    "apres": _lire_comparaison(produit),
                }

        # ── Garde : les fichiers RÉELS n'ont pas bougé ──────────────────────
        sha_sources_apres_reel = {p.name: (_sha256(p) if p.exists() else None) for p in sources_reelles}
        reel_intact = sha_sources_apres_reel == sha_sources_avant
        if not reel_intact:
            statut = STATUT_ECHEC
            erreur_code = "E_REEL_MODIFIE"
            erreur_resume = "ANOMALIE GRAVE : une source réelle a changé pendant un recalcul sur copies."

        duree = round(time.monotonic() - t0, 2)
        date_fin = datetime.now(timezone.utc).isoformat()
        run_id = _enregistrer_run(
            db_path, mode=mode, periode=mois, statut=statut,
            date_debut=date_debut, date_fin=date_fin, duree_secondes=duree,
            snapshot_id=snap["id"], git_head=_git_head(), workspace_path=str(workspace),
            sha256_sources_avant=json.dumps(sha_sources_avant, ensure_ascii=False),
            sha256_sorties_avant=json.dumps(sha_sorties_avant, ensure_ascii=False),
            sha256_sorties_apres=json.dumps(sha_sorties_apres, ensure_ascii=False),
            comparaison_json=json.dumps({"rapprochement": comparaison, "reel_intact": reel_intact},
                                        ensure_ascii=False),
            etapes_json=json.dumps(etapes, ensure_ascii=False, default=str),
            erreur_code=erreur_code, erreur_resume=erreur_resume,
        )
        log_event("MENAGE_RECALC", {"mode": mode, "mois": mois, "statut": statut,
                                    "run_id": run_id, "reel_intact": reel_intact}, db_path=db_path)
        return {
            "ok": statut == STATUT_SUCCES, "run_id": run_id, "statut": statut,
            "mode": mode, "mois": mois, "duree_secondes": duree,
            "snapshot_id": snap["id"], "workspace": str(workspace),
            "etapes": etapes, "comparaison": comparaison,
            "sha256_sorties_avant": sha_sorties_avant, "sha256_sorties_apres": sha_sorties_apres,
            "reel_intact": reel_intact,
            "erreur_code": erreur_code, "erreur_resume": erreur_resume,
        }
    except Exception as exc:
        # Défense en profondeur : toute erreur inattendue devient un run ECHEC lisible, jamais un 500
        # ni un traceback dans le navigateur. Le type d'exception seul est exposé (aucun chemin absolu).
        run_id = _enregistrer_run(
            db_path, mode=mode, periode=mois, statut=STATUT_ECHEC, date_debut=date_debut,
            git_head=_git_head(), erreur_code="E_INATTENDU",
            erreur_resume=f"{type(exc).__name__}: {exc}",
        )
        log_event("MENAGE_RECALC_ERREUR", {"mode": mode, "mois": mois,
                                           "erreur": type(exc).__name__}, db_path=db_path)
        return {"ok": False, "run_id": run_id, "statut": STATUT_ECHEC,
                "erreur_code": "E_INATTENDU",
                "message": "Une erreur inattendue a interrompu la simulation "
                           f"({type(exc).__name__}). Aucun fichier réel n'a été touché."}
    finally:
        try:
            verrou_lib.liberer_verrou(verrou)
        except verrou_lib.LiberationVerrouSaisieChargesError:
            # La libération n'a pas pu conclure : le verrou reste en place (comportement sûr).
            # On ne masque pas : l'anomalie est tracée en audit.
            log_event("MENAGE_RECALC_VERROU_NON_LIBERE", {"lock": str(_lock_path())}, db_path=db_path)


# ── Lecture historique ───────────────────────────────────────────────────────

def load_runs(limit: int = 5, db_path: Path | None = None) -> list[dict[str, Any]]:
    db_path = db_path or cfg.DB_PATH
    conn = get_db(db_path)
    try:
        rows = conn.execute(
            "SELECT id, ts, mode, periode, statut, duree_secondes, snapshot_id, git_head, "
            "erreur_code FROM menages_recalcul_runs ORDER BY id DESC LIMIT ?",
            (limit,),
        ).fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def load_run(run_id: int, db_path: Path | None = None) -> dict[str, Any] | None:
    db_path = db_path or cfg.DB_PATH
    conn = get_db(db_path)
    try:
        row = conn.execute("SELECT * FROM menages_recalcul_runs WHERE id=?", (run_id,)).fetchone()
    finally:
        conn.close()
    if not row:
        return None
    d = dict(row)
    for champ in ("etapes_json", "comparaison_json", "sha256_sources_avant",
                  "sha256_sorties_avant", "sha256_sorties_apres"):
        if d.get(champ):
            try:
                d[champ.replace("_json", "")] = json.loads(d[champ])
            except (json.JSONDecodeError, TypeError):
                pass
    d["etapes"] = d.get("etapes") or (json.loads(d["etapes_json"]) if d.get("etapes_json") else [])
    if isinstance(d.get("comparaison"), dict):
        d["reel_intact"] = d["comparaison"].get("reel_intact")
    return d
