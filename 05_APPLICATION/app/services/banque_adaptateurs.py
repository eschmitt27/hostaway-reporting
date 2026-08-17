"""Adaptateurs de source bancaire — un seul DTO en sortie, quelle que soit l'origine.

    export XLSX ─┐
    API (demain) ─┼─► DTO canonique ─► banque_mouvements_service ─► SQLite ─► classification…
    CSV, …       ─┘

Tout l'aval dépend du DTO, jamais du transport. C'est la condition pour qu'une API bancaire
remplace un classeur sans toucher à un seul traitement — et pour que la future banque n'impose pas
de réécrire la chaîne.

Aucune intégration API réelle n'est développée ici. L'adaptateur `mock` existe pour prouver que le
contrat tient, pas pour appeler quoi que ce soit.

LE DTO
    external_transaction_id  identifiant banque si fourni — sinon chaîne vide, jamais inventé
    date_operation           AAAA-MM-JJ
    date_valeur              AAAA-MM-JJ ou vide
    sens                     DEBIT|CREDIT
    montant                  positif ; le signe économique est porté par `sens`
    devise                   EUR par défaut
    libelle_brut             tel que reçu, jamais réinterprété
    contrepartie_brute       si la source la distingue
"""
from __future__ import annotations

import hashlib
from datetime import date, datetime
from pathlib import Path
from typing import Any

from app.services.banque_mouvements_service import SENS_CREDIT, SENS_DEBIT


class SourceBancaireInvalide(RuntimeError):
    """La source ne respecte pas le contrat attendu. Refuser vaut mieux qu'interpréter."""


def _date(v: Any) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = str(v).strip()
    return t[:10] if len(t) >= 10 else t


def _nombre(v: Any) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    t = str(v).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return round(float(t), 2)
    except ValueError:
        return 0.0


def sha256_fichier(chemin: Path) -> str:
    h = hashlib.sha256()
    with open(chemin, "rb") as f:
        for bloc in iter(lambda: f.read(1 << 20), b""):
            h.update(bloc)
    return h.hexdigest()


# ── Adaptateur : relevé consolidé Crédit Mutuel (format actuellement reçu) ───────────────────────

ENTETE_ATTENDUE = ["N°", "Date opération", "Date de valeur", "Libellé",
                   "Débit", "Crédit", "Montant net", "Solde consolidé"]


def depuis_xlsx_releve_consolide(chemin: Path) -> list[dict[str, Any]]:
    """Lit un relevé consolidé et rend le DTO canonique. Lecture seule.

    L'en-tête n'est pas en première ligne : le fichier commence par un titre et un rappel de
    compte. On la cherche plutôt que de la supposer — un décalage d'une ligne produirait des
    mouvements muets.
    """
    import openpyxl

    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    try:
        if "Mouvements" not in wb.sheetnames:
            raise SourceBancaireInvalide(
                f"Onglet « Mouvements » absent de {chemin.name} "
                f"(onglets : {', '.join(wb.sheetnames)}).")
        ws = wb["Mouvements"]

        entete_ligne = None
        for i, row in enumerate(ws.iter_rows(max_row=20, values_only=True), start=1):
            valeurs = [str(c).strip() if c is not None else "" for c in row]
            if valeurs[:4] == ENTETE_ATTENDUE[:4]:
                entete_ligne = i
                break
        if entete_ligne is None:
            raise SourceBancaireInvalide(
                f"En-tête introuvable dans {chemin.name} — attendu {ENTETE_ATTENDUE[:4]}.")

        lignes: list[dict[str, Any]] = []
        for row in ws.iter_rows(min_row=entete_ligne + 1, values_only=True):
            if row is None or all(c in (None, "") for c in row):
                continue
            date_op = _date(row[1])
            if not date_op:
                continue  # ligne de pied ou séparateur : pas un mouvement
            debit, credit = _nombre(row[4]), _nombre(row[5])
            # Le sens vient des colonnes de la banque, jamais du signe du montant net : c'est elle
            # qui sait si la ligne débite ou crédite.
            sens = SENS_DEBIT if debit > 0 else SENS_CREDIT
            montant = debit if debit > 0 else credit
            lignes.append({
                "external_transaction_id": "",   # ce format n'en fournit aucun
                "date_operation": date_op,
                "date_valeur": _date(row[2]),
                "sens": sens,
                "montant": montant,
                "devise": "EUR",
                "libelle_brut": str(row[3] or "").strip(),
                "contrepartie_brute": "",
            })
        if not lignes:
            raise SourceBancaireInvalide(f"Aucun mouvement exploitable dans {chemin.name}.")
        return lignes
    finally:
        wb.close()


# ── Adaptateur : API (contrat seulement) ────────────────────────────────────────────────────────

def depuis_api(charge_utile: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Transforme une réponse d'API bancaire en DTO canonique.

    Aucun appel réseau ici : la fonction prend la charge utile déjà obtenue. C'est volontaire —
    elle décrit la FORME du contrat, et permet de le tester sans banque ni secret.

    Différence essentielle avec le classeur : une API fournit en général un identifiant de
    transaction stable. Quand il est là, la déduplication devient exacte au lieu d'être heuristique.
    """
    lignes: list[dict[str, Any]] = []
    for i, t in enumerate(charge_utile, start=1):
        date_op = _date(t.get("bookingDate") or t.get("date_operation"))
        if not date_op:
            raise SourceBancaireInvalide(f"Transaction {i} sans date d'opération.")
        montant = _nombre(t.get("amount", t.get("montant")))
        sens = t.get("sens")
        if sens not in (SENS_DEBIT, SENS_CREDIT):
            # Convention API courante : montant signé. Le sens s'en déduit, le montant devient absolu.
            sens = SENS_DEBIT if montant < 0 else SENS_CREDIT
        lignes.append({
            "external_transaction_id": str(t.get("transactionId")
                                           or t.get("external_transaction_id") or "").strip(),
            "date_operation": date_op,
            "date_valeur": _date(t.get("valueDate") or t.get("date_valeur")),
            "sens": sens,
            "montant": abs(montant),
            "devise": (t.get("currency") or t.get("devise") or "EUR").strip().upper(),
            "libelle_brut": str(t.get("remittanceInformation")
                                or t.get("libelle_brut") or "").strip(),
            "contrepartie_brute": str(t.get("counterpartyName")
                                      or t.get("contrepartie_brute") or "").strip(),
        })
    return lignes
