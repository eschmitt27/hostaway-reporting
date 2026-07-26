"""Création d'un logement au référentiel (REF_Setup) — écriture gardée et validée.

Manque historique : le parc était en lecture seule (aucune route, aucune fonction de création dans
toutes les branches). Ce service ajoute la création, en réutilisant les protections déjà éprouvées :
flags d'écriture, write-guard (refus hors `data_recette` en mode recette), écriture atomique.

Deux feuilles sont alimentées de façon cohérente :
- `REF_Logements` : le bien lui-même ;
- `REF_Gestion_Logements_Hist` : le rattachement au propriétaire, **daté** (c'est cette feuille que
  lisent le moteur de charges et Lot10 pour résoudre le propriétaire d'un logement à une période).

Aucune règle métier n'est recalculée : le service valide, puis écrit.
"""
from __future__ import annotations

import os
import tempfile
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

import app.config as cfg

SH_LOG = "REF_Logements"
SH_GEST = "REF_Gestion_Logements_Hist"
SH_PROP = "REF_Proprietaires"

E_FLAGS = "E_FLAGS_DESACTIVES"
E_ID_MANQUANT = "V01_ID_MANQUANT"
E_ID_EXISTANT = "V02_ID_DEJA_UTILISE"
E_NOM_MANQUANT = "V03_NOM_MANQUANT"
E_PROP_MANQUANT = "V04_PROPRIETAIRE_MANQUANT"
E_PROP_INCONNU = "V05_PROPRIETAIRE_INCONNU"
E_DATE_INVALIDE = "V06_DATE_DEBUT_INVALIDE"
E_TYPE_INCONNU = "V07_TYPE_LOGEMENT_INCONNU"
E_ECRITURE = "E_ECRITURE_REFUSEE"

MESSAGES = {
    E_FLAGS: "Écriture désactivée sur cette installation : la création est impossible.",
    E_ID_MANQUANT: "L'identifiant du logement est obligatoire.",
    E_ID_EXISTANT: "Cet identifiant de logement existe déjà.",
    E_NOM_MANQUANT: "Le nom du logement est obligatoire.",
    E_PROP_MANQUANT: "Le propriétaire est obligatoire.",
    E_PROP_INCONNU: "Ce propriétaire n'existe pas dans le référentiel.",
    E_DATE_INVALIDE: "La date de début de gestion est invalide (format AAAA-MM-JJ attendu).",
    E_TYPE_INCONNU: "Ce type de logement n'existe pas dans le référentiel.",
    E_ECRITURE: "Écriture refusée : la cible n'est pas autorisée.",
}


def _flags_actifs() -> bool:
    return bool(cfg.CHARGES_REAL_WRITE_ENABLED and cfg.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED)


def _txt(v) -> str:
    return str(v or "").strip()


def _lire(wb, sheet: str) -> tuple[list[str], list[dict[str, Any]]]:
    if sheet not in wb.sheetnames:
        return [], []
    ws = wb[sheet]
    data = list(ws.iter_rows(values_only=True))
    if not data:
        return [], []
    hdr = [c for c in data[0]]
    return hdr, [dict(zip(hdr, r)) for r in data[1:]]


def referentiels(ref_path: Path | None = None) -> dict[str, Any]:
    """Options nécessaires au formulaire (propriétaires actifs, types, logements existants)."""
    p = Path(ref_path or cfg.REF_SETUP)
    if not p.exists():
        return {"status": "SOURCE_ABSENTE", "proprietaires": [], "types": [], "logements": []}
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        _, props = _lire(wb, SH_PROP)
        _, types = _lire(wb, "REF_Types_Logements")
        _, logs = _lire(wb, SH_LOG)
    finally:
        wb.close()
    return {
        "status": "OK",
        "proprietaires": [{"proprietaire_id": _txt(r.get("proprietaire_id")),
                           "nom": _txt(r.get("nom_proprietaire"))}
                          for r in props if _txt(r.get("proprietaire_id"))
                          and _txt(r.get("actif")).upper() == "OUI"],
        "types": [{"type_logement_id": _txt(r.get("type_logement_id")),
                   "libelle": _txt(r.get("libelle") or r.get("type_logement_libelle"))}
                  for r in types if _txt(r.get("type_logement_id"))],
        "logements": [_txt(r.get("logement_id")) for r in logs if _txt(r.get("logement_id"))],
    }


def valider(form: dict[str, Any], ref_path: Path | None = None) -> list[dict[str, str]]:
    """Contrôles métier AVANT écriture. Retourne la liste des erreurs (vide = valide)."""
    erreurs: list[dict[str, str]] = []

    def err(code: str, detail: str = ""):
        erreurs.append({"code": code, "message": MESSAGES.get(code, code), "detail": detail})

    refs = referentiels(ref_path)
    logement_id = _txt(form.get("logement_id"))
    if not logement_id:
        err(E_ID_MANQUANT)
    elif logement_id in refs["logements"]:
        err(E_ID_EXISTANT, logement_id)

    if not _txt(form.get("nom_logement_officiel")) and not _txt(form.get("nom_court")):
        err(E_NOM_MANQUANT)

    prop = _txt(form.get("proprietaire_id"))
    connus = {p["proprietaire_id"] for p in refs["proprietaires"]}
    if not prop:
        err(E_PROP_MANQUANT)
    elif connus and prop not in connus:
        err(E_PROP_INCONNU, prop)

    d = _txt(form.get("date_debut"))
    if not d:
        err(E_DATE_INVALIDE, "date manquante")
    else:
        try:
            date.fromisoformat(d)
        except ValueError:
            err(E_DATE_INVALIDE, d)

    t = _txt(form.get("type_logement_id"))
    types_connus = {x["type_logement_id"] for x in refs["types"]}
    if t and types_connus and t not in types_connus:
        err(E_TYPE_INCONNU, t)

    return erreurs


def creer(form: dict[str, Any], ref_path: Path | None = None) -> dict[str, Any]:
    """Crée le logement + son rattachement propriétaire daté. Jamais de modification d'une ligne
    existante : uniquement des ajouts."""
    if not _flags_actifs():
        return {"ok": False, "code": E_FLAGS, "message": MESSAGES[E_FLAGS], "erreurs": []}

    erreurs = valider(form, ref_path)
    if erreurs:
        return {"ok": False, "code": "V_INVALIDE", "message": "Saisie invalide.", "erreurs": erreurs}

    p = Path(ref_path or cfg.REF_SETUP)
    wb = openpyxl.load_workbook(p, keep_vba=p.suffix.lower() == ".xlsm")
    try:
        logement_id = _txt(form.get("logement_id"))
        prop = _txt(form.get("proprietaire_id"))
        actif = "OUI" if _txt(form.get("actif")).upper() != "NON" else "NON"

        ws = wb[SH_LOG]
        hdr = [c.value for c in ws[1]]
        valeurs = {
            "logement_id": logement_id,
            "hostaway_listing_id": _txt(form.get("hostaway_listing_id")) or None,
            "nom_logement_officiel": _txt(form.get("nom_logement_officiel")) or _txt(form.get("nom_court")),
            "nom_court": _txt(form.get("nom_court")) or _txt(form.get("nom_logement_officiel")),
            "adresse": _txt(form.get("adresse")) or None,
            "ville": _txt(form.get("ville")) or None,
            "type_logement_id": _txt(form.get("type_logement_id")) or None,
            "sur_hostaway": _txt(form.get("sur_hostaway")) or "NON",
            "actif": actif,
            "statut_parc": _txt(form.get("statut_parc")) or ("GERE" if actif == "OUI" else "RETIRE"),
            "commentaire": _txt(form.get("commentaire")) or None,
            "forfait_logiciel_consommables_mensuel": form.get("forfait_logiciel_consommables_mensuel") or 0,
        }
        ws.append([valeurs.get(h) for h in hdr])

        # Rattachement propriétaire DATÉ (source de résolution du propriétaire par période)
        wsg = wb[SH_GEST]
        hdrg = [c.value for c in wsg[1]]
        gest = {
            "gestion_id": f"GST_{logement_id}_{prop}",
            "logement_id": logement_id,
            "proprietaire_id": prop,
            "date_debut": _txt(form.get("date_debut")),
            "date_fin": _txt(form.get("date_fin")) or None,
            "statut_gestion": "ACTIF" if actif == "OUI" else "RETIRE",
            "source": "SAISIE_APPLICATION",
            "commentaire": _txt(form.get("commentaire")) or None,
        }
        wsg.append([gest.get(h) for h in hdrg])

        from app.services.saisie_charges_transaction_service import _remplacer_fichier
        fd, tmp = tempfile.mkstemp(suffix=p.suffix, dir=str(p.parent))
        os.close(fd)
        tmp_path = Path(tmp)
        try:
            wb.save(tmp_path)
            _remplacer_fichier(tmp_path, p)
        except Exception as exc:
            tmp_path.unlink(missing_ok=True)
            return {"ok": False, "code": E_ECRITURE,
                    "message": MESSAGES[E_ECRITURE], "erreurs": [],
                    "detail": f"{type(exc).__name__}: {exc}"}
    finally:
        wb.close()

    return {"ok": True, "logement_id": logement_id, "proprietaire_id": prop,
            "date_debut": _txt(form.get("date_debut")), "erreurs": []}
