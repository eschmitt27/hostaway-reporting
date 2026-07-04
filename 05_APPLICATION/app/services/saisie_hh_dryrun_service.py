"""APP-2c - previsualisation d'ecriture HH sur copies uniquement."""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import uuid
from typing import Any

import openpyxl

import app.config as cfg
from app.readers.saisie_hh_reader import (
    MANUAL_COL_MAP,
    find_first_empty_data_row,
    _col_index,
)
from app.services import saisie_hh_service as saisie_svc
from app.services.saisie_hh_schema_migration import (
    NEW_SAISIE_FIELDS,
    migrate_ref_setup_copy,
    migrate_saisie_copy,
)
from app.writers.saisie_hh_writer import (
    _check_formula_cells,
    _check_structural_preservation,
    _measure_structure,
)


DRYRUNS_DIR = cfg.DRYRUNS_DIR
SAISIE_COPY_NAME = "SAISIE_ReservationsHorsHostaway_copie.xlsx"
REF_COPY_NAME = "REF_Setup_copie.xlsm"
MASTER_SIM_NAME = "MASTER_FACT_MAN_ReservationsHorsHostaway_simule.xlsx"
RESULT_LOT4A_NAME = "resultat_lot4a.json"
MANIFEST_NAME = "manifest.json"
LOT4A_REQUEST_NAME = "lot4a_request.json"
LOT4A_RESPONSE_NAME = "lot4a_response.json"
STRUCTURE_BEFORE_NAME = "SAISIE_structure_avant_injection.xlsx"
LOT4A_RUNNER = Path(__file__).resolve().with_name("lot4a_dryrun_runner.py")


class DryRunError(RuntimeError):
    """Erreur bloquante de simulation APP-2c."""


def _json_default(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, Path):
        return str(value)
    return str(value)


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(str(path), "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _fingerprint(path: Path) -> dict[str, Any]:
    p = Path(path)
    if not p.exists():
        return {"exists": False}
    stat = p.stat()
    return {
        "exists": True,
        "sha256": _sha256(p),
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
    }


def _token() -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    return f"{stamp}_{uuid.uuid4().hex[:12]}"


def _safe_token(token: str) -> str:
    clean = str(token).strip()
    allowed = set("ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789_-TZ")
    if not clean or any(ch not in allowed for ch in clean):
        raise DryRunError("Identifiant de previsualisation invalide")
    return clean


def _copy_source(source: Path, destination: Path) -> None:
    if not source.exists():
        raise DryRunError(f"Source introuvable pour simulation: {source}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)


def _normalise_cell_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "OUI" if value else ""
    return value


def _build_simulated_row(preview: dict[str, Any]) -> dict[str, Any]:
    row = saisie_svc.build_row_data(preview)
    for field in NEW_SAISIE_FIELDS:
        value = preview.get(field)
        if field.startswith("confirmation_override_"):
            value = "OUI" if value else ""
        row[field] = _normalise_cell_value(value)
    return row


def _formula_cache_state(saisie_copy: Path, target_row: int | None) -> dict[str, Any]:
    if target_row is None:
        return {"target_row": None, "cached_values": {}, "requires_excel_recalc": False}
    wb = openpyxl.load_workbook(str(saisie_copy), read_only=True, data_only=True)
    try:
        ws = wb["SAISIE"]
        cells = ("B", "C", "K", "N", "O", "Q", "V", "Y", "Z")
        cached = {col: ws[f"{col}{target_row}"].value for col in cells}
    finally:
        wb.close()
    return {
        "target_row": target_row,
        "cached_values": cached,
        "excel_cache_missing": any(value in (None, "") for value in cached.values()),
        "requires_excel_recalc_for_lot4a": False,
        "lot4a_recomputes_derived_fields": True,
        "status": (
            "LOT4A_RECALCUL_PYTHON_SANS_CACHE_EXCEL"
            if any(value in (None, "") for value in cached.values())
            else "CACHE_FORMULES_DISPONIBLE"
        ),
    }


def _write_simulated_row(saisie_copy: Path, row_data: dict[str, Any], run_dir: Path) -> tuple[int, list[str]]:
    target_row = find_first_empty_data_row(saisie_copy)
    if target_row is None:
        raise DryRunError("Aucune ligne libre dans la copie SAISIE")

    formula_violations = _check_formula_cells(saisie_copy, target_row)
    if formula_violations:
        raise DryRunError("; ".join(formula_violations))

    structure_before = run_dir / STRUCTURE_BEFORE_NAME
    shutil.copy2(saisie_copy, structure_before)
    struct_avant = _measure_structure(structure_before)
    wb = openpyxl.load_workbook(str(saisie_copy), data_only=False)
    try:
        ws = wb["SAISIE"]
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        header_idx = {name: idx + 1 for idx, name in enumerate(headers) if name}

        for col_letter, field_name in MANUAL_COL_MAP.items():
            if field_name in row_data:
                ws.cell(
                    row=target_row,
                    column=_col_index(col_letter),
                    value=_normalise_cell_value(row_data[field_name]),
                )
        for field in NEW_SAISIE_FIELDS:
            col = header_idx.get(field)
            if col is not None:
                ws.cell(row=target_row, column=col, value=_normalise_cell_value(row_data.get(field)))

        wb.calculation.fullCalcOnLoad = True
        wb.save(str(saisie_copy))
    finally:
        wb.close()

    structure_violations = _check_structural_preservation(structure_before, saisie_copy, struct_avant)
    structure_violations.extend(_check_formula_cells(saisie_copy, target_row))
    structure_before.unlink(missing_ok=True)
    if structure_violations:
        raise DryRunError("; ".join(structure_violations))
    return target_row, structure_violations


def _count_acomptes(rows: list[dict[str, Any]]) -> int:
    count = 0
    for row in rows:
        value = row.get("acompte_facture")
        if value not in (None, ""):
            count += 1
    return count


def _assert_under(path: Path, allowed_root: Path) -> Path:
    resolved = Path(path).resolve()
    root = Path(allowed_root).resolve()
    if resolved != root and root not in resolved.parents:
        raise DryRunError(f"Chemin refuse hors dry-run: {resolved}")
    return resolved


def _run_lot4a_engine(
    saisie_path: Path,
    ref_path: Path,
    as_of_iso: str,
    run_dir: Path,
    *,
    master_path: Path | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, Any]]:
    run_root = Path(run_dir).resolve()
    request_path = run_root / LOT4A_REQUEST_NAME
    response_path = run_root / LOT4A_RESPONSE_NAME
    _assert_under(Path(saisie_path), run_root)
    _assert_under(Path(ref_path), run_root)
    if master_path is not None:
        _assert_under(Path(master_path), run_root)

    request = {
        "allowed_root": str(run_root),
        "project_root": str(cfg.PROJECT_ROOT.resolve()),
        "saisie_path": str(Path(saisie_path).resolve()),
        "ref_path": str(Path(ref_path).resolve()),
        "master_path": str(Path(master_path).resolve()) if master_path is not None else None,
        "as_of_iso": as_of_iso,
    }
    request_path.write_text(
        json.dumps(request, ensure_ascii=False, indent=2, default=_json_default),
        encoding="utf-8",
    )

    cmd = [str(cfg.LOT4A_ENGINE_PYTHON), str(LOT4A_RUNNER), str(request_path), str(response_path)]
    try:
        completed = subprocess.run(
            cmd,
            cwd=str(cfg.PROJECT_ROOT),
            capture_output=True,
            text=True,
            timeout=cfg.LOT4A_ENGINE_TIMEOUT_SECONDS,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise DryRunError(f"LOT4A_TIMEOUT: {exc}") from exc

    if completed.returncode != 0:
        stderr = (completed.stderr or "").strip()
        stdout = (completed.stdout or "").strip()
        raise DryRunError(f"LOT4A_SUBPROCESS_ERREUR: rc={completed.returncode}; stdout={stdout}; stderr={stderr}")
    if not response_path.exists():
        raise DryRunError("LOT4A_SUBPROCESS_SANS_REPONSE")
    response = json.loads(response_path.read_text(encoding="utf-8"))
    if not response.get("ok"):
        raise DryRunError(f"LOT4A_SUBPROCESS_REPONSE_INVALIDE: {response}")
    return response["saisie_rows"], response["result"], response.get("engine", {})


def _find_row(rows: list[dict[str, Any]], pk: str) -> dict[str, Any] | None:
    for row in rows:
        if str(row.get("reservation_hh_id", "")).strip() == pk:
            return row
    return None


def _comparatif(
    pk: str,
    saisie_before: list[dict[str, Any]],
    saisie_after: list[dict[str, Any]],
    lot4a_before: dict[str, Any],
    lot4a_after: dict[str, Any],
) -> list[dict[str, Any]]:
    master_before = lot4a_before.get("master_rows", [])
    master_after = lot4a_after.get("master_rows", [])
    return [
        {"element": "Lignes SAISIE HH", "avant": len(saisie_before), "apres": len(saisie_after)},
        {"element": "Lignes MASTER HH", "avant": len(master_before), "apres": len(master_after)},
        {
            "element": "Reservation HH simulee",
            "avant": "Presente" if _find_row(saisie_before, pk) else "Absente",
            "apres": "Presente" if _find_row(saisie_after, pk) else "Absente",
        },
        {
            "element": "Anomalies taux",
            "avant": len(lot4a_before.get("anomalies_taux", [])),
            "apres": len(lot4a_after.get("anomalies_taux", [])),
        },
        {
            "element": "Acomptes crees",
            "avant": _count_acomptes(master_before),
            "apres": _count_acomptes(master_after),
        },
    ]


def _payload_summary(preview: dict[str, Any]) -> dict[str, Any]:
    fields = [
        "reservation_hh_id", "logement_id", "proprietaire_id", "canal_id",
        "date_arrivee", "date_depart", "total_percu", "mode_paiement_id",
        "code_impact", "comptabilisation", "taux_commission_override",
        "menage_override",
    ]
    return {field: preview.get(field) for field in fields}


def run_previsualisation(
    form_data: dict[str, str],
    *,
    saisie_source: Path | None = None,
    ref_setup_source: Path | None = None,
    dryruns_root: Path | None = None,
    as_of_iso: str | None = None,
) -> dict[str, Any]:
    """Valide et execute une simulation complete sur copies.

    La fonction ne contourne pas `HH_REAL_WRITE_ENABLED` : elle n'appelle jamais
    le writer reel et n'ecrit que dans le sous-dossier de dry-run.
    """
    source_saisie = Path(saisie_source or cfg.SAISIE_RESERVATIONS_HH)
    source_ref = Path(ref_setup_source or cfg.REF_SETUP)
    root = Path(dryruns_root or DRYRUNS_DIR)
    simulation_id = _token()
    run_dir = root / simulation_id
    run_dir.mkdir(parents=True, exist_ok=False)

    validation = saisie_svc.valider(
        form_data,
        saisie_path=source_saisie,
        ref_setup_path=source_ref,
    )
    if not validation["ok"]:
        manifest = {
            "simulation_id": simulation_id,
            "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "app_version": "APP-2c",
            "hh_real_write_enabled": cfg.HH_REAL_WRITE_ENABLED,
            "status": "VALIDATION_REFUSEE",
            "errors": validation["erreurs"],
            "payload_summary": {},
        }
        (run_dir / MANIFEST_NAME).write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2, default=_json_default),
            encoding="utf-8",
        )
        return {"ok": False, "token": simulation_id, "run_dir": run_dir, "manifest": manifest}

    preview = validation["preview"]
    pk = validation["pk"]
    saisie_copy = run_dir / SAISIE_COPY_NAME
    ref_copy = run_dir / REF_COPY_NAME
    master_sim = run_dir / MASTER_SIM_NAME
    result_lot4a_path = run_dir / RESULT_LOT4A_NAME
    now_iso = as_of_iso or datetime.now(timezone.utc).isoformat()

    source_hashes = {
        "saisie": _fingerprint(source_saisie),
        "ref_setup": _fingerprint(source_ref),
    }
    errors: list[str] = []
    status = "OK"
    target_row: int | None = None
    added_fields: list[str] = []
    ref_migrated = False
    lot4a_before: dict[str, Any] = {}
    lot4a_after: dict[str, Any] = {}
    saisie_before: list[dict[str, Any]] = []
    saisie_after: list[dict[str, Any]] = []
    simulated_master_row: dict[str, Any] | None = None
    engine_info: dict[str, Any] = {}
    formula_cache: dict[str, Any] = {}

    try:
        _copy_source(source_saisie, saisie_copy)
        _copy_source(source_ref, ref_copy)
        added_fields = migrate_saisie_copy(saisie_copy)
        ref_migrated = migrate_ref_setup_copy(ref_copy)

        saisie_before, lot4a_before, engine_before = _run_lot4a_engine(
            saisie_copy, ref_copy, now_iso, run_dir
        )
        engine_info = engine_before
        row_data = _build_simulated_row(preview)
        target_row, _ = _write_simulated_row(saisie_copy, row_data, run_dir)
        formula_cache = _formula_cache_state(saisie_copy, target_row)
        saisie_after, lot4a_after, engine_after = _run_lot4a_engine(
            saisie_copy, ref_copy, now_iso, run_dir, master_path=master_sim
        )
        engine_info = engine_after or engine_info
        simulated_master_row = _find_row(lot4a_after.get("master_rows", []), pk)
        if lot4a_after.get("statut") != "ANALYSE_TERMINEE":
            status = "ERREUR_LOT4A"
            errors.append(str(lot4a_after.get("motif_blocage") or lot4a_after.get("statut")))
    except Exception as exc:
        status = "ERREUR"
        errors.append(str(exc))

    resultat_lot4a = {
        "before": lot4a_before,
        "after": lot4a_after,
        "reservation_simulee": simulated_master_row,
        "engine": engine_info,
        "formula_cache": formula_cache,
        "status": status,
        "errors": errors,
    }
    result_lot4a_path.write_text(
        json.dumps(resultat_lot4a, ensure_ascii=False, indent=2, default=_json_default),
        encoding="utf-8",
    )

    copy_hashes = {
        "saisie": _fingerprint(saisie_copy),
        "ref_setup": _fingerprint(ref_copy),
        "master_simule": _fingerprint(master_sim),
        "resultat_lot4a": _fingerprint(result_lot4a_path),
    }
    comparatif = _comparatif(pk, saisie_before, saisie_after, lot4a_before, lot4a_after)
    manifest = {
        "simulation_id": simulation_id,
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "app_version": "APP-2c",
        "hh_real_write_enabled": cfg.HH_REAL_WRITE_ENABLED,
        "source_hashes": source_hashes,
        "copy_hashes": copy_hashes,
        "paths": {
            "run_dir": str(run_dir),
            "saisie_copy": str(saisie_copy),
            "ref_setup_copy": str(ref_copy),
            "master_simule": str(master_sim),
            "resultat_lot4a": str(result_lot4a_path),
        },
        "payload_summary": _payload_summary(preview),
        "migration": {
            "saisie_fields_added": added_fields,
            "ref_direct_proprietaire_added": ref_migrated,
        },
        "engine": engine_info,
        "formula_cache": formula_cache,
        "target_row": target_row,
        "status": status,
        "errors": errors,
        "comparatif": comparatif,
        "lot4a_status": lot4a_after.get("statut"),
    }
    (run_dir / MANIFEST_NAME).write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, default=_json_default),
        encoding="utf-8",
    )

    return {
        "ok": status == "OK",
        "token": simulation_id,
        "run_dir": run_dir,
        "manifest": manifest,
        "preview": preview,
        "pk": pk,
        "target_row": target_row,
        "saisie_copy": saisie_copy,
        "ref_copy": ref_copy,
        "master_simule": master_sim,
        "resultat_lot4a": resultat_lot4a,
        "simulated_master_row": simulated_master_row,
        "comparatif": comparatif,
    }


def load_previsualisation(token: str, *, dryruns_root: Path | None = None) -> dict[str, Any]:
    simulation_id = _safe_token(token)
    root = Path(dryruns_root or DRYRUNS_DIR)
    run_dir = root / simulation_id
    manifest_path = run_dir / MANIFEST_NAME
    result_path = run_dir / RESULT_LOT4A_NAME
    if not manifest_path.exists():
        raise DryRunError("Previsualisation introuvable")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    resultat_lot4a = {}
    if result_path.exists():
        resultat_lot4a = json.loads(result_path.read_text(encoding="utf-8"))
    return {
        "token": simulation_id,
        "run_dir": run_dir,
        "manifest": manifest,
        "resultat_lot4a": resultat_lot4a,
        "preview": manifest.get("payload_summary", {}),
        "simulated_master_row": resultat_lot4a.get("reservation_simulee"),
        "comparatif": manifest.get("comparatif", []),
    }
