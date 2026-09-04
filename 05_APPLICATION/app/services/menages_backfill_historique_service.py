"""Import historique FIGÉ des coûts complets ménage (mission « backfill historique ménages figé »).

CE MODULE N'EST PAS UN MOTEUR DE CALCUL. Il ne rejoue aucune formule économique.

Contexte — pourquoi il existe
    Lot9 lisait TYPE_FLUX_018/019 dans `MASTER_CALC_CoutComplet_Menages.xlsx`. Basculer cette
    lecture vers `menages_cout_complet` (SQLite) supprimait le flux de 2026-05 : le classeur ne
    contient que ce mois-là, la table ne contenait que 2026-06/2026-07. Or 2026-05 est CLOTURÉ, et
    le recalculer est interdit — un mois arrêté ne se recalcule pas, il se conserve.

    D'où ce chemin : les valeurs déjà figées du classeur sont importées VERBATIM, telles qu'elles
    ont été arrêtées, avec leur provenance. Aucune ligne n'est recalculée, aucune règle n'est
    rejouée, `ref_cloture_mensuelle` n'est pas touchée, le mois reste CLOTURE.

Garde-fou central (`_mois_deja_calcule`)
    Un mois qui possède déjà des lignes CALCULÉES par le moteur n'est pas de l'historique : le
    backfiller écraserait un résultat courant par une valeur gelée, soit exactement la régression
    silencieuse que la mission interdit. Ces mois-là sont refusés, jamais importés.
"""
from __future__ import annotations

import hashlib
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

SOURCE_TYPE_FIGE = "LEGACY_IMPORT_FIGE"
FEUILLE = "DETAIL_COUT_COMPLET"

# Statut de contrôle du classeur legacy Lot6c — vocabulaire DISTINCT de l'énumération applicative
# `factures_service.ST_VALIDEE` ("VALIDEE"). Ici c'est la valeur historique "VALIDE", telle que lot9
# la filtrait (`men_valide`). Ne pas confondre les deux : ce sont deux référentiels différents.
STATUT_LEGACY_VALIDE = "VALIDE"

# Clé métier d'une ligne de coût complet (cf. index partiel unique, migration 0068).
CLE = ("mois", "logement_id", "intervenant_id")

E_CLASSEUR_ABSENT = "BACKFILL_CLASSEUR_ABSENT"
E_FEUILLE_ABSENTE = "BACKFILL_FEUILLE_ABSENTE"
E_MOIS_DEJA_CALCULE = "BACKFILL_MOIS_DEJA_CALCULE"
E_TABLE_ABSENTE = "BACKFILL_TABLE_ABSENTE"


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _colonnes_table(conn: sqlite3.Connection) -> list[str]:
    return [r[1] for r in conn.execute("PRAGMA table_info(menages_cout_complet)")]


def lire_classeur(chemin) -> list[dict[str, Any]]:
    """Lignes brutes de la feuille DETAIL_COUT_COMPLET, sans aucune transformation.

    Lecture seule : le classeur n'est jamais réécrit, son empreinte reste inchangée.
    """
    import openpyxl

    path = Path(chemin)
    if not path.exists():
        raise FileNotFoundError(f"{E_CLASSEUR_ABSENT}: {path}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if FEUILLE not in wb.sheetnames:
            raise KeyError(f"{E_FEUILLE_ABSENTE}: {FEUILLE}")
        ws = wb[FEUILLE]
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    finally:
        wb.close()
    if not rows:
        return []
    entetes = [str(c) for c in rows[0]]
    return [dict(zip(entetes, r)) for r in rows[1:]]


def _mois_deja_calcule(conn: sqlite3.Connection, mois: str) -> bool:
    """Vrai si le mois porte des lignes produites par le moteur (donc pas de l'historique gelé).

    Une ligne de `menages_cout_complet` SANS entrée de provenance a été écrite par lot6f. Importer
    par-dessus reviendrait à remplacer un calcul courant par une valeur figée : refusé.
    """
    n = conn.execute(
        "SELECT COUNT(*) FROM menages_cout_complet c "
        "WHERE c.mois = ? AND NOT EXISTS ("
        "  SELECT 1 FROM menages_cout_complet_provenance p "
        "  WHERE p.mois = c.mois AND p.logement_id = c.logement_id "
        "    AND p.intervenant_id = c.intervenant_id)",
        (mois,),
    ).fetchone()[0]
    return n > 0


def importer_historique_fige(*, chemin_classeur, db_path, mois_autorises=None) -> dict[str, Any]:
    """Importe VERBATIM les lignes figées du classeur legacy dans `menages_cout_complet`.

    `mois_autorises` restreint l'import (None = tous les mois présents dans le classeur). Un mois
    déjà calculé par le moteur est refusé, jamais écrasé.

    Idempotent : un rejeu remplace la ligne figée de même clé au lieu d'en ajouter une seconde.
    """
    path = Path(chemin_classeur)
    lignes = lire_classeur(path)
    empreinte = _sha256(path)
    horodatage = _now()

    conn = sqlite3.connect(str(db_path))
    try:
        colonnes = _colonnes_table(conn)
        if not colonnes:
            return {"ok": False, "code": E_TABLE_ABSENTE,
                    "message": "menages_cout_complet absente : migrations non appliquées."}

        # Seules les colonnes réellement communes au classeur et à la table sont reprises : aucune
        # valeur n'est inventée pour combler une colonne absente du classeur (methode,
        # cout_interne_ref_id, run_id... restent NULL, ce qui est l'information exacte).
        communes = [c for c in colonnes if c in (lignes[0] if lignes else {})]
        for cle in CLE:
            if cle not in communes:
                return {"ok": False, "code": E_FEUILLE_ABSENTE,
                        "message": f"Colonne clé absente du classeur : {cle}"}

        mois_classeur = sorted({str(l.get("mois")) for l in lignes if l.get("mois")})
        cibles = [m for m in mois_classeur
                  if mois_autorises is None or m in set(mois_autorises)]

        refuses = [m for m in cibles if _mois_deja_calcule(conn, m)]
        if refuses:
            return {"ok": False, "code": E_MOIS_DEJA_CALCULE, "mois_refuses": refuses,
                    "message": f"Mois déjà calculés par le moteur, import figé refusé : {refuses}. "
                               "Un résultat courant ne se remplace pas par une valeur gelée."}

        insertions = [l for l in lignes if str(l.get("mois")) in set(cibles)]
        place = ", ".join("?" for _ in communes)

        nb = 0
        for l in insertions:
            cle = (l.get("mois"), l.get("logement_id"), l.get("intervenant_id"))
            # Idempotence : on retire la ligne figée de même clé avant réinsertion. La suppression
            # est bornée aux clés DÉJÀ enregistrées comme figées — une ligne calculée par le moteur
            # n'est jamais supprimée ici (et les mois calculés ont de toute façon été refusés
            # au-dessus).
            conn.execute(
                "DELETE FROM menages_cout_complet "
                "WHERE mois = ? AND logement_id = ? AND intervenant_id = ? AND EXISTS ("
                "  SELECT 1 FROM menages_cout_complet_provenance p "
                "  WHERE p.mois = ? AND p.logement_id = ? AND p.intervenant_id = ?)",
                cle + cle,
            )
            conn.execute(
                f"INSERT INTO menages_cout_complet ({', '.join(communes)}) VALUES ({place})",
                [l.get(c) for c in communes])
            conn.execute(
                "INSERT OR REPLACE INTO menages_cout_complet_provenance "
                "(mois, logement_id, intervenant_id, source_type, source_fichier, source_hash, "
                " date_import) VALUES (?,?,?,?,?,?,?)",
                cle + (SOURCE_TYPE_FIGE, path.name, empreinte, horodatage))
            nb += 1
        conn.commit()
    finally:
        conn.close()

    return {"ok": True, "nb_lignes": nb, "mois": cibles, "source_fichier": path.name,
            "source_hash": empreinte, "date_import": horodatage,
            "colonnes_reprises": communes}


# ── Ménages externes facturés (SRC_MEN / TYPE_FLUX_014) ─────────────────────────────────────────
# Même décision, même garde-fous que ci-dessus, sur l'autre classeur legacy que lot9 lisait encore.
# La différence tient à la cible : il n'existe AUCUNE table canonique capable d'accueillir ces
# lignes sans inventer de fausses `factures` (cf. migration 0069), elles vont donc dans leur propre
# table d'historique.

FEUILLE_EXTERNES = "MASTER"

# Colonnes reprises verbatim du classeur legacy vers `menages_externes_historique`.
# `source_pk` alimente le ROW_HASH de lot9 : il doit être repris tel quel, jamais régénéré.
_COLONNES_EXTERNES = (
    "source_pk", "mois", "logement_id", "proprietaire_id", "prestataire_id",
    "date_facture", "date_menage", "montant_ligne_ttc", "type_flux_id", "sens", "code_impact",
)


def _date_iso(valeur):
    """Date de classeur -> 'AAAA-MM-JJ'. Laisse passer None et les chaînes déjà normalisées."""
    if valeur is None or isinstance(valeur, str):
        return valeur
    iso = getattr(valeur, "isoformat", None)
    return iso()[:10] if callable(iso) else str(valeur)


def lire_classeur_externes(chemin) -> list[dict[str, Any]]:
    """Lignes brutes de la feuille MASTER du classeur Lot6c, sans transformation (lecture seule)."""
    import openpyxl

    path = Path(chemin)
    if not path.exists():
        raise FileNotFoundError(f"{E_CLASSEUR_ABSENT}: {path}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if FEUILLE_EXTERNES not in wb.sheetnames:
            raise KeyError(f"{E_FEUILLE_ABSENTE}: {FEUILLE_EXTERNES}")
        ws = wb[FEUILLE_EXTERNES]
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    finally:
        wb.close()
    if not rows:
        return []
    entetes = [str(c) for c in rows[0]]
    return [dict(zip(entetes, r)) for r in rows[1:]]


def importer_externes_historique(*, chemin_classeur, db_path, mois_autorises=None,
                                 statut_cloture: str | None = None) -> dict[str, Any]:
    """Importe VERBATIM les ménages externes figés du classeur Lot6c.

    Toutes les lignes sont importées, y compris celles dont le `statut_controle` legacy n'est pas
    VALIDE : l'historique n'est jamais amputé. C'est le LECTEUR qui écarte ensuite les non-VALIDE du
    calcul économique, exactement comme le faisait lot9 (`men_valide`).

    Idempotent : la clé est `source_pk`, un rejeu remplace la ligne au lieu d'en ajouter une seconde.
    """
    path = Path(chemin_classeur)
    lignes = lire_classeur_externes(path)
    empreinte = _sha256(path)
    horodatage = _now()

    conn = sqlite3.connect(str(db_path))
    try:
        existe = conn.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name=?",
            ("menages_externes_historique",)).fetchone()[0]
        if not existe:
            return {"ok": False, "code": E_TABLE_ABSENTE,
                    "message": "menages_externes_historique absente : migrations non appliquées."}

        mois_classeur = sorted({str(l.get("mois")) for l in lignes if l.get("mois")})
        cibles = [m for m in mois_classeur if mois_autorises is None or m in set(mois_autorises)]
        insertions = [l for l in lignes if str(l.get("mois")) in set(cibles)]

        nb = nb_valide = 0
        for l in insertions:
            statut_src = str(l.get("statut_controle") or "")
            # openpyxl rend les cellules de date en `datetime` ; sqlite3 3.12 déprécie leur
            # adaptation implicite. On stocke la forme ISO courte, celle que lot9 relit.
            valeurs = [_date_iso(l.get(c)) if c in ("date_facture", "date_menage")
                       else l.get(c) for c in _COLONNES_EXTERNES]
            conn.execute(
                f"INSERT OR REPLACE INTO menages_externes_historique "
                f"({', '.join(_COLONNES_EXTERNES)}, statut_source, source_type, source_fichier, "
                f" source_hash, statut_cloture, date_import) "
                f"VALUES ({', '.join('?' for _ in _COLONNES_EXTERNES)},?,?,?,?,?,?)",
                valeurs + [statut_src, SOURCE_TYPE_FIGE, path.name, empreinte,
                           statut_cloture, horodatage])
            nb += 1
            if statut_src == STATUT_LEGACY_VALIDE:
                nb_valide += 1
        conn.commit()
    finally:
        conn.close()

    return {"ok": True, "nb_lignes": nb, "nb_valide": nb_valide,
            "nb_non_valide": nb - nb_valide, "mois": cibles,
            "source_fichier": path.name, "source_hash": empreinte, "date_import": horodatage}


def menages_externes_economiques(*, db_path) -> list[dict[str, Any]]:
    """INTERFACE UNIQUE de lot9 pour TYPE_FLUX_014 (§C) — historique figé + factures courantes.

    Deux origines, un seul contrat de sortie ; lot9 ne sait pas laquelle il consomme :

      - historique figé (`menages_externes_historique`) : mois déjà arrêtés, repris tels quels du
        classeur legacy. Filtré sur le statut legacy VALIDE — c'est exactement ce que faisait
        `men_valide` dans lot9, le comportement économique est donc inchangé.

      - factures courantes (`facture_lignes_menage` × `factures`) : restreintes aux statuts qui
        engagent l'économie. Une facture A_CONTROLER n'apparaît jamais ici — elle reste visible dans
        le rapprochement (lot6d), avec un impact économique nul.

    `source_pk` est rendu tel quel pour l'historique (clé legacy conservée, ROW_HASH stable) et
    dérivé de l'identifiant de ligne pour le courant.

    `db_path=None` DOIT résoudre vers la vraie base (`cfg.DB_PATH`), pas se connecter au fichier
    littéral « None » : c'est exactement l'appel que fait `flux_unifie_service.construire()` en
    production (l'orchestrateur ne passe jamais de `db_path` explicite). `sqlite3.connect(str(None))`
    crée silencieusement un fichier vide sans table — TYPE_FLUX_014 retombait à 0 sans la moindre
    erreur, bug trouvé lors de la bascule réelle. `get_db()` fait la même résolution « à chaud » que
    tout le reste de l'application.
    """
    from app.db.connection import get_db

    conn = get_db(db_path)
    try:
        lignes: list[dict[str, Any]] = []

        def _table(nom: str) -> bool:
            return bool(conn.execute(
                "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name=?",
                (nom,)).fetchone()[0])

        mois_figes: set[str] = set()

        if _table("menages_externes_historique"):
            for r in conn.execute(
                    "SELECT * FROM menages_externes_historique WHERE statut_source = ? "
                    "ORDER BY source_pk", (STATUT_LEGACY_VALIDE,)):
                d = dict(r)
                mois_figes.add(d["mois"])
                lignes.append({
                    "source_pk": d["source_pk"], "mois": d["mois"],
                    "logement_id": d["logement_id"], "proprietaire_id": d["proprietaire_id"],
                    "prestataire_id": d["prestataire_id"], "date_facture": d["date_facture"],
                    "date_menage": d["date_menage"],
                    "montant_ligne_ttc": d["montant_ligne_ttc"],
                    "type_flux_id": d["type_flux_id"], "sens": d["sens"],
                    "code_impact": d["code_impact"], "origine": "HISTORIQUE_FIGE",
                })

        if _table("facture_lignes_menage") and _table("factures"):
            # Source de vérité applicative des statuts qui engagent l'économie (le miroir moteur
            # vit dans `lib_db_moteur.STATUTS_FACTURE_COMPTABLES`, verrouillé par un test de sync).
            from app.services.factures_service import STATUTS_COMPTABLES

            statuts_comptables = tuple(sorted(STATUTS_COMPTABLES))
            statuts = ", ".join("?" for _ in statuts_comptables)
            sql = (
                "SELECT l.ligne_id_opaque, l.logement_id, l.montant_ttc, "
                "       f.fournisseur_id_opaque, f.date_facture, f.statut "
                "FROM facture_lignes_menage l "
                "JOIN factures f ON f.facture_id_opaque = l.facture_id_opaque "
                f"WHERE l.type_ligne = 'MENAGE_EXTERNE' AND f.statut IN ({statuts}) "
                "ORDER BY l.ligne_id_opaque")
            for r in conn.execute(sql, statuts_comptables):
                d = dict(r)
                mois = str(d["date_facture"] or "")[:7]
                if mois in mois_figes:
                    # Autorité par période (§A3) : un mois déjà couvert par l'historique figé
                    # (mois clôturé au moment du backfill) ne doit JAMAIS être complété/écrasé par
                    # une facture courante apparue depuis — même statut comptable, même mois. Faire
                    # entrer cette ligne ici recréerait en silence exactement le risque que le
                    # backfill figé existe pour éviter : une donnée d'un mois clôturé qui bouge sans
                    # passer par le workflow explicite de correction rétroactive. Cette ligne reste
                    # visible ailleurs (rapprochement, écran facture) ; elle est seulement exclue de
                    # CETTE agrégation économique.
                    continue
                lignes.append({
                    "source_pk": d["ligne_id_opaque"],
                    "mois": mois,
                    "logement_id": d["logement_id"], "proprietaire_id": None,
                    "prestataire_id": d["fournisseur_id_opaque"],
                    "date_facture": d["date_facture"], "date_menage": None,
                    "montant_ligne_ttc": d["montant_ttc"],
                    "type_flux_id": "TYPE_FLUX_014", "sens": "CHARGE",
                    "code_impact": "IC", "origine": "FACTURE_VALIDEE",
                })
        return lignes
    finally:
        conn.close()


def lignes_figees(*, db_path, mois=None) -> list[dict[str, Any]]:
    """Relit les lignes figées importées — utilisé par les contrôles de parité."""
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    try:
        sql = ("SELECT c.* FROM menages_cout_complet c "
               "JOIN menages_cout_complet_provenance p "
               "  ON p.mois = c.mois AND p.logement_id = c.logement_id "
               " AND p.intervenant_id = c.intervenant_id "
               "WHERE p.source_type = ?")
        args: list[Any] = [SOURCE_TYPE_FIGE]
        if mois:
            sql += " AND c.mois = ?"
            args.append(mois)
        return [dict(r) for r in conn.execute(sql, args)]
    finally:
        conn.close()
