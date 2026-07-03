"""Migration APP-2b sur copies de classeurs uniquement.

Ce module refuse les chemins sources reels et ne doit pas etre utilise pour
activer l'ecriture applicative. Il prepare le schema cible a tester sur copies.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

SHEET_SAISIE = "SAISIE"
SHEET_MODES_PAIEMENT = "REF_Modes_Paiement"

NEW_SAISIE_FIELDS = [
    "taux_commission_override",
    "motif_override_taux_commission",
    "confirmation_override_taux_commission",
    "menage_override",
    "motif_override_menage",
    "confirmation_override_menage",
    "source_acompte_facture",
]

DIRECT_PROPRIETAIRE_ROW = {
    "mode_paiement_id": "PAY_006",
    "mode_paiement": "DIRECT_PROPRIETAIRE",
    "impact_banque": "NON",
    "impact_caisse": "NON",
    "impact_associee": "NON",
    "actif": "OUI",
    "commentaire": "APP-2b - Direct proprietaire",
}

FORBIDDEN_PARTS = {
    "01_SOURCES_BRUTES",
    "MASTER",
    "03_EXPORTS",
}


def _assert_copy_path(path: Path) -> Path:
    resolved = Path(path).resolve()
    parts = {part.upper() for part in resolved.parts}
    forbidden = {part.upper() for part in FORBIDDEN_PARTS}
    if parts & forbidden:
        raise RuntimeError(f"Migration refusee sur chemin source ou production: {resolved}")
    return resolved


def migrate_saisie_copy(path: Path) -> list[str]:
    target = _assert_copy_path(path)
    wb = openpyxl.load_workbook(str(target), keep_vba=False)
    try:
        ws = wb[SHEET_SAISIE]
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        added: list[str] = []
        for field in NEW_SAISIE_FIELDS:
            if field in headers:
                continue
            ws.cell(row=1, column=len(headers) + 1, value=field)
            headers.append(field)
            added.append(field)
        wb.calculation.fullCalcOnLoad = True
        wb.save(str(target))
        return added
    finally:
        wb.close()


def migrate_ref_setup_copy(path: Path) -> bool:
    target = _assert_copy_path(path)
    wb = openpyxl.load_workbook(str(target), keep_vba=True)
    try:
        ws = wb[SHEET_MODES_PAIEMENT]
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        idx = {name: pos + 1 for pos, name in enumerate(headers)}
        id_col = idx.get("mode_paiement_id")
        if id_col is None:
            raise RuntimeError("REF_Modes_Paiement sans colonne mode_paiement_id")
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=id_col).value or "").strip() == "PAY_006":
                return False
        new_row = ws.max_row + 1
        for name, value in DIRECT_PROPRIETAIRE_ROW.items():
            if name in idx:
                ws.cell(row=new_row, column=idx[name], value=value)
        wb.save(str(target))
        return True
    finally:
        wb.close()
