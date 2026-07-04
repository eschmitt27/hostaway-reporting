"""APP-2e - diagnostic et preparation controlee du schema HH.

Les fonctions de migration sur copies sont utilisees par les tests et la recette.
La migration reelle future reste separee, explicite et protegee par confirmation.
"""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
from typing import Any
import zipfile

import openpyxl

import app.config as cfg
from app.readers.saisie_hh_reader import FORMULA_COLS, find_first_empty_data_row, _col_index
from app.services import saisie_hh_dryrun_service as dryrun_svc
from app.services.saisie_hh_schema_migration import (
    DIRECT_PROPRIETAIRE_ROW,
    NEW_SAISIE_FIELDS,
    migrate_ref_setup_copy,
    migrate_saisie_copy,
)
from app.writers.saisie_hh_writer import (
    _check_formula_cells,
    _check_structural_preservation,
    _measure_structure,
)


CONFIRMATION_EXECUTION = "MIGRER_SCHEMA_HH_REELLE"
MANIFEST_NAME = "manifest_schema_hh.json"
SAISIE_MIGREE_NAME = "SAISIE_ReservationsHorsHostaway_schema_prepare.xlsx"
REF_MIGREE_NAME = "REF_Setup_schema_prepare.xlsm"
SAISIE_REF_NAME = "SAISIE_ReservationsHorsHostaway_ref.xlsx"
REF_REF_NAME = "REF_Setup_ref.xlsm"

# ── Contrôle intégrité ZIP / VBA ───────────────────────────────────────────────
_VBA_PROJECT_PATH = "xl/vbaProject.bin"
_VBA_SIG_PATH = "xl/vbaProjectSignature.bin"
_CONTENT_TYPES_PATH = "[Content_Types].xml"
_WORKBOOK_RELS_PATH = "xl/_rels/workbook.xml.rels"
_VBA_MACRO_ENABLED_MARKER = "macroEnabled"
_VBA_RELATIONSHIP_TYPE = "relationships/vbaProject"

_SENSITIVE_PREFIXES = (
    "xl/activeX/",
    "xl/ctrlProps/",
    "xl/embeddings/",
    "xl/externalLinks/",
    "xl/connections.xml",
    "customUI/",
    "docProps/custom.xml",
    "xl/printerSettings/",
)


def _utc_stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def _fingerprint(path: Path) -> dict[str, Any]:
    return dryrun_svc._fingerprint(path)


def _read_headers(path: Path, sheet_name: str, *, keep_vba: bool = False) -> list[str]:
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True, keep_vba=keep_vba)
    try:
        ws = wb[sheet_name]
        return [str(cell.value or "").strip() for cell in ws[1]]
    finally:
        wb.close()


def _sheet_exists(path: Path, sheet_name: str, *, keep_vba: bool = False) -> bool:
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True, keep_vba=keep_vba)
        try:
            return sheet_name in wb.sheetnames
        finally:
            wb.close()
    except Exception:
        return False


def _mode_direct_present(ref_setup_path: Path) -> bool:
    wb = openpyxl.load_workbook(str(ref_setup_path), read_only=True, data_only=True, keep_vba=True)
    try:
        ws = wb["REF_Modes_Paiement"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return False
    headers = [str(value or "").strip() for value in rows[0]]
    if "mode_paiement_id" not in headers:
        return False
    id_idx = headers.index("mode_paiement_id")
    code_idx = headers.index("mode_paiement") if "mode_paiement" in headers else None
    for row in rows[1:]:
        mode_id = str(row[id_idx] or "").strip() if len(row) > id_idx else ""
        mode_code = str(row[code_idx] or "").strip() if code_idx is not None and len(row) > code_idx else ""
        if mode_id == DIRECT_PROPRIETAIRE_ROW["mode_paiement_id"] or mode_code == DIRECT_PROPRIETAIRE_ROW["mode_paiement"]:
            return True
    return False


def _workbook_features(path: Path, *, keep_vba: bool = False) -> dict[str, Any]:
    wb = openpyxl.load_workbook(str(path), read_only=False, data_only=False, keep_vba=keep_vba)
    try:
        validations = 0
        tables = 0
        mfc = 0
        for ws in wb.worksheets:
            validations += len(ws.data_validations.dataValidation)
            tables += len(ws.tables)
            mfc += sum(len(rules) for rules in ws.conditional_formatting._cf_rules.values())
        return {
            "sheetnames": list(wb.sheetnames),
            "defined_names": len(list(wb.defined_names.values())),
            "data_validations": validations,
            "conditional_formatting": mfc,
            "tables": tables,
            "has_vba": bool(getattr(wb, "vba_archive", None)),
            "full_calc_on_load": bool(getattr(wb.calculation, "fullCalcOnLoad", False)),
        }
    finally:
        wb.close()


def _zip_sha256_entry(zf: zipfile.ZipFile, name: str) -> str | None:
    """SHA-256 du contenu décompressé d'une entrée ZIP, ou None si absente."""
    try:
        return hashlib.sha256(zf.read(name)).hexdigest()
    except KeyError:
        return None


def _vba_snapshot(path: Path) -> dict[str, Any]:
    """Empreinte des parties VBA du package ZIP (pour manifest et comparaison)."""
    try:
        with zipfile.ZipFile(str(path), "r") as zf:
            names = set(zf.namelist())
            vba_present = _VBA_PROJECT_PATH in names
            vba_sha256 = _zip_sha256_entry(zf, _VBA_PROJECT_PATH) if vba_present else None
            sig_present = _VBA_SIG_PATH in names
            sig_sha256 = _zip_sha256_entry(zf, _VBA_SIG_PATH) if sig_present else None
            macro_enabled = False
            vba_rel_valid = False
            if _CONTENT_TYPES_PATH in names:
                ct = zf.read(_CONTENT_TYPES_PATH).decode("utf-8", errors="replace")
                macro_enabled = _VBA_MACRO_ENABLED_MARKER in ct
            if _WORKBOOK_RELS_PATH in names:
                wr = zf.read(_WORKBOOK_RELS_PATH).decode("utf-8", errors="replace")
                vba_rel_valid = _VBA_RELATIONSHIP_TYPE in wr
            return {
                "vba_project_present": vba_present,
                "vba_project_sha256": vba_sha256,
                "vba_signature_present": sig_present,
                "vba_signature_sha256": sig_sha256,
                "macro_enabled_declared": macro_enabled,
                "vba_relationship_valid": vba_rel_valid,
            }
    except Exception as exc:
        return {"error": str(exc)}


def _check_zip_vba_integrity(ref_path: Path, work_path: Path) -> list[str]:
    """Vérifie l'intégrité binaire VBA (octet pour octet) entre ref et travail."""
    violations: list[str] = []
    try:
        snap_ref = _vba_snapshot(ref_path)
        snap_work = _vba_snapshot(work_path)

        if "error" in snap_ref or "error" in snap_work:
            return [
                f"VBA_PRESERVATION_ECHEC: lecture ZIP impossible "
                f"(ref={snap_ref.get('error')}, work={snap_work.get('error')})"
            ]

        if snap_ref["vba_project_present"]:
            if not snap_work["vba_project_present"]:
                violations.append(
                    "VBA_PRESERVATION_ECHEC: xl/vbaProject.bin disparu apres migration"
                )
            elif snap_ref["vba_project_sha256"] != snap_work["vba_project_sha256"]:
                violations.append(
                    f"VBA_PRESERVATION_ECHEC: xl/vbaProject.bin modifie "
                    f"(ref={snap_ref['vba_project_sha256'][:16]}... "
                    f"work={snap_work['vba_project_sha256'][:16]}...)"
                )

        if snap_ref["vba_signature_present"]:
            if not snap_work["vba_signature_present"]:
                violations.append(
                    "VBA_PRESERVATION_ECHEC: xl/vbaProjectSignature.bin disparu apres migration"
                )
            elif snap_ref["vba_signature_sha256"] != snap_work["vba_signature_sha256"]:
                violations.append(
                    "VBA_PRESERVATION_ECHEC: xl/vbaProjectSignature.bin modifie apres migration"
                )

        if snap_ref["macro_enabled_declared"] and not snap_work["macro_enabled_declared"]:
            violations.append(
                "VBA_PRESERVATION_ECHEC: [Content_Types].xml ne declare plus de classeur macro-enabled"
            )

        if snap_ref["vba_relationship_valid"] and not snap_work["vba_relationship_valid"]:
            violations.append(
                "VBA_PRESERVATION_ECHEC: xl/_rels/workbook.xml.rels ne contient plus de relation VBA"
            )

    except Exception as exc:
        violations.append(f"VBA_PRESERVATION_ECHEC: verification impossible : {exc}")

    return violations


def _check_zip_sensitive_parts(ref_path: Path, work_path: Path) -> list[str]:
    """Vérifie que les parties ZIP sensibles n'ont pas disparu ou été altérées."""
    violations: list[str] = []
    try:
        with (
            zipfile.ZipFile(str(ref_path), "r") as zf_ref,
            zipfile.ZipFile(str(work_path), "r") as zf_work,
        ):
            ref_names = set(zf_ref.namelist())
            work_names = set(zf_work.namelist())
            for entry in sorted(ref_names):
                for prefix in _SENSITIVE_PREFIXES:
                    if entry == prefix or entry.startswith(prefix):
                        if entry not in work_names:
                            violations.append(
                                f"PACKAGE_SENSIBLE_PRESERVATION_ECHEC: {entry} disparu apres migration"
                            )
                        else:
                            if _zip_sha256_entry(zf_ref, entry) != _zip_sha256_entry(zf_work, entry):
                                violations.append(
                                    f"PACKAGE_SENSIBLE_PRESERVATION_ECHEC: {entry} modifie apres migration"
                                )
                        break
    except Exception as exc:
        violations.append(
            f"PACKAGE_SENSIBLE_PRESERVATION_ECHEC: verification impossible : {exc}"
        )
    return violations


def diagnostiquer_schema_hh(
    *,
    saisie_path: Path | None = None,
    ref_setup_path: Path | None = None,
) -> dict[str, Any]:
    """Diagnostic idempotent du schema HH actuel."""
    p_saisie = Path(saisie_path or cfg.SAISIE_RESERVATIONS_HH)
    p_ref = Path(ref_setup_path or cfg.REF_SETUP)
    errors: list[str] = []

    saisie_headers: list[str] = []
    ref_headers: list[str] = []
    try:
        saisie_headers = _read_headers(p_saisie, "SAISIE")
    except Exception as exc:
        errors.append(f"SAISIE_HEADERS_ILLISIBLES: {exc}")
    try:
        ref_headers = _read_headers(p_ref, "REF_Modes_Paiement", keep_vba=True)
    except Exception as exc:
        errors.append(f"REF_MODES_HEADERS_ILLISIBLES: {exc}")

    missing_saisie = [field for field in NEW_SAISIE_FIELDS if field not in saisie_headers]
    mode_present = False
    try:
        mode_present = _mode_direct_present(p_ref)
    except Exception as exc:
        errors.append(f"REF_MODE_DIRECT_ILLISIBLE: {exc}")

    target_row = find_first_empty_data_row(p_saisie)
    formula_violations = _check_formula_cells(p_saisie, target_row) if target_row else ["AUCUNE_LIGNE_CIBLE"]
    critical_formulas = {
        col: not any(f"col {col} " in violation for violation in formula_violations)
        for col in sorted(FORMULA_COLS)
    }

    try:
        saisie_features = _workbook_features(p_saisie)
    except Exception as exc:
        saisie_features = {"error": str(exc)}
        errors.append(f"SAISIE_FEATURES_ILLISIBLES: {exc}")
    try:
        ref_features = _workbook_features(p_ref, keep_vba=True)
    except Exception as exc:
        ref_features = {"error": str(exc)}
        errors.append(f"REF_FEATURES_ILLISIBLES: {exc}")

    migration_needed = bool(missing_saisie or not mode_present)
    return {
        "status": "OK" if not errors else "ERREUR",
        "hashes": {"saisie": _fingerprint(p_saisie), "ref_setup": _fingerprint(p_ref)},
        "paths": {"saisie": str(p_saisie), "ref_setup": str(p_ref)},
        "sheets": {
            "saisie": _sheet_exists(p_saisie, "SAISIE"),
            "ref_modes_paiement": _sheet_exists(p_ref, "REF_Modes_Paiement", keep_vba=True),
        },
        "missing_saisie_fields": missing_saisie,
        "mode_direct_proprietaire_present": mode_present,
        "missing_ref_modes": [] if mode_present else ["PAY_006 / DIRECT_PROPRIETAIRE"],
        "target_row": target_row,
        "critical_formulas": critical_formulas,
        "formula_violations": formula_violations,
        "features": {"saisie": saisie_features, "ref_setup": ref_features},
        "migration_needed": migration_needed,
        "errors": errors,
    }


def _assert_not_real_source(path: Path) -> None:
    resolved = Path(path).resolve()
    for real in (cfg.SAISIE_RESERVATIONS_HH.resolve(), cfg.REF_SETUP.resolve()):
        if resolved == real:
            raise RuntimeError(f"Operation sur fichier source reel interdite pendant APP-2e: {resolved}")


def preparer_migration_hh_sur_copies(
    *,
    saisie_source: Path | None = None,
    ref_setup_source: Path | None = None,
    output_dir: Path,
) -> dict[str, Any]:
    """Copie les classeurs dans output_dir, migre uniquement les copies de travail et valide.

    Deux copies sont créées par source :
    - copie référence (immuable, jamais migrée) — sert de base pour la comparaison avant
    - copie de travail — reçoit la migration, sert de base pour la comparaison après
    """
    source_saisie = Path(saisie_source or cfg.SAISIE_RESERVATIONS_HH)
    source_ref = Path(ref_setup_source or cfg.REF_SETUP)
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    # 1. Deux copies par source : référence immuable + copie de travail
    saisie_work = out / SAISIE_MIGREE_NAME
    saisie_ref_copy = out / SAISIE_REF_NAME
    ref_work = out / REF_MIGREE_NAME
    ref_ref_copy = out / REF_REF_NAME
    shutil.copy2(source_saisie, saisie_work)
    shutil.copy2(source_saisie, saisie_ref_copy)
    shutil.copy2(source_ref, ref_work)
    shutil.copy2(source_ref, ref_ref_copy)

    # 2. Empreinte structurelle avant (depuis les copies référence, jamais migrées)
    struct_saisie_avant = _measure_structure(saisie_ref_copy)
    struct_ref_avant = _measure_structure(ref_ref_copy)
    diag_avant = diagnostiquer_schema_hh(saisie_path=saisie_work, ref_setup_path=ref_work)

    # 3. Migrer uniquement les copies de travail
    added_fields = migrate_saisie_copy(saisie_work)
    direct_added = migrate_ref_setup_copy(ref_work)

    # 4. Comparer référence (avant) → travail (après)
    struct_saisie_errors = _check_structural_preservation(saisie_ref_copy, saisie_work, struct_saisie_avant)
    struct_ref_errors = _check_structural_preservation(ref_ref_copy, ref_work, struct_ref_avant)

    # Empreintes VBA avant/après pour manifest et contrôle binaire
    vba_snap_saisie_avant = _vba_snapshot(saisie_ref_copy)
    vba_snap_saisie_apres = _vba_snapshot(saisie_work)
    vba_snap_ref_avant = _vba_snapshot(ref_ref_copy)
    vba_snap_ref_apres = _vba_snapshot(ref_work)
    vba_errors_saisie = _check_zip_vba_integrity(saisie_ref_copy, saisie_work)
    vba_errors_ref = _check_zip_vba_integrity(ref_ref_copy, ref_work)
    sens_errors_saisie = _check_zip_sensitive_parts(saisie_ref_copy, saisie_work)
    sens_errors_ref = _check_zip_sensitive_parts(ref_ref_copy, ref_work)

    diag_apres = diagnostiquer_schema_hh(saisie_path=saisie_work, ref_setup_path=ref_work)

    errors: list[str] = []
    errors.extend(struct_saisie_errors)
    errors.extend(struct_ref_errors)
    errors.extend(vba_errors_saisie)
    errors.extend(vba_errors_ref)
    errors.extend(sens_errors_saisie)
    errors.extend(sens_errors_ref)
    errors.extend(diag_apres.get("formula_violations", []))
    if diag_apres["missing_saisie_fields"] or diag_apres["missing_ref_modes"]:
        errors.append("SCHEMA_COPIE_INCOMPLET_APRES_MIGRATION")
    status = "OK" if not errors else "ERREUR"

    manifest = {
        "operation": "APP-2e schema HH sur copies",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": status,
        "source_hashes": {"saisie": _fingerprint(source_saisie), "ref_setup": _fingerprint(source_ref)},
        "copy_hashes": {"saisie": _fingerprint(saisie_work), "ref_setup": _fingerprint(ref_work)},
        "paths": {
            "output_dir": str(out),
            "saisie_copy": str(saisie_work),
            "ref_setup_copy": str(ref_work),
        },
        "diagnostic_avant": diag_avant,
        "diagnostic_apres": diag_apres,
        "migration": {"saisie_fields_added": added_fields, "direct_proprietaire_added": direct_added},
        "structure_errors": {"saisie": struct_saisie_errors, "ref_setup": struct_ref_errors},
        "vba_snapshots": {
            "saisie": {"avant": vba_snap_saisie_avant, "apres": vba_snap_saisie_apres},
            "ref_setup": {"avant": vba_snap_ref_avant, "apres": vba_snap_ref_apres},
        },
        "errors": errors,
    }
    (out / MANIFEST_NAME).write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return manifest


def executer_migration_hh_reelle(
    *,
    confirmation: str,
    saisie_path: Path | None = None,
    ref_setup_path: Path | None = None,
    work_dir: Path | None = None,
) -> dict[str, Any]:
    """Migration reelle future. Ne pas appeler sans confirmation humaine explicite."""
    if confirmation != CONFIRMATION_EXECUTION:
        return {"status": "REFUSE", "reason": "CONFIRMATION_INCORRECTE"}
    p_saisie = Path(saisie_path or cfg.SAISIE_RESERVATIONS_HH)
    p_ref = Path(ref_setup_path or cfg.REF_SETUP)
    root = Path(work_dir or (Path(tempfile.gettempdir()) / "schema_hh_real_migrations" / _utc_stamp()))
    root.mkdir(parents=True, exist_ok=True)

    before_hashes = {"saisie": _fingerprint(p_saisie), "ref_setup": _fingerprint(p_ref)}
    backup_saisie = root / f"{p_saisie.name}.backup"
    backup_ref = root / f"{p_ref.name}.backup"
    shutil.copy2(p_saisie, backup_saisie)
    shutil.copy2(p_ref, backup_ref)

    tmp_root = root / "tmp_validation"
    manifest = preparer_migration_hh_sur_copies(
        saisie_source=p_saisie,
        ref_setup_source=p_ref,
        output_dir=tmp_root,
    )
    if manifest["status"] != "OK":
        manifest["real_status"] = "REFUSE_VALIDATION_TEMPORAIRE"
        return manifest

    tmp_saisie = Path(manifest["paths"]["saisie_copy"])
    tmp_ref = Path(manifest["paths"]["ref_setup_copy"])
    replaced_saisie = False
    replaced_ref = False
    try:
        os.replace(str(tmp_saisie), str(p_saisie))
        replaced_saisie = True
        os.replace(str(tmp_ref), str(p_ref))
        replaced_ref = True
        after_hashes = {"saisie": _fingerprint(p_saisie), "ref_setup": _fingerprint(p_ref)}
        manifest.update({
            "real_status": "OK",
            "backup_paths": {"saisie": str(backup_saisie), "ref_setup": str(backup_ref)},
            "real_hashes_before": before_hashes,
            "real_hashes_after": after_hashes,
        })
        (root / MANIFEST_NAME).write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return manifest
    except Exception as exc:
        if replaced_saisie:
            os.replace(str(backup_saisie), str(p_saisie))
        if replaced_ref:
            os.replace(str(backup_ref), str(p_ref))
        rollback_hashes = {"saisie": _fingerprint(p_saisie), "ref_setup": _fingerprint(p_ref)}
        hash_ok = (
            rollback_hashes["saisie"].get("sha256") == before_hashes["saisie"].get("sha256")
            and rollback_hashes["ref_setup"].get("sha256") == before_hashes["ref_setup"].get("sha256")
        )
        manifest.update({
            "real_status": "ROLLBACK" if hash_ok else "ROLLBACK_HASH_MISMATCH",
            "error": str(exc),
            "real_hashes_before": before_hashes,
            "real_hashes_after_rollback": rollback_hashes,
            "rollback_hash_verified": hash_ok,
        })
        (root / MANIFEST_NAME).write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return manifest
