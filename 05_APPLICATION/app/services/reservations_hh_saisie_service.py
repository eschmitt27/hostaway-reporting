"""Saisie des réservations hors Hostaway — UI → SQLite (migration 0052).

CE QUE CE MODULE REMPLACE
La saisie passait par `SAISIE_ReservationsHorsHostaway.xlsx`, qu'un rafraîchissement Power Query
transformait en `MASTER_FACT_MAN_ReservationsHorsHostaway.xlsx`, lu ensuite par l'application. La
cible est directe : **UI → service → SQLite**, sans aller-retour par Excel.

L'ancien chemin d'écriture Excel (`saisie_hh_writer`, `saisie_hh_real_write_service` et leurs
garde-fous) n'est pas supprimé : il reste disponible pour une reprise contrôlée, mais il n'est plus
nécessaire au fonctionnement quotidien.

IDENTITÉ
`reservation_hh_id` suit la convention `RESHH-{mois}-{NNN}` du moteur. Le suffixe est calculé par
rapport aux réservations DÉJÀ enregistrées pour ce mois — il n'est jamais recalculé pour les autres,
donc l'ajout d'une réservation ne renomme rien.

CE QUI N'EST PAS RECALCULÉ ICI
Le montant retenu, le code d'impact et le statut de contrôle sont SAISIS. Ce service ne dérive
aucune valeur économique : Lot4bis/Lot9 restent seuls responsables du calcul.
"""
from __future__ import annotations

import json
from datetime import datetime, timezone
from typing import Any

from app.contrats_donnees import ContratInvalideError, ReservationHH
from app.db.connection import get_db

STATUT_ACTIVE = "ACTIVE"
STATUT_ANNULEE = "ANNULEE"

EVT_CREATION = "CREATION"
EVT_MODIFICATION = "MODIFICATION"
EVT_ANNULATION = "ANNULATION"

E_CHAMP_MANQUANT = "RESHH_CHAMP_MANQUANT"
E_MOIS_INVALIDE = "RESHH_MOIS_INVALIDE"
E_DATES_INCOHERENTES = "RESHH_DATES_INCOHERENTES"
E_MONTANT_INVALIDE = "RESHH_MONTANT_INVALIDE"
E_INTROUVABLE = "RESHH_INTROUVABLE"
E_DEJA_ANNULEE = "RESHH_DEJA_ANNULEE"
E_CONTRAT_INVALIDE = "RESHH_CONTRAT_INVALIDE"

CHAMPS_SAISIE = (
    "mois", "canal_id", "source_financiere", "proprietaire_id", "logement_id",
    "reservation_id_hostaway", "date_arrivee", "date_depart", "nuits", "guest_count",
    "montant_percu", "montant_retenu", "mode_paiement_id", "code_impact", "impact_resultat_reel",
    "impact_resultat_comptable", "statut_controle", "niveau_anomalie", "code_anomalie",
    "commentaire",
)

# Table compagne `reservation_hh_overrides` (migration 0053) : dérogations menage/commission et
# montants conditionnels au mode de paiement. Voir la migration pour la sémantique exacte des
# champs `*_override` (NULL si absent ou égal au standard résolu — jamais forcé par ce service).
CHAMPS_OVERRIDES = (
    "menage", "menage_standard", "menage_standard_source", "menage_override", "motif_override_menage",
    "confirmation_override_menage", "taux_commission_standard", "taux_commission_standard_source",
    "taux_commission_override", "motif_override_taux_commission",
    "confirmation_override_taux_commission", "commentaire_taux_commission", "montant_recupere",
    "associe_id_recuperateur", "montant_reverse_proprietaire", "source_acompte_facture",
)

OBLIGATOIRES = ("mois", "logement_id", "date_arrivee", "date_depart")


def _maintenant() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _refus(code: str, message: str) -> dict[str, Any]:
    return {"ok": False, "code": code, "message": message}


def _nombre(valeur: Any) -> float | None:
    if valeur is None or valeur == "":
        return None
    try:
        return float(str(valeur).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def valider(donnees: dict[str, Any]) -> dict[str, Any]:
    """Validation métier (§17). Refus lisible, jamais d'exception."""
    for champ in OBLIGATOIRES:
        valeur = donnees.get(champ)
        if valeur is None or str(valeur).strip() == "":
            return _refus(E_CHAMP_MANQUANT, f"Champ obligatoire manquant : {champ}.")

    mois = str(donnees["mois"]).strip()
    if len(mois) != 7 or mois[4] != "-":
        return _refus(E_MOIS_INVALIDE, f"Mois attendu au format AAAA-MM, reçu : {mois!r}.")

    arrivee, depart = str(donnees["date_arrivee"]).strip(), str(donnees["date_depart"]).strip()
    if depart < arrivee:
        return _refus(E_DATES_INCOHERENTES,
                      f"Départ ({depart}) antérieur à l'arrivée ({arrivee}).")

    for champ in ("montant_percu", "montant_retenu"):
        brut = donnees.get(champ)
        if brut not in (None, "") and _nombre(brut) is None:
            return _refus(E_MONTANT_INVALIDE, f"{champ} doit être un nombre.")

    return {"ok": True, "mois": mois}


def _verifier_contrat(donnees: dict[str, Any]) -> dict[str, Any] | None:
    """Mission 11 : rejette une date_arrivee/date_depart non calendaire ou incohérente, un mois
    mal formé, un logement_id vide — avant `valider()`, seule la cohérence `depart >= arrivee` en
    tant que chaînes était vérifiée, jamais leur validité calendaire réelle (ex. "2026-02-30" aurait
    circulé jusqu'à Lot9/Lot10 sans être détecté). `montant_retenu` reste optionnel (contrat ajusté,
    cf. `ReservationHH` — un placeholder sans montant est un état réel valide)."""
    try:
        ReservationHH.from_dict(donnees)
    except ContratInvalideError as exc:
        return _refus(E_CONTRAT_INVALIDE, str(exc))
    return None


def _prochain_id(conn, mois: str) -> str:
    """`RESHH-{mois}-{NNN}` — le rang ne dépend que des réservations DÉJÀ enregistrées ce mois-là."""
    prefixe = f"RESHH-{mois}-"
    existants = [r[0] for r in conn.execute(
        "SELECT reservation_hh_id FROM reservations_hors_hostaway "
        "WHERE reservation_hh_id LIKE ?", (prefixe + "%",))]
    rangs = []
    for rid in existants:
        suffixe = rid[len(prefixe):]
        if suffixe.isdigit():
            rangs.append(int(suffixe))
    return f"{prefixe}{(max(rangs) + 1) if rangs else 1:03d}"


def _journaliser(conn, rid: str, evenement: str, acteur: str, motif: str,
                 avant: Any = None, apres: Any = None) -> None:
    conn.execute(
        "INSERT INTO reservation_hh_evenements (reservation_hh_id, evenement, acteur, motif, "
        "avant_json, apres_json) VALUES (?,?,?,?,?,?)",
        (rid, evenement, acteur or None, motif or None,
         json.dumps(avant, default=str) if avant else None,
         json.dumps(apres, default=str) if apres else None))


def _ligne(conn, rid: str) -> dict[str, Any] | None:
    r = conn.execute("SELECT * FROM reservations_hors_hostaway WHERE reservation_hh_id = ?",
                     (rid,)).fetchone()
    return dict(r) if r else None


def _overrides(conn, rid: str) -> dict[str, Any] | None:
    r = conn.execute("SELECT * FROM reservation_hh_overrides WHERE reservation_hh_id = ?",
                     (rid,)).fetchone()
    return dict(r) if r else None


def _upsert_overrides(conn, rid: str, donnees: dict[str, Any]) -> dict[str, Any]:
    """Insère ou met à jour la ligne compagne (1-1, migration 0053). Jamais de calcul métier ici :
    les valeurs sont celles déjà résolues par `saisie_hh_service.valider` (standard/override)."""
    valeurs = {c: donnees.get(c) for c in CHAMPS_OVERRIDES}
    for champ in ("menage", "menage_standard", "menage_override", "taux_commission_standard",
                  "taux_commission_override", "montant_recupere", "montant_reverse_proprietaire"):
        valeurs[champ] = _nombre(valeurs.get(champ))
    if _overrides(conn, rid) is None:
        colonnes = ["reservation_hh_id", *CHAMPS_OVERRIDES]
        conn.execute(
            f"INSERT INTO reservation_hh_overrides ({', '.join(colonnes)}) "
            f"VALUES ({', '.join(['?'] * len(colonnes))})",
            [rid, *(valeurs[c] for c in CHAMPS_OVERRIDES)])
    else:
        conn.execute(
            f"UPDATE reservation_hh_overrides SET "
            f"{', '.join(f'{c} = ?' for c in CHAMPS_OVERRIDES)} WHERE reservation_hh_id = ?",
            [*(valeurs[c] for c in CHAMPS_OVERRIDES), rid])
    return valeurs


def creer(donnees: dict[str, Any], *, acteur: str = "", db_path=None) -> dict[str, Any]:
    validation = valider(donnees)
    if not validation["ok"]:
        return validation
    contrat_refus = _verifier_contrat(donnees)
    if contrat_refus is not None:
        return contrat_refus

    valeurs = {c: donnees.get(c) for c in CHAMPS_SAISIE}
    valeurs["mois"] = validation["mois"]
    for champ in ("montant_percu", "montant_retenu"):
        valeurs[champ] = _nombre(valeurs.get(champ))

    conn = get_db(db_path)
    try:
        rid = str(donnees.get("reservation_hh_id") or "").strip() or _prochain_id(
            conn, validation["mois"])
        if _ligne(conn, rid) is not None:
            return _refus("RESHH_ID_DEJA_UTILISE", f"Identifiant déjà utilisé : {rid}.")
        colonnes = ["reservation_hh_id", *CHAMPS_SAISIE, "date_saisie", "source_module", "acteur"]
        params = [rid, *(valeurs[c] for c in CHAMPS_SAISIE), _maintenant(), "SAISIE_APP",
                  acteur or None]
        conn.execute(
            f"INSERT INTO reservations_hors_hostaway ({', '.join(colonnes)}) "
            f"VALUES ({', '.join(['?'] * len(colonnes))})", params)
        valeurs_overrides = _upsert_overrides(conn, rid, donnees)
        _journaliser(conn, rid, EVT_CREATION, acteur, "",
                     apres={**valeurs, **valeurs_overrides})
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "reservation_hh_id": rid}


def modifier(rid: str, donnees: dict[str, Any], *, acteur: str = "", motif: str = "",
             db_path=None) -> dict[str, Any]:
    validation = valider(donnees)
    if not validation["ok"]:
        return validation
    contrat_refus = _verifier_contrat(donnees)
    if contrat_refus is not None:
        return contrat_refus

    conn = get_db(db_path)
    try:
        avant = _ligne(conn, rid)
        if avant is None:
            return _refus(E_INTROUVABLE, f"Réservation inconnue : {rid}.")
        if avant["statut"] == STATUT_ANNULEE:
            return _refus(E_DEJA_ANNULEE, "Une réservation annulée ne se corrige pas.")
        avant_overrides = _overrides(conn, rid) or {}
        valeurs = {c: donnees.get(c) for c in CHAMPS_SAISIE}
        valeurs["mois"] = validation["mois"]
        for champ in ("montant_percu", "montant_retenu"):
            valeurs[champ] = _nombre(valeurs.get(champ))
        conn.execute(
            f"UPDATE reservations_hors_hostaway SET "
            f"{', '.join(f'{c} = ?' for c in CHAMPS_SAISIE)}, date_modification = ? "
            "WHERE reservation_hh_id = ?",
            [*(valeurs[c] for c in CHAMPS_SAISIE), _maintenant(), rid])
        valeurs_overrides = _upsert_overrides(conn, rid, donnees)
        _journaliser(conn, rid, EVT_MODIFICATION, acteur, motif,
                     avant={**{c: avant.get(c) for c in CHAMPS_SAISIE}, **avant_overrides},
                     apres={**valeurs, **valeurs_overrides})
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "reservation_hh_id": rid}


def annuler(rid: str, *, acteur: str = "", motif: str = "", db_path=None) -> dict[str, Any]:
    conn = get_db(db_path)
    try:
        avant = _ligne(conn, rid)
        if avant is None:
            return _refus(E_INTROUVABLE, f"Réservation inconnue : {rid}.")
        if avant["statut"] == STATUT_ANNULEE:
            return _refus(E_DEJA_ANNULEE, f"{rid} est déjà annulée.")
        conn.execute(
            "UPDATE reservations_hors_hostaway SET statut = ?, date_modification = ? "
            "WHERE reservation_hh_id = ?", (STATUT_ANNULEE, _maintenant(), rid))
        _journaliser(conn, rid, EVT_ANNULATION, acteur, motif,
                     avant={"statut": STATUT_ACTIVE}, apres={"statut": STATUT_ANNULEE})
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "reservation_hh_id": rid, "statut": STATUT_ANNULEE}


def historique(rid: str, *, db_path=None) -> list[dict[str, Any]]:
    conn = get_db(db_path)
    try:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM reservation_hh_evenements WHERE reservation_hh_id = ? ORDER BY id",
            (rid,))]
    finally:
        conn.close()


# ── Reprise contrôlée depuis le master historique (§10/§19) ─────────────────────────────────────

def reprendre_master(chemin_master, *, acteur: str = "REPRISE", db_path=None) -> dict[str, Any]:
    """Importe UNE FOIS les réservations du master Excel historique vers SQLite.

    Outil de bootstrap/reprise (§10), pas un chemin runtime : après l'import, SQLite est canonique
    et le classeur n'est plus nécessaire. Idempotent — une réservation déjà présente est ignorée,
    jamais dupliquée ni écrasée (une correction passe par `modifier`, qui laisse une trace).
    """
    from pathlib import Path

    import openpyxl

    chemin = Path(chemin_master)
    if not chemin.exists():
        return _refus("MASTER_ABSENT", f"Master introuvable : {chemin.name}")

    wb = openpyxl.load_workbook(str(chemin), read_only=True, data_only=True)
    try:
        if "MASTER" not in wb.sheetnames:
            return _refus("ONGLET_ABSENT", "Onglet MASTER absent du classeur.")
        lignes = list(wb["MASTER"].iter_rows(values_only=True))
    finally:
        wb.close()
    if len(lignes) <= 1:
        return {"ok": True, "nb_importees": 0, "nb_ignorees": 0}

    entetes = [str(c) if c is not None else "" for c in lignes[0]]
    brutes = [dict(zip(entetes, r)) for r in lignes[1:]]
    # Les lignes de gabarit Power Query ne sont pas des réservations.
    reelles = [r for r in brutes
               if str(r.get("reservation_hh_id") or "").strip().startswith("RESHH-")]

    _ALIAS = {"guestCount": "guest_count", "ROW_HASH": "row_hash"}
    importees, ignorees = 0, 0
    conn = get_db(db_path)
    try:
        colonnes_table = {c[1] for c in conn.execute(
            "PRAGMA table_info(reservations_hors_hostaway)")}
        for brute in reelles:
            rid = str(brute["reservation_hh_id"]).strip()
            if _ligne(conn, rid) is not None:
                ignorees += 1
                continue
            valeurs = {}
            for cle, valeur in brute.items():
                colonne = _ALIAS.get(cle, cle)
                if colonne in colonnes_table and colonne != "reservation_hh_id":
                    valeurs[colonne] = valeur
            colonnes = ["reservation_hh_id", *valeurs.keys(), "source_module", "acteur"]
            params = [rid, *valeurs.values(), "REPRISE_MASTER", acteur]
            conn.execute(
                f"INSERT INTO reservations_hors_hostaway ({', '.join(colonnes)}) "
                f"VALUES ({', '.join(['?'] * len(colonnes))})", params)
            _journaliser(conn, rid, EVT_CREATION, acteur, "reprise du master historique",
                         apres=valeurs)
            importees += 1
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "nb_importees": importees, "nb_ignorees": ignorees}
