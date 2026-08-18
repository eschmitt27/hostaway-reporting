"""Adaptateur d'entrée CleaningTasks : master legacy → couche RAW SQLite (`hostaway_cleaning_tasks`).

Même rôle que `hostaway_adaptateurs.py` pour les réservations : une REPRISE, pas une source. Elle
permet d'établir la parité et de peupler une installation neuve sans relancer H6 en réel — le dataset
legacy (`MASTER_FACT_HA_CleaningTasks_Discovery.xlsx`, onglet `data`) date, et l'API a déjà essuyé des
429 sur cet endpoint. Une fois le chemin API → SQLite direct en service pour CleaningTasks, ce module
ne sert plus qu'à comparer.

Lit l'onglet `data` (tâches brutes), pas `MASTER_ENRICHI` : la RAW ne reprend que ce que l'API a
renvoyé, comme pour les réservations — la résolution logement/mois reste le travail de Lot6a en aval.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import app.config as cfg

DOSSIER_LOT1 = ("02_TRAVAIL", "Lot1_Hostaway")
MASTER = "MASTER_FACT_HA_CleaningTasks_Discovery.xlsx"
ONGLET = "data"

E_MASTER_ABSENT = "CLEANING_TASKS_MASTER_ABSENT"


def _racine(racine: Path | None = None) -> Path:
    return Path(racine) if racine is not None else Path(cfg.PROJECT_ROOT)


def chemin_master(racine: Path | None = None) -> Path:
    return _racine(racine).joinpath(*DOSSIER_LOT1, MASTER)


def master_disponible(racine: Path | None = None) -> bool:
    return chemin_master(racine).exists()


def depuis_master_excel(racine: Path | None = None) -> list[dict[str, Any]]:
    import openpyxl

    chemin = chemin_master(racine)
    if not chemin.exists():
        return []
    wb = openpyxl.load_workbook(str(chemin), read_only=True, data_only=True)
    try:
        if ONGLET not in wb.sheetnames:
            return []
        lignes = list(wb[ONGLET].iter_rows(values_only=True))
    finally:
        wb.close()
    if len(lignes) <= 1:
        return []
    entetes = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(lignes[0])]
    return [dict(zip(entetes, r)) for r in lignes[1:]
            if any(v is not None and str(v).strip() for v in r)]


def reprendre(*, racine: Path | None = None, db_path=None) -> dict[str, Any]:
    from app.services import hostaway_cleaning_tasks_raw_service as raw

    if not master_disponible(racine):
        return {"ok": False, "code": E_MASTER_ABSENT,
                "message": "Master CleaningTasks introuvable : la reprise n'a rien à lire."}

    lignes = depuis_master_excel(racine)
    extraction_id = raw.ouvrir(mode=raw.MODE_REPRISE_EXCEL, db_path=db_path)
    if not extraction_id:
        return {"ok": False, "code": "MIGRATION_ABSENTE",
                "message": "Table hostaway_cleaning_tasks_extractions absente (migration 0035 non "
                           "appliquée)."}
    try:
        raw.enregistrer(extraction_id, taches=lignes, db_path=db_path)
    except Exception as exc:
        raw.cloturer(extraction_id, statut=raw.ST_ECHEC, message=f"{type(exc).__name__}: {exc}",
                     db_path=db_path)
        raise
    resultat = raw.cloturer(extraction_id, statut=raw.ST_SUCCES, db_path=db_path)
    return {"ok": True, **resultat}
