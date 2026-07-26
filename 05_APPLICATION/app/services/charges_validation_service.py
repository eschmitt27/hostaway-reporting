"""Validation humaine des charges (A_CONTROLER → VALIDE / REJETE).

Pourquoi ce parcours existe : une charge saisie est créée en `A_CONTROLER`. Lot9 n'ingère QUE les
charges `VALIDE` ; une charge non validée n'a donc **aucun impact financier** (ni résultat, ni net
propriétaire, ni préfacture). La validation était jusqu'ici simulée en modifiant le MASTER à la
main : ce module la rend applicative, tracée et sûre.

Principes :
- la vérité reste la SAISIE (`SAISIE_Charges_Flux`, onglet SAISIE) ; le MASTER est régénéré par Lot3 ;
- écriture par remplacement atomique passant par le write-guard (refus hors `data_recette` en mode
  recette) ; les flags `CHARGES_REAL_WRITE_*` restent requis ;
- chaque action est journalisée (qui, quand, quoi, commentaire) et jamais destructive : le rejet
  conserve la ligne, il change seulement son statut.
"""
from __future__ import annotations

import os
import shutil
import sqlite3
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

import app.config as cfg

SHEET = "SAISIE"
ST_A_CONTROLER = "A_CONTROLER"
ST_VALIDE = "VALIDE"
ST_REJETE = "REJETE"

E_FLAGS = "E_FLAGS_DESACTIVES"
E_INTROUVABLE = "E_CHARGE_INTROUVABLE"
E_DEJA = "E_DEJA_TRAITEE"
E_COMMENTAIRE = "E_COMMENTAIRE_OBLIGATOIRE"
E_ECRITURE = "E_ECRITURE_REFUSEE"

MESSAGES = {
    E_FLAGS: "Écriture désactivée sur cette installation : la validation est impossible.",
    E_INTROUVABLE: "Charge introuvable dans la saisie.",
    E_DEJA: "Cette charge a déjà été traitée (validée ou rejetée).",
    E_COMMENTAIRE: "Un commentaire est obligatoire pour rejeter une charge.",
    E_ECRITURE: "Écriture refusée : la cible n'est pas autorisée.",
}

AVERTISSEMENT_NON_VALIDEE = (
    "Cette charge n'aura aucun impact financier tant qu'elle ne sera pas validée."
)


def _flags_actifs() -> bool:
    return bool(cfg.CHARGES_REAL_WRITE_ENABLED and cfg.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _saisie_path() -> Path:
    return Path(cfg.SAISIE_CHARGES)


def _journal_init(conn) -> None:
    conn.execute(
        """CREATE TABLE IF NOT EXISTS charges_validation_journal (
               id INTEGER PRIMARY KEY AUTOINCREMENT,
               charge_id TEXT NOT NULL,
               ancien_statut TEXT,
               nouveau_statut TEXT NOT NULL,
               acteur TEXT,
               commentaire TEXT,
               horodatage_utc TEXT NOT NULL
           )"""
    )


def _journaliser(charge_id: str, ancien: str, nouveau: str, acteur: str, commentaire: str,
                 db_path=None) -> None:
    conn = sqlite3.connect(str(db_path or cfg.DB_PATH))
    try:
        _journal_init(conn)
        conn.execute(
            "INSERT INTO charges_validation_journal "
            "(charge_id, ancien_statut, nouveau_statut, acteur, commentaire, horodatage_utc) "
            "VALUES (?,?,?,?,?,?)",
            (charge_id, ancien, nouveau, acteur or "local", commentaire or "", _now()),
        )
        conn.commit()
    finally:
        conn.close()


def historique(charge_id: str = "", db_path=None) -> list[dict[str, Any]]:
    conn = sqlite3.connect(str(db_path or cfg.DB_PATH))
    conn.row_factory = sqlite3.Row
    try:
        _journal_init(conn)
        if charge_id:
            rows = conn.execute(
                "SELECT * FROM charges_validation_journal WHERE charge_id=? ORDER BY id DESC",
                (charge_id,)).fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM charges_validation_journal ORDER BY id DESC LIMIT 200").fetchall()
        return [dict(r) for r in rows]
    finally:
        conn.close()


def lister(statut: str = "", refacturable: str = "", sans_proprietaire: bool = False,
           code_impact: str = "") -> dict[str, Any]:
    """Liste les charges de la SAISIE, avec filtres. Lecture seule, jamais d'écriture."""
    path = _saisie_path()
    if not path.exists():
        return {"status": "SOURCE_ABSENTE", "source": path.name, "rows": [], "compteurs": {}}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    data = list(ws.iter_rows(values_only=True))
    wb.close()
    if not data:
        return {"status": "VIDE", "source": path.name, "rows": [], "compteurs": {}}
    hdr = list(data[0])
    rows = []
    for r in data[1:]:
        d = dict(zip(hdr, r))
        cid = str(d.get("charge_id") or "").strip()
        if not cid or cid[0] in "#[<←-*":
            continue
        rows.append(d)

    compteurs = {
        "total": len(rows),
        ST_A_CONTROLER: sum(1 for d in rows if str(d.get("statut_controle")) == ST_A_CONTROLER),
        ST_VALIDE: sum(1 for d in rows if str(d.get("statut_controle")) == ST_VALIDE),
        ST_REJETE: sum(1 for d in rows if str(d.get("statut_controle")) == ST_REJETE),
        "refacturables": sum(1 for d in rows if str(d.get("refacturable")).upper() == "OUI"),
        "sans_proprietaire": sum(1 for d in rows if not str(d.get("proprietaire_id") or "").strip()),
    }

    def _match(d: dict) -> bool:
        if statut and str(d.get("statut_controle")) != statut:
            return False
        if refacturable and str(d.get("refacturable")).upper() != refacturable.upper():
            return False
        if sans_proprietaire and str(d.get("proprietaire_id") or "").strip():
            return False
        if code_impact and str(d.get("code_impact")) != code_impact:
            return False
        return True

    filtres = [d for d in rows if _match(d)]
    return {"status": "OK", "source": path.name, "rows": filtres, "compteurs": compteurs,
            "avertissement": AVERTISSEMENT_NON_VALIDEE,
            "applied": {"statut": statut, "refacturable": refacturable,
                        "sans_proprietaire": sans_proprietaire, "code_impact": code_impact}}


def _refus(code: str, details: str = "") -> dict[str, Any]:
    return {"ok": False, "code": code, "message": MESSAGES.get(code, code), "details": details}


def _transition(charge_id: str, nouveau: str, *, acteur: str = "", commentaire: str = "",
                db_path=None) -> dict[str, Any]:
    charge_id = str(charge_id or "").strip()
    if nouveau == ST_REJETE and not str(commentaire or "").strip():
        return _refus(E_COMMENTAIRE)
    if not _flags_actifs():
        return _refus(E_FLAGS)

    path = _saisie_path()
    if not path.exists():
        return _refus(E_INTROUVABLE, path.name)

    wb = openpyxl.load_workbook(path)          # avec formules : jamais data_only en écriture
    try:
        ws = wb[SHEET]
        hdr = [c.value for c in ws[1]]
        if "charge_id" not in hdr or "statut_controle" not in hdr:
            return _refus(E_INTROUVABLE, "colonnes manquantes")
        ci, si = hdr.index("charge_id"), hdr.index("statut_controle")
        cible = None
        for row in ws.iter_rows(min_row=2):
            if str(row[ci].value or "").strip() == charge_id:
                cible = row
                break
        if cible is None:
            return _refus(E_INTROUVABLE, charge_id)
        ancien = str(cible[si].value or "").strip()
        if ancien in (ST_VALIDE, ST_REJETE):
            return _refus(E_DEJA, f"statut actuel {ancien}")
        cible[si].value = nouveau

        # Écriture atomique passant par le write-guard (barrière de dernière ligne).
        from app.services.saisie_charges_transaction_service import _remplacer_fichier
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(path.parent))
        os.close(fd)
        tmp_path = Path(tmp)
        try:
            wb.save(tmp_path)
            _remplacer_fichier(tmp_path, path)
        except Exception as exc:            # guard ou disque : aucune modification conservée
            if tmp_path.exists():
                tmp_path.unlink(missing_ok=True)
            return _refus(E_ECRITURE, f"{type(exc).__name__}: {exc}")
    finally:
        wb.close()

    _journaliser(charge_id, ancien, nouveau, acteur, commentaire, db_path=db_path)
    return {"ok": True, "charge_id": charge_id, "ancien_statut": ancien, "nouveau_statut": nouveau,
            "horodatage_utc": _now(), "acteur": acteur or "local", "commentaire": commentaire or "",
            "recalcul_requis": nouveau == ST_VALIDE}


def valider(charge_id: str, *, acteur: str = "", commentaire: str = "", db_path=None):
    """A_CONTROLER → VALIDE. La charge devient éligible à Lot9 (impact financier)."""
    return _transition(charge_id, ST_VALIDE, acteur=acteur, commentaire=commentaire, db_path=db_path)


def rejeter(charge_id: str, *, commentaire: str, acteur: str = "", db_path=None):
    """A_CONTROLER → REJETE. Ligne conservée, jamais supprimée ; jamais ingérée par Lot9."""
    return _transition(charge_id, ST_REJETE, acteur=acteur, commentaire=commentaire, db_path=db_path)
