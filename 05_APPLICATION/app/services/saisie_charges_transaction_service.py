"""APP-3b / commit 3 — Orchestrateur TRANSACTIONNEL de l'écriture d'une charge.

Deux classeurs sont concernés par une charge, et `os.replace` n'est atomique que **par fichier** :
il n'existe aucune atomicité entre deux fichiers. L'atomicité est donc reconstruite au niveau
applicatif — préparer tout, ne rien remplacer tant que tout n'est pas prêt, et restaurer ce qui a
déjà été remplacé si la fin échoue.

    SAISIE_Charges_Impacts.xlsx   (ordre 1)  — AFFECTATIONS / MENAGE / RESERVE
    SAISIE_Charges_Flux.xlsx      (ordre 2)  — la ligne de charge

**L'ordre n'est pas arbitraire.** Les impacts sont remplacés d'abord : si le second remplacement
échoue et que le rollback échoue lui aussi, on laisse des impacts *orphelins* (aucune charge ne les
référence, donc invisibles pour les lots) plutôt qu'une charge *sans impacts* — qui serait, elle,
comptée à tort par le résultat. Le pire cas est choisi, pas subi.

Lot7 (`MASTER_FACT_MAN_IK_Avantages`) est **hors transaction** : `file_registry` interdit à l'app
d'écrire dans `02_TRAVAIL` / `MASTER_*`. L'avantage associé est porté par la colonne
`avantage_associe_id` de la ligne de charge ; Lot7 se régénère ensuite par le moteur.

Séquence (chaque étape est une fonction distincte, monkeypatchable) :
    1. garde des flags          — refus AVANT tout verrou et toute préparation
    2. validation syntaxique de la demande
    3. ACQUISITION DU VERROU    — interprocessus, avant toute lecture concluant à l'unicité
    4. résolution du charge_id   — vérifié (ou généré) SOUS verrou, jamais avant
    5. préparation des temporaires (writer bas niveau — aucun remplacement)
    6. validation des temporaires (existence, lisibilité, empreinte, même volume)
    7. sauvegardes techniques    — créées AVANT le premier remplacement
    8. commit                    — os.replace, dans l'ordre déterministe (SEUL endroit du projet)
    9. vérification post-commit  — fichiers finaux lisibles et conformes aux temporaires
   10. rollback si échec         — restauration vérifiée par SHA256
   11. nettoyage                 — temporaires puis sauvegardes ; jamais avant la validation complète
   12. LIBÉRATION DU VERROU      — dans un finally, quoi qu'il arrive

**Le verrou est la réservation.** Il n'existe pas de registre séparé d'identifiants : l'unicité d'un
`charge_id` tient au fait que sa vérification (ou sa génération) et l'écriture qui en découle se font
sans jamais relâcher le verrou entre les deux. Sans lui, deux processus liraient tous deux « cet
identifiant est libre » avant que l'un des deux n'écrive — et aucune relecture ne rattraperait ça.

Garde-fous : tant que `CHARGES_REAL_WRITE_ENABLED` **et**
`CHARGES_REAL_WRITE_CONFIRMATION_ENABLED` ne sont pas tous deux à True, la transaction est refusée
avant qu'un seul octet d'un fichier réel ne soit touché, et **sans même poser de verrou**.
L'orchestrateur ne contourne jamais ces flags.
"""
from __future__ import annotations

import os
import re
import shutil
import uuid
from dataclasses import dataclass, field, replace as dataclass_replace
from pathlib import Path
from typing import Any, Callable

import openpyxl

import app.config as cfg
from app.services import saisie_charges_journal_service as journal
from app.services import saisie_charges_lock_service as lock
from app.writers import saisie_charges_writer as writer
from app.writers.saisie_hh_writer import _same_volume, _sha256

# Ordre déterministe des remplacements (voir docstring : impacts avant charge).
ORDRE_IMPACTS = 1
ORDRE_FLUX = 2

STATUT_OK = "OK"
STATUT_GARDE = "GARDE_SECURITE"
STATUT_ERREUR = "ERREUR"
STATUT_ROLLBACK = "ROLLBACK"

E_FLAGS_DESACTIVES = "E_FLAGS_DESACTIVES"
E_DEMANDE_INVALIDE = "E_DEMANDE_INVALIDE"
E_TOKEN_DEJA_ECRIT = "E_TOKEN_DEJA_ECRIT"
E_IDEMPOTENCE_INDISPONIBLE = "E_IDEMPOTENCE_INDISPONIBLE"
E_VERROU_DEJA_PRIS = "E_VERROU_DEJA_PRIS"
E_CHARGE_ID_INVALIDE = "E_CHARGE_ID_INVALIDE"
E_CHARGE_ID_EXISTANT = "E_CHARGE_ID_EXISTANT"
E_CHARGE_ID_EPUISE = "E_CHARGE_ID_EPUISE"
E_PREPARATION = "E_PREPARATION"
E_TEMPORAIRE_MANQUANT = "E_TEMPORAIRE_MANQUANT"
E_TEMPORAIRE_INVALIDE = "E_TEMPORAIRE_INVALIDE"
E_SAUVEGARDE = "E_SAUVEGARDE"
E_REMPLACEMENT = "E_REMPLACEMENT"
E_POST_COMMIT = "E_POST_COMMIT"

# CHG-AAAA-MM-<IMPACT>-<ASSOC_MODE>-NNN (format verrouillé, cf. generate_charge_id).
CHARGE_ID_RE = re.compile(r"^CHG-\d{4}-\d{2}-(IC|HC|HR)-[A-Z0-9_]+-\d{3}$")
SEQUENCE_MAX = 999


class ChargeIdInvalideError(ValueError):
    """Le charge_id ne respecte pas le format attendu."""


class ChargeIdDejaExistantError(RuntimeError):
    """Le charge_id est déjà utilisé. Détecté SOUS verrou, avant toute préparation."""

    def __init__(self, charge_id: str, emplacements: list[str]) -> None:
        self.charge_id = charge_id
        self.emplacements = emplacements
        super().__init__(
            f"charge_id={charge_id} déjà présent dans : {', '.join(emplacements)}. "
            f"Aucun fichier n'a été préparé."
        )


class RollbackCritiqueError(RuntimeError):
    """Le rollback n'a pas pu restaurer tous les fichiers déjà remplacés.

    État NON récupérable automatiquement. Les sauvegardes des fichiers non restaurés sont
    CONSERVÉES et listées ici : elles sont la seule voie de retour. Ne jamais les supprimer.
    """

    def __init__(self, fichiers_non_restaures: list[dict[str, str]], cause: str) -> None:
        self.fichiers_non_restaures = fichiers_non_restaures
        self.cause = cause
        # Renseigné par la journalisation : si la trace n'a pas pu être écrite, on le dit ici plutôt
        # que de le perdre (l'incident critique prime, mais son absence de trace doit se voir).
        self.journal_erreur: str | None = None
        details = "; ".join(
            f"{f['cible']} (sauvegarde conservée : {f['sauvegarde']})"
            for f in fichiers_non_restaures
        )
        super().__init__(
            f"ROLLBACK INCOMPLET — {len(fichiers_non_restaures)} fichier(s) non restauré(s) : "
            f"{details}. Cause initiale : {cause}"
        )


# ── Objets de transaction ────────────────────────────────────────────────────

@dataclass(frozen=True)
class DemandeEcritureCharge:
    """Demande déjà structurée (validée en amont par la prévisualisation).

    Deux modes d'identifiant :

    - **fourni** : `charge_id` renseigné, avec `row_data` et `persistable` cohérents. L'orchestrateur
      vérifie sa syntaxe et son absence de toutes les tables — **sous verrou**, avant préparation.
    - **généré** : `charge_id` vide, `charge_id_prefixe` et `payload_factory` renseignés.
      L'identifiant est choisi **sous verrou**, puis `payload_factory(charge_id)` construit le
      payload. La fabrique est fournie par l'appelant (ses propres constructeurs) : l'orchestrateur
      ne duplique aucune logique métier.
    """

    montant: float
    target_row: int
    saisie_path: Path
    impacts_path: Path
    charge_id: str = ""
    row_data: dict[str, Any] | None = None
    persistable: dict[str, Any] | None = None
    charge_id_prefixe: str | None = None
    payload_factory: Callable[[str], tuple[dict[str, Any], dict[str, Any]]] | None = None
    master_path: Path | None = None
    lock_path: Path | None = None
    sha256_saisie_attendu: str | None = None
    sha256_impacts_attendu: str | None = None
    # Journalisation : token du dry-run (clé d'idempotence) et base cible (injectable en test).
    token_previsualisation: str | None = None
    db_path: Path | None = None


@dataclass
class FichierPrepare:
    cible: Path
    temp: Path
    sha256_cible_avant: str
    sha256_temp: str
    ordre: int
    libelle: str


@dataclass
class FichierRemplace:
    cible: Path
    sauvegarde: Path
    sha256_avant: str
    sha256_apres: str


@dataclass
class ResultatTransaction:
    statut: str
    charge_id: str
    code: str | None = None
    details: str | None = None
    fichiers_remplaces: list[str] = field(default_factory=list)
    sauvegardes_restantes: list[str] = field(default_factory=list)
    nettoyage_incomplet: list[str] = field(default_factory=list)
    sha256_finaux: dict[str, str] = field(default_factory=dict)
    # ── Journalisation ──
    transaction_id: str | None = None
    statut_journal: str | None = None          # verdict métier écrit au journal (9 valeurs)
    journal_erreur: str | None = None          # panne du journal : signalée, jamais avalée
    rollback_tente: bool = False
    rollback_reussi: bool | None = None
    verrou_pid: int | None = None              # détenteur du verrou en cas de refus
    verrou_hostname: str | None = None

    @property
    def ok(self) -> bool:
        return self.statut == STATUT_OK


# ── 1. Garde des flags ───────────────────────────────────────────────────────

def _garde_flags() -> ResultatTransaction | None:
    """Refus AVANT toute préparation. Les deux flags sont requis (double garde-fou)."""
    manquants = []
    if not cfg.CHARGES_REAL_WRITE_ENABLED:
        manquants.append("CHARGES_REAL_WRITE_ENABLED")
    if not cfg.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED:
        manquants.append("CHARGES_REAL_WRITE_CONFIRMATION_ENABLED")
    if not manquants:
        return None
    return ResultatTransaction(
        statut=STATUT_GARDE,
        charge_id="",
        code=E_FLAGS_DESACTIVES,
        details=(
            f"Écriture réelle interdite — flag(s) désactivé(s) : {', '.join(manquants)}. "
            f"Aucun fichier n'a été touché."
        ),
    )


# ── 2. Validation de la demande ──────────────────────────────────────────────

def _valider_demande(d: DemandeEcritureCharge) -> str | None:
    """Validation PUREMENT syntaxique — aucune lecture de classeur, donc autorisée avant le verrou.

    Tout ce qui conclut à la disponibilité d'un charge_id se fait plus tard, sous verrou.
    """
    if float(d.montant) <= 0:
        return f"montant invalide : {d.montant}."
    if int(d.target_row) < 2:
        return f"target_row invalide : {d.target_row} (la ligne 1 est l'en-tête)."

    if str(d.charge_id or "").strip():
        if d.row_data is None or d.persistable is None:
            return "charge_id fourni : row_data et persistable sont obligatoires."
        return None

    # Mode génération : la fabrique de payload est indispensable (aucune logique métier ici).
    if not d.charge_id_prefixe:
        return "charge_id absent : charge_id_prefixe est obligatoire pour le générer."
    if d.payload_factory is None:
        return "charge_id absent : payload_factory est obligatoire pour construire le payload."
    return None


def _identifiants_existants(d: DemandeEcritureCharge) -> dict[str, set[str]]:
    """Lit les charge_id présents dans TOUTES les tables concernées. **Appelé sous verrou.**

    Lecture seule : SAISIE_Charges_Flux (colonne A), SAISIE_Charges_Impacts (3 onglets) et
    MASTER_FACT_MAN_Charges (sortie calculée Lot3 — un identifiant déjà consommé en aval reste
    une collision).
    """
    par_source: dict[str, set[str]] = {}

    par_source["SAISIE_Charges_Flux"] = _ids_colonne(d.saisie_path, {"SAISIE": "charge_id"})
    par_source["SAISIE_Charges_Impacts"] = _ids_colonne(
        d.impacts_path,
        {"AFFECTATIONS": "charge_id", "MENAGE": "charge_id", "RESERVE_REFACTURATION": "charge_id"},
    )
    master = Path(d.master_path) if d.master_path is not None else Path(cfg.MASTER_CHARGES)
    if master.exists():
        par_source["MASTER_FACT_MAN_Charges"] = _ids_colonne(master, {"MASTER": "charge_id"})
    return par_source


def _ids_colonne(path: Path, onglets: dict[str, str]) -> set[str]:
    """charge_id non vides d'un classeur (lignes gabarit/placeholder ignorées)."""
    ids: set[str] = set()
    p = Path(path)
    if not p.exists():
        return ids
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    try:
        for onglet, colonne in onglets.items():
            if onglet not in wb.sheetnames:
                continue
            lignes = wb[onglet].iter_rows(values_only=True)
            entetes = next(lignes, None)
            if not entetes:
                continue
            noms = [str(h).strip() if h is not None else "" for h in entetes]
            if colonne not in noms:
                continue
            idx = noms.index(colonne)
            for r in lignes:
                if idx >= len(r):
                    continue
                cid = str(r[idx] or "").strip()
                if cid and not cid.startswith("["):      # placeholder Lot3, jamais une donnée
                    ids.add(cid)
    finally:
        wb.close()
    return ids


def _resoudre_charge_id(
    d: DemandeEcritureCharge,
) -> tuple[str, dict[str, Any], dict[str, Any]]:
    """Décide l'identifiant définitif. **À n'appeler que le verrou tenu.**

    Lève ChargeIdInvalideError / ChargeIdDejaExistantError. Aucun temporaire n'est préparé avant
    que cette fonction n'ait rendu son verdict.
    """
    existants = _identifiants_existants(d)
    tous = set().union(*existants.values()) if existants else set()

    fourni = str(d.charge_id or "").strip()
    if fourni:
        if not CHARGE_ID_RE.match(fourni):
            raise ChargeIdInvalideError(
                f"charge_id={fourni!r} : format attendu CHG-AAAA-MM-<IC|HC|HR>-<MODE>-NNN."
            )
        emplacements = [source for source, ids in existants.items() if fourni in ids]
        if emplacements:
            raise ChargeIdDejaExistantError(fourni, emplacements)

        row_data, persistable = d.row_data or {}, d.persistable or {}
        if str(row_data.get("charge_id") or "").strip() != fourni:
            raise ChargeIdInvalideError("row_data.charge_id incohérent avec la demande.")
        if str(persistable.get("charge_id") or "").strip() != fourni:
            raise ChargeIdInvalideError("persistable.charge_id incohérent avec la demande.")
        return fourni, row_data, persistable

    # ── Génération sous verrou : le premier numéro libre de TOUTES les tables ──
    prefixe = str(d.charge_id_prefixe).strip().rstrip("-")
    genere = None
    for n in range(1, SEQUENCE_MAX + 1):
        candidat = f"{prefixe}-{n:03d}"
        if candidat not in tous:
            genere = candidat
            break
    if genere is None:
        raise ChargeIdDejaExistantError(f"{prefixe}-NNN", [f"séquence saturée (>{SEQUENCE_MAX})"])
    if not CHARGE_ID_RE.match(genere):
        raise ChargeIdInvalideError(
            f"charge_id généré {genere!r} invalide — préfixe {prefixe!r} non conforme."
        )

    row_data, persistable = d.payload_factory(genere)    # fabrique de l'appelant
    if str(row_data.get("charge_id") or "").strip() != genere:
        raise ChargeIdInvalideError("payload_factory : row_data.charge_id incohérent.")
    if str(persistable.get("charge_id") or "").strip() != genere:
        raise ChargeIdInvalideError("payload_factory : persistable.charge_id incohérent.")
    return genere, row_data, persistable


# ── 3. Préparation des temporaires (writer bas niveau) ───────────────────────

def _preparer(d: DemandeEcritureCharge) -> tuple[list[FichierPrepare], ResultatTransaction | None]:
    """Prépare TOUS les temporaires. Aucun fichier réel n'est modifié ici.

    Si une préparation échoue, les temporaires déjà produits sont nettoyés et rien n'est remplacé.
    """
    prepares: list[FichierPrepare] = []

    res_impacts = writer.prepare_charge_impacts_write(
        d.impacts_path, d.persistable, d.montant, sha256_attendu=d.sha256_impacts_attendu
    )
    if res_impacts["statut"] != STATUT_OK:
        return [], ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=d.charge_id, code=E_PREPARATION,
            details=f"Impacts : [{res_impacts['code']}] {res_impacts['details']}",
        )
    prepares.append(FichierPrepare(
        cible=Path(d.impacts_path), temp=Path(res_impacts["temp_path"]),
        sha256_cible_avant=res_impacts["sha256_source"], sha256_temp=res_impacts["sha256_temp"],
        ordre=ORDRE_IMPACTS, libelle="SAISIE_Charges_Impacts",
    ))

    res_flux = writer.prepare_charge_flux_write(
        d.saisie_path, d.row_data, d.target_row, sha256_attendu=d.sha256_saisie_attendu
    )
    if res_flux["statut"] != STATUT_OK:
        _nettoyer([p.temp for p in prepares])       # aucun remplacement n'a eu lieu
        return [], ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=d.charge_id, code=E_PREPARATION,
            details=f"Charge : [{res_flux['code']}] {res_flux['details']}",
        )
    prepares.append(FichierPrepare(
        cible=Path(d.saisie_path), temp=Path(res_flux["temp_path"]),
        sha256_cible_avant=res_flux["sha256_source"], sha256_temp=res_flux["sha256_temp"],
        ordre=ORDRE_FLUX, libelle="SAISIE_Charges_Flux",
    ))
    return prepares, None


# ── 4. Validation des temporaires ────────────────────────────────────────────

def _valider_temporaires(prepares: list[FichierPrepare]) -> tuple[str, str] | None:
    """(code, details) si un temporaire est inutilisable. Aucun remplacement ne doit suivre."""
    for p in prepares:
        if not p.temp.exists():
            return E_TEMPORAIRE_MANQUANT, f"{p.libelle} : temporaire absent ({p.temp})."
        if p.temp == p.cible:
            return E_TEMPORAIRE_INVALIDE, f"{p.libelle} : le temporaire est le fichier cible."
        if not _same_volume(p.temp, p.cible):
            return E_TEMPORAIRE_INVALIDE, (
                f"{p.libelle} : temporaire sur un autre volume que la cible — "
                f"remplacement atomique impossible."
            )
        if _sha256(p.temp) != p.sha256_temp:
            return E_TEMPORAIRE_INVALIDE, (
                f"{p.libelle} : le temporaire a changé depuis sa préparation "
                f"(empreinte différente)."
            )
        if not _lisible(p.temp):
            return E_TEMPORAIRE_INVALIDE, f"{p.libelle} : temporaire illisible (classeur corrompu)."
        if _sha256(p.cible) != p.sha256_cible_avant:
            return E_TEMPORAIRE_INVALIDE, (
                f"{p.libelle} : la cible a changé depuis la préparation — base périmée."
            )
    return None


def _lisible(path: Path) -> bool:
    """Le classeur s'ouvre sans réparation (aucun Excel/LibreOffice : openpyxl seul)."""
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            return bool(wb.sheetnames)
        finally:
            wb.close()
    except Exception:
        return False


# ── 5. Sauvegardes techniques ────────────────────────────────────────────────

def _sauvegarder(prepares: list[FichierPrepare], jeton: str) -> tuple[dict[Path, Path], str | None]:
    """Copie chaque cible AVANT le premier remplacement. Même dossier = même volume."""
    sauvegardes: dict[Path, Path] = {}
    for p in prepares:
        bak = p.cible.parent / f"{p.cible.stem}.{jeton}.bak{p.cible.suffix}"
        try:
            shutil.copy2(str(p.cible), str(bak))
        except Exception as exc:
            _nettoyer(list(sauvegardes.values()))
            return {}, f"{p.libelle} : sauvegarde impossible ({exc})."
        sauvegardes[p.cible] = bak
    return sauvegardes, None


# ── 6/7/8. Commit, vérification post-commit, rollback ────────────────────────

def _remplacer_fichier(temp: Path, cible: Path) -> None:
    """SEUL point de remplacement définitif du projet. Atomique (même volume garanti en amont)."""
    os.replace(str(temp), str(cible))


def _restaurer_fichier(sauvegarde: Path, cible: Path) -> None:
    """Restaure une cible depuis sa sauvegarde. Atomique si possible, sinon copie de secours.

    Le repli par copie est volontaire : si `os.replace` échoue (verrou, erreur disque), une copie
    non atomique vaut mieux qu'un fichier laissé dans un état incohérent.
    """
    try:
        os.replace(str(sauvegarde), str(cible))
    except Exception:
        shutil.copy2(str(sauvegarde), str(cible))


def _rollback(
    remplaces: list[FichierRemplace], sauvegardes: dict[Path, Path], cause: str
) -> None:
    """Restaure les fichiers DÉJÀ remplacés, en ordre inverse. Vérifie chaque restauration.

    Lève RollbackCritiqueError si un fichier ne revient pas à son empreinte initiale. Les
    sauvegardes des fichiers non restaurés sont CONSERVÉES (jamais nettoyées).
    """
    non_restaures: list[dict[str, str]] = []
    for r in reversed(remplaces):
        bak = sauvegardes.get(r.cible)
        if bak is None or not bak.exists():
            non_restaures.append({"cible": str(r.cible), "sauvegarde": str(bak or "(absente)")})
            continue
        try:
            _restaurer_fichier(bak, r.cible)
        except Exception as exc:  # jamais masquée : elle devient une erreur critique
            non_restaures.append({
                "cible": str(r.cible), "sauvegarde": str(bak), "erreur": str(exc),
            })
            continue
        if _sha256(r.cible) != r.sha256_avant:      # restauration non conforme = non restauré
            non_restaures.append({"cible": str(r.cible), "sauvegarde": str(bak)})

    if non_restaures:
        raise RollbackCritiqueError(non_restaures, cause)


def _verifier_post_commit(remplaces: list[FichierRemplace]) -> list[str]:
    """Les fichiers finaux sont-ils lisibles et conformes aux temporaires validés ?

    Étape distincte et monkeypatchable : elle s'exécute **verrou tenu**, avant toute libération.
    """
    ecarts: list[str] = []
    for r in remplaces:
        if not _lisible(r.cible):
            ecarts.append(f"{r.cible.name} : illisible après remplacement.")
        elif _sha256(r.cible) != r.sha256_apres:
            ecarts.append(f"{r.cible.name} : empreinte différente du temporaire validé.")
    return ecarts


def _nettoyer(chemins: list[Path]) -> list[str]:
    """Suppression best-effort. Retourne ce qui n'a PAS pu être supprimé (jamais silencieux)."""
    restants: list[str] = []
    for c in chemins:
        try:
            Path(c).unlink(missing_ok=True)
        except Exception:
            restants.append(str(c))
    return restants


# ── Point d'entrée ───────────────────────────────────────────────────────────

def confirmer_ecriture_charge(demande: DemandeEcritureCharge) -> ResultatTransaction:
    """Écrit une charge dans les deux classeurs SAISIE : transactionnelle, VERROUILLÉE, TRACÉE.

    Ordre non négociable : garde des flags → verrou → résolution du charge_id → préparation →
    commit → vérification → libération → **journalisation**. Le verrou est pris avant toute lecture
    concluant à la disponibilité d'un identifiant ; il n'est relâché qu'une fois la transaction
    entièrement jouée — c'est lui qui tient la réservation.

    **Toute tentative est journalisée**, y compris les refus.

    Deux règles distinctes sur les pannes du journal, à ne pas confondre :

    - **AVANT la transaction (garde d'idempotence) : FAIL-CLOSED.** Si un `token_previsualisation`
      est fourni et que le journal est illisible, on ne peut pas savoir si cette charge a déjà été
      écrite → refus (`E_IDEMPOTENCE_INDISPONIBLE`), avant toute préparation. Une saisie bloquée
      vaut mieux qu'une charge comptée deux fois.
    - **PENDANT ou APRÈS la transaction (clôture de la trace) : jamais bloquant.** Une panne à ce
      moment-là n'empêche ni l'écriture ni un rollback ; elle est signalée dans
      `ResultatTransaction.journal_erreur` (ou `RollbackCritiqueError.journal_erreur`), jamais avalée.
    """
    trace = journal.ouvrir_trace(
        token_previsualisation=demande.token_previsualisation,
        charge_id=demande.charge_id or None,
        charge_id_source=journal.ID_FOURNI if demande.charge_id else journal.ID_DEMANDE,
        cible_saisie=demande.saisie_path,
        cible_impacts=demande.impacts_path,
        db_path=demande.db_path,
    )
    trace.sha256_saisie_avant = _sha_si_existe(demande.saisie_path)
    trace.sha256_impacts_avant = _sha_si_existe(demande.impacts_path)

    try:
        resultat = _confirmer(demande)
    except RollbackCritiqueError as exc:
        # État critique : la trace est ce qui restera pour reconstituer les faits.
        trace.rollback_tente = True
        trace.rollback_reussi = False
        trace.fichiers_non_restaures = exc.fichiers_non_restaures
        exc.journal_erreur = journal.cloturer_trace_signalee(
            trace, journal.ROLLBACK_CRITIQUE, code=E_REMPLACEMENT, details=str(exc),
            sha256_saisie_apres=_sha_si_existe(demande.saisie_path),
            sha256_impacts_apres=_sha_si_existe(demande.impacts_path),
        )
        raise                                   # l'incident prime, le journal ne le masque jamais
    except Exception as exc:
        # Chemin non prévu (bug, panne système…) : on ne sait PAS conclure sur l'état des fichiers.
        # On refuse de le deviner : statut ETAT_INCOHERENT, empreintes réelles relevées telles
        # quelles. Aucun rollback n'est tenté à l'aveugle, et les sauvegardes techniques ne sont
        # PAS nettoyées (le nettoyage est en aval du point de rupture) : elles restent la voie de
        # retour. Le verrou, lui, est bien libéré — un état incohérent ne doit pas bloquer en plus.
        journal.cloturer_trace_signalee(
            trace, journal.ETAT_INCOHERENT, code=type(exc).__name__, details=str(exc),
            sha256_saisie_apres=_sha_si_existe(demande.saisie_path),
            sha256_impacts_apres=_sha_si_existe(demande.impacts_path),
        )
        raise

    _completer_trace(trace, demande, resultat)
    statut_journal = _statut_journal(resultat)
    resultat.transaction_id = trace.transaction_id
    resultat.statut_journal = statut_journal
    erreur_journal = journal.cloturer_trace_signalee(
        trace, statut_journal, code=resultat.code, details=resultat.details,
        sha256_saisie_apres=_sha_si_existe(demande.saisie_path),
        sha256_impacts_apres=_sha_si_existe(demande.impacts_path),
    )
    # Le résultat principal est conservé tel quel ; la panne de journal est ajoutée, pas substituée.
    resultat.journal_erreur = _fusionner_erreur_journal(resultat.journal_erreur, erreur_journal)
    return resultat


def _sha_si_existe(path: Path | None) -> str | None:
    try:
        p = Path(path)
        return _sha256(p) if p.exists() else None
    except Exception:
        return None


def _fusionner_erreur_journal(existante: str | None, nouvelle: str | None) -> str | None:
    if existante and nouvelle:
        return f"{existante} | {nouvelle}"
    return nouvelle or existante


def _completer_trace(
    trace: journal.TraceTransaction, demande: DemandeEcritureCharge, res: ResultatTransaction
) -> None:
    """Reporte dans la trace ce que la transaction a réellement fait."""
    trace.charge_id = res.charge_id or demande.charge_id or None
    if demande.charge_id:
        trace.charge_id_source = journal.ID_FOURNI
    elif res.charge_id:
        trace.charge_id_source = journal.ID_GENERE
    else:
        trace.charge_id_source = journal.ID_DEMANDE
    # NOMS de fichiers seulement — jamais de chemins temporaires.
    trace.fichiers_remplaces = [Path(f).name for f in res.fichiers_remplaces]
    trace.rollback_tente = res.rollback_tente
    trace.rollback_reussi = res.rollback_reussi
    trace.verrou_pid = res.verrou_pid
    trace.verrou_hostname = res.verrou_hostname
    trace.residus = [Path(r).name for r in (res.nettoyage_incomplet + res.sauvegardes_restantes)]


# Correspondance code TECHNIQUE → verdict MÉTIER. Les deux restent distincts en base.
_STATUT_PAR_CODE: dict[str, str] = {
    E_FLAGS_DESACTIVES: journal.REFUSE_FLAGS,
    E_VERROU_DEJA_PRIS: journal.REFUSE_VERROU,
    E_DEMANDE_INVALIDE: journal.REFUSE_VALIDATION,
    E_TOKEN_DEJA_ECRIT: journal.REFUSE_VALIDATION,
    E_IDEMPOTENCE_INDISPONIBLE: journal.REFUSE_VALIDATION,
    E_CHARGE_ID_INVALIDE: journal.REFUSE_CHARGE_ID,
    E_CHARGE_ID_EXISTANT: journal.REFUSE_CHARGE_ID,
    E_CHARGE_ID_EPUISE: journal.REFUSE_CHARGE_ID,
    E_PREPARATION: journal.ECHEC_PREPARATION,
    E_TEMPORAIRE_MANQUANT: journal.ECHEC_PREPARATION,
    E_TEMPORAIRE_INVALIDE: journal.ECHEC_PREPARATION,
    E_SAUVEGARDE: journal.ECHEC_PREPARATION,
    E_REMPLACEMENT: journal.ROLLBACK_REUSSI,
    E_POST_COMMIT: journal.ROLLBACK_REUSSI,
}


def _statut_journal(res: ResultatTransaction) -> str:
    if res.statut == STATUT_OK:
        return journal.SUCCES
    statut = _STATUT_PAR_CODE.get(res.code or "")
    if statut is not None:
        return statut
    return journal.ETAT_INCOHERENT      # code inconnu : on ne prétend pas savoir


def _confirmer(demande: DemandeEcritureCharge) -> ResultatTransaction:
    """Transaction proprement dite (inchangée par la journalisation)."""
    # 1. Flags — AVANT le verrou : un refus de garde ne doit même pas créer de fichier de verrou.
    if (garde := _garde_flags()) is not None:
        return garde

    # 2. Validation syntaxique — aucune lecture de classeur, donc légitime hors verrou.
    if (invalide := _valider_demande(demande)) is not None:
        return ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=demande.charge_id,
            code=E_DEMANDE_INVALIDE, details=invalide,
        )

    # 3. Verrou interprocessus — tenu jusqu'à la fin, libéré dans un finally.
    try:
        with lock.verrou_saisie_charges(
            operation=lock.OPERATION_SAISIE_CHARGE,
            charge_id=demande.charge_id or None,      # peut être None : l'id naîtra sous verrou
            lock_path=demande.lock_path,
        ):
            return _executer_sous_verrou(demande)
    except lock.VerrouSaisieChargesDejaPrisError as exc:
        # Un autre processus écrit déjà : on n'a rien lu, rien préparé, rien remplacé.
        return ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=demande.charge_id,
            code=E_VERROU_DEJA_PRIS, details=str(exc),
            verrou_pid=exc.metadonnees.get("pid"),
            verrou_hostname=exc.metadonnees.get("hostname"),
        )


def _consulter_idempotence(
    demande: DemandeEcritureCharge,
) -> tuple[dict[str, Any] | None, str | None]:
    """Ce token a-t-il déjà produit une écriture réussie ? → (trace_succes, indisponibilite).

    **FAIL-CLOSED quand un token est fourni.** Si le journal est illisible, on ne peut pas savoir si
    cette charge a déjà été écrite : on REFUSE. La disponibilité ne prime pas sur le risque d'une
    double écriture comptable — mieux vaut une saisie bloquée qu'une charge comptée deux fois.

    Sans token, aucune garde d'idempotence ne s'applique (il n'y a rien à comparer) : la transaction
    suit son cours, et c'est l'unicité du `charge_id` — vérifiée sous verrou — qui protège.
    """
    token = demande.token_previsualisation
    if not token:
        return None, None
    try:
        return journal.token_deja_ecrit(token, db_path=demande.db_path), None
    except Exception as exc:
        return None, (
            f"Garde d'idempotence indisponible ({type(exc).__name__}: {exc}) : le journal n'a pas "
            f"pu être lu, donc il est impossible de savoir si le token {token} a déjà produit une "
            f"écriture. Transaction REFUSÉE avant toute préparation — aucun fichier n'a été touché."
        )


def _executer_sous_verrou(demande: DemandeEcritureCharge) -> ResultatTransaction:
    """Corps de la transaction. Le verrou est tenu pendant TOUTE cette fonction."""
    # 3-bis. Idempotence : ce token de prévisualisation a-t-il DÉJÀ produit une écriture réussie ?
    # Vérifié sous verrou, donc sérialisé. C'est ce qui neutralise le double-clic et le retry aveugle.
    deja, indisponible = _consulter_idempotence(demande)
    if indisponible is not None:
        # Refus AVANT toute préparation et tout remplacement (fail-closed).
        return ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=demande.charge_id,
            code=E_IDEMPOTENCE_INDISPONIBLE, details=indisponible,
        )
    if deja is not None:
        return ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=str(deja.get("charge_id") or ""),
            code=E_TOKEN_DEJA_ECRIT,
            details=(
                f"Le token de prévisualisation {demande.token_previsualisation} a déjà produit "
                f"l'écriture de la charge {deja.get('charge_id')} "
                f"(transaction {deja.get('transaction_id')}, le {deja.get('fin_utc')}). "
                f"Aucune seconde écriture."
            ),
        )

    # 4. Résolution du charge_id — vérifié ou généré ICI, jamais avant le verrou.
    try:
        charge_id, row_data, persistable = _resoudre_charge_id(demande)
    except ChargeIdInvalideError as exc:
        return ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=demande.charge_id,
            code=E_CHARGE_ID_INVALIDE, details=str(exc),
        )
    except ChargeIdDejaExistantError as exc:
        return ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=exc.charge_id,
            code=E_CHARGE_ID_EXISTANT, details=str(exc),
        )

    demande = dataclass_replace(
        demande, charge_id=charge_id, row_data=row_data, persistable=persistable
    )

    # ── Préparation : rien de réel n'est touché ───────────────────────────────
    prepares, erreur = _preparer(demande)
    if erreur is not None:
        return erreur

    temps = [p.temp for p in prepares]

    if (probleme := _valider_temporaires(prepares)) is not None:
        code, details = probleme
        restants = _nettoyer(temps)
        return ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=demande.charge_id, code=code, details=details,
            nettoyage_incomplet=restants,
        )

    # ── Sauvegardes : AVANT le premier remplacement ───────────────────────────
    jeton = uuid.uuid4().hex[:12]
    sauvegardes, erreur_bak = _sauvegarder(prepares, jeton)
    if erreur_bak is not None:
        restants = _nettoyer(temps)
        return ResultatTransaction(
            statut=STATUT_ERREUR, charge_id=demande.charge_id, code=E_SAUVEGARDE,
            details=erreur_bak, nettoyage_incomplet=restants,
        )

    # ── Commit : ordre déterministe ───────────────────────────────────────────
    remplaces: list[FichierRemplace] = []
    for p in sorted(prepares, key=lambda x: x.ordre):
        try:
            _remplacer_fichier(p.temp, p.cible)
        except Exception as exc:
            cause = f"{p.libelle} : remplacement échoué ({exc})."
            _rollback(remplaces, sauvegardes, cause)   # peut lever RollbackCritiqueError
            restants = _nettoyer(temps + list(sauvegardes.values()))
            return ResultatTransaction(
                statut=STATUT_ROLLBACK, charge_id=demande.charge_id, code=E_REMPLACEMENT,
                details=f"{cause} Tous les fichiers ont été restaurés.",
                nettoyage_incomplet=restants,
                rollback_tente=True, rollback_reussi=True,   # sinon RollbackCritiqueError aurait été levée
            )
        remplaces.append(FichierRemplace(
            cible=p.cible, sauvegarde=sauvegardes[p.cible],
            sha256_avant=p.sha256_cible_avant, sha256_apres=p.sha256_temp,
        ))

    # ── Vérification post-commit (verrou toujours tenu) ───────────────────────
    ecarts = _verifier_post_commit(remplaces)
    if ecarts:
        cause = "Vérification post-commit : " + " ".join(ecarts)
        _rollback(remplaces, sauvegardes, cause)       # peut lever RollbackCritiqueError
        restants = _nettoyer(temps + list(sauvegardes.values()))
        return ResultatTransaction(
            statut=STATUT_ROLLBACK, charge_id=demande.charge_id, code=E_POST_COMMIT,
            details=f"{cause} Tous les fichiers ont été restaurés.",
            nettoyage_incomplet=restants,
            rollback_tente=True, rollback_reussi=True,
        )

    # ── Succès : nettoyage des temporaires PUIS des sauvegardes ───────────────
    restants = _nettoyer(temps)                        # les temps ont été consommés par os.replace
    restants += _nettoyer(list(sauvegardes.values()))  # jamais avant la validation complète
    return ResultatTransaction(
        statut=STATUT_OK, charge_id=demande.charge_id,
        fichiers_remplaces=[str(r.cible) for r in remplaces],
        sauvegardes_restantes=[],
        nettoyage_incomplet=restants,                  # signalé, jamais tu
        sha256_finaux={r.cible.name: r.sha256_apres for r in remplaces},
    )
