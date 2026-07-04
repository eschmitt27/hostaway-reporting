"""APP-2d - activation controlee de l'ecriture reelle HH.

Ce service ne change jamais les flags applicatifs. Il verifie un dry-run APP-2c
recent et coherent avant d'appeler l'orchestrateur d'ecriture existant.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from decimal import Decimal
import json
import os
from pathlib import Path
import shutil
from typing import Any

import openpyxl

import app.config as cfg
from app.readers.ref_setup_hh_reader import get_cloture_mois
from app.readers.saisie_hh_reader import read_existing_pks
from app.services import saisie_hh_dryrun_service as dryrun_svc
from app.services import saisie_hh_orchestrator as hh_orchestrator
from app.services import saisie_hh_service as saisie_svc
from app.services.audit_service import log_event
from app.services.saisie_hh_schema_migration import NEW_SAISIE_FIELDS


EXPIRATION_MINUTES = 30
STATUS_CANCELLED_RESTORED = "ECRITURE_REELLE_ANNULEE_ET_RESTAUREE"


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _safe_log(action: str, details: dict[str, Any], db_path: Path | None = None) -> None:
    try:
        log_event(action, details=details, db_path=db_path or cfg.DB_PATH)
    except Exception:
        pass


def _parse_dt(value: str | None) -> datetime | None:
    if not value:
        return None
    raw = str(value).strip().replace("Z", "+00:00")
    try:
        dt = datetime.fromisoformat(raw)
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def _fingerprint(path: Path) -> dict[str, Any]:
    return dryrun_svc._fingerprint(path)


def _sha256(path: Path) -> str:
    return dryrun_svc._sha256(path)


def _payload(manifest: dict[str, Any]) -> dict[str, Any]:
    return dict(manifest.get("payload_canonique") or manifest.get("payload_summary") or {})


def _pk_from_manifest(manifest: dict[str, Any]) -> str:
    payload = _payload(manifest)
    return str(payload.get("reservation_hh_id") or "").strip()


def _expected_confirmation(pk: str) -> str:
    return f"ENREGISTRER {pk}"


def _add_check(checks: list[dict[str, Any]], code: str, ok: bool, label: str, details: str = "") -> None:
    checks.append({"code": code, "ok": bool(ok), "label": label, "details": details})


def _source_hash(manifest: dict[str, Any], key: str) -> str:
    return str((manifest.get("source_hashes") or {}).get(key, {}).get("sha256") or "").lower()


def _hash_matches(manifest: dict[str, Any], key: str, path: Path) -> tuple[bool, str]:
    expected = _source_hash(manifest, key)
    if not expected:
        return False, "hash source absent du manifest"
    actual = _fingerprint(path).get("sha256")
    if not actual:
        return False, f"hash actuel illisible: {path}"
    ok = str(actual).lower() == expected
    return ok, f"manifest={expected[:12]} actuel={str(actual).lower()[:12]}"


def _schema_state(saisie_path: Path, ref_setup_path: Path) -> dict[str, Any]:
    missing_saisie: list[str] = []
    missing_ref_modes: list[str] = []
    try:
        wb = openpyxl.load_workbook(str(saisie_path), read_only=True, data_only=True)
        try:
            ws = wb["SAISIE"]
            headers = [str(cell.value or "").strip() for cell in ws[1]]
        finally:
            wb.close()
        missing_saisie = [field for field in NEW_SAISIE_FIELDS if field not in headers]
    except Exception as exc:
        return {"ok": False, "missing_saisie_fields": NEW_SAISIE_FIELDS, "missing_ref_modes": [], "error": str(exc)}

    try:
        wb_ref = openpyxl.load_workbook(str(ref_setup_path), read_only=True, data_only=True, keep_vba=True)
        try:
            ws_ref = wb_ref["REF_Modes_Paiement"]
            rows = list(ws_ref.iter_rows(values_only=True))
        finally:
            wb_ref.close()
        headers = [str(v or "").strip() for v in rows[0]]
        id_idx = headers.index("mode_paiement_id")
        ids = {str(row[id_idx] or "").strip() for row in rows[1:] if len(row) > id_idx}
        if "PAY_006" not in ids:
            missing_ref_modes.append("PAY_006 / DIRECT_PROPRIETAIRE")
    except Exception as exc:
        return {
            "ok": False,
            "missing_saisie_fields": missing_saisie,
            "missing_ref_modes": ["PAY_006 / DIRECT_PROPRIETAIRE"],
            "error": str(exc),
        }

    return {
        "ok": not missing_saisie and not missing_ref_modes,
        "missing_saisie_fields": missing_saisie,
        "missing_ref_modes": missing_ref_modes,
        "error": "",
    }


def _pk_exists(path: Path, pk: str) -> tuple[bool, str]:
    if not pk:
        return False, "reservation_hh_id absent du manifest"
    pks = set(read_existing_pks(path))
    return pk in pks, f"{pk} {'present' if pk in pks else 'absent'} dans SAISIE"


def _month_open(ref_setup_path: Path, mois: str) -> tuple[bool, str]:
    if not mois:
        return False, "mois absent du payload"
    row = get_cloture_mois(mois, ref_setup_path=ref_setup_path)
    if row is None:
        return False, f"{mois} absent de REF_Cloture_Mensuelle"
    statut = str(row.get("statut_mois", "")).strip().upper()
    return statut == "OUVERT", f"{mois} statut={statut or 'VIDE'}"


def _find_master_row(result: dict[str, Any], pk: str) -> dict[str, Any] | None:
    for row in result.get("master_rows", []) or []:
        if str(row.get("reservation_hh_id") or "").strip() == pk:
            return row
    return None


def _same_value(left: Any, right: Any) -> bool:
    if left in (None, "") and right in (None, ""):
        return True
    try:
        return Decimal(str(left)) == Decimal(str(right))
    except Exception:
        return str(left) == str(right)


def _compare_lot4a_rows(simulated: dict[str, Any], post: dict[str, Any]) -> list[str]:
    keys = (
        "reservation_hh_id", "mois", "date_arrivee", "date_depart", "nuits",
        "total_percu", "menage", "taux_commission", "taux_commission_source",
        "commission", "mode_paiement_id", "montant_recupere",
        "associe_id_recuperateur", "montant_reverse_proprietaire",
        "acompte_facture", "source_acompte_facture", "code_impact",
        "comptabilisation", "taux_commission_override", "menage_override",
    )
    diffs: list[str] = []
    for key in keys:
        if not _same_value(simulated.get(key), post.get(key)):
            diffs.append(f"{key}: simulation={simulated.get(key)!r} post={post.get(key)!r}")
    return diffs


def evaluate_real_write_prerequisites(
    token: str,
    *,
    saisie_path: Path | None = None,
    ref_setup_path: Path | None = None,
    dryruns_root: Path | None = None,
    now: datetime | None = None,
) -> dict[str, Any]:
    """Retourne l'etat complet des prerequis d'ecriture reelle sans ecrire."""
    p_saisie = Path(saisie_path or cfg.SAISIE_RESERVATIONS_HH)
    p_ref = Path(ref_setup_path or cfg.REF_SETUP)
    checks: list[dict[str, Any]] = []
    dryrun: dict[str, Any] | None = None
    manifest: dict[str, Any] = {}

    try:
        dryrun = dryrun_svc.load_previsualisation(token, dryruns_root=dryruns_root)
        manifest = dict(dryrun.get("manifest") or {})
        _add_check(checks, "SIMULATION_PRESENTE", True, "simulation presente")
    except Exception as exc:
        _add_check(checks, "SIMULATION_PRESENTE", False, "simulation absente", str(exc))
        return {
            "eligible": False,
            "checks": checks,
            "dryrun": None,
            "manifest": {},
            "pk": "",
            "expected_confirmation": "",
        }

    pk = _pk_from_manifest(manifest)
    payload = _payload(manifest)
    created = _parse_dt(manifest.get("created_at_utc"))
    age_ok = created is not None and ((now or _utcnow()) - created) <= timedelta(minutes=EXPIRATION_MINUTES)
    _add_check(
        checks,
        "SIMULATION_RECENTE",
        age_ok,
        "simulation de moins de 30 minutes",
        "Simulation expiree : relancez une previsualisation avant toute ecriture reelle." if not age_ok else "",
    )
    _add_check(checks, "SIMULATION_OK", manifest.get("status") == "OK", "simulation valide", str(manifest.get("status")))
    _add_check(
        checks,
        "LOT4A_TERMINE",
        manifest.get("lot4a_status") == "ANALYSE_TERMINEE",
        "Lot4A termine",
        str(manifest.get("lot4a_status") or ""),
    )
    errors = manifest.get("errors") or []
    _add_check(checks, "AUCUNE_ANOMALIE_BLOQUANTE", not errors, "aucune erreur ou anomalie bloquante", "; ".join(map(str, errors)))

    ok_hash_saisie, details_hash_saisie = _hash_matches(manifest, "saisie", p_saisie)
    _add_check(checks, "HASH_SAISIE_IDENTIQUE", ok_hash_saisie, "fichier SAISIE inchange depuis la simulation", details_hash_saisie)
    ok_hash_ref, details_hash_ref = _hash_matches(manifest, "ref_setup", p_ref)
    _add_check(checks, "HASH_REF_SETUP_IDENTIQUE", ok_hash_ref, "REF_Setup inchange depuis la simulation", details_hash_ref)

    exists, details_pk = _pk_exists(p_saisie, pk)
    _add_check(checks, "PK_ABSENTE", not exists and bool(pk), "reservation absente du fichier reel", details_pk)

    mois = str(payload.get("mois") or "") or str(payload.get("date_arrivee") or "")[:7]
    try:
        mois_ok, mois_details = _month_open(p_ref, mois)
    except Exception as exc:
        mois_ok, mois_details = False, str(exc)
    _add_check(checks, "MOIS_OUVERT", mois_ok, "mois ouvert dans REF_Cloture_Mensuelle", mois_details)

    schema = _schema_state(p_saisie, p_ref)
    schema_details = {
        "missing_saisie_fields": schema.get("missing_saisie_fields", []),
        "missing_ref_modes": schema.get("missing_ref_modes", []),
        "error": schema.get("error", ""),
    }
    _add_check(
        checks,
        "SCHEMA_REEL_NON_PREPARE",
        schema["ok"],
        "schema reel prepare",
        json.dumps(schema_details, ensure_ascii=False),
    )

    _add_check(checks, "FLAG_GLOBAL_ACTIVE", bool(cfg.HH_REAL_WRITE_ENABLED), "flag global active", "HH_REAL_WRITE_ENABLED")
    _add_check(
        checks,
        "FLAG_CONFIRMATION_ACTIVE",
        bool(getattr(cfg, "HH_REAL_WRITE_CONFIRMATION_ENABLED", False)),
        "confirmation reelle activee",
        "HH_REAL_WRITE_CONFIRMATION_ENABLED",
    )

    eligible = all(item["ok"] for item in checks)
    return {
        "eligible": eligible,
        "checks": checks,
        "dryrun": dryrun,
        "manifest": manifest,
        "pk": pk,
        "expected_confirmation": _expected_confirmation(pk) if pk else "",
        "schema": schema,
        "source_hashes_current": {
            "saisie": _fingerprint(p_saisie),
            "ref_setup": _fingerprint(p_ref),
        },
    }


def _restore_snapshot(saisie_path: Path, rollback_path: Path, expected_hash: str) -> tuple[str, str]:
    if not rollback_path.exists():
        return "ROLLBACK_ECHEC", "snapshot APP-2d absent"
    os.replace(str(rollback_path), str(saisie_path))
    restored = _sha256(saisie_path)
    if restored == expected_hash:
        return "ROLLBACK_REUSSI", restored
    return "ROLLBACK_ECHEC", restored


def _post_write_lot4a_check(
    dryrun: dict[str, Any],
    saisie_path: Path,
    ref_setup_path: Path,
    pk: str,
) -> dict[str, Any]:
    run_dir = Path(dryrun["run_dir"])
    check_dir = run_dir / "real_write_check"
    check_dir.mkdir(parents=True, exist_ok=True)
    saisie_copy = check_dir / "SAISIE_post_ecriture_copie.xlsx"
    ref_copy = check_dir / "REF_Setup_post_ecriture_copie.xlsm"
    master_copy = check_dir / "MASTER_FACT_MAN_post_ecriture_simule.xlsx"
    shutil.copy2(saisie_path, saisie_copy)
    shutil.copy2(ref_setup_path, ref_copy)
    rows, result, engine = dryrun_svc._run_lot4a_engine(
        saisie_copy,
        ref_copy,
        datetime.now(timezone.utc).isoformat(),
        check_dir,
        master_path=master_copy,
    )
    if result.get("statut") != "ANALYSE_TERMINEE":
        raise RuntimeError(f"LOT4A_POST_ECRITURE_NON_TERMINE: {result.get('statut')}")
    post_row = _find_master_row(result, pk)
    if not post_row:
        raise RuntimeError(f"RESERVATION_ABSENTE_MASTER_POST_ECRITURE: {pk}")
    simulated = dryrun.get("simulated_master_row") or (dryrun.get("resultat_lot4a") or {}).get("reservation_simulee") or {}
    diffs = _compare_lot4a_rows(simulated, post_row)
    if diffs:
        raise RuntimeError("DIVERGENCE_ECRITURE_REELLE_VS_SIMULATION: " + "; ".join(diffs))
    return {"rows": rows, "result": result, "engine": engine, "post_row": post_row, "master_copy": str(master_copy)}


def enregistrer_reservation_hh_reelle(
    simulation_token: str,
    confirmation_texte: str,
    *,
    saisie_path: Path | None = None,
    ref_setup_path: Path | None = None,
    dryruns_root: Path | None = None,
    db_path: Path | None = None,
    now: datetime | None = None,
) -> dict[str, Any]:
    """Execute l'ecriture reelle controlee, uniquement si tous les prerequis passent."""
    p_saisie = Path(saisie_path or cfg.SAISIE_RESERVATIONS_HH)
    p_ref = Path(ref_setup_path or cfg.REF_SETUP)
    p_db = db_path or cfg.DB_PATH
    state = evaluate_real_write_prerequisites(
        simulation_token,
        saisie_path=p_saisie,
        ref_setup_path=p_ref,
        dryruns_root=dryruns_root,
        now=now,
    )
    pk = state.get("pk") or ""
    expected = state.get("expected_confirmation") or ""

    if not state["eligible"]:
        failed = [item["code"] for item in state["checks"] if not item["ok"]]
        _safe_log("APP2D_ECRITURE_REELLE_REFUSEE", {"token": simulation_token, "pk": pk, "failed": failed}, p_db)
        return {"statut": "REFUSE", "pk": pk, "checks": state["checks"], "details": failed}

    if str(confirmation_texte or "") != expected:
        _safe_log("APP2D_CONFIRMATION_REELLE_REFUSEE", {"token": simulation_token, "pk": pk}, p_db)
        return {
            "statut": "REFUSE",
            "pk": pk,
            "checks": state["checks"],
            "details": ["TEXTE_CONFIRMATION_INCORRECT"],
            "expected_confirmation": expected,
        }

    manifest = state["manifest"]
    form_data = manifest.get("form_data")
    if not isinstance(form_data, dict):
        return {"statut": "REFUSE", "pk": pk, "checks": state["checks"], "details": ["FORM_DATA_ABSENT"]}

    validation = saisie_svc.valider(form_data, saisie_path=p_saisie, ref_setup_path=p_ref)
    if not validation.get("ok") or validation.get("pk") != pk:
        return {
            "statut": "REFUSE",
            "pk": pk,
            "checks": state["checks"],
            "details": ["VALIDATION_METIER_DIVERGENTE", validation.get("erreurs", [])],
        }

    sha_avant = _sha256(p_saisie)
    rollback_path = p_saisie.parent / f"{p_saisie.stem}.app2d.rollback.xlsx"
    shutil.copy2(p_saisie, rollback_path)
    rollback_created = True
    write_started = False
    try:
        row_data = saisie_svc.build_row_data(validation["preview"])
        write_started = True
        write_result = hh_orchestrator.confirm_write(
            row_data=row_data,
            pk=pk,
            mois=validation["preview"].get("mois", ""),
            saisie_path=p_saisie,
            db_path=p_db,
        )
        if write_result.get("statut") != "OK":
            raise RuntimeError(f"ORCHESTRATEUR_ECRITURE_REFUSEE: {write_result}")
        post_check = _post_write_lot4a_check(state["dryrun"], p_saisie, p_ref, pk)
        _safe_log(
            "APP2D_ECRITURE_REELLE_OK",
            {"token": simulation_token, "pk": pk, "sha256_avant": sha_avant, "sha256_apres": _sha256(p_saisie)},
            p_db,
        )
        rollback_path.unlink(missing_ok=True)
        return {
            "statut": "OK",
            "pk": pk,
            "checks": state["checks"],
            "write_result": write_result,
            "post_lot4a": post_check,
            "sha256_avant": sha_avant,
            "sha256_apres": _sha256(p_saisie),
        }
    except Exception as exc:
        current_hash = _sha256(p_saisie) if p_saisie.exists() else ""
        rollback_status = "ROLLBACK_NON_NECESSAIRE"
        restored_hash = current_hash
        if rollback_created and (write_started or current_hash != sha_avant):
            rollback_status, restored_hash = _restore_snapshot(p_saisie, rollback_path, sha_avant)
        _safe_log(
            "APP2D_ECRITURE_REELLE_ROLLBACK",
            {
                "statut": rollback_status,
                "erreur_initiale": str(exc),
                "hash_avant": sha_avant,
                "hash_apres_restauration": restored_hash,
                "token": simulation_token,
                "reservation_hh_id": pk,
            },
            p_db,
        )
        return {
            "statut": STATUS_CANCELLED_RESTORED if rollback_status == "ROLLBACK_REUSSI" else "ERREUR",
            "pk": pk,
            "checks": state["checks"],
            "details": str(exc),
            "rollback_status": rollback_status,
            "sha256_avant": sha_avant,
            "sha256_apres_restauration": restored_hash,
        }
    finally:
        if rollback_path.exists():
            try:
                rollback_path.unlink()
            except Exception:
                pass
