"""Reader Contrôles & clôture (APP-5A) — LECTURE SEULE.

Lit les sorties moteur consolidées du Lot11 et l'état de clôture :
  - MASTER_CTRL_Coherence.xlsx : MASTER (contrôles unifiés), BLOQUANTS_OUVERTS, A_CONTROLER_OUVERTS,
    DASHBOARD_MOIS (état de clôturabilité par mois)
  - REF_Setup.xlsm > REF_Cloture_Mensuelle : statut de clôture officiel des mois (OUVERT/EN_CONTROLE/CLOTURE)

Ne recrée aucun contrôle, ne clôture rien, n'écrit rien. Le code, le niveau et le statut viennent du
moteur. openpyxl read_only=True. Aucun chemin absolu exposé.
"""
from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

import app.config as cfg

# ── Onglets ──────────────────────────────────────────────────────────────────
ONGLET_MASTER = "MASTER"
ONGLET_BLOQUANTS = "BLOQUANTS_OUVERTS"
ONGLET_A_CONTROLER = "A_CONTROLER_OUVERTS"
ONGLET_DASHBOARD = "DASHBOARD_MOIS"
ONGLET_CLOTURE_REF = "REF_Cloture_Mensuelle"

SOURCE_COHERENCE = "MASTER_CTRL_Coherence.xlsx"
SOURCE_REF = "REF_Setup.xlsm"

# ── États ────────────────────────────────────────────────────────────────────
ETAT_OK = "OK"
ETAT_FICHIER_ABSENT = "FICHIER_ABSENT"
ETAT_ONGLET_ABSENT = "ONGLET_ABSENT"
ETAT_VIDE = "VIDE"
ETAT_ILLISIBLE = "ILLISIBLE"

_ETAT_LIBELLE = {
    ETAT_OK: "Alimentée", ETAT_FICHIER_ABSENT: "Fichier absent", ETAT_ONGLET_ABSENT: "Onglet absent",
    ETAT_VIDE: "Source vide", ETAT_ILLISIBLE: "Source illisible",
}


@dataclass(frozen=True)
class EtatSource:
    cle: str
    libelle: str
    fichier: str
    onglet: str
    etat: str
    nb_lignes: int = 0
    derniere_maj: str | None = None

    @property
    def disponible(self) -> bool:
        return self.etat == ETAT_OK

    @property
    def etat_libelle(self) -> str:
        return _ETAT_LIBELLE.get(self.etat, self.etat)


@dataclass(frozen=True)
class Source:
    etat: EtatSource
    lignes: list[dict[str, Any]] = field(default_factory=list)


_CACHE: dict[tuple, tuple[str, list[dict[str, Any]], str | None]] = {}


def vider_cache() -> None:
    _CACHE.clear()


def _lire(path: Path, sheet: str) -> tuple[str, list[dict[str, Any]], str | None]:
    p = Path(path)
    if not p.exists():
        return ETAT_FICHIER_ABSENT, [], None
    try:
        st = p.stat()
        cle = (str(p), sheet, st.st_mtime_ns, st.st_size)
        if cle in _CACHE:
            return _CACHE[cle]
        maj = _dt.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M")
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        try:
            if sheet not in wb.sheetnames:
                res = (ETAT_ONGLET_ABSENT, [], maj); _CACHE[cle] = res; return res
            ws = wb[sheet]
            rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        finally:
            wb.close()
        if len(rows) <= 1:
            res = (ETAT_VIDE, [], maj)
        else:
            hdr = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
            res = (ETAT_OK, [dict(zip(hdr, r)) for r in rows[1:]], maj)
        _CACHE[cle] = res
        return res
    except Exception:
        return ETAT_ILLISIBLE, [], None


def _fichier_coherence() -> Path:
    return Path(cfg.MASTER_CTRL_COHERENCE_FILE)


def _src_coherence(cle: str, libelle: str, sheet: str) -> Source:
    etat, lignes, maj = _lire(_fichier_coherence(), sheet)
    return Source(etat=EtatSource(cle=cle, libelle=libelle, fichier=SOURCE_COHERENCE, onglet=sheet,
                                  etat=etat, nb_lignes=len(lignes), derniere_maj=maj), lignes=lignes)


# ── Convertisseurs ───────────────────────────────────────────────────────────

def to_texte(v: Any) -> str:
    return "" if v is None else str(v).strip()


def to_nombre(v: Any) -> float | int | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def to_mois(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m")
    return str(v)[:7]


def to_date(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10]


# ── Sources ──────────────────────────────────────────────────────────────────

def controles() -> Source:
    return _src_coherence("controles", "Contrôles consolidés (Lot11)", ONGLET_MASTER)


def bloquants_ouverts() -> Source:
    return _src_coherence("bloquants", "Contrôles bloquants ouverts", ONGLET_BLOQUANTS)


def a_controler_ouverts() -> Source:
    return _src_coherence("a_controler", "Contrôles à contrôler ouverts", ONGLET_A_CONTROLER)


def dashboard_mois() -> Source:
    return _src_coherence("dashboard_mois", "Clôturabilité par mois", ONGLET_DASHBOARD)


def cloture_ref() -> Source:
    """État de clôture officiel des mois — REF_Setup.xlsm > REF_Cloture_Mensuelle."""
    etat, lignes, maj = _lire(cfg.REF_SETUP, ONGLET_CLOTURE_REF)
    return Source(etat=EtatSource(cle="cloture_ref", libelle="Statut de clôture des mois",
                                  fichier=SOURCE_REF, onglet=ONGLET_CLOTURE_REF, etat=etat,
                                  nb_lignes=len(lignes), derniere_maj=maj), lignes=lignes)


def etats_sources() -> list[EtatSource]:
    return [controles().etat, dashboard_mois().etat, cloture_ref().etat]
