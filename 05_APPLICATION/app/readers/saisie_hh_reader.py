"""Lecteur dédié SAISIE HH — lecture strictement read-only (APP-2b).

Fournit :
- listes REF_LOCALE pour alimenter le formulaire de saisie,
- PKs existants pour génération RESHH-AAAA-MM-NNN,
- détection première ligne vide pour ciblage écriture.
"""
from pathlib import Path
import re
import openpyxl

SHEET_SAISIE = "SAISIE"
SHEET_REF_LOCALE = "REF_LOCALE"


class SaisieHHReadError(RuntimeError):
    """Erreur bloquante de lecture du classeur SAISIE HH."""

FORMULA_COLS = {"B", "C", "K", "N", "O", "Q", "V", "Y", "Z"}

# Colonnes manuelles : lettre → nom de champ (confirmé sur SAISIE réelle)
MANUAL_COL_MAP: dict[str, str] = {
    "A":  "reservation_hh_id",
    "D":  "canal_id",
    "E":  "source_financiere",
    "F":  "proprietaire_id",
    "G":  "logement_id",
    "H":  "reservation_id_hostaway",
    "I":  "date_arrivee",
    "J":  "date_depart",
    "L":  "total_percu",
    "M":  "menage",
    "P":  "commentaire_taux_commission",
    "R":  "montant_recupere",
    "S":  "associe_id_recuperateur",
    "T":  "montant_reverse_proprietaire",
    "U":  "mode_paiement_id",
    "W":  "code_impact",
    "X":  "comptabilisation",
    "AA": "statut_controle",
    "AB": "niveau_anomalie",
    "AC": "code_anomalie",
    "AD": "commentaire",
}

# Champs forcés — jamais saisis par l'utilisateur
FORCED_VALUES = {
    "statut_controle": "A_CONTROLER",
    "niveau_anomalie": "A_CONTROLER",
}


def _col_index(letter: str) -> int:
    """Lettre(s) de colonne → index 1-based."""
    result = 0
    for c in letter.upper():
        result = result * 26 + (ord(c) - 64)
    return result


def _col_letter(n: int) -> str:
    """Index 1-based → lettre(s) de colonne."""
    col = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        col = chr(65 + r) + col
    return col


def read_ref_locale(saisie_path: Path) -> dict[str, list[str]]:
    """Lit REF_LOCALE et retourne un dict {nom_liste: [valeur, ...]}."""
    p = Path(saisie_path)
    if not p.exists():
        raise SaisieHHReadError(f"SAISIE introuvable pour REF_LOCALE : {p}")
    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        ws = wb[SHEET_REF_LOCALE]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception as exc:
        raise SaisieHHReadError(f"Lecture REF_LOCALE impossible : {exc}") from exc

    if not rows:
        return {}

    headers = list(rows[0])
    result: dict[str, list[str]] = {}
    for h in headers:
        if h is not None:
            result[str(h)] = []

    for row in rows[1:]:
        for i, h in enumerate(headers):
            if h is None:
                continue
            key = str(h)
            if i >= len(row) or row[i] is None:
                continue
            v = str(row[i]).strip()
            if v and v not in result[key]:
                result[key].append(v)

    return result


def read_existing_pks(saisie_path: Path) -> list[str]:
    """Retourne tous les reservation_hh_id (col A, ligne 2+) non vides."""
    p = Path(saisie_path)
    if not p.exists():
        raise SaisieHHReadError(f"SAISIE introuvable pour lecture PK : {p}")
    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        ws = wb[SHEET_SAISIE]
        rows = list(ws.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True))
        wb.close()
    except Exception as exc:
        raise SaisieHHReadError(f"Lecture PK existantes impossible : {exc}") from exc

    return [str(r[0]).strip() for r in rows if r[0] is not None and str(r[0]).strip()]


def find_first_empty_data_row(saisie_path: Path, max_scan: int = 502) -> int | None:
    """Retourne le numéro 1-based de la première ligne où col A est vide (>= ligne 2)."""
    p = Path(saisie_path)
    if not p.exists():
        return None
    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        ws = wb[SHEET_SAISIE]
        rows = list(ws.iter_rows(min_row=2, max_row=max_scan, min_col=1, max_col=1, values_only=True))
        wb.close()
    except Exception:
        return None

    if not rows:
        # Aucune ligne de données : le sheet n'a que l'en-tête → ligne 2 disponible
        return 2

    for i, r in enumerate(rows):
        if r[0] is None or str(r[0]).strip() == "":
            return i + 2  # +2 : 1-indexed + saut de l'en-tête
    return None


def generate_pk(date_arrivee_str: str, existing_pks: list[str]) -> tuple[str, str | None]:
    """
    Génère RESHH-AAAA-MM-NNN à partir de date_arrivee et des PKs existants.
    Retourne (pk, code_erreur). code_erreur est None si succès.
    Codes : 'DATE_ARRIVEE_INVALIDE', 'SEQUENCE_PK_INCOHERENTE'.
    """
    try:
        from datetime import datetime as dt
        date_str = str(date_arrivee_str).strip()[:10]
        d = dt.strptime(date_str, "%Y-%m-%d").date()
        mois = f"{d.year:04d}-{d.month:02d}"
    except (ValueError, AttributeError, TypeError):
        return "", "DATE_ARRIVEE_INVALIDE"

    prefix = f"RESHH-{mois}-"
    max_num = 0

    for pk in existing_pks:
        if not pk.startswith(prefix):
            continue
        suffix = pk[len(prefix):]
        if not re.fullmatch(r"\d+", suffix):
            return "", "SEQUENCE_PK_INCOHERENTE"
        max_num = max(max_num, int(suffix))

    new_pk = f"{prefix}{max_num + 1:03d}"
    return new_pk, None
