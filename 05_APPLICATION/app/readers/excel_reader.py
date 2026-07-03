from pathlib import Path
from typing import Any
import openpyxl


def read_sheet(
    path: Path,
    sheet_name: str | None = None,
    max_rows: int | None = 500,
) -> list[dict[str, Any]]:
    """Lit un onglet Excel en lecture seule. Retourne liste de dicts (en-tête = row 1).
    Ne transforme pas les valeurs. Jamais d'handle d'écriture."""
    p = Path(path)
    if not p.exists():
        return []

    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    except Exception:
        return []

    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return []

    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
    result = []
    for row in rows[1:max_rows]:
        result.append(dict(zip(headers, row)))
    return result


def list_sheets(path: Path) -> list[str]:
    p = Path(path)
    if not p.exists():
        return []
    try:
        wb = openpyxl.load_workbook(str(p), read_only=True)
        names = wb.sheetnames
        wb.close()
        return names
    except Exception:
        return []
