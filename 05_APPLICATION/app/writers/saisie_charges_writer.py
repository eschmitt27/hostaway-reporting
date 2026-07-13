"""APP-3b / commit 2 — Writer BAS NIVEAU des charges (préparation seule, aucun remplacement).

Ce module **prépare** une écriture et s'arrête là : il produit un fichier **temporaire validé**,
sur le même volume que la cible, prêt à être remplacé plus tard par l'orchestrateur transactionnel
(commit 3). Il ne fait JAMAIS :

  - `os.replace` sur un fichier métier réel ;
  - de sauvegarde, de rollback multi-fichiers, de journalisation transactionnelle ;
  - d'écriture réelle : les flags (`CHARGES_REAL_WRITE_ENABLED`) ne sont pas lus ici, car préparer
    un temporaire ne modifie aucune donnée. C'est l'orchestrateur qui portera la garde.

Le fichier source est ouvert en lecture seule et son empreinte SHA256 est vérifiée avant/après :
toute préparation qui aurait modifié la source est un échec.

Pattern repris de `saisie_hh_writer` (APP-2b), dont les helpers génériques (empreinte, détection
d'Excel ouvert, même volume, signature structurelle) sont réutilisés tels quels plutôt que dupliqués.

Deux cibles, deux fonctions :
  - `prepare_charge_flux_write`    → SAISIE_Charges_Flux.xlsx   (la ligne de charge)
  - `prepare_charge_impacts_write` → SAISIE_Charges_Impacts.xlsx (AFFECTATIONS / MENAGE / RESERVE)

Invariants métier gardés ici (refus explicite, jamais de correction silencieuse) :
  - les colonnes formule C (mois), I (impact_resultat_reel), J (impact_resultat_comptable) et
    AD (ROW_HASH) ne sont JAMAIS écrites — un payload qui les vise est refusé ;
  - `charge_id` unique : une même charge ne peut pas être écrite sur deux lignes ;
  - une charge = UNE ligne de charge (la réécriture au même endroit remplace, n'ajoute pas) ;
  - AFFECTATIONS : Σ quote_part = montant ;
  - MENAGE : intervenant XOR logement, et jamais de réserve pour une charge ménage ;
  - RESERVE_REFACTURATION : statut EN_ATTENTE, montants cohérents (aucune préfacture automatique).
"""
from __future__ import annotations

import shutil
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from app.readers.saisie_charges_reader import (
    FORMULA_COLS,
    FORMULA_COL_INDICES,
    MANUAL_COL_MAP,
    SAISIE_SHEET,
    _col_index,
)
from app.services.file_registry import is_writable
# Helpers génériques éprouvés (APP-2b) — réutilisés, jamais redupliqués.
from app.writers.saisie_hh_writer import (
    _detect_excel_lock,
    _same_volume,
    _sha256,
    _structure_signature,
)

AFFECT_SHEET = "AFFECTATIONS"
MENAGE_SHEET = "MENAGE"
RESERVE_SHEET = "RESERVE_REFACTURATION"

STATUT_RESERVE_ATTENDU = "EN_ATTENTE"
TOLERANCE_CENTIMES = 0.005

# Champs portés par une formule Excel : interdits en écriture (le moteur Lot3 les dérive).
CHAMPS_FORMULE: frozenset[str] = frozenset(
    {"mois", "impact_resultat_reel", "impact_resultat_comptable", "ROW_HASH"}
)
CHAMPS_MANUELS: frozenset[str] = frozenset(MANUAL_COL_MAP.values())

# Codes de refus (jamais de correction silencieuse).
E_CHEMIN_NON_AUTORISE = "E_CHEMIN_NON_AUTORISE"
E_FICHIER_ABSENT = "E_FICHIER_ABSENT"
E_FICHIER_OUVERT = "E_FICHIER_OUVERT"
E_SOURCE_MODIFIEE = "E_SOURCE_MODIFIEE"
E_COLONNE_FORMULE = "E_COLONNE_FORMULE"
E_CHAMP_INCONNU = "E_CHAMP_INCONNU"
E_CHARGE_ID_MANQUANT = "E_CHARGE_ID_MANQUANT"
E_CHARGE_ID_COLLISION = "E_CHARGE_ID_COLLISION"
E_LIGNE_CIBLE_OCCUPEE = "E_LIGNE_CIBLE_OCCUPEE"
E_FORMULE_MODELE_ABSENTE = "E_FORMULE_MODELE_ABSENTE"
E_FORMULE_MODELE_FIGEE = "E_FORMULE_MODELE_FIGEE"
E_VOLUME_DIFFERENT = "E_VOLUME_DIFFERENT"
E_VALEUR_NON_ECRITE = "E_VALEUR_NON_ECRITE"
E_DELTA_HORS_PERIMETRE = "E_DELTA_HORS_PERIMETRE"
E_STRUCTURE_ALTEREE = "E_STRUCTURE_ALTEREE"
E_SOMME_QUOTES = "E_SOMME_QUOTES"
E_MENAGE_XOR = "E_MENAGE_XOR"
E_MENAGE_AVEC_RESERVE = "E_MENAGE_AVEC_RESERVE"
E_RESERVE_STATUT = "E_RESERVE_STATUT"
E_RESERVE_MONTANT = "E_RESERVE_MONTANT"
E_ID_DOUBLON = "E_ID_DOUBLON"
E_ONGLET_ABSENT = "E_ONGLET_ABSENT"


# ── Résultats ────────────────────────────────────────────────────────────────

def _ko(code: str, details: str, **extra: Any) -> dict[str, Any]:
    return {"statut": "ERREUR", "code": code, "details": details, "temp_path": None, **extra}


def _ok(**payload: Any) -> dict[str, Any]:
    return {"statut": "OK", "code": None, "details": None, **payload}


# ── Gardes communes ──────────────────────────────────────────────────────────

def _garde_cible(path: Path) -> dict[str, Any] | None:
    """Refuse tout ce qui n'est pas un SAISIE_* écrivable et libre. None si la cible est saine."""
    if not is_writable(path):
        return _ko(
            E_CHEMIN_NON_AUTORISE,
            f"Écriture refusée sur {path} — seuls les fichiers SAISIE_* hors 02_TRAVAIL / "
            f"MASTER_* / REF_* sont écrivables (file_registry).",
        )
    if not path.exists():
        return _ko(E_FICHIER_ABSENT, f"Fichier cible introuvable : {path}")
    if _detect_excel_lock(path):
        return _ko(
            E_FICHIER_OUVERT,
            f"~${path.name} détecté — le classeur est ouvert dans Excel. Fermer Excel avant écriture.",
        )
    return None


def _creer_temp(path: Path, temp_dir: Path | None) -> tuple[Path | None, dict[str, Any] | None]:
    """Copie la cible vers un temporaire sur le MÊME volume (condition de l'os.replace atomique)."""
    dossier = Path(temp_dir) if temp_dir else path.parent
    with tempfile.NamedTemporaryFile(dir=str(dossier), suffix=path.suffix, delete=False) as f:
        temp = Path(f.name)
    if not _same_volume(path, temp):
        temp.unlink(missing_ok=True)
        return None, _ko(
            E_VOLUME_DIFFERENT,
            f"Temporaire sur un volume différent de {path} — remplacement atomique impossible.",
        )
    shutil.copy2(str(path), str(temp))
    return temp, None


def _structure_preservee(source: Path, temp: Path, *, exiger_full_calc: bool = False) -> list[str]:
    """Onglets, plages nommées, validations, MFC, tables, protections, liens externes.

    `exiger_full_calc` : uniquement pour SAISIE_Charges_Flux, dont les colonnes C/I/J/AD sont des
    formules à recalculer à l'ouverture. Le classeur d'impacts n'en porte aucune.
    """
    try:
        avant, apres = _structure_signature(source), _structure_signature(temp)
    except Exception as exc:  # classeur illisible = corruption
        return [f"structure illisible après écriture : {exc}"]
    ecarts = [
        cle for cle in (
            "sheetnames", "defined_names", "data_validations", "mfc_rules",
            "tables", "protections", "external_links",
        )
        if avant.get(cle) != apres.get(cle)
    ]
    if exiger_full_calc and not apres.get("full_calc_on_load", False):
        ecarts.append("fullCalcOnLoad perdu")
    return ecarts


# ── 1. Ligne de charge (SAISIE_Charges_Flux) ─────────────────────────────────

def _valider_payload_flux(row_data: dict[str, Any]) -> dict[str, Any] | None:
    """Refuse tout champ formule et tout champ inconnu. None si le payload est propre."""
    vises = {str(k) for k in row_data}
    formules = sorted(vises & CHAMPS_FORMULE)
    if formules:
        cols = ", ".join(sorted(FORMULA_COLS))
        return _ko(
            E_COLONNE_FORMULE,
            f"Champs portés par une formule Excel — écriture interdite : {', '.join(formules)} "
            f"(colonnes {cols}). Ces valeurs sont dérivées par le moteur Lot3, jamais saisies.",
        )
    inconnus = sorted(vises - CHAMPS_MANUELS)
    if inconnus:
        return _ko(
            E_CHAMP_INCONNU,
            f"Champs hors schéma SAISIE_Charges_Flux : {', '.join(inconnus)}",
        )
    if not str(row_data.get("charge_id") or "").strip():
        return _ko(E_CHARGE_ID_MANQUANT, "charge_id obligatoire.")
    return None


def _controler_formules_ligne(temp: Path, target_row: int) -> dict[str, Any] | None:
    """La ligne cible doit porter des formules VIVANTES en C/I/J/AD (jamais des valeurs figées)."""
    wb = openpyxl.load_workbook(str(temp), data_only=False)
    try:
        ws = wb[SAISIE_SHEET]
        for col in sorted(FORMULA_COLS):
            v = ws.cell(row=target_row, column=_col_index(col)).value
            if isinstance(v, str) and v.startswith("="):
                continue
            if v is None or (isinstance(v, str) and not v.strip()):
                return _ko(
                    E_FORMULE_MODELE_ABSENTE,
                    f"Colonne {col} ligne {target_row} : formule absente — la ligne modèle "
                    f"n'est pas exploitable.",
                )
            return _ko(
                E_FORMULE_MODELE_FIGEE,
                f"Colonne {col} ligne {target_row} : valeur figée {v!r} au lieu d'une formule.",
            )
    finally:
        wb.close()
    return None


def _delta_flux(source: Path, temp: Path, target_row: int) -> list[str]:
    """Seules les colonnes manuelles de la ligne cible ont le droit d'avoir changé."""
    autorisees = {_col_index(c) for c in MANUAL_COL_MAP if c not in FORMULA_COLS}
    ecarts: list[str] = []
    wb_s = openpyxl.load_workbook(str(source), data_only=False)
    wb_t = openpyxl.load_workbook(str(temp), data_only=False)
    try:
        ws_s, ws_t = wb_s[SAISIE_SHEET], wb_t[SAISIE_SHEET]
        max_row = max(ws_s.max_row or 1, ws_t.max_row or 1)
        max_col = max(ws_s.max_column or 1, ws_t.max_column or 1)
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                v_s = ws_s.cell(row=r, column=c).value
                v_t = ws_t.cell(row=r, column=c).value
                if v_s == v_t:
                    continue
                if r == target_row and c in autorisees:
                    continue
                ecarts.append(f"{get_column_letter(c)}{r} : {v_s!r} → {v_t!r}")
                if len(ecarts) >= 10:
                    return ecarts
    finally:
        wb_s.close()
        wb_t.close()
    return ecarts


def prepare_charge_flux_write(
    saisie_path: Path,
    row_data: dict[str, Any],
    target_row: int,
    *,
    sha256_attendu: str | None = None,
    temp_dir: Path | None = None,
) -> dict[str, Any]:
    """Prépare l'écriture d'UNE ligne de charge dans un temporaire validé. Ne remplace rien.

    `sha256_attendu` : empreinte lue au moment de la prévisualisation. Si la source a bougé depuis,
    la préparation est refusée (E_SOURCE_MODIFIEE) — jamais d'écriture sur une base périmée.

    Réécrire le MÊME charge_id sur la MÊME ligne est autorisé (remplacement idempotent) ; le même
    charge_id sur une AUTRE ligne est un refus (E_CHARGE_ID_COLLISION).

    Retour : {"statut": "OK", "temp_path", "sha256_source", "sha256_temp", "charge_id", "target_row"}
             ou {"statut": "ERREUR", "code", "details"}.
    """
    saisie_path = Path(saisie_path)
    if (garde := _garde_cible(saisie_path)) is not None:
        return garde
    if (invalide := _valider_payload_flux(row_data)) is not None:
        return invalide

    sha_avant = _sha256(saisie_path)
    if sha256_attendu and sha256_attendu != sha_avant:
        return _ko(
            E_SOURCE_MODIFIEE,
            f"La source a changé depuis la prévisualisation "
            f"(attendu {sha256_attendu[:12]}…, lu {sha_avant[:12]}…).",
        )

    charge_id = str(row_data["charge_id"]).strip()

    # Unicité du charge_id + cohérence de la ligne cible (lecture seule sur la source).
    wb = openpyxl.load_workbook(str(saisie_path), read_only=True, data_only=True)
    try:
        ws = wb[SAISIE_SHEET]
        lignes_par_id: dict[str, int] = {}
        for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            if idx == 1 or not row:
                continue
            cid = str(row[0] or "").strip()
            if cid:
                lignes_par_id.setdefault(cid, idx)
        cible_actuelle = lignes_par_id.get(charge_id)
    finally:
        wb.close()

    if cible_actuelle is not None and cible_actuelle != target_row:
        return _ko(
            E_CHARGE_ID_COLLISION,
            f"charge_id={charge_id} déjà présent ligne {cible_actuelle} — une charge = UNE ligne.",
        )
    occupant = next((cid for cid, r in lignes_par_id.items() if r == target_row), None)
    if occupant is not None and occupant != charge_id:
        return _ko(
            E_LIGNE_CIBLE_OCCUPEE,
            f"Ligne {target_row} déjà occupée par charge_id={occupant}.",
        )

    temp, erreur = _creer_temp(saisie_path, temp_dir)
    if erreur is not None:
        return erreur

    try:
        if (formules_ko := _controler_formules_ligne(temp, target_row)) is not None:
            temp.unlink(missing_ok=True)
            return formules_ko

        # Écriture : colonnes manuelles uniquement. Les indices formule sont sautés — double
        # verrou avec la validation du payload (une régression du mapping ne peut pas passer).
        wb = openpyxl.load_workbook(str(temp), data_only=False)
        try:
            ws = wb[SAISIE_SHEET]
            for col, champ in MANUAL_COL_MAP.items():
                idx = _col_index(col)
                if idx in FORMULA_COL_INDICES:
                    continue
                if champ in row_data:
                    ws.cell(row=target_row, column=idx, value=row_data[champ])
            wb.calculation.fullCalcOnLoad = True
            wb.save(str(temp))
        finally:
            wb.close()

        # Revalidation des valeurs effectivement écrites.
        wb = openpyxl.load_workbook(str(temp), data_only=False)
        try:
            ws = wb[SAISIE_SHEET]
            ecarts = []
            for col, champ in MANUAL_COL_MAP.items():
                if champ not in row_data or _col_index(col) in FORMULA_COL_INDICES:
                    continue
                attendu, lu = row_data[champ], ws.cell(row=target_row, column=_col_index(col)).value
                if attendu is None and lu is None:
                    continue
                if attendu == "" and lu is None:
                    continue
                if lu != attendu:
                    ecarts.append(f"{col}/{champ} : attendu {attendu!r}, lu {lu!r}")
        finally:
            wb.close()
        if ecarts:
            temp.unlink(missing_ok=True)
            return _ko(E_VALEUR_NON_ECRITE, "; ".join(ecarts))

        if delta := _delta_flux(saisie_path, temp, target_row):
            temp.unlink(missing_ok=True)
            return _ko(E_DELTA_HORS_PERIMETRE, "; ".join(delta))

        if ecarts_struct := _structure_preservee(saisie_path, temp, exiger_full_calc=True):
            temp.unlink(missing_ok=True)
            return _ko(E_STRUCTURE_ALTEREE, "; ".join(ecarts_struct))

        if _sha256(saisie_path) != sha_avant:
            temp.unlink(missing_ok=True)
            return _ko(E_SOURCE_MODIFIEE, "La source a été modifiée pendant la préparation.")

        return _ok(
            temp_path=temp,
            sha256_source=sha_avant,
            sha256_temp=_sha256(temp),
            charge_id=charge_id,
            target_row=target_row,
            remplacement=cible_actuelle is not None,   # True = réécriture idempotente
        )
    except Exception:
        temp.unlink(missing_ok=True)
        raise


# ── 2. Impacts (SAISIE_Charges_Impacts) ──────────────────────────────────────

def _valider_persistable(persistable: dict[str, Any], montant: float) -> dict[str, Any] | None:
    """Règles métier des impacts. Refus explicite — jamais de correction silencieuse."""
    charge_id = str(persistable.get("charge_id") or "").strip()
    if not charge_id:
        return _ko(E_CHARGE_ID_MANQUANT, "charge_id obligatoire dans le persistable.")

    affectations = persistable.get("affectations") or []
    menage = persistable.get("menage") or []
    reserve = persistable.get("reserve") or []

    for nom, rows in (("affectations", affectations), ("menage", menage), ("reserve", reserve)):
        for r in rows:
            if str(r.get("charge_id") or "").strip() != charge_id:
                return _ko(
                    E_CHARGE_ID_MANQUANT,
                    f"{nom} : une ligne porte charge_id={r.get('charge_id')!r} "
                    f"au lieu de {charge_id!r}.",
                )

    # AFFECTATIONS : Σ quote_part = montant (jamais une charge répliquée intégralement).
    if affectations:
        somme = round(sum(float(r.get("quote_part") or 0) for r in affectations), 2)
        if abs(somme - float(montant)) > TOLERANCE_CENTIMES:
            return _ko(
                E_SOMME_QUOTES,
                f"AFFECTATIONS : Σ quote_part = {somme} ≠ montant {montant}.",
            )

    # MENAGE : intervenant XOR logement, ligne par ligne.
    for r in menage:
        a_int = bool(str(r.get("intervenant_id") or "").strip())
        a_log = bool(str(r.get("logement_id") or "").strip())
        if a_int == a_log:
            return _ko(
                E_MENAGE_XOR,
                f"MENAGE {r.get('menage_impact_id')!r} : intervenant_id et logement_id doivent "
                f"être renseignés l'un OU l'autre, jamais les deux ni aucun.",
            )

    # Une charge ménage n'a JAMAIS de réserve refacturable.
    if menage and reserve:
        return _ko(
            E_MENAGE_AVEC_RESERVE,
            "Une charge ménage ne peut pas porter de réserve de refacturation.",
        )

    # RESERVE : statut EN_ATTENTE (aucune application automatique en préfacture), montants cohérents.
    for r in reserve:
        if str(r.get("statut_traitement") or "").strip().upper() != STATUT_RESERVE_ATTENDU:
            return _ko(
                E_RESERVE_STATUT,
                f"RESERVE {r.get('reserve_id')!r} : statut_traitement="
                f"{r.get('statut_traitement')!r} — attendu {STATUT_RESERVE_ATTENDU} "
                f"(aucune refacturation automatique).",
            )
        m = float(r.get("montant_refacturable") or 0)
        if m <= 0:
            return _ko(
                E_RESERVE_MONTANT,
                f"RESERVE {r.get('reserve_id')!r} : montant_refacturable={m} — doit être > 0.",
            )
    if reserve:
        somme_res = round(sum(float(r.get("montant_refacturable") or 0) for r in reserve), 2)
        if somme_res - float(montant) > TOLERANCE_CENTIMES:
            return _ko(
                E_RESERVE_MONTANT,
                f"RESERVE : Σ montant_refacturable = {somme_res} > montant de la charge {montant}.",
            )
    return None


def _lignes(ws) -> tuple[list[str], list[dict[str, Any]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = [
        dict(zip(headers, r))
        for r in rows[1:]
        if any(c is not None for c in r)
    ]
    return headers, out


def _remplacer_lignes(ws, charge_id: str, rows: list[dict[str, Any]]) -> None:
    """Idempotent : supprime les lignes du charge_id puis ré-ajoute. Jamais d'accumulation."""
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col_charge = headers.index("charge_id") + 1 if "charge_id" in headers else None
    if col_charge:
        for r in range(ws.max_row, 1, -1):
            if str(ws.cell(r, col_charge).value or "").strip() == charge_id:
                ws.delete_rows(r, 1)
    for row in rows:
        ws.append([row.get(h) for h in headers])


def prepare_charge_impacts_write(
    impacts_path: Path,
    persistable: dict[str, Any],
    montant: float,
    *,
    sha256_attendu: str | None = None,
    temp_dir: Path | None = None,
) -> dict[str, Any]:
    """Prépare l'écriture des impacts dans un temporaire validé. Ne remplace rien.

    Idempotent : rejouer le même persistable remplace les lignes du charge_id, ne les duplique
    jamais. Les lignes des AUTRES charges doivent rester strictement inchangées.
    """
    impacts_path = Path(impacts_path)
    if (garde := _garde_cible(impacts_path)) is not None:
        return garde
    if (invalide := _valider_persistable(persistable, montant)) is not None:
        return invalide

    charge_id = str(persistable["charge_id"]).strip()
    attendus = {
        AFFECT_SHEET: (persistable.get("affectations") or [], "affectation_id"),
        MENAGE_SHEET: (persistable.get("menage") or [], "menage_impact_id"),
        RESERVE_SHEET: (persistable.get("reserve") or [], "reserve_id"),
    }

    sha_avant = _sha256(impacts_path)
    if sha256_attendu and sha256_attendu != sha_avant:
        return _ko(
            E_SOURCE_MODIFIEE,
            f"La source a changé depuis la prévisualisation "
            f"(attendu {sha256_attendu[:12]}…, lu {sha_avant[:12]}…).",
        )

    # État des autres charges avant écriture — elles ne doivent pas bouger d'un iota.
    wb = openpyxl.load_workbook(str(impacts_path), read_only=True, data_only=True)
    try:
        manquants = [s for s in attendus if s not in wb.sheetnames]
        autres_avant = {}
        if not manquants:
            for sheet in attendus:
                _, rows = _lignes(wb[sheet])
                autres_avant[sheet] = [
                    r for r in rows if str(r.get("charge_id") or "").strip() != charge_id
                ]
    finally:
        wb.close()
    if manquants:
        return _ko(E_ONGLET_ABSENT, f"Onglets absents de {impacts_path.name} : {', '.join(manquants)}")

    temp, erreur = _creer_temp(impacts_path, temp_dir)
    if erreur is not None:
        return erreur

    try:
        wb = openpyxl.load_workbook(str(temp))
        try:
            for sheet, (rows, _id_col) in attendus.items():
                _remplacer_lignes(wb[sheet], charge_id, rows)
            wb.save(str(temp))
        finally:
            wb.close()

        # Relecture : comptes exacts, identifiants uniques, autres charges intactes.
        ecrits: dict[str, int] = {}
        wb = openpyxl.load_workbook(str(temp), read_only=True, data_only=True)
        try:
            for sheet, (rows_attendues, id_col) in attendus.items():
                _, rows = _lignes(wb[sheet])
                mine = [r for r in rows if str(r.get("charge_id") or "").strip() == charge_id]
                autres = [r for r in rows if str(r.get("charge_id") or "").strip() != charge_id]
                ecrits[sheet] = len(mine)

                if len(mine) != len(rows_attendues):
                    temp.unlink(missing_ok=True)
                    return _ko(
                        E_ID_DOUBLON,
                        f"{sheet} : {len(mine)} ligne(s) pour charge_id={charge_id}, "
                        f"{len(rows_attendues)} attendue(s) — duplication.",
                    )
                ids = [str(r.get(id_col) or "").strip() for r in rows if str(r.get(id_col) or "").strip()]
                if len(ids) != len(set(ids)):
                    temp.unlink(missing_ok=True)
                    return _ko(E_ID_DOUBLON, f"{sheet} : {id_col} en doublon.")
                if autres != autres_avant[sheet]:
                    temp.unlink(missing_ok=True)
                    return _ko(
                        E_DELTA_HORS_PERIMETRE,
                        f"{sheet} : les lignes d'autres charges ont été modifiées.",
                    )
        finally:
            wb.close()

        if ecarts_struct := _structure_preservee(impacts_path, temp):
            temp.unlink(missing_ok=True)
            return _ko(E_STRUCTURE_ALTEREE, "; ".join(ecarts_struct))

        if _sha256(impacts_path) != sha_avant:
            temp.unlink(missing_ok=True)
            return _ko(E_SOURCE_MODIFIEE, "La source a été modifiée pendant la préparation.")

        return _ok(
            temp_path=temp,
            sha256_source=sha_avant,
            sha256_temp=_sha256(temp),
            charge_id=charge_id,
            affectations_ecrites=ecrits[AFFECT_SHEET],
            menage_ecrits=ecrits[MENAGE_SHEET],
            reserve_ecrites=ecrits[RESERVE_SHEET],
        )
    except Exception:
        temp.unlink(missing_ok=True)
        raise
