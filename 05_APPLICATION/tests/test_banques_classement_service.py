"""File humaine de classement Banque (A_ENVOYER_IA, migration 0026). Fixtures fictives uniquement."""
from __future__ import annotations

import openpyxl
import pytest

import app.config as cfg
from app.services import banques_classement_service as svc
from app.services import banques_controle_service as ctrl_svc

NORM_HDR = [
    "mouvement_id", "ROW_HASH", "import_id", "ligne_source", "date_operation", "date_valeur",
    "libelle", "libelle_brut", "montant", "sens", "devise", "compte_id", "tiers_detecte",
    "categorie", "type_flux_id", "code_impact", "source_classification", "source_economique",
    "statut_controle", "niveau_risque", "codes_anomalie", "date_integration", "commentaire",
    "statut_classification", "niveau_anomalie", "regle_id_appliquee",
]


def _ligne(mid, montant=8.64, sens="DEBIT", categorie="", statut_classification="A_ENVOYER_IA",
          date_operation="2026-01-05"):
    return [mid, f"HASH-{mid}", "IMP1", 1, date_operation, date_operation,
           "PAIEMENT CB TEST", "PAIEMENT CB TEST", montant, sens, "EUR", "CM_TEST", None,
           categorie, None, "", "REGLE_DETERMINISTE", "", "A_CONTROLER", "MOYEN", "", "2026-01-31",
           "", statut_classification, "A_CONTROLER", "R_099"]


@pytest.fixture
def source(tmp_path, monkeypatch):
    p = tmp_path / "BANQUE_LOT8_IMPORT.xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("NORM_Banque")
    ws.append(NORM_HDR)
    ws.append(_ligne("MVT-TEST-001"))
    ws.append(_ligne("MVT-TEST-002", montant=120.0, sens="CREDIT"))
    ws.append(_ligne("MVT-TEST-003", montant=45.0, statut_classification="CLASSE", categorie="LOGICIEL_GESTION"))
    wb.save(p); wb.close()
    monkeypatch.setattr(cfg, "MASTER_BANQUE", p)
    ctrl_svc.vider_cache()
    yield p
    ctrl_svc.vider_cache()


def _opq(mid):
    return ctrl_svc.id_opaque(mid)


def test_01_lister_ne_retient_que_a_envoyer_ia(source, tmp_db):
    lignes = svc.lister(db_path=tmp_db)
    assert len(lignes) == 2
    assert all(l["statut_humain"] == "SANS_DECISION" for l in lignes)


def test_02_compter(source):
    assert svc.compter() == 2


@pytest.fixture
def source_avec_doublon(tmp_path, monkeypatch):
    """Reproduit un doublon physique de mouvement_id (même ligne source dupliquée en amont, ex.
    lot8a) — garantie requise : jamais deux lignes de file ni un double comptage pour ce cas."""
    p = tmp_path / "BANQUE_LOT8_IMPORT.xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("NORM_Banque")
    ws.append(NORM_HDR)
    ws.append(_ligne("MVT-DUP-001", montant=120.0))
    ws.append(_ligne("MVT-DUP-001", montant=120.0))  # doublon exact du mouvement_id ci-dessus
    ws.append(_ligne("MVT-TEST-002", montant=45.0))
    wb.save(p); wb.close()
    monkeypatch.setattr(cfg, "MASTER_BANQUE", p)
    ctrl_svc.vider_cache()
    yield p
    ctrl_svc.vider_cache()


def test_02b_doublon_mouvement_id_compte_une_seule_fois(source_avec_doublon):
    assert svc.compter() == 2  # MVT-DUP-001 (une fois) + MVT-TEST-002


def test_02c_doublon_mouvement_id_une_seule_ligne_dans_la_file(source_avec_doublon, tmp_db):
    lignes = svc.lister(db_path=tmp_db)
    assert len(lignes) == 2
    ids_opaques = [l["id_opaque"] for l in lignes]
    assert len(ids_opaques) == len(set(ids_opaques)), "aucun id_opaque ne doit apparaître deux fois"


def test_02d_doublon_mouvement_id_une_seule_decision_economique_possible(source_avec_doublon, tmp_db):
    """Une décision prise sur l'opaque du mouvement dupliqué s'applique une seule fois — pas de
    second enregistrement indépendant possible pour la même ligne source dupliquée."""
    opq = _opq("MVT-DUP-001")
    svc.decider(opq, "MAINTENIR_A_CONTROLER", db_path=tmp_db)
    lignes = svc.lister(db_path=tmp_db)
    ligne = next(l for l in lignes if l["id_opaque"] == opq)
    assert ligne["statut_humain"] == "MAINTENIR_A_CONTROLER"
    # Toujours une seule ligne pour cet opaque, même après décision.
    assert sum(1 for l in lignes if l["id_opaque"] == opq) == 1


def test_03_filtre_montant(source, tmp_db):
    assert len(svc.lister(montant_min=100, db_path=tmp_db)) == 1


def test_04_filtre_sens(source, tmp_db):
    assert len(svc.lister(sens="CREDIT", db_path=tmp_db)) == 1


def test_05_charger_detail(source, tmp_db):
    d = svc.charger(_opq("MVT-TEST-001"), db_path=tmp_db)
    assert d is not None and d["categorie_moteur"] == ""
    assert d["historique"] == []


def test_06_mouvement_introuvable(source):
    assert svc.charger("MVT-INEXISTANT") is None


def test_07_previsualiser_categoriser(source):
    r = svc.previsualiser(_opq("MVT-TEST-001"), "CATEGORISER", nouvelle_categorie="LOGICIEL_GESTION")
    assert r["ok"] and r["apercu"]["nouvelle_categorie"] == "LOGICIEL_GESTION"


def test_08_categorie_invalide_refusee(source):
    r = svc.previsualiser(_opq("MVT-TEST-001"), "CATEGORISER", nouvelle_categorie="INVENTEE")
    assert not r["ok"] and r["code"] == svc.E_CATEGORIE_INCONNUE


def test_09_categorie_manquante_refusee(source):
    r = svc.previsualiser(_opq("MVT-TEST-001"), "CATEGORISER")
    assert not r["ok"] and r["code"] == svc.E_CATEGORIE_MANQUANTE


def test_10_type_decision_invalide(source):
    r = svc.previsualiser(_opq("MVT-TEST-001"), "TYPE_INVENTE")
    assert not r["ok"] and r["code"] == svc.E_TYPE_DECISION_INCONNU


def test_11_decider_maintenir_a_controler(source, tmp_db):
    r = svc.decider(_opq("MVT-TEST-001"), "MAINTENIR_A_CONTROLER", db_path=tmp_db)
    assert r["ok"]
    d = svc.charger(_opq("MVT-TEST-001"), db_path=tmp_db)
    assert d["statut_humain"] == "MAINTENIR_A_CONTROLER"


def test_12_decider_non_classe(source, tmp_db):
    r = svc.decider(_opq("MVT-TEST-001"), "NON_CLASSE", db_path=tmp_db)
    assert r["ok"]
    assert svc.charger(_opq("MVT-TEST-001"), db_path=tmp_db)["statut_humain"] == "NON_CLASSE"


def test_13_decider_reporter(source, tmp_db):
    r = svc.decider(_opq("MVT-TEST-001"), "REPORTER", db_path=tmp_db)
    assert r["ok"]


def test_14_decider_avec_justification(source, tmp_db):
    svc.decider(_opq("MVT-TEST-001"), "REPORTER", justification="à vérifier avec le relevé suivant",
               db_path=tmp_db)
    h = svc.historique(_opq("MVT-TEST-001"), db_path=tmp_db)
    assert h[0]["justification"] == "à vérifier avec le relevé suivant"


def test_15_signaler_anomalie_moteur(source, tmp_db):
    r = svc.decider(_opq("MVT-TEST-001"), "REPORTER", anomalie_moteur="libellé tronqué", db_path=tmp_db)
    assert r["ok"]
    assert svc.historique(_opq("MVT-TEST-001"), db_path=tmp_db)[0]["anomalie_moteur"] == "libellé tronqué"


def test_16_proposer_future_regle(source, tmp_db):
    svc.decider(_opq("MVT-TEST-001"), "REPORTER", future_regle="détecter PAIEMENT CB + montant < 10€",
               db_path=tmp_db)
    assert svc.historique(_opq("MVT-TEST-001"), db_path=tmp_db)[0]["future_regle"]


def test_17_historique_append_only(source, tmp_db):
    svc.decider(_opq("MVT-TEST-001"), "REPORTER", db_path=tmp_db)
    svc.decider(_opq("MVT-TEST-001"), "NON_CLASSE", db_path=tmp_db)
    h = svc.historique(_opq("MVT-TEST-001"), db_path=tmp_db)
    assert len(h) == 2   # les deux décisions restent, rien supprimé


def test_18_seconde_decision_devient_active(source, tmp_db):
    svc.decider(_opq("MVT-TEST-001"), "REPORTER", db_path=tmp_db)
    svc.decider(_opq("MVT-TEST-001"), "NON_CLASSE", db_path=tmp_db)
    d = svc.charger(_opq("MVT-TEST-001"), db_path=tmp_db)
    assert d["statut_humain"] == "NON_CLASSE"   # la plus récente fait foi
    h = svc.historique(_opq("MVT-TEST-001"), db_path=tmp_db)
    actifs = [x for x in h if x["actif"] == 1]
    assert len(actifs) == 1 and actifs[0]["type_decision"] == "NON_CLASSE"


def test_19_aucune_ia_externe_aucun_reseau(source):
    import inspect
    src = inspect.getsource(svc)
    for interdit in ("requests.", "httpx.", "openai", "anthropic", "urllib.request"):
        assert interdit not in src


def test_20_aucune_pii_dans_les_decisions(source, tmp_db):
    svc.decider(_opq("MVT-TEST-001"), "CATEGORISER", nouvelle_categorie="LOGICIEL_GESTION",
               db_path=tmp_db)
    h = svc.historique(_opq("MVT-TEST-001"), db_path=tmp_db)
    assert "iban" not in str(h).lower() and "compte_id" not in str(h)


def test_21_aucun_libelle_complet_dans_la_liste(source, tmp_db):
    lignes = svc.lister(db_path=tmp_db)
    assert all("PAIEMENT CB TEST" != l.get("libelle_brut") for l in lignes)  # champ jamais exposé


def test_22_aucune_ecriture_excel(source, tmp_db):
    import os
    mtime_avant = os.path.getmtime(source)
    svc.decider(_opq("MVT-TEST-001"), "NON_CLASSE", db_path=tmp_db)
    assert os.path.getmtime(source) == mtime_avant   # le fichier Excel n'a jamais bougé


def test_23_idempotence_meme_decision_repetee(source, tmp_db):
    svc.decider(_opq("MVT-TEST-001"), "REPORTER", db_path=tmp_db)
    svc.decider(_opq("MVT-TEST-001"), "REPORTER", db_path=tmp_db)
    h = svc.historique(_opq("MVT-TEST-001"), db_path=tmp_db)
    assert len(h) == 2   # deux événements distincts, pas une fusion — mais un seul actif
    assert sum(1 for x in h if x["actif"] == 1) == 1
