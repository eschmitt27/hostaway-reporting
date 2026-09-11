"""§29 — lot6f ne LIT plus aucun classeur.

Mission « FIN DU LEGACY CHARGES / MÉNAGES », DÉCISION 1 : « ZÉRO EXCEL OPÉRATIONNEL ». lot6f
lisait quatre sources Excel — REF_Setup.xlsm (référentiels), MASTER_FACT_MEN_MenagesExternes.xlsx
(ménages externes), la Google Sheet M04 (déclarations internes + lavage) et
SAISIE_Charges_Flux.xlsx (pools courses/consommables). Le chemin `--source EXCEL` a été supprimé
après preuve d'équivalence sur 2026-06/07/08 ; ces tests empêchent qu'il revienne.

Deux niveaux, comme `test_lot6b_anti_excel` :
  - STRUCTUREL : le code ne contient plus ni branche EXCEL, ni lecture de classeur, ni appel
    réseau vers la feuille.
  - COMPORTEMENTAL : un run réel n'OUVRE EN LECTURE aucun `.xlsx/.xlsm/.xls`, prouvé par un
    audit hook CPython qui voit tous les `open()` — y compris ceux faits en C par `zipfile`
    sous `openpyxl`. L'écriture du classeur de SORTIE reste permise : lot11, lot13/PowerBI et
    `menages_reader` le consomment encore.
"""
from __future__ import annotations

import hashlib
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pytest

import app.config as cfg

_TRAVAIL = Path(cfg.PROJECT_ROOT) / "02_TRAVAIL"
LOT6F = _TRAVAIL / "lot6f_cout_complet_menages.py"

pytestmark = pytest.mark.skipif(not LOT6F.exists(), reason="lot6f absent")

SOURCE = LOT6F.read_text(encoding="utf-8", errors="replace")


def _code_seul(source: str) -> str:
    """Le source privé de sa docstring de module et de ses commentaires.

    Le fichier explique en prose les classeurs qu'il ne lit PLUS ; viser le code seul évite qu'une
    documentation honnête fasse échouer le garde-fou — et évite surtout la tentation inverse, qui
    serait d'effacer l'explication pour faire passer un test."""
    import ast

    lignes = source.splitlines()
    arbre = ast.parse(source)
    if (arbre.body and isinstance(arbre.body[0], ast.Expr)
            and isinstance(arbre.body[0].value, ast.Constant)
            and isinstance(arbre.body[0].value.value, str)):
        debut = arbre.body[0].lineno - 1
        fin = arbre.body[0].end_lineno
        lignes = lignes[:debut] + lignes[fin:]
    return "\n".join(l for l in lignes if not l.lstrip().startswith("#"))


CODE_TXT = _code_seul(SOURCE)

EXTENSIONS = (".xlsx", ".xlsm", ".xls")

#: L'interpréteur qui exécute les tests : il a les dépendances des moteurs et il existe, ce que
#: `cfg.LOT4A_ENGINE_PYTHON` (chemin Windows codé en dur, surchargeable par variable
#: d'environnement) ne garantit sur aucun poste.
PYTHON = sys.executable


# ── Niveau STRUCTUREL ────────────────────────────────────────────────────────

def test_le_mode_excel_n_existe_plus():
    """`--source` n'admet plus qu'une valeur : un `--source EXCEL` résiduel doit ÉCHOUER.

    Le drapeau survit parce que des appelants le passent (orchestrateur_moteur,
    menages_chaine_service, run_menages_pipeline) ; c'est `choices` qui garantit qu'aucun chemin
    Excel ne peut être ré-emprunté en silence."""
    assert '_ap.add_argument("--source", choices=("SQLITE",)' in CODE_TXT
    assert '"EXCEL"' not in CODE_TXT, "plus aucune valeur EXCEL dans le code"
    assert "args.source ==" not in CODE_TXT, "plus aucun aiguillage sur la source"


def test_aucune_lecture_de_classeur_dans_le_code():
    """`load_workbook` servait à lire ; seule l'écriture (`Workbook()` + `save`) doit subsister."""
    assert "load_workbook" not in CODE_TXT
    assert "def sh(" not in CODE_TXT, "l'aide de lecture de feuille doit avoir disparu"
    assert "openpyxl.Workbook()" in CODE_TXT and "wb.save(OUT)" in CODE_TXT


def test_aucun_chemin_vers_les_sources_excel():
    """Les constantes de chemin vers les classeurs sources ne doivent plus être construites."""
    for nom in ("REF_Setup", "SAISIE_Charges_Flux", "MASTER_FACT_MEN_MenagesExternes",
                "01_SOURCES_BRUTES"):
        assert nom not in CODE_TXT, f"{nom} est encore référencé par le code"


def test_aucun_appel_reseau_vers_la_google_sheet():
    """La feuille M04 n'est plus lue ici : lot6b la résout et écrit `menages_declarations_internes`.

    Garder un second lecteur de la feuille, c'était garder deux mappings de libellés et de prénoms
    qui finiraient par diverger — c'est exactement l'écart trouvé lors de la bascule."""
    assert "fetch_sheet_csv" not in CODE_TXT
    assert "SHEET_URL" not in CODE_TXT
    assert "INTMAP" not in CODE_TXT, "la table de prénoms en dur doit avoir disparu"
    assert "menages_declarations_internes" in CODE_TXT


def test_pas_de_provenance_sheet_posee_par_lot6f():
    """lot6f ne peut plus attester une lecture réseau qu'il n'effectue pas.

    `begin_step`/`commit_step` décrivaient un appel à la feuille. Les laisser produirait une
    provenance mensongère ; les retirer sans retirer lot6f du contrôle lot11 laisserait
    `SOURCE_SHEET_PROVENANCE_INCOMPLETE` rouge en permanence — les deux vont ensemble."""
    assert "begin_step" not in CODE_TXT and "commit_step" not in CODE_TXT
    from app.services import controles_lot11_service as c11
    src11 = Path(c11.__file__).read_text(encoding="utf-8", errors="replace")
    assert 'etapes = {"lot6b": "menages_declarations_internes"}' in src11, \
        "lot6f ne doit plus être une étape à provenance propre"


# ── Niveau COMPORTEMENTAL ────────────────────────────────────────────────────

_SITECUSTOMIZE = '''\
"""Espion d'ouvertures de fichiers, chargé automatiquement par CPython au démarrage."""
import atexit, os, sys

_LOG = os.environ.get("AUDIT_OPEN_LOG")
if _LOG:
    _vus = []
    _reentrant = False

    def _hook(evenement, args):
        # Le hook ouvre lui-même un fichier pour écrire son journal : sans garde, il se
        # rappellerait indéfiniment.
        global _reentrant
        if evenement != "open" or _reentrant:
            return
        _reentrant = True
        try:
            _vus.append("%s\\t%s" % (args[0], args[1]))
        except Exception:
            pass
        finally:
            _reentrant = False

    def _vider():
        global _reentrant
        _reentrant = True
        with open(_LOG, "w", encoding="utf-8", errors="replace") as fh:
            fh.write("\\n".join(str(x) for x in _vus))

    atexit.register(_vider)
    sys.addaudithook(_hook)
'''


@pytest.fixture(scope="module")
def racine_isolee(tmp_path_factory) -> Path:
    """Une racine projet ne contenant QUE les moteurs — aucun classeur, aucun `01_SOURCES_BRUTES`.

    lot6f déduit sa racine de `__file__`. L'y exécuter prouve deux choses d'un coup : qu'il aboutit
    quand les sources Excel sont matériellement absentes (§9 de la mission), et qu'il n'écrase pas
    le classeur de sortie versionné du dépôt pendant les tests.
    """
    import shutil

    racine = tmp_path_factory.mktemp("racine_sans_excel")
    (racine / "02_TRAVAIL").mkdir()
    for module in _TRAVAIL.glob("*.py"):
        shutil.copy2(module, racine / "02_TRAVAIL" / module.name)

    restants = [p for p in racine.rglob("*") if p.suffix.lower() in EXTENSIONS]
    assert not restants, f"la racine isolée doit être vierge de classeurs : {restants}"
    return racine


def _base_migree(tmp_path: Path) -> Path:
    from app.db.connection import apply_migrations

    db = tmp_path / "app.db"
    apply_migrations(db)
    return db


def _run_lot6f_espionne(racine: Path, tmp_path: Path, db: Path, *args: str):
    """Lance lot6f en sous-processus sous audit hook, depuis la racine isolée.

    Rend (process, ouvertures)."""
    espion = tmp_path / "espion"
    espion.mkdir(exist_ok=True)
    (espion / "sitecustomize.py").write_text(_SITECUSTOMIZE, encoding="utf-8")
    journal = tmp_path / "ouvertures.txt"

    travail = racine / "02_TRAVAIL"
    env = dict(os.environ)
    env["PYTHONPATH"] = str(espion)
    env["AUDIT_OPEN_LOG"] = str(journal)
    env["PYTHONIOENCODING"] = "utf-8"
    env.pop("PILOTAGE_DB_PATH", None)
    env.pop("APP_DATA_DIR", None)

    proc = subprocess.run(
        [PYTHON, str(travail / LOT6F.name), "--db", str(db), *args],
        cwd=str(travail), env=env, capture_output=True, text=True, timeout=600)

    lignes = journal.read_text(encoding="utf-8", errors="replace").splitlines() \
        if journal.exists() else []
    ouvertures = []
    for ligne in lignes:
        chemin, _, mode = ligne.partition("\t")
        ouvertures.append((chemin, mode))
    return proc, ouvertures


def _classeurs(ouvertures, *, en_lecture: bool):
    resultat = []
    for chemin, mode in ouvertures:
        if not chemin.lower().endswith(EXTENSIONS):
            continue
        lecture = ("r" in mode and "+" not in mode) if mode and mode != "None" else True
        if lecture is en_lecture:
            resultat.append((chemin, mode))
    return resultat


def test_un_run_reel_n_ouvre_aucun_classeur_en_lecture(racine_isolee, tmp_path):
    """PREUVE CENTRALE (§29) — aucun `.xlsx/.xlsm/.xls` n'est ouvert en lecture.

    L'audit hook voit les ouvertures faites en C par `zipfile` sous `openpyxl` : une lecture de
    classeur ne peut pas lui échapper. Le run tourne sur une base migrée VIDE — le calcul n'a rien
    à produire, mais toutes les résolutions de source sont traversées, et c'est ce qui est
    observé."""
    db = _base_migree(tmp_path)
    proc, ouvertures = _run_lot6f_espionne(racine_isolee, tmp_path, db, "--mois", "2026-07")

    assert ouvertures, "l'espion n'a rien enregistré — le hook n'a pas été installé"
    assert proc.returncode == 0, (proc.stdout + proc.stderr)[-3000:]

    lus = _classeurs(ouvertures, en_lecture=True)
    assert not lus, f"lot6f a lu un classeur : {lus}"


def test_le_classeur_de_sortie_reste_ecrit(racine_isolee, tmp_path):
    """Contre-épreuve : l'espion fonctionne bien, puisqu'il voit l'ÉCRITURE de la sortie.

    Sans cette assertion, un hook cassé rendrait le test précédent vert pour une mauvaise raison —
    « aucune lecture » serait vrai parce qu'aucune ouverture n'est vue du tout."""
    db = _base_migree(tmp_path)
    _, ouvertures = _run_lot6f_espionne(racine_isolee, tmp_path, db, "--mois", "2026-07")

    ecrits = _classeurs(ouvertures, en_lecture=False)
    assert any("CoutComplet" in c for c, _ in ecrits), \
        f"la sortie lot6f devrait être écrite ; vu : {ecrits}"


def test_sans_excel_n_ecrit_meme_pas_la_sortie(racine_isolee, tmp_path):
    """`--sans-excel` porte sur la SORTIE : plus aucun classeur n'est touché, ni lu ni écrit."""
    db = _base_migree(tmp_path)
    proc, ouvertures = _run_lot6f_espionne(
        racine_isolee, tmp_path, db, "--mois", "2026-07", "--sans-excel")

    assert proc.returncode == 0, (proc.stdout + proc.stderr)[-3000:]
    tous = _classeurs(ouvertures, en_lecture=True) + _classeurs(ouvertures, en_lecture=False)
    assert not tous, f"--sans-excel ne doit toucher aucun classeur ; vu : {tous}"


def test_les_classeurs_sources_restent_intacts(racine_isolee, tmp_path):
    """Les classeurs sources ne sont ni lus ni modifiés — leur empreinte ne bouge pas.

    Ils ne sont PAS supprimés : REF_Setup et SAISIE_Charges_Flux restent des pièces d'origine et
    des sources d'import d'autres modules (cf. RESTANT_EXCEL_OPERATIONNEL.md). Ce qui disparaît
    est la LECTURE runtime par lot6f, pas le fichier."""
    sources = [
        Path(cfg.PROJECT_ROOT) / "01_SOURCES_BRUTES" / "REF_Setup" / "REF_Setup.xlsm",
        Path(cfg.PROJECT_ROOT) / "01_SOURCES_BRUTES" / "Charges" / "SAISIE_Charges_Flux.xlsx",
    ]
    avant = {p: (hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None)
             for p in sources}

    db = _base_migree(tmp_path)
    _run_lot6f_espionne(racine_isolee, tmp_path, db, "--mois", "2026-07", "--sans-excel")

    for chemin, empreinte in avant.items():
        apres = hashlib.sha256(chemin.read_bytes()).hexdigest() if chemin.exists() else None
        assert apres == empreinte, f"{chemin.name} a été modifié"


def test_une_base_absente_echoue_clairement(tmp_path):
    """FAIL-CLOSED — sans base, lot6f s'arrête en nommant ce qu'il attend, il ne replie pas.

    L'ancien comportement était de basculer sur le classeur ; ce repli est précisément ce que la
    mission interdit de conserver."""
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    for cle in ("PILOTAGE_DB_PATH", "APP_DATA_DIR"):
        env.pop(cle, None)

    proc = subprocess.run(
        [PYTHON, str(LOT6F), "--mois", "2026-07"],
        cwd=str(_TRAVAIL), env=env, capture_output=True, text=True, timeout=300)

    sortie = proc.stdout + proc.stderr
    assert proc.returncode != 0, "une base absente doit être bloquante"
    assert "PILOTAGE_DB_PATH" in sortie, sortie[-2000:]


def test_les_pools_nomment_leur_vraie_source(racine_isolee, tmp_path):
    """La feuille POOLS annonçait « SAISIE_Charges_Flux » : elle doit nommer la base.

    Un libellé de source qui ment est pire qu'absent — c'est ce que l'utilisateur lit pour savoir
    où corriger une donnée."""
    pytest.importorskip("openpyxl")
    import openpyxl

    db = _base_migree(tmp_path)
    proc, _ = _run_lot6f_espionne(racine_isolee, tmp_path, db, "--mois", "2026-07")
    assert proc.returncode == 0, (proc.stdout + proc.stderr)[-3000:]

    sortie = (racine_isolee / "02_TRAVAIL" / "Lot6f_CoutComplet_Menages"
              / "MASTER_CALC_CoutComplet_Menages.xlsx")
    classeur = openpyxl.load_workbook(sortie, read_only=True, data_only=True)
    try:
        lignes = [r for r in classeur["POOLS_CHARGES_MENAGE"].iter_rows(values_only=True)
                  if any(c is not None for c in r)]
    finally:
        classeur.close()

    entetes = [str(c) for c in lignes[0]]
    sources = {dict(zip(entetes, r))["source"] for r in lignes[1:]}
    assert not any("SAISIE_Charges_Flux" in str(s) for s in sources), sources
    assert any("charges" in str(s) for s in sources), sources


def test_menages_cout_complet_est_ecrit_en_base(racine_isolee, tmp_path):
    """La vraie sortie métier est le dataset SQLite, pas le classeur : il doit exister."""
    db = _base_migree(tmp_path)
    proc, _ = _run_lot6f_espionne(
        racine_isolee, tmp_path, db, "--mois", "2026-07", "--sans-excel")
    assert proc.returncode == 0, (proc.stdout + proc.stderr)[-3000:]

    conn = sqlite3.connect(str(db))
    try:
        # Base vide : zéro ligne attendu, mais la table doit avoir été traitée sans erreur.
        assert conn.execute(
            "SELECT COUNT(*) FROM menages_cout_complet WHERE mois='2026-07'").fetchone()[0] == 0
    finally:
        conn.close()
