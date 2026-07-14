"""APP-3b — Audit de clôture : scénarios non couverts par la preuve E2E initiale.

Ces tests ferment les trous trouvés en revue adversariale :
  - la validation métier n'était PAS rejouée à la confirmation (un mois clôturé entre la
    prévisualisation et la confirmation passait au travers — les empreintes ne couvrent pas
    les référentiels) ;
  - les flags étaient vérifiés en DERNIER, ce qui pouvait produire un faux message ;
  - un manifest corrompu était annoncé comme « token inconnu » ;
  - aucune notion de fraîcheur de la prévisualisation ;
  - Lot7 / Lot11 en échec, runner cassé : jamais testés.

Plus les tests de concurrence (vrais processus) et l'E2E HTTP de bout en bout.
Environnement 100 % isolé : aucun Excel réel, aucune app.db réelle.
"""
from __future__ import annotations

import hashlib
import json
import multiprocessing as mp
import shutil
import sqlite3
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

import app.config as cfg
from app.main import app
from app.services import charges_confirmation_service as confirmation
from app.services import charges_post_write_service as aval
from app.services import charges_preview_service as prev
from app.services import saisie_charges_journal_service as jrn
from app.services import saisie_charges_lock_service as lk
from app.services import saisie_charges_transaction_service as tx

REELS = [cfg.SAISIE_CHARGES, cfg.SAISIE_CHARGES_IMPACTS, cfg.MASTER_CHARGES, cfg.SAISIE_IK_AVANTAGES]
PYTHON_MOTEUR = cfg.LOT4A_ENGINE_PYTHON


def _sha(p) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


@pytest.fixture
def flags_on(monkeypatch):
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)


@pytest.fixture
def env(tmp_path: Path, monkeypatch):
    e = tmp_path / "env"
    e.mkdir()
    saisie = e / "SAISIE_Charges_Flux.xlsx"
    impacts = e / "SAISIE_Charges_Impacts.xlsx"
    master = e / "MASTER_FACT_MAN_Charges.xlsx"
    lot7 = e / "MASTER_FACT_MAN_IK_Avantages.xlsx"
    shutil.copy2(cfg.SAISIE_CHARGES, saisie)
    shutil.copy2(cfg.SAISIE_CHARGES_IMPACTS, impacts)
    shutil.copy2(cfg.MASTER_CHARGES, master)
    shutil.copy2(cfg.SAISIE_IK_AVANTAGES, lot7)
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", saisie)
    monkeypatch.setattr(cfg, "SAISIE_CHARGES_IMPACTS", impacts)
    return {
        "root": e, "saisie": saisie, "impacts": impacts, "master": master, "lot7": lot7,
        "ref": cfg.REF_SETUP, "dryruns": e / "dryruns",
        "lock": e / "verrous" / lk.LOCK_NAME, "db": e / "journal" / "journal.db",
    }


def _form(**kw) -> dict:
    base = {
        "date_charge": "2026-06-15", "montant": "100.00", "categorie_charge_id": "CHG_025",
        "code_impact": "IC", "mode_paiement_id": "PAY_001",
        "impact_menage": "NON", "refacturable": "NON",
    }
    base.update(kw)
    return base


def _previsualiser(env, form=None) -> dict:
    res = prev.previsualiser(form or _form(), saisie_source=env["saisie"],
                             dryruns_root=env["dryruns"], ref_path=env["ref"])
    assert res["ok"], res["manifest"].get("errors")
    return res


def _confirmer(env, **kw):
    params = dict(
        dryruns_root=env["dryruns"], saisie_path=env["saisie"], impacts_path=env["impacts"],
        master_path=env["master"], lot7_path=env["lot7"], ref_path=env["ref"],
        lock_path=env["lock"], db_path=env["db"], python_moteur=PYTHON_MOTEUR,
    )
    params.update(kw)
    return confirmation.confirmer(params.pop("token"), **params)


def _charges(env, charge_id: str) -> list:
    wb = openpyxl.load_workbook(env["saisie"], read_only=True, data_only=True)
    try:
        rows = [r for r in wb["SAISIE"].iter_rows(values_only=True) if r[0] == charge_id]
    finally:
        wb.close()
    return rows


# ═════════════════════════════════════════════════════════════════════════════
# DÉFAUT 1 — la validation métier n'était pas rejouée à la confirmation
# ═════════════════════════════════════════════════════════════════════════════

def test_mois_cloture_entre_previsualisation_et_confirmation(env, flags_on, monkeypatch):
    """LE trou principal : les empreintes ne couvrent PAS les référentiels.

    On prévisualise sur un mois ouvert, le mois est clôturé, puis on confirme. Sans revalidation,
    la charge s'écrirait dans un mois fermé — aucune empreinte n'aurait bougé.
    """
    p = _previsualiser(env)
    avant = {k: _sha(env[k]) for k in ("saisie", "impacts")}

    # Le mois 2026-06 est clôturé après coup, dans le référentiel (pas dans les fichiers hashés).
    refs_reels = prev.load_form_refs(env["ref"])
    refs_clos = dict(refs_reels)
    refs_clos["cloture"] = [
        {**r, "statut_mois": "CLOTURE"} if str(r.get("mois")) == "2026-06" else r
        for r in refs_reels["cloture"]
    ]
    monkeypatch.setattr(prev, "load_form_refs", lambda *a, **kw: refs_clos)

    res = _confirmer(env, token=p["token"])

    assert res.statut == confirmation.REFUSE
    assert res.code == confirmation.E_VALIDATION_PERIMEE
    assert "V02_MOIS_CLOTURE" in res.transaction["details"]
    assert "clôturé" in res.message
    assert {k: _sha(env[k]) for k in ("saisie", "impacts")} == avant   # rien écrit
    assert not env["lock"].exists()


def test_categorie_devenue_invalide_est_refusee(env, flags_on, monkeypatch):
    """Toute règle métier qui casse entre-temps bloque la confirmation, pas seulement la clôture."""
    p = _previsualiser(env)
    refs = prev.load_form_refs(env["ref"])
    sans_categorie = dict(refs)
    sans_categorie["categories_all"] = [
        c for c in refs["categories_all"]
        if str(c.get("categorie_charge_id")).strip() != "CHG_025"
    ]
    monkeypatch.setattr(prev, "load_form_refs", lambda *a, **kw: sans_categorie)

    res = _confirmer(env, token=p["token"])
    assert res.code == confirmation.E_VALIDATION_PERIMEE
    assert "V04" in res.transaction["details"]


# ═════════════════════════════════════════════════════════════════════════════
# DÉFAUT 2 — flags vérifiés trop tard (faux message)
# ═════════════════════════════════════════════════════════════════════════════

def test_flags_off_prime_sur_toute_autre_cause_de_refus(env):
    """Flags off ET source modifiée : le message doit parler des FLAGS, pas envoyer ressaisir."""
    p = _previsualiser(env)

    wb = openpyxl.load_workbook(env["saisie"])       # la source bouge aussi
    wb["SAISIE"]["AC500"] = "modification externe"
    wb.save(env["saisie"])
    wb.close()

    res = _confirmer(env, token=p["token"])          # aucun flag actif

    assert res.code == tx.E_FLAGS_DESACTIVES         # et NON E_SOURCE_MODIFIEE
    assert "pas activée" in res.message
    assert not env["lock"].exists()

    # Le refus reste tracé, alors même que la transaction n'est jamais atteinte.
    traces = jrn.traces_du_token(p["token"], db_path=env["db"])
    assert len(traces) == 1
    assert traces[0]["statut"] == jrn.REFUSE_FLAGS
    assert traces[0]["sha256_saisie_avant"] == traces[0]["sha256_saisie_apres"]   # rien touché


def test_un_seul_flag_ne_suffit_pas(env, monkeypatch):
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)   # le second reste False
    p = _previsualiser(env)
    res = _confirmer(env, token=p["token"])
    assert res.code == tx.E_FLAGS_DESACTIVES


# ═════════════════════════════════════════════════════════════════════════════
# DÉFAUT 3 — manifest illisible ≠ token inconnu.  DÉFAUT 4 — fraîcheur
# ═════════════════════════════════════════════════════════════════════════════

def test_manifest_illisible_nest_pas_un_token_inconnu(env, flags_on):
    p = _previsualiser(env)
    (Path(p["run_dir"]) / prev.MANIFEST_NAME).write_text("{ ceci n'est pas du JSON",
                                                         encoding="utf-8")
    res = _confirmer(env, token=p["token"])

    assert res.code == confirmation.E_MANIFEST_ILLISIBLE      # et NON E_TOKEN_INCONNU
    assert "corrompu" in res.message


def test_manifest_expire_nest_plus_confirmable(env, flags_on):
    p = _previsualiser(env)
    chemin = Path(p["run_dir"]) / prev.MANIFEST_NAME
    manifest = json.loads(chemin.read_text(encoding="utf-8"))
    vieux = datetime.now(timezone.utc) - timedelta(hours=confirmation.DUREE_VIE_MANIFEST_HEURES + 1)
    manifest["created_at_utc"] = vieux.isoformat()
    # `created_at_utc` n'est pas scellé : on peut le vieillir sans casser le sceau.
    chemin.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

    avant = {k: _sha(env[k]) for k in ("saisie", "impacts")}
    res = _confirmer(env, token=p["token"])

    assert res.code == confirmation.E_MANIFEST_EXPIRE
    assert {k: _sha(env[k]) for k in ("saisie", "impacts")} == avant


def test_copie_de_flux_absente_est_refusee(env, flags_on):
    p = _previsualiser(env)
    (Path(p["run_dir"]) / prev.SAISIE_COPY_NAME).unlink()
    res = _confirmer(env, token=p["token"])
    assert res.code == confirmation.E_COPIE_ALTEREE


def test_copie_de_flux_modifiee_est_refusee(env, flags_on):
    p = _previsualiser(env)
    (Path(p["run_dir"]) / prev.SAISIE_COPY_NAME).write_bytes(b"corrompu")
    res = _confirmer(env, token=p["token"])
    assert res.code == confirmation.E_COPIE_ALTEREE


def test_source_impacts_modifiee_est_refusee(env, flags_on):
    """L'empreinte du fichier d'IMPACTS est vérifiée, pas seulement celle du flux."""
    p = _previsualiser(env)
    wb = openpyxl.load_workbook(env["impacts"])
    wb["AFFECTATIONS"].append(["AUTRE-001", "CHG-AUTRE", "2026-05", "LOG_0001", None,
                               10.0, "A_CONTROLER", "X", None, "h"])
    wb.save(env["impacts"])
    wb.close()

    res = _confirmer(env, token=p["token"])
    assert res.code == confirmation.E_SOURCE_MODIFIEE
    assert "Impacts" in res.transaction["details"]


# ═════════════════════════════════════════════════════════════════════════════
# Post-écriture : Lot7 et Lot11 en échec, runner cassé
# ═════════════════════════════════════════════════════════════════════════════

def test_echec_lot7_apres_ecriture_reussie(env, flags_on):
    """Lot7 casse : la charge reste écrite, Lot3 reste OK, et on le dit sans mentir."""
    refs = prev.load_form_refs(env["ref"])
    associe = str(refs["associes"][0]["personne_id"]).strip()
    p = _previsualiser(env, _form(avantage_associe="OUI", avantage_associe_id=associe))

    lot7_casse = env["root"] / "inexistant" / "LOT7.xlsx"
    res = _confirmer(env, token=p["token"], lot7_path=lot7_casse)

    assert res.statut == confirmation.SUCCES              # l'écriture a bien eu lieu
    assert len(_charges(env, res.charge_id)) == 1
    post = res.post_ecriture
    assert post["lot3"]["statut"] == aval.OK              # Lot3 a réussi
    assert post["lot7"]["statut"] == aval.ECHEC           # Lot7 a échoué
    assert post["ok"] is False
    assert "recalcul aval a échoué" in res.message


def test_echec_lot11_apres_ecriture_reussie(env, flags_on, monkeypatch):
    """Lot11 casse (moteur introuvable) : la charge reste écrite, l'incident est rendu."""
    p = _previsualiser(env)
    res = _confirmer(env, token=p["token"], python_moteur=Path("python_inexistant.exe"))

    assert res.statut == confirmation.SUCCES              # écriture OK malgré tout
    assert len(_charges(env, res.charge_id)) == 1
    post = res.post_ecriture
    assert post["ok"] is False
    assert post["erreur"] and "moteur" in post["erreur"].lower()
    assert post["lot3"]["statut"] == aval.NON_LANCE       # rien n'a tourné
    assert "recalcul aval a échoué" in res.message


def test_reponse_runner_invalide_est_rendue_comme_incident_aval(env, flags_on, tmp_path):
    """Réponse JSON illisible : incident de recalcul, jamais un échec d'écriture."""
    faux_runner = tmp_path / "faux_runner.py"
    faux_runner.write_text(
        "import sys, pathlib\n"
        "pathlib.Path(sys.argv[2]).write_text('pas du json', encoding='utf-8')\n",
        encoding="utf-8",
    )
    monkey = aval.RUNNER
    try:
        aval.RUNNER = faux_runner
        p = _previsualiser(env)
        res = _confirmer(env, token=p["token"])
    finally:
        aval.RUNNER = monkey

    assert res.statut == confirmation.SUCCES
    assert res.post_ecriture["ok"] is False
    assert "illisible" in res.post_ecriture["erreur"]


def test_runner_qui_meurt_brutalement(env, flags_on, tmp_path):
    """Le runner sort en catastrophe sans réponse : incident aval explicite."""
    faux_runner = tmp_path / "runner_qui_meurt.py"
    faux_runner.write_text("import sys\nsys.exit(3)\n", encoding="utf-8")
    monkey = aval.RUNNER
    try:
        aval.RUNNER = faux_runner
        p = _previsualiser(env)
        res = _confirmer(env, token=p["token"])
    finally:
        aval.RUNNER = monkey

    assert res.statut == confirmation.SUCCES              # la charge est écrite
    assert res.post_ecriture["ok"] is False
    assert "aucune réponse" in res.post_ecriture["erreur"]


# ═════════════════════════════════════════════════════════════════════════════
# Concurrence — vrais processus
# ═════════════════════════════════════════════════════════════════════════════

def _worker_confirmation(params: dict, ev_pret, ev_go, resultats) -> None:
    """Processus enfant (spawn). Flags forcés EN MÉMOIRE uniquement."""
    import app.config as cfg_enfant
    cfg_enfant.CHARGES_REAL_WRITE_ENABLED = True
    cfg_enfant.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED = True
    cfg_enfant.SAISIE_CHARGES = Path(params["saisie"])
    cfg_enfant.SAISIE_CHARGES_IMPACTS = Path(params["impacts"])

    from app.services import charges_confirmation_service as conf
    from app.services import saisie_charges_transaction_service as tx_enfant

    if params["role"] == "A":
        vrai = tx_enfant._preparer

        def preparer_bloquant(d):
            ev_pret.set()                      # « je tiens le verrou »
            ev_go.wait(timeout=60)             # B a fini d'essayer
            return vrai(d)

        tx_enfant._preparer = preparer_bloquant
    else:
        ev_pret.wait(timeout=60)               # B n'essaie que pendant que A tient le verrou

    res = conf.confirmer(
        params["token"],
        dryruns_root=Path(params["dryruns"]), saisie_path=Path(params["saisie"]),
        impacts_path=Path(params["impacts"]), master_path=Path(params["master"]),
        lot7_path=Path(params["lot7"]), ref_path=Path(params["ref"]),
        lock_path=Path(params["lock"]), db_path=Path(params["db"]),
        post_ecriture=False,                   # on teste la concurrence, pas le recalcul
    )
    resultats.put({"role": params["role"], "statut": res.statut, "code": res.code,
                   "charge_id": res.charge_id})
    if params["role"] == "B":
        ev_go.set()


def test_deux_processus_confirment_le_meme_token(env, flags_on):
    """Deux VRAIS processus, même token : un seul écrit. Aucune charge en double."""
    p = _previsualiser(env)
    ctx = mp.get_context("spawn")
    ev_pret, ev_go, resultats = ctx.Event(), ctx.Event(), ctx.Queue()

    base = {
        "token": p["token"], "dryruns": str(env["dryruns"]), "saisie": str(env["saisie"]),
        "impacts": str(env["impacts"]), "master": str(env["master"]), "lot7": str(env["lot7"]),
        "ref": str(env["ref"]), "lock": str(env["lock"]), "db": str(env["db"]),
    }
    pa = ctx.Process(target=_worker_confirmation, args=({**base, "role": "A"}, ev_pret, ev_go, resultats))
    pb = ctx.Process(target=_worker_confirmation, args=({**base, "role": "B"}, ev_pret, ev_go, resultats))
    pa.start()
    pb.start()
    recus = [resultats.get(timeout=90), resultats.get(timeout=90)]
    pa.join(timeout=90)
    pb.join(timeout=90)
    assert pa.exitcode == 0 and pb.exitcode == 0

    par_role = {r["role"]: r for r in recus}
    assert par_role["A"]["statut"] == confirmation.SUCCES, par_role["A"]
    assert par_role["B"]["statut"] == confirmation.REFUSE, par_role["B"]
    assert par_role["B"]["code"] == tx.E_VERROU_DEJA_PRIS

    # Une seule charge écrite, un seul succès au journal, aucun verrou résiduel.
    assert len(_charges(env, par_role["A"]["charge_id"])) == 1
    traces = jrn.traces_du_token(p["token"], db_path=env["db"])
    assert [t["statut"] for t in traces].count(jrn.SUCCES) == 1
    assert not env["lock"].exists()


def test_deux_tokens_distincts_ne_produisent_jamais_le_meme_charge_id(env, flags_on):
    """Deux prévisualisations concurrentes : la seconde est périmée dès que la première écrit."""
    p1 = _previsualiser(env)
    p2 = _previsualiser(env)                   # même état de départ → même charge_id indicatif
    assert p1["manifest"]["charge_id"] == p2["manifest"]["charge_id"]

    r1 = _confirmer(env, token=p1["token"])
    assert r1.statut == confirmation.SUCCES

    r2 = _confirmer(env, token=p2["token"])
    assert r2.statut == confirmation.REFUSE
    assert r2.code == confirmation.E_SOURCE_MODIFIEE   # la base a bougé sous ses pieds
    assert len(_charges(env, r1.charge_id)) == 1       # une seule charge, jamais deux


def test_le_verrou_est_libere_meme_apres_un_refus_de_validation(env, flags_on, monkeypatch):
    """Aucun blocage durable : un refus tardif ne laisse pas de verrou derrière lui."""
    p = _previsualiser(env)
    refs = prev.load_form_refs(env["ref"])
    clos = dict(refs)
    clos["cloture"] = [{**r, "statut_mois": "CLOTURE"} if str(r.get("mois")) == "2026-06" else r
                       for r in refs["cloture"]]
    monkeypatch.setattr(prev, "load_form_refs", lambda *a, **kw: clos)

    assert _confirmer(env, token=p["token"]).code == confirmation.E_VALIDATION_PERIMEE
    assert not env["lock"].exists()


# ═════════════════════════════════════════════════════════════════════════════
# E2E HTTP complet
# ═════════════════════════════════════════════════════════════════════════════

def test_e2e_http_complet_sur_copies(env, flags_on, tmp_db, monkeypatch):
    """Parcours HTTP entier : formulaire → prévisualisation → confirmation → 303 → résultat.

    Les routes n'acceptent aucune injection de chemin : on redirige donc la configuration (en
    mémoire) vers l'environnement isolé. Aucun fichier réel n'est touché.
    """
    empreintes_reelles = {p: _sha(p) for p in REELS if p.exists()}

    monkeypatch.setattr(prev, "DRYRUNS_DIR", env["dryruns"])
    monkeypatch.setattr(cfg, "DRYRUNS_DIR", env["dryruns"])
    monkeypatch.setattr(cfg, "MASTER_CHARGES", env["master"])
    monkeypatch.setattr(cfg, "SAISIE_IK_AVANTAGES", env["lot7"])
    monkeypatch.setattr(lk, "chemin_verrou_par_defaut", lambda: env["lock"])

    with TestClient(app) as client:
        # 1. Formulaire
        assert client.get("/fournisseurs/nouvelle").status_code == 200

        # 2. Prévisualisation
        r = client.post("/fournisseurs/nouvelle/previsualiser", data=_form(),
                        follow_redirects=False)
        assert r.status_code == 303
        token = r.headers["location"].rsplit("/", 1)[-1]

        # 3. Écran de prévisualisation : le formulaire de confirmation est offert (flags actifs).
        html = client.get(f"/fournisseurs/nouvelle/previsualisation/{token}").text
        assert f'action="/fournisseurs/nouvelle/confirmer/{token}"' in html
        assert 'data-testid="avertissement-ecriture"' in html
        assert "C:\\Users" not in html and "OneDrive" not in html     # aucun chemin interne

        # 4. Confirmation — AUCUNE donnée métier envoyée (et ce qu'on envoie est ignoré).
        r = client.post(f"/fournisseurs/nouvelle/confirmer/{token}",
                        data={"montant": "999999", "charge_id": "CHG-PIRATE"},
                        follow_redirects=False)
        assert r.status_code == 303                                   # POST-Redirect-Get
        assert r.headers["location"] == f"/fournisseurs/nouvelle/resultat/{token}"

        # 5. Résultat : écriture et lots aval affichés SÉPARÉMENT.
        html = client.get(f"/fournisseurs/nouvelle/resultat/{token}").text
        assert 'data-statut="SUCCES"' in html
        assert 'data-testid="statut-lot3"' in html
        assert 'data-testid="statut-lot7"' in html
        assert 'data-testid="statut-lot11"' in html
        assert "PIRATE" not in html                                   # le payload client n'a rien fait
        assert "C:\\Users" not in html

        # 6. Rafraîchir : simple lecture, aucune réécriture.
        for _ in range(3):
            assert client.get(f"/fournisseurs/nouvelle/resultat/{token}").status_code == 200

        # 7. Second POST : redirigé, jamais rejoué.
        r = client.post(f"/fournisseurs/nouvelle/confirmer/{token}", follow_redirects=False)
        assert r.status_code == 303

    # La charge existe UNE fois, dans la copie isolée.
    resultat = confirmation.charger_resultat(token, env["dryruns"])
    assert resultat["statut"] == "SUCCES"
    assert len(_charges(env, resultat["charge_id"])) == 1

    # Les fichiers réels n'ont pas bougé d'un octet.
    assert {p: _sha(p) for p in REELS if p.exists()} == empreintes_reelles
    assert not env["lock"].exists()


def test_aucun_residu_apres_succes_refus_et_rollback(env, flags_on, monkeypatch):
    """Ni verrou, ni .bak, ni temporaire dans les dossiers métier — quel que soit le dénouement."""
    import os

    # Succès
    p = _previsualiser(env)
    assert _confirmer(env, token=p["token"]).statut == confirmation.SUCCES

    # Refus
    p2 = _previsualiser(env)
    (Path(p2["run_dir"]) / prev.SAISIE_COPY_NAME).unlink()
    assert _confirmer(env, token=p2["token"]).code == confirmation.E_COPIE_ALTEREE

    # Rollback
    p3 = _previsualiser(env)
    vrai = os.replace
    etat = {"n": 0}

    def faux(src, dst, *a, **kw):
        etat["n"] += 1
        if etat["n"] == 2:
            raise OSError("échec simulé")
        return vrai(src, dst, *a, **kw)

    monkeypatch.setattr(os, "replace", faux)
    assert _confirmer(env, token=p3["token"]).statut == confirmation.ROLLBACK
    monkeypatch.undo()

    residus = [f.name for f in env["root"].iterdir()
               if f.is_file() and (".bak" in f.name or f.name.startswith("~$")
                                   or f.name.startswith("tmp"))]
    assert residus == []
    assert not env["lock"].exists()
