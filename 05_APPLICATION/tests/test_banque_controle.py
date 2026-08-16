"""APP-4B — Contrôle & catégorisation bancaire sur copies (34 points de recette).

Aucune écriture bancaire réelle. Excel réel + app.db réelle intacts. Flags False.
Les tests qui touchent la vraie banque sont skipés si BANQUE_LOT8_IMPORT.xlsx est absent.
"""
import hashlib
from pathlib import Path

import pytest

import app.config as cfg
from app.services import banques_controle_service as ctrl
from app.services import banques_controle_writer as writer

REAL_BANQUE = Path(cfg.MASTER_BANQUE)
banque_requise = pytest.mark.skipif(not REAL_BANQUE.exists(), reason="BANQUE_LOT8_IMPORT.xlsx absent")


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


@pytest.fixture(autouse=True)
def _isoler_workspace(tmp_path, monkeypatch):
    """Isole le workspace de copie et les snapshots sous tmp (jamais dans le worktree)."""
    monkeypatch.setattr(cfg, "BANQUE_CONTROLE_WORKSPACE", tmp_path / "banque_ws")
    from app.services import snapshot_service
    monkeypatch.setattr(cfg, "SNAPSHOTS_DIR", tmp_path / "snapshots")


@pytest.fixture
def opaque():
    ctrl.vider_cache()
    idx = ctrl.index_opaque()
    assert idx, "Aucun mouvement bancaire"
    return list(idx)[0]


# ── 1-2 : identifiant opaque ─────────────────────────────────────────────────

def test_01_id_opaque_stable():
    a = ctrl.id_opaque("MVT-CM_02211_00021321603-20260225-DEBIT-5013-77DFEC")
    b = ctrl.id_opaque("MVT-CM_02211_00021321603-20260225-DEBIT-5013-77DFEC")
    assert a == b and a.startswith("MVT-")


def test_02_id_opaque_sans_numero_compte():
    opq = ctrl.id_opaque("MVT-CM_02211_00021321603-20260225-DEBIT-5013-77DFEC")
    assert "00021321603" not in opq and "CM_" not in opq and "02211" not in opq


# ── 3-4 : fiche ──────────────────────────────────────────────────────────────

@banque_requise
def test_03_fiche_disponible(opaque, tmp_db):
    f = ctrl.load_fiche(opaque, db_path=tmp_db)
    assert f is not None and f["id_opaque"] == opaque
    assert "••••" in f["compte_masque"]


def test_04_mouvement_inconnu_none(tmp_db):
    assert ctrl.load_fiche("MVT-000000000000", db_path=tmp_db) is None


# ── 5-11 : décisions sur copie (journal) ─────────────────────────────────────

@banque_requise
def test_05_accepter_proposition(opaque, tmp_db):
    f = ctrl.load_fiche(opaque, db_path=tmp_db)
    cat = f["proposition_moteur"]["categorie"]
    r = ctrl.enregistrer_decision(opaque, categorie=cat, statut_controle="EN_COURS", db_path=tmp_db)
    assert r["ok"] and r["statut_controle"] == "EN_COURS"


@banque_requise
def test_06_modifier_categorie(opaque, tmp_db):
    r = ctrl.enregistrer_decision(opaque, categorie="FRAIS_BANCAIRES", statut_controle="EN_COURS", db_path=tmp_db)
    f = ctrl.load_fiche(opaque, db_path=tmp_db)
    assert f["decision"]["categorie_validee"] == "FRAIS_BANCAIRES"


@banque_requise
def test_07_08_rattacher_proprietaire_logement(opaque, tmp_db):
    opts = ctrl.options_reference()
    pid = opts["proprietaires"][0]["id"]
    lid = opts["logements"][0]["id"]
    ctrl.enregistrer_decision(opaque, proprietaire_id=pid, logement_id=lid,
                              statut_controle="EN_COURS", db_path=tmp_db)
    f = ctrl.load_fiche(opaque, db_path=tmp_db)
    assert f["decision"]["proprietaire_id"] == pid and f["decision"]["logement_id"] == lid


@banque_requise
def test_09_10_rattacher_reservation_facture(opaque, tmp_db):
    ctrl.enregistrer_decision(opaque, reservation_id="RES_0001", facture_id="FAC_0001",
                              statut_controle="EN_COURS", db_path=tmp_db)
    f = ctrl.load_fiche(opaque, db_path=tmp_db)
    assert f["decision"]["reservation_id"] == "RES_0001" and f["decision"]["facture_id"] == "FAC_0001"


@banque_requise
def test_11_commentaire_enregistre(opaque, tmp_db):
    ctrl.enregistrer_decision(opaque, commentaire="vérifié le 17/07", statut_controle="EN_COURS", db_path=tmp_db)
    f = ctrl.load_fiche(opaque, db_path=tmp_db)
    assert f["decision"]["commentaire"] == "vérifié le 17/07"


# ── 12-16 : workflow de statuts ──────────────────────────────────────────────

@banque_requise
def test_12_ignore_exige_justification(opaque, tmp_db):
    with pytest.raises(ctrl.DecisionRefusee):
        ctrl.enregistrer_decision(opaque, statut_controle="IGNORE", db_path=tmp_db)
    r = ctrl.enregistrer_decision(opaque, statut_controle="IGNORE", justification="hors périmètre", db_path=tmp_db)
    assert r["statut_controle"] == "IGNORE"


@banque_requise
def test_13_14_15_16_transitions(opaque, tmp_db):
    ctrl.enregistrer_decision(opaque, statut_controle="EN_COURS", db_path=tmp_db)        # A_CONTROLER->EN_COURS
    ctrl.enregistrer_decision(opaque, statut_controle="CONTROLE", db_path=tmp_db)        # EN_COURS->CONTROLE
    ctrl.enregistrer_decision(opaque, statut_controle="RAPPROCHE", db_path=tmp_db)       # CONTROLE->RAPPROCHE
    r = ctrl.enregistrer_decision(opaque, statut_controle="ROUVERT", db_path=tmp_db)     # RAPPROCHE->ROUVERT
    assert r["statut_controle"] == "ROUVERT"


@banque_requise
def test_17_statut_invalide_refuse(opaque, tmp_db):
    with pytest.raises(ctrl.DecisionRefusee):
        ctrl.enregistrer_decision(opaque, statut_controle="NIMPORTEQUOI", db_path=tmp_db)


@banque_requise
def test_18_entite_inconnue_refusee(opaque, tmp_db):
    with pytest.raises(ctrl.DecisionRefusee):
        ctrl.enregistrer_decision(opaque, proprietaire_id="PROP_INEXISTANT", statut_controle="EN_COURS", db_path=tmp_db)


# ── 19-21 : proposition préservée, décision séparée, historique ──────────────

@banque_requise
def test_19_20_proposition_preservee_decision_separee(opaque, tmp_db):
    f0 = ctrl.load_fiche(opaque, db_path=tmp_db)
    prop_avant = f0["proposition_moteur"]["categorie"]
    ctrl.enregistrer_decision(opaque, categorie="FRAIS_BANCAIRES", statut_controle="EN_COURS", db_path=tmp_db)
    f1 = ctrl.load_fiche(opaque, db_path=tmp_db)
    assert f1["proposition_moteur"]["categorie"] == prop_avant   # jamais écrasée
    assert f1["decision"]["categorie_validee"] == "FRAIS_BANCAIRES"   # séparée


@banque_requise
def test_21_historique_conserve(opaque, tmp_db):
    ctrl.enregistrer_decision(opaque, statut_controle="EN_COURS", db_path=tmp_db)
    ctrl.enregistrer_decision(opaque, statut_controle="CONTROLE", db_path=tmp_db)
    f = ctrl.load_fiche(opaque, db_path=tmp_db)
    assert len(f["historique"]) >= 2
    assert sum(1 for h in f["historique"] if h["actif"]) == 1   # une seule décision active


# ── 22 : conflit de version ──────────────────────────────────────────────────

@banque_requise
def test_22_conflit_version_refuse(opaque, tmp_db):
    ctrl.enregistrer_decision(opaque, statut_controle="EN_COURS", db_path=tmp_db)   # version 1
    with pytest.raises(ctrl.DecisionRefusee):
        ctrl.enregistrer_decision(opaque, statut_controle="CONTROLE", version_attendue=0, db_path=tmp_db)


# ── 23-24, 27-28, 31 : writer copie, réel intact, impact contrôle ────────────

@banque_requise
def test_23_24_27_28_writer_copie_reel_intact_impact(opaque, tmp_db):
    sha_avant = _sha(REAL_BANQUE)
    ctrl.enregistrer_decision(opaque, statut_controle="IGNORE", justification="test", db_path=tmp_db)
    res = writer.enregistrer_sur_copie(db_path=tmp_db)
    assert res["statut"] == "SUCCES" and res["reel_intact"] is True
    # 23 : copie modifiée (onglet override + NORM), pas le réel
    copie = Path(res["copie"])
    import openpyxl
    wb = openpyxl.load_workbook(copie, read_only=True)
    assert cfg.BANQUE_OVERRIDE_SHEET in wb.sheetnames
    assert all(s in wb.sheetnames for s in ("BRUT_Banque", "NORM_Banque", "CTRL_A_CONTROLER"))
    wb.close()
    # 24 : réel intact
    assert _sha(REAL_BANQUE) == sha_avant
    # 27-28 : le contrôle évolue (statut NORM), l'anomalie n'est pas masquée dans SQLite
    assert res["controles_avant"] != res["controles_apres"]
    assert res["controles_apres"].get("IGNORE", 0) >= 1


# ── 32 : flags ───────────────────────────────────────────────────────────────

def test_32_flags_false():
    assert cfg.BANQUE_REAL_WRITE_ENABLED is False
    assert cfg.BANQUE_REAL_WRITE_CONFIRMATION_ENABLED is False


@banque_requise
def test_32b_mode_reel_refuse(tmp_db):
    res = writer.enregistrer_sur_copie(mode="REEL", db_path=tmp_db)
    assert res["statut"] == "BLOQUE" and res["erreur_code"] == "E_REEL_DESACTIVE"


# ── 33-34 : intégrité réelle ─────────────────────────────────────────────────

@banque_requise
def test_33_34_app_db_et_excel_reels_intacts(opaque, tmp_db):
    sha_banque = _sha(REAL_BANQUE)
    real_db = Path(cfg.PROJECT_ROOT) / "05_APPLICATION" / "data" / "app.db"
    sha_db = _sha(real_db) if real_db.exists() else None
    ctrl.enregistrer_decision(opaque, statut_controle="EN_COURS", db_path=tmp_db)
    writer.enregistrer_sur_copie(db_path=tmp_db)
    assert _sha(REAL_BANQUE) == sha_banque
    if sha_db is not None:
        assert _sha(real_db) == sha_db


# ── 29-30 : routes sans 500, sans chemin absolu ──────────────────────────────

@banque_requise
def test_29_30_31_routes_lisibles(opaque, client):
    for url in ("/banques-caisse/controle",
                f"/banques-caisse/mouvements/{opaque}",
                f"/banques-caisse/mouvements/{opaque}/modifier"):
        r = client.get(url)
        assert r.status_code == 200
        assert "C:\\" not in r.text and "OneDrive" not in r.text
    # compte masqué, pas de numéro de compte brut
    t = client.get(f"/banques-caisse/mouvements/{opaque}").text
    assert "••••" in t and "00021321603" not in t
    # 404 lisible
    assert client.get("/banques-caisse/mouvements/MVT-000000000000").status_code == 404
