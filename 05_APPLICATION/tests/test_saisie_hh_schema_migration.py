"""Tests migration APP-2b sur copie uniquement."""
from pathlib import Path
import shutil

import openpyxl
import pytest

from app.services.saisie_hh_schema_migration import (
    NEW_SAISIE_FIELDS,
    migrate_saisie_copy,
)


ARTIFACT = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "dapp05c_formules"
    / "20260702T143503Z"
    / "SAISIE_copie.xlsx"
)

FORMULA_CELLS = ("B2", "C2", "K2", "N2", "O2", "Q2", "V2", "Y2", "Z2")


def _formula_values(path: Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(str(path), data_only=False)
    try:
        ws = wb["SAISIE"]
        return {cell: ws[cell].value for cell in FORMULA_CELLS}
    finally:
        wb.close()


def test_migration_saisie_app2b_sur_copie_preserve_formules(tmp_path):
    assert ARTIFACT.exists(), f"Artefact D-APP-05C absent: {ARTIFACT}"
    target = tmp_path / "SAISIE_copie_migree.xlsx"
    shutil.copy2(ARTIFACT, target)

    before = _formula_values(target)
    added = migrate_saisie_copy(target)
    after = _formula_values(target)

    assert added == NEW_SAISIE_FIELDS
    assert after == before
    assert all(isinstance(value, str) and value.startswith("=") for value in after.values())

    wb = openpyxl.load_workbook(str(target), data_only=False)
    try:
        ws = wb["SAISIE"]
        headers = [cell.value for cell in ws[1]]
        assert headers[-len(NEW_SAISIE_FIELDS):] == NEW_SAISIE_FIELDS
        assert wb.calculation.fullCalcOnLoad is True
    finally:
        wb.close()


def test_migration_refuse_chemin_source_reel():
    with pytest.raises(RuntimeError):
        migrate_saisie_copy(
            Path("01_SOURCES_BRUTES")
            / "ReservationsHH"
            / "SAISIE_ReservationsHorsHostaway.xlsx"
        )
