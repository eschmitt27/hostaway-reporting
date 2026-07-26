"""Impact Lot9 : vérifie que les lignes produites par l'import applicatif se comportent
correctement face au filtre RÉEL de `lot9_construire_flux.py` (ligne 126 :
`bnq_valide = [r for r in bnq_all if r.get('type_flux_id') == 'TYPE_FLUX_016' and
r.get('statut_controle') == 'VALIDE']`) — reproduit ici verbatim (pas de pandas requis pour cette
ligne, donc exécutable directement dans cet environnement, contrairement au script complet).

Ce que ça prouve :
- un mouvement fraîchement importé (statut EN_ATTENTE_CLASSIFICATION) n'est JAMAIS repris par
  Lot9 tant qu'il n'a pas été validé — pas d'impact financier silencieux d'un import brut ;
- seuls les frais bancaires (TYPE_FLUX_016) VALIDE entrent dans Lot9 — un encaissement réservation,
  un payout plateforme, un paiement propriétaire, etc. n'y entrent jamais via ce chemin (ce sont
  d'autres modules — Hostaway/HH — qui portent ces flux, jamais la banque : pas de double comptage).
"""
from __future__ import annotations

import openpyxl
import pytest

import app.config as cfg
from app.services import banques_import_service as svc

CSV_MIXTE = (
    "Date operation;Libelle;Debit;Credit\n"
    "05/06/2026;VIR HOSTAWAY PAYOUT;;850,00\n"      # encaissement réservation/payout -> jamais BNQ
    "10/06/2026;VIR PROPRIETAIRE PROP A;400,00;\n"   # reversement propriétaire -> jamais BNQ
    "15/06/2026;FRAIS TENUE DE COMPTE;8,90;\n"       # frais bancaires -> seul cas BNQ
).encode("utf-8")


def _bnq_valide(rows: list[dict]) -> list[dict]:
    """Filtre copié verbatim depuis lot9_construire_flux.py:126 — jamais réécrit ici."""
    return [r for r in rows if r.get("type_flux_id") == "TYPE_FLUX_016"
            and r.get("statut_controle") == "VALIDE"]


@pytest.fixture
def ref(tmp_path, monkeypatch):
    p = tmp_path / "BANQUE_LOT8_IMPORT.xlsx"
    monkeypatch.setattr(cfg, "MASTER_BANQUE", p)
    monkeypatch.setattr(cfg, "RECETTE_MODE", True)
    monkeypatch.setattr(cfg, "RECETTE_ROOT", tmp_path.resolve())
    monkeypatch.setattr(cfg, "BANQUE_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "BANQUE_REAL_WRITE_CONFIRMATION_ENABLED", True)
    return p


def _lignes(p):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb["NORM_Banque"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = rows[0]
    return [dict(zip(hdr, r)) for r in rows[1:]]


def test_import_brut_nentre_jamais_dans_lot9(ref, tmp_path):
    """Un mouvement fraîchement importé (statut EN_ATTENTE_CLASSIFICATION) est exclu du filtre
    Lot9 quel que soit son type_flux_id — aucun impact financier avant classification/validation."""
    from app.db.connection import apply_migrations
    db = tmp_path / "test.db"; apply_migrations(db)
    prev = svc.previsualiser(CSV_MIXTE, "releve.csv", "CM_TEST", ref_path=ref)
    svc.confirmer(prev["token"], acteur="recette", db_path=db)

    rows = _lignes(ref)
    assert len(rows) == 3
    assert all(r["statut_controle"] == "EN_ATTENTE_CLASSIFICATION" for r in rows)
    assert _bnq_valide(rows) == []                # rien n'entre dans Lot9 avant validation


def test_seuls_les_frais_bancaires_valides_entrent_dans_lot9(ref, tmp_path):
    """Une fois classifiés (simulateur d'un passage lot8b + validation), seul le mouvement
    TYPE_FLUX_016 VALIDE est repris par Lot9 — jamais le payout ni le reversement propriétaire."""
    from app.db.connection import apply_migrations
    db = tmp_path / "test.db"; apply_migrations(db)
    prev = svc.previsualiser(CSV_MIXTE, "releve.csv", "CM_TEST", ref_path=ref)
    svc.confirmer(prev["token"], acteur="recette", db_path=db)

    # Simule la classification + validation (normalement faite par lot8b + validation humaine) :
    wb = openpyxl.load_workbook(ref)
    ws = wb["NORM_Banque"]
    hdr = [c.value for c in ws[1]]
    i_libelle = hdr.index("libelle")
    i_type = hdr.index("type_flux_id")
    i_statut = hdr.index("statut_controle")
    for row in ws.iter_rows(min_row=2):
        libelle = row[i_libelle].value
        if libelle == "FRAIS TENUE DE COMPTE":
            row[i_type].value = "TYPE_FLUX_016"
        elif libelle == "VIR HOSTAWAY PAYOUT":
            row[i_type].value = "TYPE_FLUX_017"          # jamais BNQ : porté par Hostaway
        elif libelle == "VIR PROPRIETAIRE PROP A":
            row[i_type].value = "TYPE_FLUX_007"          # reversement propriétaire : jamais BNQ
        row[i_statut].value = "VALIDE"
    wb.save(ref); wb.close()

    rows = _lignes(ref)
    bnq = _bnq_valide(rows)
    assert len(bnq) == 1
    assert bnq[0]["libelle"] == "FRAIS TENUE DE COMPTE"
    assert bnq[0]["montant"] == 8.90

    # Le payout et le reversement propriétaire, même VALIDE, n'entrent jamais dans Lot9 par ce
    # chemin : aucun risque de double comptage avec les flux Hostaway/HH ou les règlements propriétaire.
    libelles_bnq = {r["libelle"] for r in bnq}
    assert "VIR HOSTAWAY PAYOUT" not in libelles_bnq
    assert "VIR PROPRIETAIRE PROP A" not in libelles_bnq
