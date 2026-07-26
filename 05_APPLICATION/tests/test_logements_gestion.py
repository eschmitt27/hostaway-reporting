"""Cycle de vie d'un logement existant : modifier / archiver / réactiver / changer_proprietaire /
changer_taux_commission — validations + écriture gardée."""
from datetime import date

import openpyxl
import pytest

import app.config as cfg
from app.services import logements_gestion_service as svc


def _ref(tmp_path):
    p = tmp_path / "REF_Setup.xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("REF_Logements")
    ws.append(["logement_id", "hostaway_listing_id", "nom_logement_officiel", "nom_court", "adresse",
               "ville", "type_logement_id", "sur_hostaway", "actif", "statut_parc", "commentaire",
               "forfait_logiciel_consommables_mensuel"])
    ws.append(["LOG_A1", 900001, "Fictif A1", "A1", "1 rue", "RECETTE", "TYPE_001", "OUI", "OUI",
               "GERE", None, 0])
    ws.append(["LOG_INACTIF", 900005, "Fictif Inactif", "INACTIF", "9 rue", "RECETTE", "TYPE_001",
               "NON", "NON", "RETIRE", None, 0])
    ws = wb.create_sheet("REF_Gestion_Logements_Hist")
    ws.append(["gestion_id", "logement_id", "proprietaire_id", "date_debut", "date_fin",
               "statut_gestion", "source", "commentaire"])
    ws.append(["GST_1", "LOG_A1", "PROP_A", "2026-01-01", None, "ACTIF", "FICTIF", None])
    ws.append(["GST_2", "LOG_INACTIF", "PROP_C", "2026-01-01", "2026-05-31", "RETIRE", "FICTIF", None])
    ws = wb.create_sheet("REF_Proprietaires")
    ws.append(["proprietaire_id", "nom_proprietaire", "actif"])
    for pid in ("PROP_A", "PROP_B", "PROP_C"):
        ws.append([pid, f"Nom {pid}", "OUI"])
    ws = wb.create_sheet("REF_Types_Logements")
    ws.append(["type_logement_id", "libelle"]); ws.append(["TYPE_001", "STUDIO"])
    ws = wb.create_sheet("REF_Taux_Commission")
    ws.append(["taux_commission_id", "proprietaire_id", "logement_id", "taux_commission",
               "date_debut", "date_fin", "actif", "justification", "commentaire"])
    ws.append(["TX_A1", "PROP_A", "LOG_A1", 0.19, "2026-01-01", None, "OUI", "FICTIF", ""])
    wb.save(p); wb.close()
    return p


@pytest.fixture
def ref(tmp_path, monkeypatch):
    p = _ref(tmp_path)
    monkeypatch.setattr(cfg, "REF_SETUP", p)
    monkeypatch.setattr(cfg, "RECETTE_MODE", True)
    monkeypatch.setattr(cfg, "RECETTE_ROOT", tmp_path.resolve())
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)
    return p


def _lignes(p, sheet):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    hdr = [c for c in next(wb[sheet].iter_rows(max_row=1, values_only=True))]
    rows = [dict(zip(hdr, r)) for r in wb[sheet].iter_rows(min_row=2, values_only=True)]
    wb.close()
    return rows


# ── modifier ──────────────────────────────────────────────────────────────

def test_modifier_met_a_jour_les_champs_descriptifs(ref):
    res = svc.modifier("LOG_A1", {"nom_court": "A1-bis", "ville": "NOUVELLE_VILLE"})
    assert res["ok"], res
    logs = _lignes(ref, "REF_Logements")
    assert logs[0]["nom_court"] == "A1-bis"
    assert logs[0]["ville"] == "NOUVELLE_VILLE"
    assert logs[0]["logement_id"] == "LOG_A1"        # jamais touché


def test_modifier_logement_inconnu_refuse(ref):
    res = svc.modifier("LOG_ZZZ", {"nom_court": "X"})
    assert res["ok"] is False and res["code"] == svc.E_LOGEMENT_INCONNU


def test_modifier_type_inconnu_refuse(ref):
    res = svc.modifier("LOG_A1", {"type_logement_id": "TYPE_ZZZ"})
    assert res["ok"] is False


def test_modifier_refuse_si_flags_desactives(ref, monkeypatch):
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", False)
    res = svc.modifier("LOG_A1", {"nom_court": "X"})
    assert res["ok"] is False and res["code"] == svc.E_FLAGS


# ── archiver / reactiver ──────────────────────────────────────────────────

def test_archiver_desactive_et_cloture_la_gestion(ref):
    res = svc.archiver("LOG_A1", "2026-06-30")
    assert res["ok"], res
    logs = {r["logement_id"]: r for r in _lignes(ref, "REF_Logements")}
    assert logs["LOG_A1"]["actif"] == "NON" and logs["LOG_A1"]["statut_parc"] == "RETIRE"
    gest = [r for r in _lignes(ref, "REF_Gestion_Logements_Hist") if r["logement_id"] == "LOG_A1"]
    assert gest[0]["date_fin"] == "2026-06-30" and gest[0]["statut_gestion"] == "RETIRE"


def test_archiver_deja_archive_refuse(ref):
    res = svc.archiver("LOG_INACTIF", "2026-06-30")
    assert res["ok"] is False and res["code"] == svc.E_DEJA_ARCHIVE


def test_archiver_date_invalide_refusee(ref):
    res = svc.archiver("LOG_A1", "30/06/2026")
    assert res["ok"] is False and res["code"] == svc.E_DATE_INVALIDE


def test_reactiver_reactive_et_ouvre_un_rattachement(ref):
    res = svc.reactiver("LOG_INACTIF", "2026-07-01", "PROP_B")
    assert res["ok"], res
    logs = {r["logement_id"]: r for r in _lignes(ref, "REF_Logements")}
    assert logs["LOG_INACTIF"]["actif"] == "OUI" and logs["LOG_INACTIF"]["statut_parc"] == "GERE"
    gest = [r for r in _lignes(ref, "REF_Gestion_Logements_Hist")
            if r["logement_id"] == "LOG_INACTIF"]
    assert len(gest) == 2                            # ancienne conservée + nouvelle
    nouvelle = [g for g in gest if not g["date_fin"]][0]
    assert nouvelle["proprietaire_id"] == "PROP_B" and nouvelle["date_debut"] == "2026-07-01"


def test_reactiver_deja_actif_refuse(ref):
    res = svc.reactiver("LOG_A1", "2026-07-01", "PROP_A")
    assert res["ok"] is False and res["code"] == svc.E_DEJA_ACTIF


def test_reactiver_proprietaire_inconnu_refuse(ref):
    res = svc.reactiver("LOG_INACTIF", "2026-07-01", "PROP_ZZ")
    assert res["ok"] is False and res["code"] == svc.E_PROP_INCONNU


# ── changer_proprietaire ──────────────────────────────────────────────────

def test_changer_proprietaire_cloture_et_ouvre(ref):
    res = svc.changer_proprietaire("LOG_A1", "PROP_B", "2026-07-01")
    assert res["ok"], res
    gest = [r for r in _lignes(ref, "REF_Gestion_Logements_Hist") if r["logement_id"] == "LOG_A1"]
    assert len(gest) == 2
    ancienne = [g for g in gest if g["date_fin"]][0]
    assert ancienne["proprietaire_id"] == "PROP_A" and ancienne["date_fin"] == "2026-06-30"
    nouvelle = [g for g in gest if not g["date_fin"]][0]
    assert nouvelle["proprietaire_id"] == "PROP_B" and nouvelle["date_debut"] == "2026-07-01"


def test_changer_proprietaire_ne_modifie_jamais_la_ligne_close(ref):
    svc.changer_proprietaire("LOG_A1", "PROP_B", "2026-07-01")
    svc.changer_proprietaire("LOG_A1", "PROP_C", "2026-08-01")
    gest = [r for r in _lignes(ref, "REF_Gestion_Logements_Hist") if r["logement_id"] == "LOG_A1"]
    assert len(gest) == 3
    closes = sorted([g for g in gest if g["date_fin"]], key=lambda g: g["date_debut"])
    assert closes[0]["proprietaire_id"] == "PROP_A" and closes[0]["date_fin"] == "2026-06-30"
    assert closes[1]["proprietaire_id"] == "PROP_B" and closes[1]["date_fin"] == "2026-07-31"


def test_changer_proprietaire_inconnu_refuse(ref):
    res = svc.changer_proprietaire("LOG_A1", "PROP_ZZ", "2026-07-01")
    assert res["ok"] is False and res["code"] == svc.E_PROP_INCONNU


def test_changer_proprietaire_manquant_refuse(ref):
    res = svc.changer_proprietaire("LOG_A1", "", "2026-07-01")
    assert res["ok"] is False and res["code"] == svc.E_PROP_MANQUANT


# ── changer_taux_commission ───────────────────────────────────────────────

def test_changer_taux_cloture_ancien_et_ouvre_nouveau(ref):
    res = svc.changer_taux_commission("LOG_A1", 0.15, "2026-07-01")
    assert res["ok"], res
    taux = [r for r in _lignes(ref, "REF_Taux_Commission") if r["logement_id"] == "LOG_A1"]
    assert len(taux) == 2
    ancien = [t for t in taux if t["date_fin"]][0]
    assert ancien["taux_commission"] == 0.19 and ancien["date_fin"] == "2026-06-30"
    nouveau = [t for t in taux if not t["date_fin"]][0]
    assert nouveau["taux_commission"] == 0.15 and nouveau["date_debut"] == "2026-07-01"
    assert nouveau["proprietaire_id"] == "PROP_A"     # hérité du rattachement actif


def test_changer_taux_grain_logement_pas_de_regression_sur_taux_historique(ref):
    """Une ligne close (date_fin déjà renseignée) n'est jamais modifiée par un nouveau
    changement — seule la ligne active (date_fin vide) est close."""
    svc.changer_taux_commission("LOG_A1", 0.15, "2026-07-01")
    svc.changer_taux_commission("LOG_A1", 0.12, "2026-08-01")
    taux = [r for r in _lignes(ref, "REF_Taux_Commission") if r["logement_id"] == "LOG_A1"]
    assert len(taux) == 3
    initial = [t for t in taux if t["taux_commission"] == 0.19][0]
    assert initial["date_fin"] == "2026-06-30"        # jamais retouché par le 2e changement


def test_changer_taux_invalide_refuse(ref):
    res = svc.changer_taux_commission("LOG_A1", 1.5, "2026-07-01")
    assert res["ok"] is False and res["code"] == svc.E_TAUX_INVALIDE


def test_changer_taux_logement_inconnu_refuse(ref):
    res = svc.changer_taux_commission("LOG_ZZZ", 0.15, "2026-07-01")
    assert res["ok"] is False and res["code"] == svc.E_LOGEMENT_INCONNU


def test_changer_taux_refuse_hors_racine_recette(ref, monkeypatch, tmp_path):
    monkeypatch.setattr(cfg, "RECETTE_ROOT", (tmp_path / "ailleurs").resolve())
    res = svc.changer_taux_commission("LOG_A1", 0.15, "2026-07-01")
    assert res["ok"] is False and res["code"] == svc.E_ECRITURE
    taux = [r for r in _lignes(ref, "REF_Taux_Commission") if r["logement_id"] == "LOG_A1"]
    assert len(taux) == 1                             # write-guard : aucune écriture
