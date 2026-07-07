"""Tests persistance durable des impacts charges (source Excel, écriture sur copie, flags off)."""
from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import openpyxl
import pytest

import app.config as cfg
from app.services import charges_impacts_persist_service as persist
from app.services.charges_preview_service import load_form_refs, compute_guidee, previsualiser
from tools import creer_saisie_charges_impacts as gen


def _sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ── Schéma source ────────────────────────────────────────────────────────────

def test_source_impacts_existe_et_conforme():
    assert cfg.SAISIE_CHARGES_IMPACTS.exists()
    assert gen.verifier(cfg.SAISIE_CHARGES_IMPACTS)["conforme"] is True


def test_source_impacts_vide():
    wb = openpyxl.load_workbook(str(cfg.SAISIE_CHARGES_IMPACTS), read_only=True)
    try:
        for sheet in ("AFFECTATIONS", "MENAGE", "RESERVE_REFACTURATION"):
            ws = wb[sheet]
            data = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(c not in (None, "") for c in r)]
            assert data == [], sheet
    finally:
        wb.close()


def test_generateur_idempotent(tmp_path: Path):
    p = tmp_path / "impacts.xlsx"
    assert gen.creer(p)["cree"] is True
    assert gen.creer(p)["cree"] is False  # 2e passe : aucun changement


# ── build_persistable ────────────────────────────────────────────────────────

@pytest.fixture()
def refs():
    return load_form_refs()


def _guide(form):
    return compute_guidee(form, load_form_refs(), form.get("date_charge", "2026-06-15")[:7])


def _base(cat="CHG_005", **extra):
    f = {"date_charge": "2026-06-15", "montant": "100.00", "categorie_charge_id": cat,
         "code_impact": "IC", "mode_paiement_id": "PAY_001"}
    f.update(extra)
    return f


def test_persistable_affectations_somme_egale_montant():
    refs = load_form_refs()
    logs = [str(l["logement_id"]).strip() for l in refs["logements"]][:3]
    form = _base(logements=logs)
    g = _guide(form)
    pers = persist.build_persistable("CHG-T-001", "2026-06", 100.0, g, form)
    assert len(pers["affectations"]) == 3
    total = round(sum(a["quote_part"] for a in pers["affectations"]), 2)
    assert total == 100.0  # jamais répliqué, somme exacte
    # une seule ligne par logement
    lids = [a["logement_id"] for a in pers["affectations"]]
    assert len(lids) == len(set(lids))


def test_persistable_global_zero_affectation():
    form = _base()
    g = _guide(form)
    pers = persist.build_persistable("CHG-T-002", "2026-06", 100.0, g, form)
    assert pers["affectations"] == []
    assert pers["reserve"] == []


def test_persistable_menage_jamais_reserve():
    form = _base(cat="CHG_004", code_impact="HC", commentaire="ok",
                 menage_mode="INTERVENANT", menage_intervenants=["INT_0001", "INT_0002"])
    g = _guide(form)
    pers = persist.build_persistable("CHG-T-003", "2026-06", 100.0, g, form)
    assert len(pers["menage"]) == 2
    assert pers["reserve"] == []  # charge ménage jamais refacturable
    assert pers["affectations"] == []


def test_persistable_reserve_somme_egale_montant():
    refs = load_form_refs()
    logs = [str(l["logement_id"]).strip() for l in refs["logements"]][:2]
    form = _base(logements=logs, refacturable="OUI")
    g = _guide(form)
    pers = persist.build_persistable("CHG-T-004", "2026-06", 100.0, g, form)
    assert len(pers["reserve"]) == 2
    total = round(sum(r["montant_refacturable"] for r in pers["reserve"]), 2)
    assert total == 100.0


def test_persistable_avantage_une_seule_ligne():
    refs = load_form_refs()
    assoc = [str(a["personne_id"]).strip() for a in refs["associes"]][:1]
    form = _base(cat="CHG_009", avantage_associe="OUI", avantage_associe_id=assoc[0])
    g = _guide(form)
    pers = persist.build_persistable("CHG-T-005", "2026-06", 100.0, g, form)
    assert pers["avantage"] is not None
    assert pers["avantage"]["lien_origine"] == "CHG-T-005"
    assert pers["avantage"]["nature"] == "AVANTAGE_CHARGE"


# ── persister_sur_copie (idempotence, fichier réel intouché) ─────────────────

def test_persister_sur_copie_idempotent(tmp_path: Path):
    copy = tmp_path / "impacts.xlsx"
    shutil.copy2(str(cfg.SAISIE_CHARGES_IMPACTS), str(copy))
    refs = load_form_refs()
    logs = [str(l["logement_id"]).strip() for l in refs["logements"]][:2]
    form = _base(logements=logs, refacturable="OUI")
    g = _guide(form)
    pers = persist.build_persistable("CHG-IDEM-001", "2026-06", 100.0, g, form)
    persist.persister_sur_copie(pers, copy)
    persist.persister_sur_copie(pers, copy)  # 2e passe : remplace, pas de doublon
    wb = openpyxl.load_workbook(str(copy))
    try:
        aff = [r for r in wb["AFFECTATIONS"].iter_rows(min_row=2, values_only=True) if r[0]]
        res = [r for r in wb["RESERVE_REFACTURATION"].iter_rows(min_row=2, values_only=True) if r[0]]
    finally:
        wb.close()
    assert len(aff) == 2  # pas 4 : idempotent
    assert len(res) == 2


def test_persister_fichier_reel_intouche(tmp_path: Path):
    h_av = _sha(cfg.SAISIE_CHARGES_IMPACTS)
    copy = tmp_path / "impacts.xlsx"
    shutil.copy2(str(cfg.SAISIE_CHARGES_IMPACTS), str(copy))
    refs = load_form_refs()
    logs = [str(l["logement_id"]).strip() for l in refs["logements"]][:2]
    form = _base(logements=logs, refacturable="OUI")
    g = _guide(form)
    persist.persister_sur_copie(persist.build_persistable("CHG-R-001", "2026-06", 100.0, g, form), copy)
    assert _sha(cfg.SAISIE_CHARGES_IMPACTS) == h_av  # source réelle inchangée


# ── Garde-fou écriture réelle ────────────────────────────────────────────────

def test_persister_reel_interdit_flags_off():
    assert cfg.CHARGES_REAL_WRITE_ENABLED is False
    with pytest.raises(PermissionError):
        persist.persister_reel()


# ── Intégration previsualiser (persistable dans manifest, source réelle intouchée) ──

def test_previsualiser_expose_persistable(tmp_path: Path):
    refs = load_form_refs()
    logs = [str(l["logement_id"]).strip() for l in refs["logements"]][:2]
    form = _base(logements=logs, refacturable="OUI")
    h_av = _sha(cfg.SAISIE_CHARGES_IMPACTS)
    r = previsualiser(form, dryruns_root=tmp_path / "d")
    assert r["ok"], r["manifest"].get("errors")
    m = r["manifest"]
    assert m["persistable"]["charge_id"] == m["charge_id"]
    assert len(m["persistable"]["affectations"]) == 2
    assert m["persist_report"]["affectations_ecrites"] == 2
    # copie impacts écrite dans le dryrun
    assert (Path(m["paths"]["run_dir"]) / "SAISIE_Charges_Impacts_copie.xlsx").exists()
    # source réelle impacts inchangée
    assert _sha(cfg.SAISIE_CHARGES_IMPACTS) == h_av
