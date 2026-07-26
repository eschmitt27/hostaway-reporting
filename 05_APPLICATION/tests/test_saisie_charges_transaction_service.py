"""APP-3b / commit 3 — orchestrateur transactionnel : atomicité applicative sur deux classeurs.

Aucun fichier métier réel n'est touché : tout se joue sur des fixtures en tmp_path (nommées comme
les vrais fichiers, car file_registry n'autorise que les `SAISIE_*`). Les flags réels restent False
dans le dépôt — ils ne sont forcés à True que par monkeypatch, dans le seul processus de test.

`os.replace` est monkeypatché pour provoquer un échec à chaque position de la séquence de commit,
et vérifier que le rollback ramène les fichiers à leur empreinte initiale.
"""
from __future__ import annotations

import hashlib
import multiprocessing as mp
import os
from pathlib import Path

import openpyxl
import pytest

import app.config as cfg
from app.readers.saisie_charges_reader import MANUAL_COL_MAP, _col_index
from app.services import saisie_charges_lock_service as lk
from app.services import saisie_charges_transaction_service as svc
from app.writers import saisie_charges_writer as writer

MONTANT = 100.0
CHARGE_ID = "CHG-2026-06-IC-BANQUE-001"
LIGNE = 2

FORMULES = {
    "C": '=IF(B{r}="","",TEXT(B{r},"YYYY-MM"))',
    "I": '=IF(H{r}="IC","OUI",IF(H{r}="HC","OUI",IF(H{r}="HR","NON","A_CONTROLER")))',
    "J": '=IF(H{r}="IC","OUI",IF(H{r}="HC","NON",IF(H{r}="HR","NON","A_CONTROLER")))',
    "AD": '=IF(A{r}="","",TEXT(B{r},"YYYYMMDD")&"|"&TEXT(D{r},"0.00")&"|"&F{r}&"|"&M{r})',
}

IMPACTS_HEADERS = {
    "AFFECTATIONS": ["affectation_id", "charge_id", "mois", "logement_id", "proprietaire_id",
                     "quote_part", "statut", "origine", "commentaire", "ROW_HASH"],
    "MENAGE": ["menage_impact_id", "charge_id", "mois", "mode", "intervenant_id", "logement_id",
               "statut", "origine", "commentaire", "ROW_HASH"],
    "RESERVE_REFACTURATION": ["reserve_id", "charge_id", "mois", "logement_id", "proprietaire_id",
                              "montant_refacturable", "libelle", "justificatif",
                              "statut_traitement", "trace_decision", "origine", "commentaire",
                              "ROW_HASH"],
}

REELS = [cfg.SAISIE_CHARGES, cfg.SAISIE_CHARGES_IMPACTS, cfg.MASTER_CHARGES, cfg.SAISIE_IK_AVANTAGES]


def _sha(p) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ── Fixtures ─────────────────────────────────────────────────────────────────

@pytest.fixture
def flags_on(monkeypatch):
    """Les deux flags à True — UNIQUEMENT en mémoire, dans ce processus de test."""
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)


@pytest.fixture
def cibles(tmp_path: Path) -> tuple[Path, Path]:
    saisie = tmp_path / "SAISIE_Charges_Flux.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    entetes = {**MANUAL_COL_MAP, "C": "mois", "I": "impact_resultat_reel",
               "J": "impact_resultat_comptable", "AD": "ROW_HASH"}
    for col, champ in entetes.items():
        ws.cell(row=1, column=_col_index(col), value=champ)
    for r in range(2, 7):
        for col, f in FORMULES.items():
            ws.cell(row=r, column=_col_index(col), value=f.format(r=r))
    wb.save(saisie)
    wb.close()

    impacts = tmp_path / "SAISIE_Charges_Impacts.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, headers in IMPACTS_HEADERS.items():
        wb.create_sheet(sheet).append(headers)
    wb["AFFECTATIONS"].append(["AUTRE-AFF-001", "CHG-AUTRE", "2026-05", "LOG_0009", "PROP_01",
                              42.0, "A_CONTROLER", "NOUVELLE_CHARGE_GUIDEE", None, "hash"])
    wb.save(impacts)
    wb.close()
    return saisie, impacts


def _row_data(charge_id: str = CHARGE_ID) -> dict:
    return {
        "charge_id": charge_id, "date_charge": "2026-06-15", "montant": MONTANT,
        "sens_flux": "DEPENSE", "categorie_charge_id": "CHG_025", "type_flux_id": "TYPE_FLUX_020",
        "code_impact": "IC", "prise_en_compta": "OUI", "mode_paiement_id": "PAY_001",
        "affectation_type": "GLOBAL", "refacturable": "NON", "source_flux": "SAISIE_MANUELLE",
        "statut_controle": "A_CONTROLER", "niveau_anomalie": "INFO",
        "avantage_associe_id": "PERS_EWAN",
    }


def _persistable(charge_id: str = CHARGE_ID) -> dict:
    return {
        "charge_id": charge_id,
        "affectations": [{
            "affectation_id": f"{charge_id}-AFF-001", "charge_id": charge_id, "mois": "2026-06",
            "logement_id": "LOG_0001", "proprietaire_id": "PROP_01", "quote_part": MONTANT,
            "statut": "A_CONTROLER", "origine": "NOUVELLE_CHARGE_GUIDEE", "commentaire": None,
            "ROW_HASH": "h1",
        }],
        "menage": [], "reserve": [],
    }


def _demande(cibles, **kw) -> svc.DemandeEcritureCharge:
    """Verrou, master ET journal isolés sous tmp.

    `db_path` est obligatoire ici : sans lui, la journalisation écrirait dans la VRAIE base
    applicative (05_APPLICATION/data/app.db) et polluerait le journal de production.
    """
    saisie, impacts = cibles
    tmp = saisie.parent
    params = {
        "charge_id": CHARGE_ID, "montant": MONTANT, "target_row": LIGNE,
        "row_data": _row_data(), "persistable": _persistable(),
        "saisie_path": saisie, "impacts_path": impacts,
        "master_path": tmp / "verrous" / "MASTER_absent.xlsx",   # absent : aucune collision
        "lock_path": tmp / "verrous" / lk.LOCK_NAME,
        # Journal isolé, dans un sous-dossier DISTINCT : sa base (+ -wal/-shm) ne doit pas polluer
        # le dossier des classeurs, que les tests de résidus inspectent au fichier près.
        "db_path": tmp / "journal" / "journal_test.db",
    }
    params.update(kw)
    return svc.DemandeEcritureCharge(**params)


def _fichiers_du_dossier(p: Path) -> list[str]:
    return sorted(f.name for f in p.iterdir() if f.is_file())


def _replace_qui_echoue_a(appel_ko: int):
    """Fabrique un os.replace qui échoue au N-ième appel et délègue au vrai sinon."""
    vrai = os.replace
    etat = {"n": 0}

    def faux(src, dst, *a, **kw):
        etat["n"] += 1
        if etat["n"] == appel_ko:
            raise OSError(f"échec simulé au remplacement n°{appel_ko}")
        return vrai(src, dst, *a, **kw)

    return faux, etat


# ── 1-2 : succès et garde des flags ──────────────────────────────────────────

def test_1_transaction_reussie_sur_tous_les_fichiers(cibles, flags_on):
    saisie, impacts = cibles
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.statut == svc.STATUT_OK, res
    assert res.ok
    assert sorted(res.fichiers_remplaces) == sorted([str(impacts), str(saisie)])
    assert res.nettoyage_incomplet == []

    wb = openpyxl.load_workbook(saisie, read_only=True, data_only=True)
    lignes = [r for r in wb["SAISIE"].iter_rows(values_only=True) if r[0] == CHARGE_ID]
    wb.close()
    assert len(lignes) == 1

    wb = openpyxl.load_workbook(impacts, read_only=True, data_only=True)
    aff = [r for r in wb["AFFECTATIONS"].iter_rows(values_only=True) if r[1] == CHARGE_ID]
    wb.close()
    assert len(aff) == 1 and aff[0][5] == MONTANT

    # Les empreintes finales annoncées correspondent aux fichiers sur disque.
    assert res.sha256_finaux[saisie.name] == _sha(saisie)
    assert res.sha256_finaux[impacts.name] == _sha(impacts)


def test_2_refus_si_les_flags_sont_desactives(cibles):
    """Aucun monkeypatch : les flags réels (False) doivent bloquer AVANT toute préparation."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    dossier = _fichiers_du_dossier(saisie.parent)

    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.statut == svc.STATUT_GARDE
    assert res.code == svc.E_FLAGS_DESACTIVES
    assert "CHARGES_REAL_WRITE_ENABLED" in res.details
    assert {p: _sha(p) for p in (saisie, impacts)} == avant       # rien touché
    assert _fichiers_du_dossier(saisie.parent) == dossier         # aucun temporaire créé


@pytest.mark.parametrize("flag_on", ["CHARGES_REAL_WRITE_ENABLED",
                                     "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED"])
def test_2b_un_seul_flag_ne_suffit_pas(cibles, monkeypatch, flag_on):
    monkeypatch.setattr(cfg, flag_on, True)          # l'autre reste False
    res = svc.confirmer_ecriture_charge(_demande(cibles))
    assert res.statut == svc.STATUT_GARDE
    assert res.code == svc.E_FLAGS_DESACTIVES


# ── 3-4 : échec de préparation → aucun fichier réel modifié ──────────────────

def test_3_echec_preparation_du_premier_temporaire(cibles, flags_on):
    """Impacts (préparé en premier) invalide : Σ quote_part ≠ montant."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    dossier = _fichiers_du_dossier(saisie.parent)

    d = _demande(cibles)
    d.persistable["affectations"][0]["quote_part"] = 60.0        # Σ = 60 ≠ 100
    res = svc.confirmer_ecriture_charge(d)

    assert res.statut == svc.STATUT_ERREUR
    assert res.code == svc.E_PREPARATION
    assert res.fichiers_remplaces == []
    assert {p: _sha(p) for p in (saisie, impacts)} == avant
    assert _fichiers_du_dossier(saisie.parent) == dossier         # aucun temporaire laissé


def test_4_echec_preparation_du_temporaire_intermediaire(cibles, flags_on):
    """Impacts OK, puis la charge est refusée (colonne formule visée) : rien ne doit rester."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    dossier = _fichiers_du_dossier(saisie.parent)

    d = _demande(cibles)
    d.row_data["mois"] = "TRICHE"                                 # colonne formule → refus writer
    res = svc.confirmer_ecriture_charge(d)

    assert res.statut == svc.STATUT_ERREUR
    assert res.code == svc.E_PREPARATION
    assert "E_COLONNE_FORMULE" in res.details
    assert {p: _sha(p) for p in (saisie, impacts)} == avant
    # Le temporaire des impacts, déjà produit, a bien été nettoyé (test 15).
    assert _fichiers_du_dossier(saisie.parent) == dossier


def test_18_erreur_de_validation_writer_ne_touche_aucun_fichier(cibles, flags_on):
    """Charge ménage portant une réserve : refus métier du writer, aucun remplacement."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}

    d = _demande(cibles)
    d.persistable["affectations"] = []
    d.persistable["menage"] = [{
        "menage_impact_id": f"{CHARGE_ID}-MEN-001", "charge_id": CHARGE_ID, "mois": "2026-06",
        "mode": "INTERVENANT", "intervenant_id": "INT_0001", "logement_id": None,
        "statut": "A_CONTROLER", "origine": "NOUVELLE_CHARGE_GUIDEE", "commentaire": None,
        "ROW_HASH": "m1",
    }]
    d.persistable["reserve"] = [{
        "reserve_id": f"{CHARGE_ID}-RES-001", "charge_id": CHARGE_ID, "mois": "2026-06",
        "logement_id": "LOG_0001", "proprietaire_id": "PROP_01", "montant_refacturable": MONTANT,
        "libelle": "x", "justificatif": None, "statut_traitement": "EN_ATTENTE",
        "trace_decision": None, "origine": "NOUVELLE_CHARGE_GUIDEE", "commentaire": None,
        "ROW_HASH": "r1",
    }]
    res = svc.confirmer_ecriture_charge(d)

    assert res.statut == svc.STATUT_ERREUR
    assert "E_MENAGE_AVEC_RESERVE" in res.details
    assert {p: _sha(p) for p in (saisie, impacts)} == avant


# ── 5-6 : temporaires manquants ou invalides → aucun remplacement ────────────

def test_5_aucun_remplacement_si_un_temporaire_est_manquant(cibles, flags_on, monkeypatch):
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    vrai = writer.prepare_charge_flux_write

    def prepare_puis_supprime(*a, **kw):
        res = vrai(*a, **kw)
        Path(res["temp_path"]).unlink()               # le temporaire disparaît après préparation
        return res

    monkeypatch.setattr(writer, "prepare_charge_flux_write", prepare_puis_supprime)
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.code == svc.E_TEMPORAIRE_MANQUANT
    assert res.fichiers_remplaces == []
    assert {p: _sha(p) for p in (saisie, impacts)} == avant


def test_6_aucun_remplacement_si_un_temporaire_est_invalide(cibles, flags_on, monkeypatch):
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    vrai = writer.prepare_charge_impacts_write

    def prepare_puis_corrompt(*a, **kw):
        res = vrai(*a, **kw)
        Path(res["temp_path"]).write_bytes(b"ceci n'est pas un classeur")
        return res

    monkeypatch.setattr(writer, "prepare_charge_impacts_write", prepare_puis_corrompt)
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.code == svc.E_TEMPORAIRE_INVALIDE
    assert res.fichiers_remplaces == []
    assert {p: _sha(p) for p in (saisie, impacts)} == avant
    assert _fichiers_du_dossier(saisie.parent) == ["SAISIE_Charges_Flux.xlsx",
                                                   "SAISIE_Charges_Impacts.xlsx"]


# ── 7-8 : sauvegardes avant remplacement, ordre déterministe ─────────────────

def test_7_sauvegardes_creees_avant_le_premier_remplacement(cibles, flags_on, monkeypatch):
    saisie, impacts = cibles
    vu: list[list[str]] = []
    vrai = svc._remplacer_fichier

    def espion(temp, cible):
        vu.append(sorted(f.name for f in cible.parent.iterdir() if ".bak" in f.name))
        return vrai(temp, cible)

    monkeypatch.setattr(svc, "_remplacer_fichier", espion)
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.ok, res
    # Dès le PREMIER remplacement, les deux sauvegardes existent déjà.
    assert len(vu[0]) == 2
    assert any("SAISIE_Charges_Flux" in n for n in vu[0])
    assert any("SAISIE_Charges_Impacts" in n for n in vu[0])


def test_8_ordre_deterministe_des_remplacements(cibles, flags_on, monkeypatch):
    """Impacts d'abord, charge ensuite : des impacts orphelins sont préférables à une charge
    sans impacts (qui, elle, serait comptée à tort par le résultat)."""
    saisie, impacts = cibles
    ordre: list[str] = []
    vrai = svc._remplacer_fichier

    def espion(temp, cible):
        ordre.append(cible.name)
        return vrai(temp, cible)

    monkeypatch.setattr(svc, "_remplacer_fichier", espion)
    assert svc.confirmer_ecriture_charge(_demande(cibles)).ok
    assert ordre == ["SAISIE_Charges_Impacts.xlsx", "SAISIE_Charges_Flux.xlsx"]


# ── 9-12 : rollback à chaque position + vérification des empreintes ──────────

def test_9_rollback_si_le_premier_remplacement_echoue(cibles, flags_on, monkeypatch):
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}

    faux, _ = _replace_qui_echoue_a(1)
    monkeypatch.setattr(os, "replace", faux)
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.statut == svc.STATUT_ROLLBACK
    assert res.code == svc.E_REMPLACEMENT
    assert {p: _sha(p) for p in (saisie, impacts)} == avant       # test 12 : empreintes initiales
    assert res.nettoyage_incomplet == []


def test_10_11_rollback_si_le_dernier_remplacement_echoue(cibles, flags_on, monkeypatch):
    """La transaction porte deux cibles : la position « intermédiaire » et la position « dernière »
    coïncident. Ici les impacts SONT déjà remplacés — le rollback doit les défaire."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}

    faux, etat = _replace_qui_echoue_a(2)
    monkeypatch.setattr(os, "replace", faux)
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.statut == svc.STATUT_ROLLBACK
    assert res.code == svc.E_REMPLACEMENT
    assert etat["n"] >= 2                                    # le 1er remplacement avait bien eu lieu
    # test 12 : les DEUX fichiers sont revenus à leur empreinte initiale (impacts restaurés).
    assert {p: _sha(p) for p in (saisie, impacts)} == avant
    assert "restaurés" in res.details


def test_11b_rollback_si_la_verification_post_commit_echoue(cibles, flags_on, monkeypatch):
    """Troisième position d'échec : après les deux remplacements, un fichier final est illisible."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}

    # Illisible UNIQUEMENT pour les cibles (les temporaires, eux, sont validés normalement) :
    # l'échec se produit donc bien après les deux remplacements.
    monkeypatch.setattr(svc, "_lisible", lambda p: not Path(p).name.startswith("SAISIE_"))
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.statut == svc.STATUT_ROLLBACK
    assert res.code == svc.E_POST_COMMIT
    assert {p: _sha(p) for p in (saisie, impacts)} == avant       # tout est revenu en arrière


# ── 13 / 17 : rollback impossible → erreur critique, sauvegardes conservées ──

def test_13_17_erreur_critique_si_une_restauration_echoue(cibles, flags_on, monkeypatch):
    saisie, impacts = cibles

    faux, _ = _replace_qui_echoue_a(2)                # le 2e remplacement échoue
    monkeypatch.setattr(os, "replace", faux)

    def restauration_ko(sauvegarde, cible):
        raise OSError("restauration impossible (disque en lecture seule)")

    monkeypatch.setattr(svc, "_restaurer_fichier", restauration_ko)

    with pytest.raises(svc.RollbackCritiqueError) as exc:
        svc.confirmer_ecriture_charge(_demande(cibles))

    err = exc.value
    assert len(err.fichiers_non_restaures) == 1
    assert "SAISIE_Charges_Impacts" in err.fichiers_non_restaures[0]["cible"]
    assert "ROLLBACK INCOMPLET" in str(err)

    # test 17 : la sauvegarde du fichier non restauré est CONSERVÉE — seule voie de retour.
    bak = Path(err.fichiers_non_restaures[0]["sauvegarde"])
    assert bak.exists()
    assert _sha(bak) != _sha(impacts)          # elle porte bien l'état d'AVANT (impacts remplacé)


# ── 14-16 : nettoyage ────────────────────────────────────────────────────────

def test_14_16_nettoyage_des_temporaires_et_sauvegardes_apres_succes(cibles, flags_on):
    saisie, _ = cibles
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.ok, res
    restants = _fichiers_du_dossier(saisie.parent)
    assert restants == ["SAISIE_Charges_Flux.xlsx", "SAISIE_Charges_Impacts.xlsx"]
    assert not any(".bak" in n for n in restants)        # test 16 : plus aucune sauvegarde
    assert res.sauvegardes_restantes == []
    assert res.nettoyage_incomplet == []


def test_15_nettoyage_des_temporaires_apres_echec_avant_commit(cibles, flags_on, monkeypatch):
    saisie, _ = cibles
    vrai = writer.prepare_charge_flux_write

    def prepare_puis_corrompt(*a, **kw):
        res = vrai(*a, **kw)
        Path(res["temp_path"]).write_bytes(b"corrompu")
        return res

    monkeypatch.setattr(writer, "prepare_charge_flux_write", prepare_puis_corrompt)
    res = svc.confirmer_ecriture_charge(_demande(cibles))

    assert res.code == svc.E_TEMPORAIRE_INVALIDE
    # Les deux temporaires (impacts + charge) ont été nettoyés, aucune sauvegarde n'a été créée.
    assert _fichiers_du_dossier(saisie.parent) == ["SAISIE_Charges_Flux.xlsx",
                                                   "SAISIE_Charges_Impacts.xlsx"]
    assert res.nettoyage_incomplet == []


# ── 19-20 : fichiers réels et flags du dépôt ─────────────────────────────────

def test_19_les_vrais_fichiers_du_projet_ne_sont_jamais_touches(cibles, flags_on):
    presents = [p for p in REELS if p.exists()]
    avant = {p: _sha(p) for p in presents}

    assert svc.confirmer_ecriture_charge(_demande(cibles)).ok

    assert {p: _sha(p) for p in presents} == avant


def test_20_les_flags_du_depot_restent_a_false():
    """Les flags ne sont forcés qu'en mémoire (monkeypatch). Le dépôt, lui, reste sûr par défaut :
    aucune valeur littérale True codée en dur, activation possible uniquement via RECETTE_MODE ET
    variable d'environnement dédiée (double verrou)."""
    source = Path(cfg.__file__).read_text(encoding="utf-8")
    assert "CHARGES_REAL_WRITE_ENABLED = RECETTE_MODE and _env_flag(" in source
    assert "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED = RECETTE_MODE and _env_flag(" in source
    assert cfg.CHARGES_REAL_WRITE_ENABLED is False        # hors monkeypatch/env
    assert cfg.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED is False


def test_20b_os_replace_uniquement_dans_l_orchestrateur():
    """Le writer bas niveau ne remplace jamais : l'unique os.replace du flux est ici."""
    assert "os.replace(" not in Path(writer.__file__).read_text(encoding="utf-8")
    assert "os.replace(" in Path(svc.__file__).read_text(encoding="utf-8")


# ═════════════════════════════════════════════════════════════════════════════
# Commit 4 — verrou et réservation du charge_id
# ═════════════════════════════════════════════════════════════════════════════

def _verrou(d: svc.DemandeEcritureCharge) -> Path:
    return Path(d.lock_path)


# ── 15-18 : collision de charge_id détectée SOUS verrou, avant préparation ───

def test_15_identifiant_fourni_libre_accepte(cibles, flags_on):
    res = svc.confirmer_ecriture_charge(_demande(cibles))
    assert res.ok, res
    assert res.charge_id == CHARGE_ID


def test_16_identifiant_deja_present_dans_saisie_refuse(cibles, flags_on):
    saisie, impacts = cibles
    assert svc.confirmer_ecriture_charge(_demande(cibles)).ok      # 1re écriture
    avant = {p: _sha(p) for p in (saisie, impacts)}

    res = svc.confirmer_ecriture_charge(_demande(cibles, target_row=3))   # même charge_id
    assert res.statut == svc.STATUT_ERREUR
    assert res.code == svc.E_CHARGE_ID_EXISTANT
    assert "SAISIE_Charges_Flux" in res.details
    assert {p: _sha(p) for p in (saisie, impacts)} == avant        # rien préparé, rien remplacé


def test_17_identifiant_present_uniquement_dans_les_impacts_refuse(cibles, flags_on):
    saisie, impacts = cibles
    wb = openpyxl.load_workbook(impacts)
    wb["AFFECTATIONS"].append([f"{CHARGE_ID}-AFF-009", CHARGE_ID, "2026-06", "LOG_0001",
                               "PROP_01", 10.0, "A_CONTROLER", "AUTRE", None, "h"])
    wb.save(impacts)
    wb.close()
    avant = {p: _sha(p) for p in (saisie, impacts)}

    res = svc.confirmer_ecriture_charge(_demande(cibles))
    assert res.code == svc.E_CHARGE_ID_EXISTANT
    assert "SAISIE_Charges_Impacts" in res.details
    assert {p: _sha(p) for p in (saisie, impacts)} == avant


def test_18_identifiant_present_uniquement_dans_le_master_refuse(cibles, flags_on, tmp_path):
    """Un identifiant déjà consommé en aval (MASTER Lot3) reste une collision."""
    saisie, impacts = cibles
    master = tmp_path / "MASTER_FACT_MAN_Charges.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MASTER"
    ws.append(["charge_id", "mois", "montant"])
    ws.append(["[Charge par Power Query — placeholder]", None, None])   # gabarit : jamais une donnée
    ws.append([CHARGE_ID, "2026-06", MONTANT])
    wb.save(master)
    wb.close()
    avant = {p: _sha(p) for p in (saisie, impacts)}

    res = svc.confirmer_ecriture_charge(_demande(cibles, master_path=master))
    assert res.code == svc.E_CHARGE_ID_EXISTANT
    assert "MASTER_FACT_MAN_Charges" in res.details
    assert {p: _sha(p) for p in (saisie, impacts)} == avant


def test_18b_identifiant_de_syntaxe_invalide_refuse(cibles, flags_on):
    res = svc.confirmer_ecriture_charge(_demande(
        cibles, charge_id="CHG-TRUC", row_data={"charge_id": "CHG-TRUC"},
        persistable={"charge_id": "CHG-TRUC", "affectations": [], "menage": [], "reserve": []},
    ))
    assert res.code == svc.E_CHARGE_ID_INVALIDE


# ── 19 : génération de l'identifiant APRÈS acquisition du verrou ─────────────

def test_19_identifiant_genere_sous_verrou(cibles, flags_on):
    """La fabrique de payload n'est appelée qu'une fois le verrou tenu : c'est ce qui rend
    la séquence « lire les ids → en choisir un libre → écrire » réellement sûre."""
    saisie, impacts = cibles
    d_ref = _demande(cibles)
    verrou = _verrou(d_ref)
    vu: list[bool] = []

    def fabrique(charge_id: str):
        vu.append(verrou.exists())                 # le verrou est-il tenu à cet instant ?
        return _row_data(charge_id), _persistable(charge_id)

    d = _demande(cibles, charge_id="", row_data=None, persistable=None,
                 charge_id_prefixe="CHG-2026-06-IC-BANQUE", payload_factory=fabrique)
    res = svc.confirmer_ecriture_charge(d)

    assert res.ok, res
    assert res.charge_id == "CHG-2026-06-IC-BANQUE-001"   # premier numéro libre
    assert vu == [True]                                   # verrou TENU pendant la génération
    assert not verrou.exists()                            # et libéré ensuite


def test_19b_generation_saute_les_identifiants_deja_pris(cibles, flags_on):
    """Deux charges successives : la seconde reçoit -002, jamais -001 (relecture sous verrou)."""
    def fabrique(charge_id: str):
        return _row_data(charge_id), _persistable(charge_id)

    base = dict(charge_id="", row_data=None, persistable=None,
                charge_id_prefixe="CHG-2026-06-IC-BANQUE", payload_factory=fabrique)

    r1 = svc.confirmer_ecriture_charge(_demande(cibles, **base))
    assert r1.charge_id == "CHG-2026-06-IC-BANQUE-001", r1
    r2 = svc.confirmer_ecriture_charge(_demande(cibles, target_row=3, **base))
    assert r2.ok, r2
    assert r2.charge_id == "CHG-2026-06-IC-BANQUE-002"


# ── 23-24 : le verrou couvre les remplacements ET la vérification post-commit ─

def test_23_verrou_tenu_pendant_tous_les_remplacements(cibles, flags_on, monkeypatch):
    d = _demande(cibles)
    verrou = _verrou(d)
    presence: list[bool] = []
    vrai = svc._remplacer_fichier

    def espion(temp, cible):
        presence.append(verrou.exists())
        return vrai(temp, cible)

    monkeypatch.setattr(svc, "_remplacer_fichier", espion)
    assert svc.confirmer_ecriture_charge(d).ok
    assert presence == [True, True]        # verrou tenu aux DEUX os.replace


def test_24_verrou_libere_seulement_apres_la_verification_post_commit(cibles, flags_on, monkeypatch):
    d = _demande(cibles)
    verrou = _verrou(d)
    presence: list[bool] = []
    vrai = svc._verifier_post_commit

    def espion(remplaces):
        presence.append(verrou.exists())
        return vrai(remplaces)

    monkeypatch.setattr(svc, "_verifier_post_commit", espion)
    res = svc.confirmer_ecriture_charge(d)

    assert res.ok, res
    assert presence == [True]              # encore tenu pendant la vérification
    assert not verrou.exists()             # libéré seulement après


# ── 25 : flags désactivés → aucun verrou n'est même créé ─────────────────────

def test_25_flags_desactives_aucun_verrou_cree(cibles):
    d = _demande(cibles)
    res = svc.confirmer_ecriture_charge(d)
    assert res.statut == svc.STATUT_GARDE
    assert not _verrou(d).exists()
    assert not _verrou(d).parent.exists()   # même le dossier du verrou n'a pas été créé


# ── 27-28 : aucun résidu de verrou ───────────────────────────────────────────

def test_27_aucun_residu_de_verrou_apres_succes(cibles, flags_on):
    d = _demande(cibles)
    assert svc.confirmer_ecriture_charge(d).ok
    assert not _verrou(d).exists()


def test_28_aucun_residu_de_verrou_apres_erreur_recuperee(cibles, flags_on):
    """Erreur de préparation (Σ quotes ≠ montant) : le verrou est libéré dans le finally."""
    d = _demande(cibles)
    d.persistable["affectations"][0]["quote_part"] = 60.0
    res = svc.confirmer_ecriture_charge(d)
    assert res.code == svc.E_PREPARATION
    assert not _verrou(d).exists()


def test_28b_verrou_libere_apres_rollback(cibles, flags_on, monkeypatch):
    d = _demande(cibles)
    faux, _ = _replace_qui_echoue_a(2)
    monkeypatch.setattr(os, "replace", faux)

    res = svc.confirmer_ecriture_charge(d)
    assert res.statut == svc.STATUT_ROLLBACK
    assert not _verrou(d).exists()          # libéré malgré le rollback


def test_28c_verrou_libere_apres_rollback_critique(cibles, flags_on, monkeypatch):
    """Même un état critique ne laisse pas de verrou bloqué : l'incident est porté par l'exception."""
    d = _demande(cibles)
    faux, _ = _replace_qui_echoue_a(2)
    monkeypatch.setattr(os, "replace", faux)
    monkeypatch.setattr(svc, "_restaurer_fichier",
                        lambda s, c: (_ for _ in ()).throw(OSError("restauration impossible")))

    with pytest.raises(svc.RollbackCritiqueError):
        svc.confirmer_ecriture_charge(d)
    assert not _verrou(d).exists()


# ── 20-22 : deux VRAIS processus concurrents ─────────────────────────────────

def _worker_transaction(params: dict, ev_verrou_pris, ev_continuer, resultats) -> None:
    """Exécuté dans un processus séparé (spawn). Les flags ne sont forcés qu'EN MÉMOIRE ici."""
    import app.config as cfg_enfant
    cfg_enfant.CHARGES_REAL_WRITE_ENABLED = True
    cfg_enfant.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED = True

    from app.services import saisie_charges_transaction_service as svc_enfant

    saisie = Path(params["saisie"])
    impacts = Path(params["impacts"])

    if params["role"] == "A":
        # A s'arrête AU MILIEU de la section critique, verrou tenu, avant toute préparation.
        vrai_preparer = svc_enfant._preparer

        def preparer_bloquant(d):
            ev_verrou_pris.set()               # « je tiens le verrou »
            ev_continuer.wait(timeout=30)      # B a fini d'essayer
            return vrai_preparer(d)

        svc_enfant._preparer = preparer_bloquant
    else:
        ev_verrou_pris.wait(timeout=30)        # B n'essaie QUE pendant que A tient le verrou

    fichiers_avant = sorted(f.name for f in saisie.parent.iterdir() if f.is_file())
    hashes_avant = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (saisie, impacts)}

    demande = svc_enfant.DemandeEcritureCharge(
        charge_id=params["charge_id"], montant=params["montant"], target_row=params["target_row"],
        row_data=params["row_data"], persistable=params["persistable"],
        saisie_path=saisie, impacts_path=impacts,
        master_path=Path(params["master"]), lock_path=Path(params["lock"]),
        db_path=Path(params["db"]),                 # journal isolé, jamais la base de production
    )
    res = svc_enfant.confirmer_ecriture_charge(demande)

    fichiers_apres = sorted(f.name for f in saisie.parent.iterdir() if f.is_file())
    hashes_apres = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (saisie, impacts)}

    resultats.put({
        "role": params["role"], "statut": res.statut, "code": res.code,
        "charge_id": res.charge_id, "fichiers_remplaces": res.fichiers_remplaces,
        "fichiers_avant": fichiers_avant, "fichiers_apres": fichiers_apres,
        "hashes_avant": hashes_avant, "hashes_apres": hashes_apres,
    })
    if params["role"] == "B":
        ev_continuer.set()                     # B a fini : A peut reprendre


def test_20_21_22_deux_processus_concurrents(cibles, tmp_path):
    """Deux VRAIS processus (multiprocessing, spawn) tentent d'écrire la même charge.

    Synchronisation déterministe par Event : B n'essaie que pendant que A tient le verrou —
    aucun sleep fragile. B doit être refusé, sans préparer ni modifier quoi que ce soit.
    """
    saisie, impacts = cibles
    ctx = mp.get_context("spawn")              # comportement Windows réel, pas un fork
    ev_verrou_pris = ctx.Event()
    ev_continuer = ctx.Event()
    resultats = ctx.Queue()

    base = {
        "saisie": str(saisie), "impacts": str(impacts),
        "master": str(tmp_path / "verrous" / "MASTER_absent.xlsx"),
        "lock": str(tmp_path / "verrous" / lk.LOCK_NAME),
        "db": str(tmp_path / "journal" / "journal_test.db"),
        "charge_id": CHARGE_ID, "montant": MONTANT, "target_row": LIGNE,
        "row_data": _row_data(), "persistable": _persistable(),
    }
    pa = ctx.Process(target=_worker_transaction,
                     args=({**base, "role": "A"}, ev_verrou_pris, ev_continuer, resultats))
    pb = ctx.Process(target=_worker_transaction,
                     args=({**base, "role": "B"}, ev_verrou_pris, ev_continuer, resultats))
    pa.start()
    pb.start()
    recus = [resultats.get(timeout=60), resultats.get(timeout=60)]
    pa.join(timeout=60)
    pb.join(timeout=60)
    assert pa.exitcode == 0 and pb.exitcode == 0

    par_role = {r["role"]: r for r in recus}
    a, b = par_role["A"], par_role["B"]

    # 20. Les deux ne sont jamais entrés ensemble : A écrit, B est refusé par le VERROU.
    assert a["statut"] == svc.STATUT_OK, a
    assert b["statut"] == svc.STATUT_ERREUR, b
    assert b["code"] == svc.E_VERROU_DEJA_PRIS, b

    # 21. B n'a préparé AUCUN temporaire (le contenu du dossier est identique avant/après son essai).
    assert b["fichiers_avant"] == b["fichiers_apres"]
    assert b["fichiers_apres"] == ["SAISIE_Charges_Flux.xlsx", "SAISIE_Charges_Impacts.xlsx"]

    # 22. B n'a modifié AUCUN fichier.
    assert b["hashes_avant"] == b["hashes_apres"]
    assert b["fichiers_remplaces"] == []

    # État final : une seule charge écrite, par A, et aucun verrou résiduel.
    wb = openpyxl.load_workbook(saisie, read_only=True, data_only=True)
    lignes = [r for r in wb["SAISIE"].iter_rows(values_only=True) if r[0] == CHARGE_ID]
    wb.close()
    assert len(lignes) == 1
    assert not (tmp_path / "verrous" / lk.LOCK_NAME).exists()
