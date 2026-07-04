"""Tests APP-2c - previsualisation HH sur copies uniquement."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path
import shutil
from unittest.mock import patch

import openpyxl
import pytest

from app.services import saisie_hh_dryrun_service as dryrun_svc
from app.writers.saisie_hh_writer import _check_formula_cells


ARTIFACT = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "dapp05c_formules"
    / "20260702T143503Z"
    / "SAISIE_copie.xlsx"
)


def _copy_saisie(tmp_path: Path) -> Path:
    if not ARTIFACT.exists():
        pytest.fail(f"Artefact D-APP-05C absent: {ARTIFACT}")
    target = tmp_path / "source_saisie.xlsx"
    shutil.copy2(ARTIFACT, target)
    return target


def _minimal_ref(tmp_path: Path) -> Path:
    ref = tmp_path / "REF_Setup_test.xlsm"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "REF_Modes_Paiement"
    ws.append(["mode_paiement_id", "mode_paiement", "impact_banque", "impact_caisse", "impact_associee", "actif"])
    ws.append(["PAY_001", "BANQUE_PRO", "OUI", "NON", "NON", "OUI"])
    ws.append(["PAY_002", "ESPECES_CAISSE", "NON", "OUI", "NON", "OUI"])
    ws.append(["PAY_003", "CARTE_ASSOCIEE", "NON", "NON", "OUI", "OUI"])
    ws.append(["PAY_004", "COMPTE_PERSO_ASSOCIEE", "NON", "NON", "OUI", "OUI"])
    wb.create_sheet("REF_Taux_Commission").append([
        "proprietaire_id", "logement_id", "taux_commission", "date_debut", "date_fin", "actif",
    ])
    wb.save(ref)
    wb.close()
    return ref


def _preview(**overrides):
    data = {
        "reservation_hh_id": "RESHH-2026-08-001",
        "canal_id": "CANAL_001",
        "source_financiere": "SAISIE_MANUELLE",
        "proprietaire_id": "PROP_0001",
        "logement_id": "LOG_0001",
        "reservation_id_hostaway": None,
        "date_arrivee": "2026-08-15",
        "date_depart": "2026-08-18",
        "total_percu": Decimal("450.00"),
        "menage": Decimal("60.00"),
        "menage_standard": Decimal("60.00"),
        "menage_override": None,
        "motif_override_menage": None,
        "confirmation_override_menage": False,
        "taux_commission_standard": Decimal("0.15"),
        "taux_commission_override": None,
        "motif_override_taux_commission": None,
        "confirmation_override_taux_commission": False,
        "commentaire_taux_commission": None,
        "montant_recupere": None,
        "associe_id_recuperateur": None,
        "montant_reverse_proprietaire": None,
        "mode_paiement_id": "PAY_001",
        "code_impact": "HC",
        "comptabilisation": "NON",
        "statut_controle": "A_CONTROLER",
        "niveau_anomalie": "A_CONTROLER",
        "code_anomalie": None,
        "commentaire": None,
        "mois": "2026-08",
    }
    data.update(overrides)
    return data


def _validation(preview):
    return {"ok": True, "erreurs": [], "pk": preview["reservation_hh_id"], "preview": preview}


class FakeLot4A:
    MASTER_COLUMNS = [
        "reservation_hh_id", "mois", "canal_id", "source_financiere",
        "proprietaire_id", "logement_id", "date_arrivee", "date_depart", "nuits",
        "total_percu", "menage", "taux_commission", "taux_commission_source",
        "commission", "montant_recupere", "associe_id_recuperateur",
        "montant_reverse_proprietaire", "mode_paiement_id", "acompte_facture",
        "source_acompte_facture", "code_impact", "comptabilisation",
        "impact_resultat_reel", "impact_resultat_comptable", "statut_controle",
        "niveau_anomalie", "taux_commission_override", "motif_override_taux_commission",
        "confirmation_override_taux_commission", "menage_override",
        "motif_override_menage", "confirmation_override_menage",
    ]

    def __init__(self, fail=False):
        self.fail = fail
        self.saisie_paths = []
        self.ref_paths = []

    def read_saisie_values(self, path):
        self.saisie_paths.append(Path(path))
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            rows = list(wb["SAISIE"].iter_rows(values_only=True))
        finally:
            wb.close()
        headers = [str(h or "").strip() for h in rows[0]]
        out = []
        for row in rows[1:]:
            rec = dict(zip(headers, row))
            if str(rec.get("reservation_hh_id") or "").startswith("RESHH-"):
                out.append(rec)
        return out

    def read_taux_rows(self, path):
        self.ref_paths.append(Path(path))
        return []

    def build_master(self, saisie_list, taux_rows, as_of_iso):
        if self.fail:
            return {
                "statut": "ANALYSE_BLOQUEE_TAUX",
                "master_rows": [],
                "vue_active_rows": [],
                "anomalies_donnees": [],
                "anomalies_taux": [{"reservation_hh_id": "RESHH-2026-08-001", "statut": "TAUX_ABSENT"}],
                "motif_blocage": "TAUX_ABSENT",
            }
        rows = [self._master_row(row) for row in saisie_list]
        return {
            "statut": "ANALYSE_TERMINEE",
            "master_rows": rows,
            "vue_active_rows": rows,
            "anomalies_donnees": [],
            "anomalies_taux": [],
            "motif_blocage": None,
        }

    def _master_row(self, row):
        total = float(row.get("total_percu") or 0)
        raw_menage = row.get("menage_override")
        menage = float(row.get("menage") or 0) if raw_menage in (None, "") else float(raw_menage)
        raw_taux = row.get("taux_commission_override")
        taux = 0.15 if raw_taux in (None, "") else float(raw_taux)
        mode = str(row.get("mode_paiement_id") or "")
        reverse = float(row.get("montant_reverse_proprietaire") or 0)
        if mode in ("PAY_006", "DIRECT_PROPRIETAIRE"):
            acompte, source = 0.0, "DIRECT_PROPRIETAIRE"
        elif mode in ("PAY_002", "ESPECES_CAISSE", "ESPECES"):
            acompte, source = total - reverse, "TOTAL_PERCU_MOINS_REVERSE_ESPECES"
        elif mode in ("PAY_003", "CARTE_ASSOCIEE", "PAY_004", "COMPTE_PERSO_ASSOCIEE"):
            acompte, source = total, "TOTAL_PERCU_ASSOCIE"
        else:
            acompte, source = total, "TOTAL_PERCU"
        return {
            "reservation_hh_id": row.get("reservation_hh_id"),
            "mois": "2026-08",
            "canal_id": row.get("canal_id"),
            "source_financiere": row.get("source_financiere"),
            "proprietaire_id": row.get("proprietaire_id"),
            "logement_id": row.get("logement_id"),
            "date_arrivee": row.get("date_arrivee"),
            "date_depart": row.get("date_depart"),
            "nuits": 3,
            "total_percu": total,
            "menage": menage,
            "taux_commission": taux,
            "taux_commission_source": "OVERRIDE_CONFIRME" if row.get("taux_commission_override") not in (None, "") else "REF_PROPRIETAIRE",
            "commission": round((total - menage) * taux, 2),
            "montant_recupere": row.get("montant_recupere"),
            "associe_id_recuperateur": row.get("associe_id_recuperateur"),
            "montant_reverse_proprietaire": reverse,
            "mode_paiement_id": mode,
            "acompte_facture": acompte,
            "source_acompte_facture": source,
            "code_impact": row.get("code_impact"),
            "comptabilisation": row.get("comptabilisation"),
            "impact_resultat_reel": "OUI",
            "impact_resultat_comptable": "NON",
            "statut_controle": row.get("statut_controle"),
            "niveau_anomalie": row.get("niveau_anomalie"),
            "taux_commission_override": row.get("taux_commission_override"),
            "motif_override_taux_commission": row.get("motif_override_taux_commission"),
            "confirmation_override_taux_commission": row.get("confirmation_override_taux_commission"),
            "menage_override": row.get("menage_override"),
            "motif_override_menage": row.get("motif_override_menage"),
            "confirmation_override_menage": row.get("confirmation_override_menage"),
        }


def _run(tmp_path, preview=None, lot4a=None):
    saisie = _copy_saisie(tmp_path)
    ref = _minimal_ref(tmp_path)
    fake = lot4a or FakeLot4A()
    data = preview or _preview()
    with (
        patch("app.services.saisie_hh_dryrun_service.saisie_svc.valider", return_value=_validation(data)),
        patch("app.services.saisie_hh_dryrun_service._load_lot4a", return_value=fake),
    ):
        result = dryrun_svc.run_previsualisation(
            {
                "logement_id": data["logement_id"],
                "date_arrivee": data["date_arrivee"],
                "total_percu": str(data["total_percu"]),
            },
            saisie_source=saisie,
            ref_setup_source=ref,
            dryruns_root=tmp_path / "dryruns",
            as_of_iso="2026-07-04T00:00:00+00:00",
        )
    return result, fake, saisie, ref


def test_app2c_cree_dryrun_manifest_hashes_et_copies(tmp_path):
    result, fake, saisie_source, ref_source = _run(tmp_path)

    assert result["ok"] is True
    run_dir = result["run_dir"]
    assert run_dir.exists()
    assert (run_dir / dryrun_svc.MANIFEST_NAME).exists()
    assert (run_dir / dryrun_svc.SAISIE_COPY_NAME).exists()
    assert (run_dir / dryrun_svc.REF_COPY_NAME).exists()
    assert (run_dir / dryrun_svc.MASTER_SIM_NAME).exists()
    assert (run_dir / dryrun_svc.RESULT_LOT4A_NAME).exists()

    manifest = result["manifest"]
    assert manifest["source_hashes"]["saisie"]["sha256"]
    assert manifest["copy_hashes"]["saisie"]["sha256"]
    assert manifest["source_hashes"]["saisie"]["sha256"] != manifest["copy_hashes"]["saisie"]["sha256"]
    assert manifest["hh_real_write_enabled"] is False
    assert manifest["migration"]["saisie_fields_added"]
    assert result["saisie_copy"] in fake.saisie_paths
    assert result["ref_copy"] in fake.ref_paths
    assert saisie_source.read_bytes() != result["saisie_copy"].read_bytes()
    assert ref_source.exists()


def test_app2c_migration_et_formules_sur_copie_uniquement(tmp_path):
    result, _fake, saisie_source, _ref_source = _run(tmp_path)
    copied = result["saisie_copy"]

    assert _check_formula_cells(copied, result["target_row"]) == []
    assert _check_formula_cells(saisie_source, result["target_row"]) == []

    wb = openpyxl.load_workbook(str(copied), data_only=False)
    try:
        ws = wb["SAISIE"]
        headers = [cell.value for cell in ws[1]]
        assert headers[-len(dryrun_svc.NEW_SAISIE_FIELDS):] == dryrun_svc.NEW_SAISIE_FIELDS
        assert ws.cell(row=result["target_row"], column=1).value == "RESHH-2026-08-001"
    finally:
        wb.close()

    wb_source = openpyxl.load_workbook(str(saisie_source), read_only=True, data_only=True)
    try:
        assert wb_source["SAISIE"].cell(row=result["target_row"], column=1).value in (None, "")
    finally:
        wb_source.close()


@pytest.mark.parametrize(
    ("mode", "reverse", "expected_acompte", "expected_source"),
    [
        ("PAY_001", None, 450.0, "TOTAL_PERCU"),
        ("PAY_003", None, 450.0, "TOTAL_PERCU_ASSOCIE"),
        ("PAY_002", Decimal("120.00"), 330.0, "TOTAL_PERCU_MOINS_REVERSE_ESPECES"),
        ("PAY_006", None, 0.0, "DIRECT_PROPRIETAIRE"),
    ],
)
def test_app2c_lot4a_acomptes_sur_copie(tmp_path, mode, reverse, expected_acompte, expected_source):
    preview = _preview(mode_paiement_id=mode, montant_reverse_proprietaire=reverse)
    result, _fake, _saisie, _ref = _run(tmp_path, preview=preview)

    row = result["simulated_master_row"]
    assert row["acompte_facture"] == expected_acompte
    assert row["source_acompte_facture"] == expected_source
    assert row["mois"] == "2026-08"


@pytest.mark.parametrize(
    ("override", "expected"),
    [
        (Decimal("0.00"), 0.0),
        (Decimal("0.005"), 0.005),
    ],
)
def test_app2c_taux_override_canonique_propage_a_lot4a(tmp_path, override, expected):
    preview = _preview(
        taux_commission_override=override,
        motif_override_taux_commission="Derogation test",
        confirmation_override_taux_commission=True,
    )
    result, _fake, _saisie, _ref = _run(tmp_path, preview=preview)

    row = result["simulated_master_row"]
    assert row["taux_commission"] == expected
    assert row["taux_commission_source"] == "OVERRIDE_CONFIRME"


def test_app2c_echec_lot4a_visible_dans_manifest(tmp_path):
    result, _fake, _saisie, _ref = _run(tmp_path, lot4a=FakeLot4A(fail=True))

    assert result["ok"] is False
    assert result["manifest"]["status"] == "ERREUR_LOT4A"
    assert result["manifest"]["lot4a_status"] == "ANALYSE_BLOQUEE_TAUX"
    assert "TAUX_ABSENT" in ";".join(result["manifest"]["errors"])


def test_route_previsualisation_affiche_resultat_et_aucune_ecriture(client):
    dryrun = {
        "token": "20260704T120000Z_abcdef123456",
        "manifest": {
            "status": "OK",
            "errors": [],
            "target_row": 2,
            "lot4a_status": "ANALYSE_TERMINEE",
            "paths": {"run_dir": "tmp/dryruns/20260704T120000Z_abcdef123456"},
            "payload_summary": {
                "reservation_hh_id": "RESHH-2026-08-001",
                "logement_id": "LOG_0001",
                "proprietaire_id": "PROP_0001",
                "canal_id": "CANAL_001",
                "date_arrivee": "2026-08-15",
                "date_depart": "2026-08-18",
                "total_percu": "450.00",
                "mode_paiement_id": "PAY_001",
                "code_impact": "HC",
                "comptabilisation": "NON",
                "taux_commission_override": None,
                "menage_override": None,
            },
            "migration": {"saisie_fields_added": ["taux_commission_override"]},
        },
        "resultat_lot4a": {"reservation_simulee": {}},
        "simulated_master_row": {
            "nuits": 3,
            "taux_commission": 0.15,
            "taux_commission_source": "REF_PROPRIETAIRE",
            "menage": 60.0,
            "commission": 58.5,
            "acompte_facture": 450.0,
            "source_acompte_facture": "TOTAL_PERCU",
            "mois": "2026-08",
            "impact_resultat_reel": "OUI",
            "impact_resultat_comptable": "NON",
        },
        "comparatif": [
            {"element": "Lignes SAISIE HH", "avant": 0, "apres": 1},
            {"element": "Lignes MASTER HH", "avant": 0, "apres": 1},
        ],
    }
    with patch("app.routes.reservations.dryrun_svc.load_previsualisation", return_value=dryrun):
        resp = client.get("/reservations/nouvelle/previsualisation/20260704T120000Z_abcdef123456")

    assert resp.status_code == 200
    assert "Simulation sur copie" in resp.text
    assert "aucune donnée réelle n’a été modifiée" in resp.text
    assert "Résultat Lot4A" in resp.text
    assert "ANALYSE_TERMINEE" in resp.text
    assert "/reservations/nouvelle/confirmer" not in resp.text
    assert "Sauvegarder" not in resp.text


def test_route_post_previsualiser_redirige_vers_token(client):
    result = {
        "ok": True,
        "token": "20260704T120000Z_abcdef123456",
        "manifest": {"status": "OK"},
    }
    with patch("app.routes.reservations.dryrun_svc.run_previsualisation", return_value=result) as run:
        resp = client.post("/reservations/nouvelle/previsualiser", data={"total_percu": "450.00"}, follow_redirects=False)

    assert resp.status_code == 303
    assert resp.headers["location"].endswith("/reservations/nouvelle/previsualisation/20260704T120000Z_abcdef123456")
    run.assert_called_once()
