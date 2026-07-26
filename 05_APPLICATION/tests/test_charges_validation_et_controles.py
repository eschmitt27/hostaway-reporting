"""Parcours de validation des charges + contrôles d'intégrité (charge refacturable orpheline).

Tests unitaires isolés : chaque test construit ses propres classeurs en tmp_path (aucune dépendance
au jeu de recette, aucune écriture réelle).
"""
import openpyxl
import pytest

import app.config as cfg
from app.services import charges_controles_integrite_service as ctrl
from app.services import charges_validation_service as val

SAISIE_HDR = ["charge_id", "date_charge", "mois", "montant", "logement_id", "proprietaire_id",
              "refacturable", "code_impact", "statut_controle", "commentaire"]
MASTER_HDR = SAISIE_HDR


def _saisie(tmp_path, rows):
    p = tmp_path / "SAISIE_Charges_Flux.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SAISIE"; ws.append(SAISIE_HDR)
    for r in rows:
        ws.append([r.get(h) for h in SAISIE_HDR])
    wb.save(p); wb.close()
    return p


def _master(tmp_path, rows):
    p = tmp_path / "MASTER_FACT_MAN_Charges.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "MASTER"; ws.append(MASTER_HDR)
    for r in rows:
        ws.append([r.get(h) for h in MASTER_HDR])
    wb.save(p); wb.close()
    return p


def _ref(tmp_path, proprios=("PROP_A",), gestion=(("LOG_A1", "PROP_A", "2026-01-01", None),)):
    p = tmp_path / "REF_Setup.xlsm"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("REF_Proprietaires"); ws.append(["proprietaire_id", "actif"])
    for x in proprios:
        ws.append([x, "OUI"])
    ws = wb.create_sheet("REF_Gestion_Logements_Hist")
    ws.append(["gestion_id", "logement_id", "proprietaire_id", "date_debut", "date_fin", "statut_gestion"])
    for i, (lg, pr, d, f) in enumerate(gestion, 1):
        ws.append([f"G{i}", lg, pr, d, f, "ACTIF"])
    wb.save(p); wb.close()
    return p


CHARGE = {"charge_id": "CHG-1", "date_charge": "2026-06-15", "mois": "2026-06", "montant": 100,
          "logement_id": "LOG_A1", "proprietaire_id": "PROP_A", "refacturable": "OUI",
          "code_impact": "IC", "statut_controle": "A_CONTROLER"}


@pytest.fixture
def _recette(tmp_path, monkeypatch):
    """Mode recette + writers actifs, cantonnés à tmp_path."""
    monkeypatch.setattr(cfg, "RECETTE_MODE", True)
    monkeypatch.setattr(cfg, "RECETTE_ROOT", tmp_path.resolve())
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)
    monkeypatch.setattr(cfg, "DB_PATH", tmp_path / "app.db")
    return tmp_path


# ── Parcours de validation ───────────────────────────────────────────────────

def test_liste_expose_les_charges_a_controler(tmp_path, monkeypatch):
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", _saisie(tmp_path, [CHARGE]))
    res = val.lister(statut="A_CONTROLER")
    assert res["status"] == "OK"
    assert len(res["rows"]) == 1
    assert res["compteurs"]["A_CONTROLER"] == 1
    assert "aucun impact financier" in res["avertissement"]


def test_validation_passe_le_statut_a_valide_et_journalise(_recette, monkeypatch):
    p = _saisie(_recette, [CHARGE])
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", p)
    r = val.valider("CHG-1", acteur="testeur", commentaire="ok")
    assert r["ok"] and r["nouveau_statut"] == "VALIDE" and r["recalcul_requis"] is True
    # persistance sur disque
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ligne = list(wb["SAISIE"].iter_rows(min_row=2, values_only=True))[0]
    wb.close()
    assert ligne[SAISIE_HDR.index("statut_controle")] == "VALIDE"
    hist = val.historique("CHG-1")
    assert hist and hist[0]["nouveau_statut"] == "VALIDE" and hist[0]["acteur"] == "testeur"


def test_rejet_exige_un_commentaire(_recette, monkeypatch):
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", _saisie(_recette, [CHARGE]))
    assert val.rejeter("CHG-1", commentaire="")["code"] == val.E_COMMENTAIRE


def test_rejet_conserve_la_ligne(_recette, monkeypatch):
    p = _saisie(_recette, [CHARGE])
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", p)
    assert val.rejeter("CHG-1", commentaire="montant erroné")["ok"]
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    lignes = list(wb["SAISIE"].iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(lignes) == 1                      # jamais supprimée
    assert lignes[0][SAISIE_HDR.index("statut_controle")] == "REJETE"


def test_double_validation_refusee(_recette, monkeypatch):
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", _saisie(_recette, [CHARGE]))
    assert val.valider("CHG-1")["ok"]
    assert val.valider("CHG-1")["code"] == val.E_DEJA


def test_charge_inexistante_refusee(_recette, monkeypatch):
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", _saisie(_recette, [CHARGE]))
    assert val.valider("CHG-INCONNUE")["code"] == val.E_INTROUVABLE


def test_validation_refusee_si_flags_desactives(tmp_path, monkeypatch):
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", _saisie(tmp_path, [CHARGE]))
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", False)
    assert val.valider("CHG-1")["code"] == val.E_FLAGS


def test_validation_refusee_hors_racine_recette(tmp_path, monkeypatch):
    """Write-guard : cible hors RECETTE_ROOT → refus, fichier inchangé."""
    p = _saisie(tmp_path, [CHARGE])
    monkeypatch.setattr(cfg, "SAISIE_CHARGES", p)
    monkeypatch.setattr(cfg, "RECETTE_MODE", True)
    monkeypatch.setattr(cfg, "RECETTE_ROOT", (tmp_path / "ailleurs").resolve())
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)
    monkeypatch.setattr(cfg, "DB_PATH", tmp_path / "app.db")
    assert val.valider("CHG-1")["code"] == val.E_ECRITURE
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ligne = list(wb["SAISIE"].iter_rows(min_row=2, values_only=True))[0]
    wb.close()
    assert ligne[SAISIE_HDR.index("statut_controle")] == "A_CONTROLER"


# ── Contrôles d'intégrité (esprit Lot11) ─────────────────────────────────────

def test_controle_bloque_charge_refacturable_orpheline(tmp_path):
    orpheline = dict(CHARGE, proprietaire_id=None, statut_controle="VALIDE")
    res = ctrl.controler(master_charges=_master(tmp_path, [orpheline]), ref_path=_ref(tmp_path))
    assert res["statut"] == ctrl.BLOQUANT and res["nb_bloquants"] == 1
    a = res["anomalies"][0]
    assert a["code"] == ctrl.C_REFAC_ORPHELINE
    assert "préfacture" in a["message"]
    assert a["charge_id"] == "CHG-1" and a["logement_id"] == "LOG_A1" and a["montant"] == 100


def test_controle_ok_quand_proprietaire_materialise(tmp_path):
    ok = dict(CHARGE, statut_controle="VALIDE")
    res = ctrl.controler(master_charges=_master(tmp_path, [ok]), ref_path=_ref(tmp_path))
    assert res["statut"] == "OK" and res["recalcul_fiable"] is True


def test_controle_bloque_proprietaire_inconnu(tmp_path):
    mauvais = dict(CHARGE, proprietaire_id="PROP_ZZ", statut_controle="VALIDE")
    res = ctrl.controler(master_charges=_master(tmp_path, [mauvais]), ref_path=_ref(tmp_path))
    codes = [a["code"] for a in res["anomalies"]]
    assert ctrl.C_PROP_INCONNU in codes and res["nb_bloquants"] >= 1


def test_controle_bloque_proprietaire_incoherent_avec_logement(tmp_path):
    incoherent = dict(CHARGE, proprietaire_id="PROP_B", statut_controle="VALIDE")
    ref = _ref(tmp_path, proprios=("PROP_A", "PROP_B"),
               gestion=(("LOG_A1", "PROP_A", "2026-01-01", None),))
    res = ctrl.controler(master_charges=_master(tmp_path, [incoherent]), ref_path=ref)
    codes = [a["code"] for a in res["anomalies"]]
    assert ctrl.C_PROP_INCOHERENT in codes


def test_controle_ignore_charge_non_refacturable_sans_proprietaire(tmp_path):
    globale = dict(CHARGE, proprietaire_id=None, logement_id=None,
                   refacturable="NON", statut_controle="VALIDE")
    res = ctrl.controler(master_charges=_master(tmp_path, [globale]), ref_path=_ref(tmp_path))
    assert res["statut"] == "OK"


def test_controle_ignore_charge_non_validee(tmp_path):
    """Une charge refacturable sans propriétaire mais NON validée n'est pas bloquante."""
    en_attente = dict(CHARGE, proprietaire_id=None, statut_controle="A_CONTROLER")
    res = ctrl.controler(master_charges=_master(tmp_path, [en_attente]), ref_path=_ref(tmp_path))
    assert res["nb_bloquants"] == 0
