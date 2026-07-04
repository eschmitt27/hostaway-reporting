"""Tests APP-2d - activation controlee de l'ecriture reelle HH."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

import app.config as cfg
from app.services import saisie_hh_real_write_service as real_svc
from app.services.saisie_hh_schema_migration import NEW_SAISIE_FIELDS


TOKEN = "20260704T120000Z_app2dtest001"
PK = "RESHH-2026-08-001"


def _sha(path: Path) -> str:
    return real_svc._sha256(path)


def _make_saisie(path: Path, *, schema: bool = True, existing_pk: str | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    headers = [
        "reservation_hh_id", "canal_id", "source_financiere", "proprietaire_id",
        "logement_id", "date_arrivee", "date_depart", "total_percu",
    ]
    if schema:
        headers.extend(NEW_SAISIE_FIELDS)
    ws.append(headers)
    if existing_pk:
        ws.append([existing_pk])
    wb.save(path)
    wb.close()
    return path


def _make_ref(path: Path, *, schema: bool = True, statut_mois: str = "OUVERT") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "REF_Modes_Paiement"
    ws.append(["mode_paiement_id", "mode_paiement"])
    ws.append(["PAY_001", "BANQUE_PRO"])
    if schema:
        ws.append(["PAY_006", "DIRECT_PROPRIETAIRE"])
    ws_cloture = wb.create_sheet("REF_Cloture_Mensuelle")
    ws_cloture.append(["mois", "statut_mois"])
    ws_cloture.append(["2026-08", statut_mois])
    wb.save(path)
    wb.close()
    return path


def _write_manifest(
    root: Path,
    saisie: Path,
    ref: Path,
    *,
    created_at: datetime | None = None,
    status: str = "OK",
    lot4a_status: str = "ANALYSE_TERMINEE",
    errors: list[str] | None = None,
) -> Path:
    run_dir = root / TOKEN
    run_dir.mkdir(parents=True)
    payload = {
        "reservation_hh_id": PK,
        "mois": "2026-08",
        "date_arrivee": "2026-08-15",
        "date_depart": "2026-08-18",
        "logement_id": "LOG_0001",
        "proprietaire_id": "PROP_0001",
        "canal_id": "CANAL_001",
        "source_financiere": "SAISIE_MANUELLE",
        "total_percu": "450.00",
    }
    manifest = {
        "simulation_id": TOKEN,
        "created_at_utc": (created_at or datetime.now(timezone.utc)).isoformat(),
        "status": status,
        "lot4a_status": lot4a_status,
        "errors": errors or [],
        "source_hashes": {
            "saisie": {"sha256": _sha(saisie)},
            "ref_setup": {"sha256": _sha(ref)},
        },
        "payload_summary": payload,
        "payload_canonique": payload,
        "form_data": {
            "date_arrivee": "2026-08-15",
            "date_depart": "2026-08-18",
            "logement_id": "LOG_0001",
            "total_percu": "450.00",
        },
        "paths": {"run_dir": str(run_dir)},
    }
    (run_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    result = {
        "reservation_simulee": {
            "reservation_hh_id": PK,
            "mois": "2026-08",
            "taux_commission": 0.15,
            "commission": 58.5,
            "menage": 60,
            "acompte_facture": 450,
            "source_acompte_facture": "TOTAL_PERCU",
        }
    }
    (run_dir / "resultat_lot4a.json").write_text(json.dumps(result), encoding="utf-8")
    return run_dir


@pytest.fixture
def app2d_env(tmp_path):
    saisie = _make_saisie(tmp_path / "SAISIE_test.xlsx")
    ref = _make_ref(tmp_path / "REF_Setup_test.xlsm")
    dryruns = tmp_path / "dryruns"
    _write_manifest(dryruns, saisie, ref)
    return {"saisie": saisie, "ref": ref, "dryruns": dryruns}


def _eval(env):
    return real_svc.evaluate_real_write_prerequisites(
        TOKEN,
        saisie_path=env["saisie"],
        ref_setup_path=env["ref"],
        dryruns_root=env["dryruns"],
    )


def _codes(state):
    return {item["code"]: item["ok"] for item in state["checks"]}


def test_app2d_refuse_si_flag_global_desactive(app2d_env):
    with patch.object(cfg, "HH_REAL_WRITE_ENABLED", False), patch.object(cfg, "HH_REAL_WRITE_CONFIRMATION_ENABLED", True):
        state = _eval(app2d_env)
    assert state["eligible"] is False
    assert _codes(state)["FLAG_GLOBAL_ACTIVE"] is False


def test_app2d_refuse_si_flag_confirmation_desactive(app2d_env):
    with patch.object(cfg, "HH_REAL_WRITE_ENABLED", True), patch.object(cfg, "HH_REAL_WRITE_CONFIRMATION_ENABLED", False):
        state = _eval(app2d_env)
    assert state["eligible"] is False
    assert _codes(state)["FLAG_CONFIRMATION_ACTIVE"] is False


def test_app2d_refuse_simulation_expiree(tmp_path):
    saisie = _make_saisie(tmp_path / "SAISIE_test.xlsx")
    ref = _make_ref(tmp_path / "REF_Setup_test.xlsm")
    dryruns = tmp_path / "dryruns"
    _write_manifest(dryruns, saisie, ref, created_at=datetime.now(timezone.utc) - timedelta(minutes=31))
    state = real_svc.evaluate_real_write_prerequisites(
        TOKEN, saisie_path=saisie, ref_setup_path=ref, dryruns_root=dryruns
    )
    assert _codes(state)["SIMULATION_RECENTE"] is False
    assert "Simulation expiree" in next(c["details"] for c in state["checks"] if c["code"] == "SIMULATION_RECENTE")


def test_app2d_refuse_hash_saisie_modifie_depuis_simulation(app2d_env):
    wb = openpyxl.load_workbook(app2d_env["saisie"])
    wb["SAISIE"].append(["RESHH-OTHER"])
    wb.save(app2d_env["saisie"])
    wb.close()
    state = _eval(app2d_env)
    assert _codes(state)["HASH_SAISIE_IDENTIQUE"] is False


def test_app2d_refuse_hash_ref_modifie_depuis_simulation(app2d_env):
    wb = openpyxl.load_workbook(app2d_env["ref"])
    wb["REF_Cloture_Mensuelle"].append(["2026-09", "OUVERT"])
    wb.save(app2d_env["ref"])
    wb.close()
    state = _eval(app2d_env)
    assert _codes(state)["HASH_REF_SETUP_IDENTIQUE"] is False


def test_app2d_refuse_pk_deja_presente(tmp_path):
    saisie = _make_saisie(tmp_path / "SAISIE_test.xlsx", existing_pk=PK)
    ref = _make_ref(tmp_path / "REF_Setup_test.xlsm")
    dryruns = tmp_path / "dryruns"
    _write_manifest(dryruns, saisie, ref)
    state = real_svc.evaluate_real_write_prerequisites(TOKEN, saisie_path=saisie, ref_setup_path=ref, dryruns_root=dryruns)
    assert _codes(state)["PK_ABSENTE"] is False


def test_app2d_refuse_mois_devenu_cloture(tmp_path):
    saisie = _make_saisie(tmp_path / "SAISIE_test.xlsx")
    ref = _make_ref(tmp_path / "REF_Setup_test.xlsm", statut_mois="CLOTURE")
    dryruns = tmp_path / "dryruns"
    _write_manifest(dryruns, saisie, ref)
    state = real_svc.evaluate_real_write_prerequisites(TOKEN, saisie_path=saisie, ref_setup_path=ref, dryruns_root=dryruns)
    assert _codes(state)["MOIS_OUVERT"] is False


def test_app2d_refuse_schema_reel_incomplet(tmp_path):
    saisie = _make_saisie(tmp_path / "SAISIE_test.xlsx", schema=False)
    ref = _make_ref(tmp_path / "REF_Setup_test.xlsm", schema=False)
    dryruns = tmp_path / "dryruns"
    _write_manifest(dryruns, saisie, ref)
    state = real_svc.evaluate_real_write_prerequisites(TOKEN, saisie_path=saisie, ref_setup_path=ref, dryruns_root=dryruns)
    schema = next(c for c in state["checks"] if c["code"] == "SCHEMA_REEL_NON_PREPARE")
    assert schema["ok"] is False
    assert "taux_commission_override" in schema["details"]
    assert "PAY_006" in schema["details"]


def test_app2d_confirmation_incorrecte_nappelle_pas_writer(app2d_env, tmp_db):
    with (
        patch.object(cfg, "HH_REAL_WRITE_ENABLED", True),
        patch.object(cfg, "HH_REAL_WRITE_CONFIRMATION_ENABLED", True),
        patch("app.services.saisie_hh_real_write_service.hh_orchestrator.confirm_write") as writer,
    ):
        result = real_svc.enregistrer_reservation_hh_reelle(
            TOKEN, "ENREGISTRER AUTRE", saisie_path=app2d_env["saisie"],
            ref_setup_path=app2d_env["ref"], dryruns_root=app2d_env["dryruns"], db_path=tmp_db,
        )
    assert result["statut"] == "REFUSE"
    assert "TEXTE_CONFIRMATION_INCORRECT" in result["details"]
    writer.assert_not_called()


def test_app2d_confirmation_exacte_declenche_snapshot_writer_et_lot4a(app2d_env, tmp_db):
    validation = {
        "ok": True,
        "pk": PK,
        "erreurs": [],
        "preview": {"reservation_hh_id": PK, "mois": "2026-08", "total_percu": 450},
    }

    def _writer_side_effect(**kwargs):
        rollback = app2d_env["saisie"].parent / "SAISIE_test.app2d.rollback.xlsx"
        assert rollback.exists()
        return {"statut": "OK", "pk": PK, "ligne_cible": 2}

    with (
        patch.object(cfg, "HH_REAL_WRITE_ENABLED", True),
        patch.object(cfg, "HH_REAL_WRITE_CONFIRMATION_ENABLED", True),
        patch("app.services.saisie_hh_real_write_service.saisie_svc.valider", return_value=validation),
        patch("app.services.saisie_hh_real_write_service.hh_orchestrator.confirm_write", side_effect=_writer_side_effect) as writer,
        patch("app.services.saisie_hh_real_write_service._post_write_lot4a_check", return_value={"post_row": {"reservation_hh_id": PK}}) as lot4a,
    ):
        result = real_svc.enregistrer_reservation_hh_reelle(
            TOKEN, f"ENREGISTRER {PK}", saisie_path=app2d_env["saisie"],
            ref_setup_path=app2d_env["ref"], dryruns_root=app2d_env["dryruns"], db_path=tmp_db,
        )
    assert result["statut"] == "OK"
    writer.assert_called_once()
    lot4a.assert_called_once()


def test_app2d_rollback_automatique_si_writer_modifie_puis_echoue(app2d_env, tmp_db):
    sha_avant = _sha(app2d_env["saisie"])
    validation = {
        "ok": True,
        "pk": PK,
        "erreurs": [],
        "preview": {"reservation_hh_id": PK, "mois": "2026-08", "total_percu": 450},
    }

    def _writer_side_effect(**kwargs):
        app2d_env["saisie"].write_bytes(b"corruption apres debut ecriture")
        return {"statut": "ERREUR", "details": "writer force"}

    with (
        patch.object(cfg, "HH_REAL_WRITE_ENABLED", True),
        patch.object(cfg, "HH_REAL_WRITE_CONFIRMATION_ENABLED", True),
        patch("app.services.saisie_hh_real_write_service.saisie_svc.valider", return_value=validation),
        patch("app.services.saisie_hh_real_write_service.hh_orchestrator.confirm_write", side_effect=_writer_side_effect),
    ):
        result = real_svc.enregistrer_reservation_hh_reelle(
            TOKEN, f"ENREGISTRER {PK}", saisie_path=app2d_env["saisie"],
            ref_setup_path=app2d_env["ref"], dryruns_root=app2d_env["dryruns"], db_path=tmp_db,
        )
    assert result["statut"] == real_svc.STATUS_CANCELLED_RESTORED
    assert result["rollback_status"] == "ROLLBACK_REUSSI"
    assert _sha(app2d_env["saisie"]) == sha_avant


def test_app2d_audit_sqlite_cree_sur_refus(app2d_env, tmp_db):
    with patch.object(cfg, "HH_REAL_WRITE_ENABLED", False):
        real_svc.enregistrer_reservation_hh_reelle(
            TOKEN, f"ENREGISTRER {PK}", saisie_path=app2d_env["saisie"],
            ref_setup_path=app2d_env["ref"], dryruns_root=app2d_env["dryruns"], db_path=tmp_db,
        )
    import sqlite3
    conn = sqlite3.connect(tmp_db)
    try:
        row = conn.execute("SELECT action FROM audit_events WHERE action='APP2D_ECRITURE_REELLE_REFUSEE'").fetchone()
    finally:
        conn.close()
    assert row is not None


def test_route_app2c_affiche_prerequis_ecriture_reelle(client):
    dryrun = {
        "token": TOKEN,
        "manifest": {
            "status": "OK",
            "errors": [],
            "target_row": 2,
            "lot4a_status": "ANALYSE_TERMINEE",
            "paths": {"run_dir": "tmp/dryruns/" + TOKEN},
            "payload_summary": {"reservation_hh_id": PK},
            "migration": {"saisie_fields_added": []},
        },
        "resultat_lot4a": {"reservation_simulee": {}},
        "simulated_master_row": {},
        "comparatif": [],
    }
    state = {
        "eligible": False,
        "checks": [
            {"code": "FLAG_GLOBAL_ACTIVE", "ok": False, "label": "flag global active", "details": "HH_REAL_WRITE_ENABLED"},
            {"code": "SCHEMA_REEL_NON_PREPARE", "ok": False, "label": "schema reel prepare", "details": "colonnes manquantes"},
        ],
        "pk": PK,
        "expected_confirmation": f"ENREGISTRER {PK}",
    }
    with (
        patch("app.routes.reservations.dryrun_svc.load_previsualisation", return_value=dryrun),
        patch("app.routes.reservations.real_write_svc.evaluate_real_write_prerequisites", return_value=state),
    ):
        resp = client.get(f"/reservations/nouvelle/previsualisation/{TOKEN}")
    assert resp.status_code == 200
    assert "Écriture réelle désactivée" in resp.text
    assert "FLAG_GLOBAL_ACTIVE" in resp.text
    assert "SCHEMA_REEL_NON_PREPARE" in resp.text
    assert 'id="real_write_submit"' not in resp.text
