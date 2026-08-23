"""BLOC E — le système démarre, s'alimente, se recalcule et exporte SANS AUCUN Excel interne.

C'EST LE TEST QUI DÉFINIT « ZERO EXCEL OPÉRATIONNEL »
Un environnement neuf ne contient AUCUN master calculé, AUCUN `REF_Setup.xlsm`, AUCUN fichier de
saisie Excel interne. Seulement : le code, une base SQLite migrée, et ce que l'utilisateur saisit.
Si le pipeline tourne dans ces conditions, Excel n'est plus un prérequis de fonctionnement.

La preuve est ACTIVE : `openpyxl.load_workbook` lève pour tout classeur interne. Un chemin de code
qui tenterait d'en ouvrir un pour s'amorcer ferait échouer le test immédiatement, plutôt que de
retomber silencieusement sur un fichier présent par hasard sur la machine de développement.

Les classeurs EXTERNES (relevé bancaire, facture PDF d'un prestataire) ne sont pas concernés : ils
entrent par un import explicite, ce qui est leur rôle.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from app.db.connection import apply_migrations, get_db
from app.services import orchestrateur_dag as dag
from app.services import orchestrateur_service as orch

# Tout classeur INTERNE au système : masters calculés, référentiel, saisies.
_INTERDITS = ("MASTER_", "REF_Setup", "SAISIE_", "M04_MENAGES", "BANQUE_LOT8")


@pytest.fixture(autouse=True)
def _aucun_classeur_interne(monkeypatch):
    import openpyxl

    original = openpyxl.load_workbook
    original_save = openpyxl.Workbook.save

    def garde(chemin, *a, **kw):
        nom = Path(str(chemin)).name
        if any(marqueur in nom for marqueur in _INTERDITS):
            raise AssertionError(
                f"Bootstrap sans Excel : ouverture d'un classeur interne ({nom}). "
                "Le système doit démarrer sans aucun fichier Excel préexistant.")
        return original(chemin, *a, **kw)

    def garde_save(self, chemin, *a, **kw):
        # Un classeur intermédiaire CRÉÉ pendant « Actualiser toute l'activité » est aussi
        # incompatible avec ZERO EXCEL OPÉRATIONNEL que son ouverture (mission nettoyage 2026-08-22).
        nom = Path(str(chemin)).name
        if any(marqueur in nom for marqueur in _INTERDITS):
            raise AssertionError(
                f"Bootstrap sans Excel : écriture d'un classeur interne ({nom}). "
                "L'orchestrateur ne doit créer aucun intermédiaire Excel pendant un run normal.")
        return original_save(self, chemin, *a, **kw)

    monkeypatch.setattr(openpyxl.Workbook, "save", garde_save)

    monkeypatch.setattr(openpyxl, "load_workbook", garde)


@pytest.fixture
def env_neuf(tmp_path, monkeypatch):
    """Une base migrée, vierge, dans un APP_DATA_DIR isolé — et rien d'autre."""
    import app.config as cfg

    base = tmp_path / "neuf" / "app.db"
    base.parent.mkdir(parents=True, exist_ok=True)
    apply_migrations(base)
    monkeypatch.setattr(cfg, "DB_PATH", base)
    monkeypatch.setattr(cfg, "BACKUPS_DIR", tmp_path / "backups")
    return base


# ── §26 : bootstrap ─────────────────────────────────────────────────────────

def test_migrations_creent_le_socle_complet(env_neuf):
    conn = get_db(env_neuf)
    try:
        tables = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    finally:
        conn.close()
    # Une table par maillon de la chaîne : entrées saisies, référentiel, calculs, orchestration.
    for attendue in ("charges", "reservations_hors_hostaway", "ref_logements", "ref_proprietaires",
                     "banque_mouvements", "hostaway_extractions", "menages_taches_enrichies",
                     "flux_unifies", "lot10_runs", "controles_lot11_constats",
                     "lot12_prefactures_entete", "orchestrateur_datasets"):
        assert attendue in tables, attendue


def test_saisies_passent_par_les_services_pas_par_excel(env_neuf):
    """§7/§9 — une saisie quotidienne se fait UI → service → SQLite."""
    from app.readers import charges_reader
    from app.readers import reservations_hh_reader
    from app.services import charges_saisie_service
    from app.services import reservations_hh_saisie_service

    charge = charges_saisie_service.creer(
        {"date_charge": "2026-06-12", "montant": 149.90, "categorie_charge_id": "CAT_ENTRETIEN",
         "logement_id": "LOG_0001", "type_flux_id": "TYPE_FLUX_020", "code_impact": "IC",
         "statut_controle": "VALIDE"},
        acteur="recette", db_path=env_neuf)
    assert charge["ok"], charge

    reservation = reservations_hh_saisie_service.creer(
        {"mois": "2026-06", "logement_id": "LOG_0001", "date_arrivee": "2026-06-02",
         "date_depart": "2026-06-06", "montant_retenu": 420.0, "statut_controle": "VALIDE"},
        acteur="recette", db_path=env_neuf)
    assert reservation["ok"], reservation

    assert len(charges_reader.read_charges(db_path=env_neuf)) == 1
    assert len(reservations_hh_reader.read_reservations(db_path=env_neuf)) == 1


def test_etats_vides_propres_sans_classeur(env_neuf):
    """§11 — sans aucune saisie, chaque source rend un état vide, jamais une erreur."""
    from app.readers import charges_reader
    from app.readers import menages_reader
    from app.readers import reservations_hh_reader

    assert charges_reader.read_charges(db_path=env_neuf) == []
    assert reservations_hh_reader.read_reservations(db_path=env_neuf) == []
    # Les sources ménages rendent un ÉTAT, pas une exception.
    assert menages_reader.rapprochement().etat.etat in ("OK", "VIDE", "FICHIER_ABSENT")


# ── §27 : actualiser toute l'activité ───────────────────────────────────────

def test_actualiser_toute_activite_sur_environnement_neuf(env_neuf, monkeypatch):
    """Le pipeline complet tourne sur une base neuve, sans aucun classeur interne.

    Lot10 (moteur pandas en sous-processus) est remplacé : ce test prouve l'AMORÇAGE, pas
    l'exécution d'un sous-processus — celle-ci est couverte par la parité réelle de Lot10.
    """
    from app.services import charges_saisie_service

    charges_saisie_service.creer(
        {"date_charge": "2026-06-12", "montant": 80.0, "categorie_charge_id": "CAT_ENTRETIEN",
         "logement_id": "LOG_0001", "statut_controle": "VALIDE"},
        acteur="recette", db_path=env_neuf)

    vrai_appel = orch._appeler_service

    def appel(chemin, db_path):
        if "orchestrateur_moteur" in chemin:
            return {"ok": True, "simule": True}
        return vrai_appel(chemin, db_path)

    monkeypatch.setattr(orch, "_appeler_service", appel)
    resultat = orch.actualiser(db_path=env_neuf)

    assert resultat["statut"] in (orch.RUN_SUCCES, orch.RUN_PARTIEL), resultat
    etats = {d["dataset"]: d["statut"] for d in resultat["datasets"]}
    for calcule in (dag.FLUX_LOT9, dag.LOT10, dag.LOT11, dag.LOT12):
        assert etats[calcule] == orch.ST_A_JOUR, (calcule, etats[calcule])


def test_export_lot13_regenere_depuis_une_base_neuve(env_neuf, tmp_path):
    """§26 — l'export terminal se produit lui aussi sans classeur d'entrée."""
    from app.services import lot13_export_service

    destination = tmp_path / "exports"
    resultat = lot13_export_service.exporter(db_path=env_neuf, destination=destination)
    assert resultat.get("ok"), resultat
    assert list(destination.glob("*.csv")), "aucun export produit"


# ── §28 : ce qui reste autorisé ─────────────────────────────────────────────

def test_aucun_classeur_interne_requis_pour_demarrer(env_neuf):
    """Reformulation directe de l'objectif : démarrer, alimenter, recalculer, exporter.

    Garanti par le fixture `_aucun_classeur_interne` : si l'un des chemins ci-dessus ouvrait un
    master, un référentiel ou une saisie Excel, le test aurait déjà échoué.
    """
    from app.services import controles_lot11_service as lot11
    from app.services import flux_unifie_service as lot9
    from app.services import lot12_prefactures_service as lot12

    assert lot9.construire(db_path=env_neuf)["ok"]
    assert lot11.construire(db_path=env_neuf)["ok"]
    assert lot12.construire(db_path=env_neuf)["ok"]
