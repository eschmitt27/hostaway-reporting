"""Cross-module : un logement archivé via `logements_gestion_service` ne doit plus être
sélectionnable dans les NOUVEAUX traitements (réservations, charges) mais doit rester visible dans
l'historique.

Ces modules lisent tous `REF_Logements`/`REF_Gestion_Logements_Hist` directement dans REF_Setup —
la même feuille que `logements_gestion_service` écrit. Pas de cache intermédiaire à invalider :
l'exclusion est immédiate, sans attendre un cycle de pipeline (Lot13/PBI), contrairement à la vue
fiche/liste (`logements_service.py`, basée sur le CSV PBI — voir `28_CYCLE_DE_VIE_LOGEMENT.md`).
"""
import openpyxl
import pytest

import app.config as cfg
from app.readers.saisie_charges_reader import read_ref_logements
from app.readers.ref_setup_hh_reader import get_all_logements_hh, get_gestion_hist
from app.services import logements_gestion_service as svc
from app.services.saisie_hh_service import _is_managed_active_logement


def _ref(tmp_path):
    p = tmp_path / "REF_Setup.xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("REF_Logements")
    ws.append(["logement_id", "hostaway_listing_id", "nom_logement_officiel", "nom_court", "adresse",
               "ville", "type_logement_id", "sur_hostaway", "actif", "statut_parc", "commentaire",
               "forfait_logiciel_consommables_mensuel"])
    ws.append(["LOG_A1", 900001, "Fictif A1", "A1", "1 rue", "RECETTE", "TYPE_001", "OUI", "OUI",
               "GERE", None, 0])
    ws = wb.create_sheet("REF_Gestion_Logements_Hist")
    ws.append(["gestion_id", "logement_id", "proprietaire_id", "date_debut", "date_fin",
               "statut_gestion", "source", "commentaire"])
    ws.append(["GST_1", "LOG_A1", "PROP_A", "2026-01-01", None, "ACTIF", "FICTIF", None])
    ws = wb.create_sheet("REF_Proprietaires")
    ws.append(["proprietaire_id", "nom_proprietaire", "actif"])
    ws.append(["PROP_A", "Nom PROP_A", "OUI"])
    ws = wb.create_sheet("REF_Types_Logements")
    ws.append(["type_logement_id", "libelle"]); ws.append(["TYPE_001", "STUDIO"])
    ws = wb.create_sheet("REF_Taux_Commission")
    ws.append(["taux_commission_id", "proprietaire_id", "logement_id", "taux_commission",
               "date_debut", "date_fin", "actif", "justification", "commentaire"])
    wb.save(p); wb.close()
    return p


@pytest.fixture
def ref(tmp_path, monkeypatch):
    p = _ref(tmp_path)
    monkeypatch.setattr(cfg, "REF_SETUP", p)
    monkeypatch.setattr(cfg, "RECETTE_MODE", True)
    monkeypatch.setattr(cfg, "RECETTE_ROOT", tmp_path.resolve())
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)
    return p


def _eligible_charges(p) -> set[str]:
    """Reproduit le filtre réel de `charges_preview_service.load_form_refs` (ligne 432) :
    `[l for l in logements if is_active(l)]` avec `is_active = actif == 'OUI'`."""
    logs = read_ref_logements(p)
    return {str(l.get("logement_id", "")).strip() for l in logs
            if str(l.get("actif", "")).upper() == "OUI"}


def _eligible_reservations(p) -> set[str]:
    """Reproduit le filtre réel de `saisie_hh_service.load_form_refs` (ligne 493-497) :
    `_is_managed_active_logement` (actif=OUI ET statut_parc=GERE)."""
    logs = get_all_logements_hh(ref_setup_path=p)
    return {str(l.get("logement_id", "")).strip() for l in logs
            if str(l.get("logement_id", "")).strip() and _is_managed_active_logement(l)}


def test_logement_actif_est_selectionnable_partout(ref):
    assert "LOG_A1" in _eligible_charges(ref)
    assert "LOG_A1" in _eligible_reservations(ref)


def test_archiver_exclut_immediatement_des_nouveaux_traitements(ref):
    res = svc.archiver("LOG_A1", "2026-06-30", ref_path=ref)
    assert res["ok"], res

    assert "LOG_A1" not in _eligible_charges(ref)
    assert "LOG_A1" not in _eligible_reservations(ref)


def test_archiver_reste_visible_dans_lhistorique(ref):
    svc.archiver("LOG_A1", "2026-06-30", ref_path=ref)

    # Le logement reste lisible : la ligne REF_Logements existe toujours, jamais supprimée.
    logs = read_ref_logements(ref)
    assert any(str(l.get("logement_id")) == "LOG_A1" for l in logs)

    # Le rattachement de gestion clôturé reste dans l'historique (jamais supprimé).
    gest = get_gestion_hist(ref_setup_path=ref)
    ligne = next(g for g in gest if str(g.get("logement_id")) == "LOG_A1")
    assert ligne["date_fin"] == "2026-06-30"
    assert ligne["statut_gestion"] == "RETIRE"

    # Historique applicatif (gestion_svc.historique) le retrouve aussi.
    histo = svc.historique("LOG_A1", ref_path=ref)
    assert any(g["logement_id"] == "LOG_A1" for g in histo["gestion"])


def test_reactiver_rend_de_nouveau_selectionnable(ref):
    svc.archiver("LOG_A1", "2026-06-30", ref_path=ref)
    assert "LOG_A1" not in _eligible_charges(ref)

    res = svc.reactiver("LOG_A1", "2026-07-01", "PROP_A", ref_path=ref)
    assert res["ok"], res

    assert "LOG_A1" in _eligible_charges(ref)
    assert "LOG_A1" in _eligible_reservations(ref)


def test_changer_taux_resolu_correctement_a_toutes_les_dates(ref):
    """Le taux de commission historisé doit se résoudre correctement à n'importe quelle date,
    y compris avant/après un changement — cohérence avec `resolve_commission_rate` (Lot10)."""
    svc.changer_taux_commission("LOG_A1", 0.19, "2026-01-01", "PROP_A", ref_path=ref)
    svc.changer_taux_commission("LOG_A1", 0.15, "2026-03-01", "PROP_A", ref_path=ref)

    histo = svc.historique("LOG_A1", ref_path=ref)["taux"]
    par_debut = sorted(histo, key=lambda t: t["date_debut"])
    assert par_debut[0]["date_debut"] == "2026-01-01" and par_debut[0]["date_fin"] == "2026-02-28"
    assert par_debut[0]["taux_commission"] == 0.19
    assert par_debut[1]["date_debut"] == "2026-03-01" and par_debut[1]["date_fin"] is None
    assert par_debut[1]["taux_commission"] == 0.15

    # Résolution à une date antérieure au changement : c'est bien l'ancien taux (ligne close)
    # qui couvre cette période — jamais le nouveau.
    ancien = next(t for t in histo if t["date_debut"] == "2026-01-01")
    assert ancien["date_fin"] == "2026-02-28"        # couvre janvier-février, pas mars
    nouveau = next(t for t in histo if t["date_debut"] == "2026-03-01")
    assert nouveau["date_fin"] is None                # couvre mars et au-delà, jusqu'à révision
