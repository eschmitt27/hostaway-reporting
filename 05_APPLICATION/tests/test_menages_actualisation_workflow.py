"""« Actualiser le rapprochement des ménages » — bouton unique, workflow synchrone.

Ce que ces tests protègent :
  1. les TROIS sources sont réellement sollicitées (PDF, Google Sheet via lot6b, Hostaway) ;
  2. l'ordre est celui qui rend la chaîne cohérente ;
  3. seuls les mois OUVERTS réellement impactés sont recalculés — un mois CLÔTURÉ est TRACÉ, jamais
     recalculé, et ses chiffres économiques restent prouvablement identiques ;
  4. sans changement de source, aucun recalcul superflu (idempotence) ;
  5. la route est SYNCHRONE : l'état est complet au retour de la requête, sans second appel ;
  6. zéro Excel, à aucune étape.

Les moteurs (lot6b/6d/6e/6f) et l'API Hostaway sont remplacés par des espions : ce qui est vérifié
ici est l'ORCHESTRATION, pas le calcul des lots — déjà couvert par leurs propres tests.
"""
import pytest

import app.config as cfg
from app.db.connection import apply_migrations, get_db
from app.services import menages_actualisation_service as workflow


@pytest.fixture()
def db(tmp_path, monkeypatch):
    chemin = tmp_path / "app.db"
    monkeypatch.setattr(cfg, "DB_PATH", chemin, raising=False)
    apply_migrations(chemin)
    return chemin


@pytest.fixture()
def espions(monkeypatch, db):
    """Remplace les frontières externes et enregistre l'ordre exact des appels."""
    from app.services import hostaway_cleaning_tasks_actualisation_service as hostaway_ct
    from app.services import menages_actualisation_service as workflow_svc
    from app.services import menages_declarations_service as decl
    from app.services import menages_pdf_import_service as pdf_import
    from app.services import menages_service as svc
    from app.services import orchestrateur_moteur as moteur
    from app.services import orchestrateur_service as orch

    journal: list[str] = []
    etat = {"pdf": {"ok": True, "nb_detectes": 0, "nb_importees": 0, "nb_remplacees": 0,
                    "nb_deja_importees": 0, "mois_impactes": [], "details": []},
            "clotures": set(), "menages_ok": True}

    # Préflights (mission « spinner infini ») : réputés OK par défaut dans ce fixture, testés
    # explicitement KO dans les tests dédiés A/B ci-dessous (override du monkeypatch au cas par cas).
    monkeypatch.setattr(hostaway_ct, "credentials_disponibles", lambda: True)
    monkeypatch.setattr(workflow_svc, "_google_sheet_config_ok", lambda **k: True)

    def _pdf(**kw):
        journal.append("PDF")
        return etat["pdf"]

    def _lot6b(**kw):
        journal.append("GOOGLE_SHEET")
        return {"ok": True}

    def _hostaway(dataset, **kw):
        journal.append(f"HOSTAWAY:{dataset}")
        return {"ok": True}

    def _menages(*, mois, **kw):
        journal.append(f"MENAGES:{mois}")
        return {"ok": etat["menages_ok"], "mois_traite": mois,
                "code": "" if etat["menages_ok"] else "ECHEC"}

    monkeypatch.setattr(pdf_import, "importer_nouveaux", _pdf)
    monkeypatch.setattr(moteur, "executer_declarations_internes", _lot6b)
    monkeypatch.setattr(orch, "recalculer_dataset", _hostaway)
    monkeypatch.setattr(moteur, "executer_menages_cible", _menages)
    monkeypatch.setattr(orch, "marquer_dataset", lambda *a, **k: None)
    monkeypatch.setattr(orch, "actualiser",
                        lambda **k: journal.append("CASCADE") or {"ok": True})
    monkeypatch.setattr(decl, "mois_cloture", lambda mois, db_path=None: mois in etat["clotures"])
    monkeypatch.setattr(svc, "invalidate_menages_cache", lambda: None)
    monkeypatch.setattr(svc, "periode_par_defaut", lambda: "2026-06")
    return journal, etat


# ── Les trois sources, dans le bon ordre ────────────────────────────────────────────────────────

def test_les_trois_sources_sont_sollicitees_dans_l_ordre(db, espions):
    journal, _ = espions
    workflow.actualiser(mois_affiche="2026-06", db_path=db)
    assert journal[0] == "PDF"
    assert journal[1] == "GOOGLE_SHEET"
    assert journal[2] == "HOSTAWAY:HOSTAWAY_CLEANING_TASKS"
    assert "MENAGES:2026-06" in journal
    # Le rapprochement aval vient APRÈS le recalcul, jamais avant.
    assert journal.index("CASCADE") > journal.index("MENAGES:2026-06")


def test_google_sheet_reellement_branche(db, espions):
    """Le maillon lot6b n'était câblé dans AUCUN parcours applicatif avant cette mission."""
    journal, _ = espions
    workflow.actualiser(mois_affiche="2026-06", db_path=db)
    assert "GOOGLE_SHEET" in journal


# ── Ciblage des mois ────────────────────────────────────────────────────────────────────────────

def test_seuls_les_mois_impactes_sont_recalcules(db, espions):
    journal, etat = espions
    etat["pdf"] = {**etat["pdf"], "nb_importees": 1, "mois_impactes": ["2026-05"]}
    resultat = workflow.actualiser(mois_affiche="2026-06", db_path=db)
    assert sorted(resultat["mois_recalcules"]) == ["2026-05", "2026-06"]
    assert "MENAGES:2026-04" not in journal


def test_sans_changement_aucun_recalcul_superflu(db, espions):
    """Idempotence : deux clics d'affilée sur des sources inchangées font le même travail minimal."""
    journal, _ = espions
    premier = workflow.actualiser(mois_affiche="2026-06", db_path=db)
    journal.clear()
    second = workflow.actualiser(mois_affiche="2026-06", db_path=db)
    assert premier["mois_recalcules"] == second["mois_recalcules"] == ["2026-06"]
    # Aucun mois supplémentaire n'apparaît au second passage.
    assert [e for e in journal if e.startswith("MENAGES:")] == ["MENAGES:2026-06"]


# ── Mois clôturés ───────────────────────────────────────────────────────────────────────────────

def _photo_economique(db):
    conn = get_db(db)
    try:
        return {t: [tuple(r) for r in conn.execute(f"SELECT * FROM {t} ORDER BY id")]
                for t in ("lot10_resultats", "lot10_commissions", "lot12_prefactures_entete")}
    finally:
        conn.close()


def test_mois_cloture_trace_mais_jamais_recalcule(db, espions):
    journal, etat = espions
    etat["clotures"] = {"2026-05"}
    etat["pdf"] = {**etat["pdf"], "mois_impactes": ["2026-05"]}
    photo = _photo_economique(db)

    resultat = workflow.actualiser(mois_affiche="2026-06", db_path=db)

    assert resultat["mois_clotures_signales"] == ["2026-05"]
    assert "MENAGES:2026-05" not in journal, "un mois clôturé n'est jamais recalculé"
    assert _photo_economique(db) == photo, "l'économie d'un mois clôturé reste intacte"

    signalements = workflow.changements_mois_clotures(db_path=db)
    assert [s["mois"] for s in signalements] == ["2026-05"]
    assert signalements[0]["origine"] == workflow.ORIGINE_PDF
    assert signalements[0]["statut"] == "SIGNALE"


def test_statistiques_reelles_jamais_inventees(db, espions):
    _, etat = espions
    etat["pdf"] = {**etat["pdf"], "nb_importees": 1, "mois_impactes": ["2026-06"]}
    resultat = workflow.actualiser(mois_affiche="2026-06", db_path=db)
    assert "1 nouvelle facture" in resultat["statistiques"]
    assert any("1 mois mis à jour" in s for s in resultat["statistiques"])

    etat["pdf"] = {**etat["pdf"], "nb_importees": 0, "mois_impactes": []}
    suivant = workflow.actualiser(mois_affiche="2026-06", db_path=db)
    assert not any("nouvelle facture" in s for s in suivant["statistiques"])


# ── Zéro Excel ──────────────────────────────────────────────────────────────────────────────────

def test_zero_excel(db, espions, monkeypatch):
    """Aucun classeur n'est ouvert : le piège échoue le test au premier `load_workbook`."""
    import openpyxl

    def _piege(*a, **k):
        raise AssertionError("openpyxl.load_workbook appelé : le workflow doit rester SQLite pur")

    monkeypatch.setattr(openpyxl, "load_workbook", _piege)
    workflow.actualiser(mois_affiche="2026-06", db_path=db)


# ── Route : synchrone, bouton unique, retour au contexte ────────────────────────────────────────

@pytest.fixture()
def client(db, espions, monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from app.services import orchestrateur_service as orch
    monkeypatch.setattr(cfg, "DATA_DIR", tmp_path, raising=False)
    monkeypatch.setattr(orch, "marquer_runs_interrompus", lambda **k: [])
    from app.main import app as application
    with TestClient(application) as c:
        yield c


def test_route_est_synchrone(client, espions):
    """Le travail est FAIT au retour de la requête — aucun second appel n'est nécessaire."""
    journal, _ = espions
    reponse = client.post("/menages/actualiser", data={"mois": "2026-06"},
                          follow_redirects=False)
    assert reponse.status_code == 303
    assert "MENAGES:2026-06" in journal, "le recalcul a eu lieu PENDANT la requête"
    assert "actualisation=terminee" in reponse.headers["location"]


def test_route_actualiser_affichage_supprimee(client):
    """Le bouton de repli d'un traitement asynchrone n'a plus de raison d'être."""
    assert client.post("/menages/actualiser-affichage", data={}).status_code == 404


def test_ecran_ne_propose_qu_une_action_d_actualisation(client):
    reponse = client.get("/menages")
    assert reponse.status_code == 200
    assert reponse.text.count('action="/menages/actualiser"') == 1
    # Les parcours de saisie manuelle de déclaration ont été retirés : ils ne doivent pas revenir.
    assert "/menages/declarations/nouvelle" not in reponse.text
    assert "Actualisation en cours…" in reponse.text, "le bouton se désactive côté client"


def test_action_de_ligne_revient_au_contexte_filtre(client, db, monkeypatch):
    """Après une correction, l'utilisateur retrouve son mois ET ses filtres."""
    from app.services import menages_declarations_service as declarations
    monkeypatch.setattr(declarations, "modifier", lambda **k: {"ok": True})

    retour = "/menages?mois=2026-06&statut=A_CONTROLER&tri=anomalie"
    reponse = client.post(
        "/menages/2026-06/LOG_1/INT_1/modifier-declaration",
        data={"nb_menages": "2", "retour": retour}, follow_redirects=False)
    assert reponse.status_code == 303
    assert reponse.headers["location"].startswith(retour)
    assert "statut=A_CONTROLER" in reponse.headers["location"]
