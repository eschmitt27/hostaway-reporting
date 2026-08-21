"""Fabrique de datasets Lot12 en SQLite, pour les tests.

Remplace le classeur `MASTER_FACT_Proprietaires.xlsx` que les tests devaient fabriquer : depuis la
migration 0047, les lecteurs applicatifs lisent les tables `lot12_*` du run ACTIF, sans repli Excel.
Même principe que `fixtures_lot10.py`.

Insertion en SQL direct (pas via `lot12_prefactures_service`) : ces tests vérifient les LECTEURS et
les services applicatifs, pas le calcul Lot12 lui-même — celui-ci est couvert par la parité réelle
OLD/NEW (`02_TRAVAIL/lot12_generer_factures.py` vs `lot12_prefactures_service.py`).
"""
from __future__ import annotations

from typing import Any, Iterable

from app.db.connection import get_db

RUN_TEST = "L12-TEST"


def _inserer(conn, table: str, run_id: str, lignes: Iterable[dict[str, Any]]) -> None:
    for l in lignes:
        cols = ["run_id"] + list(l.keys())
        valeurs = [run_id] + [l[c] for c in l]
        conn.execute(
            f"INSERT INTO {table} ({', '.join(cols)}) VALUES ({', '.join(['?'] * len(cols))})",
            valeurs)


def seeder(db_path, *, entetes=(), lignes=(), controle=(), dashboard=(), a_controler=(),
          run_id: str = RUN_TEST, statut: str = "SUCCES", actif: bool = True) -> str:
    """Crée un run Lot12 et son dataset."""
    conn = get_db(db_path)
    try:
        conn.execute("UPDATE lot12_runs SET actif = 0 WHERE actif = 1")
        conn.execute(
            "INSERT INTO lot12_runs (run_id, statut, actif, nb_entetes, nb_lignes, nb_controle, "
            "nb_dashboard, nb_a_controler) VALUES (?,?,?,?,?,?,?,?)",
            (run_id, statut, 1 if actif else 0, len(list(entetes)), len(list(lignes)),
             len(list(controle)), len(list(dashboard)), len(list(a_controler))))
        _inserer(conn, "lot12_prefactures_entete", run_id, entetes)
        _inserer(conn, "lot12_prefactures_lignes", run_id, lignes)
        _inserer(conn, "lot12_controle_mensuel", run_id, controle)
        _inserer(conn, "lot12_dashboard_facturation", run_id, dashboard)
        _inserer(conn, "lot12_a_controler", run_id, a_controler)
        conn.commit()
    finally:
        conn.close()
    return run_id


def reprendre_master_reel(db_path, chemin_master, *, run_id: str = RUN_TEST) -> str | None:
    """Charge le VRAI `MASTER_FACT_Proprietaires.xlsx` (lecture seule) dans les tables `lot12_*`.

    Même statut que `fixtures_lot10.reprendre_master_reel` : réservé aux tests qui portent sur des
    DONNÉES RÉELLES (PROP_0001, mois 2025-07, préfactures multi-logements…) — leur valeur vient
    précisément de ce qu'ils vérifient des chiffres réels. Outil de TEST, jamais un chemin runtime :
    l'application ne lit plus ce classeur.

    Rend None si le classeur est absent (environnement sans données réelles) — l'appelant skippe.
    """
    from pathlib import Path

    import openpyxl

    chemin = Path(chemin_master)
    if not chemin.exists():
        return None

    wb = openpyxl.load_workbook(str(chemin), read_only=True, data_only=True)
    try:
        def _onglet(nom: str) -> list[dict[str, Any]]:
            if nom not in wb.sheetnames:
                return []
            lignes = list(wb[nom].iter_rows(values_only=True))
            if len(lignes) <= 1:
                return []
            entetes = [str(c) if c is not None else "" for c in lignes[0]]
            return [dict(zip(entetes, r)) for r in lignes[1:] if any(v is not None for v in r)]

        entetes = _onglet("FACT_FACTURE_ENTETE")
        lignes = _onglet("FACT_FACTURE_LIGNES")
        dashboard = _onglet("DASHBOARD_FACTURATION")
        a_controler = _onglet("A_CONTROLER")
    finally:
        wb.close()

    # Seules les colonnes réellement portées par les tables 0047 sont reprises : une colonne du
    # classeur qui n'existe pas en base signalerait une divergence de schéma, pas une donnée à
    # forcer.
    def _filtrer(rows, colonnes):
        return [{k: v for k, v in r.items() if k in colonnes} for r in rows]

    COLS_ENTETE = {"facture_id", "mois", "proprietaire_id", "nom_proprietaire",
                   "adresse_proprietaire", "logement_id", "nom_logement", "periode_debut",
                   "periode_fin", "nb_reservations", "total_exploitation_net",
                   "total_reglement_du", "reste_a_payer", "credit_a_traiter", "mode_facturation",
                   "statut_facture", "statut_generation", "balises", "date_generation"}
    COLS_LIGNES = {"facture_id", "ligne_num", "type_ligne", "libelle", "montant", "bloc",
                   "commentaire"}
    COLS_DASH = {"mois", "proprietaire_id", "nb_logements", "nb_bloquants_mois",
                 "nb_a_controler_mois", "facturation_lot12_ok", "mode_facturation",
                 "statut_facture", "balises_non_resolues"}
    COLS_AC = {"mois", "proprietaire_id", "logement_id", "reservation", "code_anomalie",
               "severite", "impact_facturation", "message"}

    seeder(db_path,
           entetes=_filtrer(entetes, COLS_ENTETE),
           lignes=_filtrer(lignes, COLS_LIGNES),
           dashboard=_filtrer(dashboard, COLS_DASH),
           a_controler=_filtrer(a_controler, COLS_AC),
           run_id=run_id)
    return run_id
