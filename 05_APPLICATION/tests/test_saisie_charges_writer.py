"""APP-3b / commit 2 — writer BAS NIVEAU des charges : préparation de temporaires validés.

Aucun fichier métier réel n'est touché : chaque test construit ses fixtures en tmp (nommées comme
les vrais fichiers, car file_registry n'autorise que les `SAISIE_*`). Le writer ne fait aucun
`os.replace` — il produit un temporaire, et les tests vérifient que la SOURCE reste intacte (SHA256).
"""
from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import openpyxl
import pytest

import app.config as cfg
from app.readers.saisie_charges_reader import MANUAL_COL_MAP, _col_index
from app.writers import saisie_charges_writer as w

MONTANT = 100.0
CHARGE_ID = "CHG-2026-06-IC-BANQUE-001"

# Formules réelles de SAISIE_Charges_Flux (colonnes C, I, J, AD).
FORMULES = {
    "C": '=IF(B{r}="","",TEXT(B{r},"YYYY-MM"))',
    "I": '=IF(H{r}="IC","OUI",IF(H{r}="HC","OUI",IF(H{r}="HR","NON","A_CONTROLER")))',
    "J": '=IF(H{r}="IC","OUI",IF(H{r}="HC","NON",IF(H{r}="HR","NON","A_CONTROLER")))',
    "AD": '=IF(A{r}="","",TEXT(B{r},"YYYYMMDD")&"|"&TEXT(D{r},"0.00")&"|"&F{r}&"|"&M{r})',
}

IMPACTS_HEADERS = {
    "AFFECTATIONS": ["affectation_id", "charge_id", "mois", "logement_id", "proprietaire_id",
                     "quote_part", "statut", "origine", "commentaire", "ROW_HASH"],
    "MENAGE": ["menage_impact_id", "charge_id", "mois", "mode", "intervenant_id", "logement_id",
               "statut", "origine", "commentaire", "ROW_HASH"],
    "RESERVE_REFACTURATION": ["reserve_id", "charge_id", "mois", "logement_id", "proprietaire_id",
                              "montant_refacturable", "libelle", "justificatif",
                              "statut_traitement", "trace_decision", "origine", "commentaire",
                              "ROW_HASH"],
}


def _sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ── Fixtures ─────────────────────────────────────────────────────────────────

@pytest.fixture
def saisie(tmp_path: Path) -> Path:
    """SAISIE_Charges_Flux : en-têtes + 5 lignes modèle portant les formules vivantes."""
    p = tmp_path / "SAISIE_Charges_Flux.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    entetes = {**MANUAL_COL_MAP, "C": "mois", "I": "impact_resultat_reel",
               "J": "impact_resultat_comptable", "AD": "ROW_HASH"}
    for col, champ in entetes.items():
        ws.cell(row=1, column=_col_index(col), value=champ)
    for r in range(2, 7):                       # lignes modèle : formules vivantes, charge_id vide
        for col, f in FORMULES.items():
            ws.cell(row=r, column=_col_index(col), value=f.format(r=r))
    wb.create_sheet("REF_LOCALE")
    wb.save(p)
    wb.close()
    return p


@pytest.fixture
def impacts(tmp_path: Path) -> Path:
    """SAISIE_Charges_Impacts : 3 onglets + 1 ligne d'une AUTRE charge (témoin de non-régression)."""
    p = tmp_path / "SAISIE_Charges_Impacts.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, headers in IMPACTS_HEADERS.items():
        ws = wb.create_sheet(sheet)
        ws.append(headers)
    wb["AFFECTATIONS"].append(["AUTRE-AFF-001", "CHG-AUTRE", "2026-05", "LOG_0009", "PROP_01",
                              42.0, "A_CONTROLER", "NOUVELLE_CHARGE_GUIDEE", None, "hash"])
    wb.create_sheet("README")
    wb.save(p)
    wb.close()
    return p


def _row_data(**kw) -> dict:
    base = {
        "charge_id": CHARGE_ID, "date_charge": "2026-06-15", "montant": MONTANT,
        "sens_flux": "DEPENSE", "categorie_charge_id": "CHG_025", "type_flux_id": "TYPE_FLUX_020",
        "code_impact": "IC", "prise_en_compta": "OUI", "mode_paiement_id": "PAY_001",
        "affectation_type": "GLOBAL", "refacturable": "NON", "source_flux": "SAISIE_MANUELLE",
        "statut_controle": "A_CONTROLER", "niveau_anomalie": "INFO",
        "statut_rapprochement": "NON_RAPPROCHE", "date_saisie": "2026-07-13",
        "affectable_menage": "NON", "avantage_associe_id": "PERS_EWAN",
    }
    base.update(kw)
    return base


def _affectation(i: int, logement: str, quote: float) -> dict:
    return {"affectation_id": f"{CHARGE_ID}-AFF-{i:03d}", "charge_id": CHARGE_ID, "mois": "2026-06",
            "logement_id": logement, "proprietaire_id": "PROP_01", "quote_part": quote,
            "statut": "A_CONTROLER", "origine": "NOUVELLE_CHARGE_GUIDEE", "commentaire": None,
            "ROW_HASH": f"h{i}"}


def _menage(i: int, intervenant=None, logement=None) -> dict:
    return {"menage_impact_id": f"{CHARGE_ID}-MEN-{i:03d}", "charge_id": CHARGE_ID, "mois": "2026-06",
            "mode": "INTERVENANT" if intervenant else "LOGEMENT", "intervenant_id": intervenant,
            "logement_id": logement, "statut": "A_CONTROLER",
            "origine": "NOUVELLE_CHARGE_GUIDEE", "commentaire": None, "ROW_HASH": f"m{i}"}


def _reserve(i: int, montant: float, statut: str = "EN_ATTENTE") -> dict:
    return {"reserve_id": f"{CHARGE_ID}-RES-{i:03d}", "charge_id": CHARGE_ID, "mois": "2026-06",
            "logement_id": "LOG_0001", "proprietaire_id": "PROP_01",
            "montant_refacturable": montant, "libelle": "Maintenance", "justificatif": None,
            "statut_traitement": statut, "trace_decision": None,
            "origine": "NOUVELLE_CHARGE_GUIDEE", "commentaire": None, "ROW_HASH": f"r{i}"}


def _persistable(affectations=(), menage=(), reserve=()) -> dict:
    return {"charge_id": CHARGE_ID, "affectations": list(affectations),
            "menage": list(menage), "reserve": list(reserve)}


def _lire(path: Path, sheet: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(c is not None for c in r)]


# ── 1-6 : ligne de charge (SAISIE_Charges_Flux) ──────────────────────────────

def test_1_ecrit_ligne_sur_temporaire_sans_toucher_la_source(saisie, tmp_path):
    avant = _sha(saisie)
    res = w.prepare_charge_flux_write(saisie, _row_data(), 2, temp_dir=tmp_path)

    assert res["statut"] == "OK", res
    assert _sha(saisie) == avant                      # la SOURCE n'a pas bougé
    temp = Path(res["temp_path"])
    assert temp.exists() and temp != saisie
    assert res["remplacement"] is False

    lignes = _lire(temp, "SAISIE")
    mine = [r for r in lignes if r.get("charge_id") == CHARGE_ID]
    assert len(mine) == 1
    assert mine[0]["montant"] == MONTANT
    assert mine[0]["categorie_charge_id"] == "CHG_025"
    assert mine[0]["avantage_associe_id"] == "PERS_EWAN"


def test_2_preserve_les_formules_c_i_j_ad(saisie, tmp_path):
    res = w.prepare_charge_flux_write(saisie, _row_data(), 2, temp_dir=tmp_path)
    assert res["statut"] == "OK", res

    wb = openpyxl.load_workbook(res["temp_path"], data_only=False)
    try:
        ws = wb["SAISIE"]
        for col in ("C", "I", "J", "AD"):
            v = ws.cell(row=2, column=_col_index(col)).value
            assert isinstance(v, str) and v.startswith("="), f"{col} n'est plus une formule : {v!r}"
        assert wb.calculation.fullCalcOnLoad is True
    finally:
        wb.close()


@pytest.mark.parametrize("champ", ["mois", "impact_resultat_reel",
                                   "impact_resultat_comptable", "ROW_HASH"])
def test_3_refuse_un_payload_visant_une_colonne_formule(saisie, tmp_path, champ):
    res = w.prepare_charge_flux_write(saisie, _row_data(**{champ: "TRICHE"}), 2, temp_dir=tmp_path)
    assert res["statut"] == "ERREUR"
    assert res["code"] == w.E_COLONNE_FORMULE
    assert champ in res["details"]
    assert res["temp_path"] is None


def test_3b_refuse_un_champ_hors_schema(saisie, tmp_path):
    res = w.prepare_charge_flux_write(saisie, _row_data(colonne_inventee="x"), 2, temp_dir=tmp_path)
    assert res["code"] == w.E_CHAMP_INCONNU


def test_4_refuse_master_et_02_travail(tmp_path):
    # Nom MASTER_* → jamais écrivable (file_registry).
    master = tmp_path / "MASTER_FACT_MAN_Charges.xlsx"
    master.write_bytes(b"")
    res = w.prepare_charge_flux_write(master, _row_data(), 2, temp_dir=tmp_path)
    assert res["code"] == w.E_CHEMIN_NON_AUTORISE

    # Chemin sous 02_TRAVAIL → refusé même avec un nom SAISIE_*.
    sous_travail = cfg.TRAVAIL / "Lot3_Charges" / "SAISIE_faux.xlsx"
    res = w.prepare_charge_flux_write(sous_travail, _row_data(), 2, temp_dir=tmp_path)
    assert res["code"] == w.E_CHEMIN_NON_AUTORISE
    assert not sous_travail.exists()        # rien n'a été créé dans 02_TRAVAIL


def test_5_refuse_si_classeur_ouvert_dans_excel(saisie, tmp_path):
    verrou = saisie.parent / f"~${saisie.name}"
    verrou.write_bytes(b"")
    try:
        res = w.prepare_charge_flux_write(saisie, _row_data(), 2, temp_dir=tmp_path)
        assert res["code"] == w.E_FICHIER_OUVERT
    finally:
        verrou.unlink()


def test_6_preserve_onglets_validations_et_tables(saisie, tmp_path):
    from app.writers.saisie_hh_writer import _structure_signature

    avant = _structure_signature(saisie)
    res = w.prepare_charge_flux_write(saisie, _row_data(), 2, temp_dir=tmp_path)
    assert res["statut"] == "OK", res
    apres = _structure_signature(Path(res["temp_path"]))

    for cle in ("sheetnames", "defined_names", "data_validations", "mfc_rules",
                "tables", "protections", "external_links"):
        assert apres[cle] == avant[cle], f"{cle} altéré"


def test_6b_refuse_une_source_modifiee_depuis_la_previsualisation(saisie, tmp_path):
    res = w.prepare_charge_flux_write(saisie, _row_data(), 2,
                                      sha256_attendu="0" * 64, temp_dir=tmp_path)
    assert res["code"] == w.E_SOURCE_MODIFIEE


# ── 7-11 : impacts (SAISIE_Charges_Impacts) ──────────────────────────────────

def test_7_ecrit_affectations(impacts, tmp_path):
    avant = _sha(impacts)
    p = _persistable(affectations=[_affectation(1, "LOG_0001", 60.0),
                                   _affectation(2, "LOG_0002", 40.0)])
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)

    assert res["statut"] == "OK", res
    assert _sha(impacts) == avant                     # source intacte
    assert res["affectations_ecrites"] == 2

    rows = _lire(Path(res["temp_path"]), "AFFECTATIONS")
    mine = [r for r in rows if r["charge_id"] == CHARGE_ID]
    assert len(mine) == 2
    assert round(sum(r["quote_part"] for r in mine), 2) == MONTANT
    # La ligne de l'autre charge est toujours là, inchangée.
    autres = [r for r in rows if r["charge_id"] == "CHG-AUTRE"]
    assert len(autres) == 1 and autres[0]["quote_part"] == 42.0


def test_8_ecrit_menage(impacts, tmp_path):
    p = _persistable(menage=[_menage(1, intervenant="INT_0001")])
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)

    assert res["statut"] == "OK", res
    assert res["menage_ecrits"] == 1
    assert res["reserve_ecrites"] == 0
    row = _lire(Path(res["temp_path"]), "MENAGE")[0]
    assert row["intervenant_id"] == "INT_0001"
    assert row["logement_id"] is None                 # intervenant XOR logement


def test_8b_refuse_menage_avec_intervenant_et_logement(impacts, tmp_path):
    p = _persistable(menage=[_menage(1, intervenant="INT_0001", logement="LOG_0001")])
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)
    assert res["code"] == w.E_MENAGE_XOR


def test_8c_refuse_menage_sans_intervenant_ni_logement(impacts, tmp_path):
    p = _persistable(menage=[_menage(1)])
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)
    assert res["code"] == w.E_MENAGE_XOR


def test_9_ecrit_reserve_refacturation(impacts, tmp_path):
    p = _persistable(affectations=[_affectation(1, "LOG_0001", MONTANT)],
                     reserve=[_reserve(1, MONTANT)])
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)

    assert res["statut"] == "OK", res
    assert res["reserve_ecrites"] == 1
    row = _lire(Path(res["temp_path"]), "RESERVE_REFACTURATION")[0]
    assert row["statut_traitement"] == "EN_ATTENTE"   # aucune refacturation automatique
    assert row["montant_refacturable"] == MONTANT


def test_9b_refuse_une_reserve_deja_traitee(impacts, tmp_path):
    p = _persistable(reserve=[_reserve(1, MONTANT, statut="APPLIQUE")])
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)
    assert res["code"] == w.E_RESERVE_STATUT


def test_9c_refuse_une_reserve_superieure_au_montant(impacts, tmp_path):
    p = _persistable(reserve=[_reserve(1, 150.0)])
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)
    assert res["code"] == w.E_RESERVE_MONTANT


def test_10_refuse_somme_quotes_parts_differente_du_montant(impacts, tmp_path):
    p = _persistable(affectations=[_affectation(1, "LOG_0001", 60.0),
                                   _affectation(2, "LOG_0002", 30.0)])   # Σ = 90 ≠ 100
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)
    assert res["code"] == w.E_SOMME_QUOTES
    assert "90" in res["details"]


def test_11_charge_menage_jamais_en_reserve(impacts, tmp_path):
    p = _persistable(menage=[_menage(1, intervenant="INT_0001")], reserve=[_reserve(1, MONTANT)])
    res = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)
    assert res["code"] == w.E_MENAGE_AVEC_RESERVE


# ── 12-13 : idempotence et collision ─────────────────────────────────────────

def test_12_idempotence_impacts_pas_de_doublon(impacts, tmp_path):
    """Rejouer le même charge_id remplace ses lignes : jamais d'accumulation."""
    p = _persistable(affectations=[_affectation(1, "LOG_0001", MONTANT)])

    r1 = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)
    assert r1["statut"] == "OK", r1
    # On simule le remplacement : le temporaire devient la nouvelle source.
    shutil.copy2(r1["temp_path"], impacts)

    r2 = w.prepare_charge_impacts_write(impacts, p, MONTANT, temp_dir=tmp_path)
    assert r2["statut"] == "OK", r2
    assert r2["affectations_ecrites"] == 1

    rows = _lire(Path(r2["temp_path"]), "AFFECTATIONS")
    assert len([r for r in rows if r["charge_id"] == CHARGE_ID]) == 1     # toujours 1, pas 2
    assert len([r for r in rows if r["charge_id"] == "CHG-AUTRE"]) == 1


def test_12b_idempotence_ligne_charge_meme_ligne(saisie, tmp_path):
    """Réécrire la même charge sur la MÊME ligne = remplacement, jamais une seconde ligne."""
    r1 = w.prepare_charge_flux_write(saisie, _row_data(), 2, temp_dir=tmp_path)
    assert r1["statut"] == "OK", r1
    shutil.copy2(r1["temp_path"], saisie)

    r2 = w.prepare_charge_flux_write(saisie, _row_data(montant=120.0), 2, temp_dir=tmp_path)
    assert r2["statut"] == "OK", r2
    assert r2["remplacement"] is True

    lignes = _lire(Path(r2["temp_path"]), "SAISIE")
    mine = [r for r in lignes if r.get("charge_id") == CHARGE_ID]
    assert len(mine) == 1 and mine[0]["montant"] == 120.0


def test_13_collision_charge_id_refus_explicite(saisie, tmp_path):
    """Le même charge_id sur une AUTRE ligne = refus. Une charge = UNE ligne."""
    r1 = w.prepare_charge_flux_write(saisie, _row_data(), 2, temp_dir=tmp_path)
    assert r1["statut"] == "OK", r1
    shutil.copy2(r1["temp_path"], saisie)

    r2 = w.prepare_charge_flux_write(saisie, _row_data(), 3, temp_dir=tmp_path)   # ligne différente
    assert r2["code"] == w.E_CHARGE_ID_COLLISION
    assert "ligne 2" in r2["details"]
    assert r2["temp_path"] is None


def test_13b_ligne_cible_occupee_par_une_autre_charge(saisie, tmp_path):
    r1 = w.prepare_charge_flux_write(saisie, _row_data(), 2, temp_dir=tmp_path)
    shutil.copy2(r1["temp_path"], saisie)

    autre = _row_data(charge_id="CHG-2026-06-IC-BANQUE-002")
    res = w.prepare_charge_flux_write(saisie, autre, 2, temp_dir=tmp_path)   # ligne 2 déjà prise
    assert res["code"] == w.E_LIGNE_CIBLE_OCCUPEE


def test_13c_refuse_une_ligne_cible_sans_formules(saisie, tmp_path):
    """Une ligne hors zone modèle (sans formules) n'est pas une cible valide."""
    res = w.prepare_charge_flux_write(saisie, _row_data(), 50, temp_dir=tmp_path)
    assert res["code"] == w.E_FORMULE_MODELE_ABSENTE
    assert res["temp_path"] is None


# ── 14 : les fichiers métier RÉELS ne sont jamais touchés ────────────────────

def test_14_fichiers_metier_reels_inchanges(saisie, impacts, tmp_path):
    """Le writer n'écrit que dans les fixtures tmp — les vrais classeurs ne bougent pas."""
    reels = [p for p in (cfg.SAISIE_CHARGES, cfg.SAISIE_CHARGES_IMPACTS,
                         cfg.MASTER_CHARGES, cfg.SAISIE_IK_AVANTAGES) if p.exists()]
    avant = {p: _sha(p) for p in reels}

    w.prepare_charge_flux_write(saisie, _row_data(), 2, temp_dir=tmp_path)
    w.prepare_charge_impacts_write(
        impacts, _persistable(affectations=[_affectation(1, "LOG_0001", MONTANT)]),
        MONTANT, temp_dir=tmp_path,
    )

    assert {p: _sha(p) for p in reels} == avant


def test_14b_aucune_ecriture_reelle_les_flags_restent_off():
    """Le writer bas niveau ne remplace rien : les flags ne sont ni lus ni modifiés."""
    assert cfg.CHARGES_REAL_WRITE_ENABLED is False
    assert cfg.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED is False
    # Aucun remplacement à ce niveau : le module n'appelle pas os.replace (il ne l'importe même pas).
    source = Path(w.__file__).read_text(encoding="utf-8")
    assert "os.replace(" not in source
    assert "import os\n" not in source
