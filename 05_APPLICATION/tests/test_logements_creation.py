"""Création d'un logement au référentiel : validations + écriture gardée."""
import openpyxl
import pytest

import app.config as cfg
from app.services import logements_creation_service as svc

FORM = {"logement_id": "LOG_NEW", "nom_logement_officiel": "Fictif Nouveau", "nom_court": "NEW",
        "adresse": "1 rue Test", "ville": "RECETTE", "type_logement_id": "TYPE_001",
        "proprietaire_id": "PROP_A", "date_debut": "2026-06-01", "actif": "OUI"}


def _ref(tmp_path):
    p = tmp_path / "REF_Setup.xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("REF_Logements")
    ws.append(["logement_id", "hostaway_listing_id", "nom_logement_officiel", "nom_court", "adresse",
               "ville", "type_logement_id", "sur_hostaway", "actif", "statut_parc", "commentaire",
               "forfait_logiciel_consommables_mensuel"])
    ws.append(["LOG_A1", 900001, "Fictif A1", "A1", "1 rue", "RECETTE", "TYPE_001", "OUI", "OUI",
               "GERE", None, 0])
    ws = wb.create_sheet("REF_Gestion_Logements_Hist")
    ws.append(["gestion_id", "logement_id", "proprietaire_id", "date_debut", "date_fin",
               "statut_gestion", "source", "commentaire"])
    ws.append(["GST_1", "LOG_A1", "PROP_A", "2026-01-01", None, "ACTIF", "FICTIF", None])
    ws = wb.create_sheet("REF_Proprietaires")
    ws.append(["proprietaire_id", "nom_proprietaire", "actif"])
    for pid in ("PROP_A", "PROP_B"):
        ws.append([pid, f"Nom {pid}", "OUI"])
    ws = wb.create_sheet("REF_Types_Logements")
    ws.append(["type_logement_id", "libelle"]); ws.append(["TYPE_001", "STUDIO"])
    wb.save(p); wb.close()
    return p


@pytest.fixture
def ref(tmp_path, monkeypatch):
    p = _ref(tmp_path)
    monkeypatch.setattr(cfg, "REF_SETUP", p)
    monkeypatch.setattr(cfg, "RECETTE_MODE", True)
    monkeypatch.setattr(cfg, "RECETTE_ROOT", tmp_path.resolve())
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)
    return p


def _lignes(p, sheet):
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows


def test_referentiels_exposent_les_options(ref):
    r = svc.referentiels()
    assert r["status"] == "OK"
    assert {p["proprietaire_id"] for p in r["proprietaires"]} == {"PROP_A", "PROP_B"}
    assert "LOG_A1" in r["logements"]


def test_creation_ajoute_logement_et_rattachement_date(ref):
    res = svc.creer(dict(FORM))
    assert res["ok"], res
    logs = _lignes(ref, "REF_Logements")
    assert len(logs) == 2 and logs[1][0] == "LOG_NEW"
    gest = _lignes(ref, "REF_Gestion_Logements_Hist")
    assert len(gest) == 2
    assert gest[1][1] == "LOG_NEW" and gest[1][2] == "PROP_A" and gest[1][3] == "2026-06-01"
    assert gest[1][5] == "ACTIF"


def test_identifiant_deja_utilise_refuse(ref):
    codes = [e["code"] for e in svc.valider(dict(FORM, logement_id="LOG_A1"))]
    assert svc.E_ID_EXISTANT in codes


def test_identifiant_obligatoire(ref):
    codes = [e["code"] for e in svc.valider(dict(FORM, logement_id=""))]
    assert svc.E_ID_MANQUANT in codes


def test_nom_obligatoire(ref):
    codes = [e["code"] for e in svc.valider(dict(FORM, nom_logement_officiel="", nom_court=""))]
    assert svc.E_NOM_MANQUANT in codes


def test_proprietaire_inconnu_refuse(ref):
    codes = [e["code"] for e in svc.valider(dict(FORM, proprietaire_id="PROP_ZZ"))]
    assert svc.E_PROP_INCONNU in codes


def test_proprietaire_obligatoire(ref):
    codes = [e["code"] for e in svc.valider(dict(FORM, proprietaire_id=""))]
    assert svc.E_PROP_MANQUANT in codes


def test_date_invalide_refusee(ref):
    codes = [e["code"] for e in svc.valider(dict(FORM, date_debut="15/06/2026"))]
    assert svc.E_DATE_INVALIDE in codes


def test_type_inconnu_refuse(ref):
    codes = [e["code"] for e in svc.valider(dict(FORM, type_logement_id="TYPE_ZZ"))]
    assert svc.E_TYPE_INCONNU in codes


def test_creation_refusee_si_flags_desactives(ref, monkeypatch):
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", False)
    res = svc.creer(dict(FORM))
    assert res["ok"] is False and res["code"] == svc.E_FLAGS
    assert len(_lignes(ref, "REF_Logements")) == 1        # rien écrit


def test_creation_refusee_hors_racine_recette(ref, monkeypatch, tmp_path):
    monkeypatch.setattr(cfg, "RECETTE_ROOT", (tmp_path / "ailleurs").resolve())
    res = svc.creer(dict(FORM))
    assert res["ok"] is False and res["code"] == svc.E_ECRITURE
    assert len(_lignes(ref, "REF_Logements")) == 1        # write-guard : aucune écriture


def test_logement_inactif_cree_un_rattachement_retire(ref):
    assert svc.creer(dict(FORM, actif="NON"))["ok"]
    gest = _lignes(ref, "REF_Gestion_Logements_Hist")
    assert gest[1][5] == "RETIRE"
