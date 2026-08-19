"""Lot9 — test bloquant : `flux_unifie_service` n'ouvre jamais `MASTER_CALC_Flux.xlsx` (§27).

`MASTER_CALC_Flux.xlsx` est le classeur produit par le script legacy `lot9_construire_flux.py` —
un outil de comparaison (`LEGACY_PARITE_TEMPORAIRE`), jamais une source lue par le NEW. Preuve
active par interception `openpyxl.load_workbook`, même principe que
`test_menages_sans_excel.py` pour Ménages.
"""
from __future__ import annotations

from pathlib import Path

import pytest

INTERDIT = "MASTER_CALC_Flux.xlsx"


@pytest.fixture(autouse=True)
def _interdire_master_calc_flux(monkeypatch):
    import openpyxl

    original = openpyxl.load_workbook

    def garde(chemin, *a, **kw):
        if Path(chemin).name == INTERDIT:
            raise AssertionError(f"MASTER_CALC_Flux.xlsx interdit rouvert par le NEW : {chemin}")
        return original(chemin, *a, **kw)

    monkeypatch.setattr(openpyxl, "load_workbook", garde)


def test_construire_sans_master_calc_flux(tmp_db):
    from app.services import flux_unifie_service as svc

    resultat = svc.construire(db_path=tmp_db)
    assert resultat["ok"], resultat
    flux = svc.lire(db_path=tmp_db)
    assert flux == []  # tmp_db vide : aucune source alimentée, mais aucun crash ni ouverture Excel
