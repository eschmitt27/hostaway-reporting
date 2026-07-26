"""APP-3b / commit 5 — journal durable des tentatives d'écriture de charges.

Chaque test écrit dans une base SQLite **isolée sous tmp_path** : la base applicative réelle
(`05_APPLICATION/data/app.db`) n'est jamais touchée. Aucun fichier Excel métier n'est utilisé.
"""
from __future__ import annotations

import hashlib
import json
import os
import sqlite3
from pathlib import Path

import openpyxl
import pytest

import app.config as cfg
from app.readers.saisie_charges_reader import MANUAL_COL_MAP, _col_index
from app.services import saisie_charges_journal_service as jrn
from app.services import saisie_charges_lock_service as lk
from app.services import saisie_charges_transaction_service as svc

MONTANT = 100.0
CHARGE_ID = "CHG-2026-06-IC-BANQUE-001"
LIGNE = 2
TOKEN = "20260713T120000Z_abcdef123456"

FORMULES = {
    "C": '=IF(B{r}="","",TEXT(B{r},"YYYY-MM"))',
    "I": '=IF(H{r}="IC","OUI",IF(H{r}="HC","OUI",IF(H{r}="HR","NON","A_CONTROLER")))',
    "J": '=IF(H{r}="IC","OUI",IF(H{r}="HC","NON",IF(H{r}="HR","NON","A_CONTROLER")))',
    "AD": '=IF(A{r}="","",TEXT(B{r},"YYYYMMDD")&"|"&TEXT(D{r},"0.00")&"|"&F{r}&"|"&M{r})',
}
IMPACTS_HEADERS = {
    "AFFECTATIONS": ["affectation_id", "charge_id", "mois", "logement_id", "proprietaire_id",
                     "quote_part", "statut", "origine", "commentaire", "ROW_HASH"],
    "MENAGE": ["menage_impact_id", "charge_id", "mois", "mode", "intervenant_id", "logement_id",
               "statut", "origine", "commentaire", "ROW_HASH"],
    "RESERVE_REFACTURATION": ["reserve_id", "charge_id", "mois", "logement_id", "proprietaire_id",
                              "montant_refacturable", "libelle", "justificatif",
                              "statut_traitement", "trace_decision", "origine", "commentaire",
                              "ROW_HASH"],
}
REELS = [cfg.SAISIE_CHARGES, cfg.SAISIE_CHARGES_IMPACTS, cfg.MASTER_CHARGES, cfg.SAISIE_IK_AVANTAGES]


def _sha(p) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


@pytest.fixture
def flags_on(monkeypatch):
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)


@pytest.fixture
def db(tmp_path: Path) -> Path:
    """Base de journal isolée, dans un sous-dossier distinct de celui des classeurs.

    La base de production (05_APPLICATION/data/app.db) n'est jamais ouverte.
    """
    return tmp_path / "journal" / "journal_test.db"


@pytest.fixture
def cibles(tmp_path: Path) -> tuple[Path, Path]:
    saisie = tmp_path / "SAISIE_Charges_Flux.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    entetes = {**MANUAL_COL_MAP, "C": "mois", "I": "impact_resultat_reel",
               "J": "impact_resultat_comptable", "AD": "ROW_HASH"}
    for col, champ in entetes.items():
        ws.cell(row=1, column=_col_index(col), value=champ)
    for r in range(2, 7):
        for col, f in FORMULES.items():
            ws.cell(row=r, column=_col_index(col), value=f.format(r=r))
    wb.save(saisie)
    wb.close()

    impacts = tmp_path / "SAISIE_Charges_Impacts.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, headers in IMPACTS_HEADERS.items():
        wb.create_sheet(sheet).append(headers)
    wb.save(impacts)
    wb.close()
    return saisie, impacts


def _row_data(charge_id: str = CHARGE_ID) -> dict:
    return {
        "charge_id": charge_id, "date_charge": "2026-06-15", "montant": MONTANT,
        "sens_flux": "DEPENSE", "categorie_charge_id": "CHG_025", "type_flux_id": "TYPE_FLUX_020",
        "code_impact": "IC", "mode_paiement_id": "PAY_001", "statut_controle": "A_CONTROLER",
    }


def _persistable(charge_id: str = CHARGE_ID) -> dict:
    return {
        "charge_id": charge_id,
        "affectations": [{
            "affectation_id": f"{charge_id}-AFF-001", "charge_id": charge_id, "mois": "2026-06",
            "logement_id": "LOG_0001", "proprietaire_id": "PROP_01", "quote_part": MONTANT,
            "statut": "A_CONTROLER", "origine": "NOUVELLE_CHARGE_GUIDEE", "commentaire": None,
            "ROW_HASH": "h1",
        }],
        "menage": [], "reserve": [],
    }


def _demande(cibles, db, **kw) -> svc.DemandeEcritureCharge:
    saisie, impacts = cibles
    tmp = saisie.parent
    params = {
        "charge_id": CHARGE_ID, "montant": MONTANT, "target_row": LIGNE,
        "row_data": _row_data(), "persistable": _persistable(),
        "saisie_path": saisie, "impacts_path": impacts,
        "master_path": tmp / "verrous" / "MASTER_absent.xlsx",
        "lock_path": tmp / "verrous" / lk.LOCK_NAME,
        "db_path": db, "token_previsualisation": TOKEN,
    }
    params.update(kw)
    return svc.DemandeEcritureCharge(**params)


def _traces(db: Path) -> list[dict]:
    return list(reversed(jrn.traces_recentes(db_path=db)))     # ordre chronologique


# ── 1-8 : un statut métier distinct par situation ────────────────────────────

def test_1_transaction_reussie_trace_succes(cibles, db, flags_on):
    saisie, impacts = cibles
    res = svc.confirmer_ecriture_charge(_demande(cibles, db))
    assert res.ok, res

    traces = _traces(db)
    assert len(traces) == 1
    t = traces[0]
    assert t["statut"] == jrn.SUCCES
    assert t["code"] is None
    assert t["charge_id"] == CHARGE_ID
    assert t["charge_id_source"] == jrn.ID_FOURNI
    assert t["token_previsualisation"] == TOKEN
    assert t["transaction_id"] == res.transaction_id
    assert t["debut_utc"] and t["fin_utc"]

    # Les empreintes racontent bien ce qui s'est passé : avant ≠ après, et après = fichier réel.
    assert t["sha256_saisie_avant"] != t["sha256_saisie_apres"]
    assert t["sha256_saisie_apres"] == _sha(saisie)
    assert t["sha256_impacts_apres"] == _sha(impacts)

    assert json.loads(t["fichiers_remplaces"]) == ["SAISIE_Charges_Impacts.xlsx",
                                                   "SAISIE_Charges_Flux.xlsx"]
    assert t["rollback_tente"] == 0
    assert t["rollback_reussi"] is None
    assert t["fichiers_non_restaures"] is None
    assert t["residus"] is None
    assert t["app_pid"] == os.getpid()
    assert res.journal_erreur is None


def test_2_flags_desactives_trace_refuse_flags(cibles, db):
    res = svc.confirmer_ecriture_charge(_demande(cibles, db))
    assert res.statut == svc.STATUT_GARDE

    t = _traces(db)[0]
    assert t["statut"] == jrn.REFUSE_FLAGS
    assert t["code"] == svc.E_FLAGS_DESACTIVES
    assert t["fichiers_remplaces"] is None          # rien remplacé
    assert t["sha256_saisie_avant"] == t["sha256_saisie_apres"]   # rien touché


def test_3_verrou_deja_pris_trace_refuse_verrou(cibles, db, flags_on):
    d = _demande(cibles, db)
    tenu = lk.acquerir_verrou(charge_id="CHG-AUTRE", lock_path=Path(d.lock_path))
    try:
        res = svc.confirmer_ecriture_charge(d)
    finally:
        lk.liberer_verrou(tenu)

    assert res.code == svc.E_VERROU_DEJA_PRIS
    t = _traces(db)[0]
    assert t["statut"] == jrn.REFUSE_VERROU
    assert t["verrou_pid"] == os.getpid()           # le détenteur qui a bloqué
    assert t["verrou_hostname"]
    assert t["fichiers_remplaces"] is None


def test_4_validation_invalide_trace_refuse_validation(cibles, db, flags_on):
    res = svc.confirmer_ecriture_charge(_demande(cibles, db, montant=0.0))
    assert res.code == svc.E_DEMANDE_INVALIDE

    t = _traces(db)[0]
    assert t["statut"] == jrn.REFUSE_VALIDATION
    assert t["code"] == svc.E_DEMANDE_INVALIDE


def test_5_charge_id_deja_utilise_trace_refuse_charge_id(cibles, db, flags_on):
    assert svc.confirmer_ecriture_charge(_demande(cibles, db)).ok
    res = svc.confirmer_ecriture_charge(
        _demande(cibles, db, target_row=3, token_previsualisation="autre-token")
    )
    assert res.code == svc.E_CHARGE_ID_EXISTANT

    traces = _traces(db)
    assert [t["statut"] for t in traces] == [jrn.SUCCES, jrn.REFUSE_CHARGE_ID]
    assert traces[1]["charge_id"] == CHARGE_ID


def test_6_echec_avant_remplacement_trace_echec_preparation(cibles, db, flags_on):
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    d = _demande(cibles, db)
    d.persistable["affectations"][0]["quote_part"] = 60.0      # Σ ≠ montant
    res = svc.confirmer_ecriture_charge(d)
    assert res.code == svc.E_PREPARATION

    t = _traces(db)[0]
    assert t["statut"] == jrn.ECHEC_PREPARATION
    assert t["rollback_tente"] == 0
    assert t["fichiers_remplaces"] is None
    assert {p: _sha(p) for p in (saisie, impacts)} == avant


def test_7_rollback_reussi_trace_rollback_reussi(cibles, db, flags_on, monkeypatch):
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    vrai = os.replace
    etat = {"n": 0}

    def faux(src, dst, *a, **kw):
        etat["n"] += 1
        if etat["n"] == 2:                          # le 2e remplacement échoue
            raise OSError("échec simulé")
        return vrai(src, dst, *a, **kw)

    monkeypatch.setattr(os, "replace", faux)
    res = svc.confirmer_ecriture_charge(_demande(cibles, db))
    assert res.statut == svc.STATUT_ROLLBACK

    t = _traces(db)[0]
    assert t["statut"] == jrn.ROLLBACK_REUSSI
    assert t["code"] == svc.E_REMPLACEMENT
    assert t["rollback_tente"] == 1
    assert t["rollback_reussi"] == 1
    assert t["fichiers_non_restaures"] is None
    # La trace prouve le retour à l'état initial.
    assert t["sha256_saisie_apres"] == t["sha256_saisie_avant"]
    assert t["sha256_impacts_apres"] == t["sha256_impacts_avant"]
    assert {p: _sha(p) for p in (saisie, impacts)} == avant


def test_8_rollback_critique_trace_avec_fichiers_non_restaures(cibles, db, flags_on, monkeypatch):
    """Le scénario qui justifie tout ce commit : sans trace, l'incident n'existerait nulle part."""
    vrai = os.replace
    etat = {"n": 0}

    def faux(src, dst, *a, **kw):
        etat["n"] += 1
        if etat["n"] == 2:
            raise OSError("échec simulé")
        return vrai(src, dst, *a, **kw)

    monkeypatch.setattr(os, "replace", faux)
    monkeypatch.setattr(svc, "_restaurer_fichier",
                        lambda s, c: (_ for _ in ()).throw(OSError("restauration impossible")))

    with pytest.raises(svc.RollbackCritiqueError) as exc:
        svc.confirmer_ecriture_charge(_demande(cibles, db))

    t = _traces(db)[0]
    assert t["statut"] == jrn.ROLLBACK_CRITIQUE
    assert t["rollback_tente"] == 1
    assert t["rollback_reussi"] == 0
    non_restaures = json.loads(t["fichiers_non_restaures"])
    assert len(non_restaures) == 1
    assert "SAISIE_Charges_Impacts" in non_restaures[0]["cible"]
    assert non_restaures[0]["sauvegarde"]           # la voie de retour est tracée
    assert exc.value.journal_erreur is None         # la trace a bien été écrite


# ── 9 : double tentative avec le même token ─────────────────────────────────

def test_9_double_tentative_meme_token_est_idempotente(cibles, db, flags_on):
    """Le token du dry-run est la clé d'idempotence : une 2e confirmation n'écrit JAMAIS 2 fois.

    Comportement documenté : la 2e tentative est REFUSÉE (E_TOKEN_DEJA_ECRIT), tracée comme
    REFUSE_VALIDATION, et renvoie le charge_id déjà écrit. Elle ne touche aucun fichier.
    """
    saisie, impacts = cibles
    r1 = svc.confirmer_ecriture_charge(_demande(cibles, db))
    assert r1.ok, r1
    apres_1 = {p: _sha(p) for p in (saisie, impacts)}

    r2 = svc.confirmer_ecriture_charge(_demande(cibles, db))     # même token, même charge
    assert r2.statut == svc.STATUT_ERREUR
    assert r2.code == svc.E_TOKEN_DEJA_ECRIT
    assert r2.charge_id == CHARGE_ID                             # renvoie la charge déjà écrite
    assert "déjà produit" in r2.details
    assert {p: _sha(p) for p in (saisie, impacts)} == apres_1    # aucun second effet

    traces = _traces(db)
    assert [t["statut"] for t in traces] == [jrn.SUCCES, jrn.REFUSE_VALIDATION]
    assert len(jrn.traces_du_token(TOKEN, db_path=db)) == 2      # les 2 tentatives sont tracées
    assert jrn.token_deja_ecrit(TOKEN, db_path=db)["charge_id"] == CHARGE_ID


# ── 10 : panne du journal ───────────────────────────────────────────────────

def test_10_panne_du_journal_ne_perd_pas_le_resultat(cibles, db, flags_on, monkeypatch):
    """Le journal tombe : la transaction garde son résultat, et la panne est SIGNALÉE (jamais avalée)."""
    saisie, impacts = cibles

    def journal_ko(*a, **kw):
        raise sqlite3.OperationalError("base verrouillée (panne simulée)")

    monkeypatch.setattr(jrn, "cloturer_trace", journal_ko)
    res = svc.confirmer_ecriture_charge(_demande(cibles, db))

    # Résultat principal intact : la charge est bien écrite.
    assert res.ok, res
    assert res.fichiers_remplaces
    wb = openpyxl.load_workbook(saisie, read_only=True, data_only=True)
    lignes = [r for r in wb["SAISIE"].iter_rows(values_only=True) if r[0] == CHARGE_ID]
    wb.close()
    assert len(lignes) == 1

    # Mais la panne est explicite — pas de `except: pass`.
    assert res.journal_erreur is not None
    assert "JOURNALISATION ÉCHOUÉE" in res.journal_erreur
    assert res.transaction_id in res.journal_erreur


def test_10b_panne_du_journal_nempeche_pas_le_rollback(cibles, db, flags_on, monkeypatch):
    """Contrainte dure : une erreur de journal ne doit JAMAIS empêcher un rollback."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    vrai = os.replace
    etat = {"n": 0}

    def faux(src, dst, *a, **kw):
        etat["n"] += 1
        if etat["n"] == 2:
            raise OSError("échec simulé")
        return vrai(src, dst, *a, **kw)

    monkeypatch.setattr(os, "replace", faux)
    monkeypatch.setattr(jrn, "cloturer_trace",
                        lambda *a, **kw: (_ for _ in ()).throw(sqlite3.OperationalError("panne")))

    res = svc.confirmer_ecriture_charge(_demande(cibles, db))

    assert res.statut == svc.STATUT_ROLLBACK
    assert {p: _sha(p) for p in (saisie, impacts)} == avant      # le rollback a bien eu lieu
    assert "JOURNALISATION ÉCHOUÉE" in res.journal_erreur


# ── FAIL-CLOSED : token fourni + journal indisponible ⇒ REFUS ───────────────

def test_10c_token_fourni_et_base_indisponible_refuse_tout(cibles, db, flags_on, monkeypatch):
    """Impossible de savoir si ce token a déjà écrit ⇒ on REFUSE. Une saisie bloquée vaut mieux
    qu'une charge comptée deux fois. La disponibilité ne prime pas sur le risque comptable."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}
    dossier = sorted(f.name for f in saisie.parent.iterdir() if f.is_file())

    monkeypatch.setattr(jrn, "token_deja_ecrit",
                        lambda *a, **kw: (_ for _ in ()).throw(sqlite3.OperationalError("base HS")))
    d = _demande(cibles, db)
    res = svc.confirmer_ecriture_charge(d)

    assert res.statut == svc.STATUT_ERREUR
    assert res.code == svc.E_IDEMPOTENCE_INDISPONIBLE
    assert "aucun fichier n'a été touché" in res.details.lower()

    # Aucun fichier modifié, aucun temporaire préparé, aucun verrou résiduel.
    assert {p: _sha(p) for p in (saisie, impacts)} == avant
    assert sorted(f.name for f in saisie.parent.iterdir() if f.is_file()) == dossier
    assert not Path(d.lock_path).exists()

    # Le refus est tracé (la clôture, elle, fonctionne).
    t = _traces(db)[0]
    assert t["statut"] == jrn.REFUSE_VALIDATION
    assert t["code"] == svc.E_IDEMPOTENCE_INDISPONIBLE
    assert t["fichiers_remplaces"] is None


def test_10d_token_fourni_et_journal_totalement_hs_refuse_quand_meme(cibles, db, flags_on, monkeypatch):
    """Même si la clôture du journal tombe aussi : le refus tient, et la panne est signalée."""
    saisie, impacts = cibles
    avant = {p: _sha(p) for p in (saisie, impacts)}

    def hs(*a, **kw):
        raise sqlite3.OperationalError("base HS")

    monkeypatch.setattr(jrn, "token_deja_ecrit", hs)
    monkeypatch.setattr(jrn, "cloturer_trace", hs)
    res = svc.confirmer_ecriture_charge(_demande(cibles, db))

    assert res.code == svc.E_IDEMPOTENCE_INDISPONIBLE      # refus maintenu
    assert {p: _sha(p) for p in (saisie, impacts)} == avant
    assert "JOURNALISATION ÉCHOUÉE" in res.journal_erreur  # et la panne n'est pas cachée


def test_10e_sans_token_la_garde_ne_sapplique_pas(cibles, db, flags_on, monkeypatch):
    """Comportement documenté : sans token, il n'y a rien à comparer — la garde d'idempotence ne
    s'applique pas et une base indisponible n'empêche pas d'écrire. C'est alors l'unicité du
    charge_id, vérifiée sous verrou, qui protège du doublon."""
    saisie, _ = cibles

    def hs(*a, **kw):
        raise sqlite3.OperationalError("base HS")

    appels: list[str] = []
    monkeypatch.setattr(jrn, "token_deja_ecrit",
                        lambda *a, **kw: appels.append("consulte") or hs())
    monkeypatch.setattr(jrn, "cloturer_trace", hs)

    res = svc.confirmer_ecriture_charge(_demande(cibles, db, token_previsualisation=None))

    assert res.ok, res                       # l'écriture a bien lieu
    assert appels == []                      # la garde n'a même pas été consultée
    assert "JOURNALISATION ÉCHOUÉE" in res.journal_erreur   # panne de clôture signalée

    wb = openpyxl.load_workbook(saisie, read_only=True, data_only=True)
    lignes = [r for r in wb["SAISIE"].iter_rows(values_only=True) if r[0] == CHARGE_ID]
    wb.close()
    assert len(lignes) == 1


# ── ETAT_INCOHERENT : chemin non prévu, on refuse de deviner ─────────────────

def test_etat_incoherent_exception_inattendue(cibles, db, flags_on, monkeypatch):
    """Une exception imprévue APRÈS les remplacements : on ne sait pas conclure, on le DIT.

    Aucun rollback n'est tenté à l'aveugle ; les sauvegardes restent la voie de retour ;
    le verrou est malgré tout libéré (un état incohérent ne doit pas bloquer en plus).
    """
    saisie, impacts = cibles
    d = _demande(cibles, db)

    def boum(remplaces):
        raise RuntimeError("panne imprévue pendant la vérification post-commit")

    monkeypatch.setattr(svc, "_verifier_post_commit", boum)

    with pytest.raises(RuntimeError, match="panne imprévue"):
        svc.confirmer_ecriture_charge(d)

    t = _traces(db)[0]
    assert t["statut"] == jrn.ETAT_INCOHERENT
    assert t["code"] == "RuntimeError"                    # code technique explicite
    assert "panne imprévue" in t["details"]               # détail conservé
    assert t["rollback_tente"] == 0                       # aucun rollback à l'aveugle
    assert t["rollback_reussi"] is None

    # Les empreintes relevées décrivent l'état RÉEL des fichiers, sans le maquiller.
    assert t["sha256_saisie_apres"] == _sha(saisie)
    assert t["sha256_impacts_apres"] == _sha(impacts)
    assert t["sha256_saisie_apres"] != t["sha256_saisie_avant"]   # les remplacements ont eu lieu

    # Aucune donnée métier dans la trace.
    contenu = json.dumps(t, ensure_ascii=False)
    for interdit in ("PROP_01", "LOG_0001", "PERS_EWAN"):
        assert interdit not in contenu

    assert not Path(d.lock_path).exists()                 # verrou libéré
    # Les sauvegardes techniques SUBSISTENT : elles sont la seule voie de retour dans cet état.
    baks = [f.name for f in saisie.parent.iterdir() if ".bak" in f.name]
    assert len(baks) == 2


# ── 11-12 : hygiène ─────────────────────────────────────────────────────────

def test_11_aucune_donnee_metier_ni_fichier_reel_dans_le_journal(cibles, db, flags_on):
    """On journalise la TENTATIVE, pas la charge : ni montant, ni logement, ni propriétaire."""
    presents = [p for p in REELS if p.exists()]
    avant = {p: _sha(p) for p in presents}

    assert svc.confirmer_ecriture_charge(_demande(cibles, db)).ok

    assert {p: _sha(p) for p in presents} == avant               # Excel réels intacts
    assert not cfg.DB_PATH.samefile(db) if cfg.DB_PATH.exists() else True

    t = _traces(db)[0]
    contenu = json.dumps(t, ensure_ascii=False)
    for interdit in ("PROP_01", "LOG_0001", "100.0", "PERS_EWAN", "quote_part", "affectation_id"):
        assert interdit not in contenu, f"donnée métier fuitée dans le journal : {interdit}"

    # Cibles = NOMS de fichiers, jamais des chemins temporaires complets.
    assert t["cible_saisie"] == "SAISIE_Charges_Flux.xlsx"
    assert t["cible_impacts"] == "SAISIE_Charges_Impacts.xlsx"
    assert "\\" not in (t["cible_saisie"] or "") and "/" not in (t["cible_saisie"] or "")


def test_12_les_flags_du_depot_restent_a_false():
    source = Path(cfg.__file__).read_text(encoding="utf-8")
    assert "CHARGES_REAL_WRITE_ENABLED = RECETTE_MODE and _env_flag(" in source
    assert "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED = RECETTE_MODE and _env_flag(" in source
    assert cfg.CHARGES_REAL_WRITE_ENABLED is False
    assert cfg.CHARGES_REAL_WRITE_CONFIRMATION_ENABLED is False


# ── Le journal lui-même ─────────────────────────────────────────────────────

def test_statut_inconnu_refuse(db):
    trace = jrn.ouvrir_trace(db_path=db)
    with pytest.raises(ValueError, match="statut inconnu"):
        jrn.cloturer_trace(trace, "PAS_UN_STATUT")


def test_les_neuf_statuts_metier_existent():
    assert jrn.STATUTS == {
        "REFUSE_FLAGS", "REFUSE_VERROU", "REFUSE_VALIDATION", "REFUSE_CHARGE_ID",
        "ECHEC_PREPARATION", "SUCCES", "ROLLBACK_REUSSI", "ROLLBACK_CRITIQUE", "ETAT_INCOHERENT",
    }


def test_statut_metier_et_code_technique_sont_distincts(cibles, db, flags_on):
    """Une colonne pour le verdict métier, une autre pour le code technique. Jamais confondus."""
    svc.confirmer_ecriture_charge(_demande(cibles, db, montant=-5.0))
    t = _traces(db)[0]
    assert t["statut"] == jrn.REFUSE_VALIDATION      # métier
    assert t["code"] == svc.E_DEMANDE_INVALIDE       # technique
    assert t["statut"] != t["code"]
