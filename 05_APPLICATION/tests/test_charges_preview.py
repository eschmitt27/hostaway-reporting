"""APP-3b-2 — Prévisualisation de saisie charge, entièrement SQLite.

Remplace l'ancienne suite Excel (copie SAISIE_Charges_Flux, injection de ligne, empreintes de
fichiers) : la prévisualisation ne lit plus que les référentiels `ref_*` (migration 0029) et ne
produit qu'un manifest JSON. La preuve « zéro classeur » est active : `openpyxl.load_workbook` lève
si un classeur interne est ouvert.
"""
from __future__ import annotations

from pathlib import Path

import pytest

import app.config as cfg
import fixtures_referentiel as fx
from app.services import charges_preview_service as prev


@pytest.fixture(autouse=True)
def _aucun_classeur(tmp_db, tmp_path, monkeypatch):
    import openpyxl

    fx.semer_parc_standard(tmp_db)
    fx.semer_referentiel_charges(tmp_db)
    monkeypatch.setattr(cfg, "REF_SETUP", tmp_path / "REF_Setup_ABSENT.xlsm")

    original = openpyxl.load_workbook

    def garde(chemin, *a, **kw):
        raise AssertionError(f"Classeur ouvert au runtime de la previsualisation : {chemin}")

    monkeypatch.setattr(openpyxl, "load_workbook", garde)
    return tmp_db


@pytest.fixture
def dryruns(tmp_path: Path) -> Path:
    return tmp_path / "dryruns"


def _form(**overrides) -> dict:
    base = {
        "date_charge": "2026-06-15", "montant": "100.00", "categorie_charge_id": "CHG_017",
        "code_impact": "IC", "mode_paiement_id": "PAY_001", "refacturable": "NON",
    }
    base.update(overrides)
    return base


def test_load_form_refs_lit_les_referentiels_sqlite(tmp_db):
    refs = prev.load_form_refs(db_path=tmp_db)
    assert {c["categorie_charge_id"] for c in refs["categories_all"]} == {"CHG_017"}
    assert {m["mode_paiement_id"] for m in refs["modes_paiement"]} == {"PAY_001"}
    assert refs["cloture"]


def test_previsualisation_ok_sans_classeur(tmp_db, dryruns):
    res = prev.previsualiser(_form(), db_path=tmp_db, dryruns_root=dryruns)
    assert res["ok"], res["manifest"].get("errors")
    manifest = res["manifest"]
    assert manifest["status"] == "OK"
    assert manifest["charge_id"].startswith("CHG-")
    assert manifest["mois_charge"] == "2026-06"
    assert manifest["row_data"]["montant"] == 100.0
    assert manifest["row_data"]["statut_controle"] == "A_CONTROLER"
    assert manifest["integrite"] == prev.sceller_manifest(manifest)


def test_previsualisation_refuse_montant_manquant(tmp_db, dryruns):
    res = prev.previsualiser(_form(montant=""), db_path=tmp_db, dryruns_root=dryruns)
    assert not res["ok"]
    codes = {e["code"] for e in res["manifest"]["errors"]}
    assert "V03_MONTANT_MANQUANT" in codes


def test_previsualisation_refuse_mois_cloture(tmp_db, dryruns):
    res = prev.previsualiser(_form(date_charge="2026-01-10"), db_path=tmp_db, dryruns_root=dryruns)
    assert not res["ok"]
    codes = {e["code"] for e in res["manifest"]["errors"]}
    assert "V02_MOIS_CLOTURE" in codes


def test_previsualisation_refuse_categorie_inconnue(tmp_db, dryruns):
    res = prev.previsualiser(_form(categorie_charge_id="CHG_INEXISTANTE"),
                             db_path=tmp_db, dryruns_root=dryruns)
    assert not res["ok"]
    codes = {e["code"] for e in res["manifest"]["errors"]}
    assert "V04_CATEGORIE_INVALIDE" in codes


def test_load_previsualisation_relit_le_manifest(tmp_db, dryruns):
    res = prev.previsualiser(_form(), db_path=tmp_db, dryruns_root=dryruns)
    relu = prev.load_previsualisation(res["token"], dryruns_root=dryruns)
    assert relu["manifest"]["charge_id"] == res["manifest"]["charge_id"]


def test_token_inconnu_leve(tmp_db, dryruns):
    with pytest.raises(prev.ChargesPreviewError):
        prev.load_previsualisation("jamais-vu", dryruns_root=dryruns)
