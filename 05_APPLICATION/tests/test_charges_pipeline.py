"""Chaîne Charges pilotée depuis l'application : orchestration, scénarios métier, double comptage.

Ces tests portent sur ce que le tour précédent n'avait pas exercé : la chaîne `charges` du pilotage
et la cohérence des scénarios métier une fois passés dans Lot9/Lot10.

Deux catégories :
- **structure** — toujours exécutés, ils vérifient la déclaration de la chaîne et le verdict runner ;
- **réconciliation** — exécutés seulement si une chaîne a réellement tourné sur la racine courante,
  ignorés sinon (ils constatent des sorties, ils ne les fabriquent pas).
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import pytest

import app.config as cfg
from app.services import calculs_executeur_service as ex

FLUX = "02_TRAVAIL/Lot9_FluxUnifie/MASTER_CALC_Flux.xlsx"
MASTER_CHARGES = "02_TRAVAIL/Lot3_Charges/MASTER_FACT_MAN_Charges.xlsx"
RESULTATS = "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_Resultats.xlsx"
NET = "02_TRAVAIL/Lot10_Resultats/MASTER_CALC_NetProprietaire.xlsx"


def _racine() -> Path:
    return Path(cfg.PROJECT_ROOT)


def _feuille(rel: str, onglet: str) -> list[dict]:
    import openpyxl
    p = _racine() / rel
    if not p.exists():
        pytest.skip(f"Sortie absente sur cette racine : {rel}")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        if onglet not in wb.sheetnames:
            pytest.skip(f"Onglet {onglet} absent de {rel}")
        it = wb[onglet].iter_rows(values_only=True)
        hdr = [str(c) if c is not None else "" for c in next(it)]
        return [dict(zip(hdr, r)) for r in it if any(c is not None for c in r)]
    finally:
        wb.close()


def _charges_du_flux() -> list[dict]:
    return [f for f in _feuille(FLUX, "MASTER") if str(f.get("sens")) == "CHARGE"]


# ── Structure de la chaîne ───────────────────────────────────────────────────

def test_lot3_est_pilote_par_le_runner_existant():
    """Lot3 est une BIBLIOTHÈQUE : aucun bloc `__main__`. Le lancer comme un script rendait code 0
    sans rien produire. Le pilotage réutilise l'orchestrateur qui existait déjà plutôt que d'en
    écrire un second."""
    lot = ex.TOUS_LES_LOTS["lot3"]
    assert lot.runner == "charges_post_write_runner.py"
    assert (ex.RUNNERS_DIR / lot.runner).exists()


def test_le_moteur_lot3_n_a_toujours_pas_de_point_d_entree():
    """Verrouille la raison d'être du runner : si Lot3 gagnait un `__main__`, ce test le signale."""
    moteur = _racine() / "02_TRAVAIL" / "lot3_generateur_charges.py"
    if not moteur.exists():
        pytest.skip("Moteur Lot3 absent de cette racine.")
    assert '__main__' not in moteur.read_text(encoding="utf-8")


def test_le_verdict_runner_refuse_une_reponse_absente(tmp_path, monkeypatch):
    """Le runner rend TOUJOURS le code 0 et porte l'échec dans son JSON. Sans lecture de ce JSON,
    une étape en échec passerait pour un succès dès qu'une sortie d'un run précédent traîne."""
    monkeypatch.setattr(cfg, "DRYRUNS_DIR", tmp_path / "dryruns")
    ok, message = ex._verdict_runner(ex.TOUS_LES_LOTS["lot3"], tmp_path)
    assert not ok and "aucune réponse" in message


def test_le_verdict_runner_refuse_une_etape_en_echec(tmp_path, monkeypatch):
    import json
    monkeypatch.setattr(cfg, "DRYRUNS_DIR", tmp_path / "dryruns")
    dossier = tmp_path / "dryruns" / "calculs_runners"
    dossier.mkdir(parents=True)
    (dossier / "lot3_reponse.json").write_text(json.dumps({
        "ok": False,
        "lot3": {"statut": "ECHEC", "details": "FileNotFoundError"},
        "lot7": {"statut": "NON_APPLICABLE"},
        "lot11": {"statut": "OK"},
    }), encoding="utf-8")
    ok, message = ex._verdict_runner(ex.TOUS_LES_LOTS["lot3"], tmp_path)
    assert not ok
    assert "lot3=ECHEC" in message


def test_le_verdict_runner_accepte_non_applicable(tmp_path, monkeypatch):
    """NON_APPLICABLE n'est pas un échec : le pilotage régénère le MASTER, il ne rejoue pas
    l'écriture d'une charge portant un avantage associé."""
    import json
    monkeypatch.setattr(cfg, "DRYRUNS_DIR", tmp_path / "dryruns")
    dossier = tmp_path / "dryruns" / "calculs_runners"
    dossier.mkdir(parents=True)
    (dossier / "lot3_reponse.json").write_text(json.dumps({
        "ok": True,
        "lot3": {"statut": "OK"},
        "lot7": {"statut": "NON_APPLICABLE"},
        "lot11": {"statut": "OK"},
    }), encoding="utf-8")
    ok, message = ex._verdict_runner(ex.TOUS_LES_LOTS["lot3"], tmp_path)
    assert ok and message == ""


# ── Absence de double comptage ───────────────────────────────────────────────

def test_pas_de_double_comptage_frais_bancaire():
    """Un frais bancaire injecté par Lot9 depuis NORM_Banque ne doit PAS être aussi saisi en charge
    manuelle. Le jeu de recette précédent le faisait : 8,90 € comptés deux fois, sans que rien ne
    le détecte."""
    par_cle: dict[tuple, list[str]] = defaultdict(list)
    for f in _charges_du_flux():
        montant = round(float(f.get("montant") or 0), 2)
        par_cle[(str(f.get("mois")), montant)].append(str(f.get("source_table")))

    doubles = {
        cle: sources for cle, sources in par_cle.items()
        if "BANQUE_LOT8_IMPORT_NORM_Banque" in sources
        and "MASTER_FACT_MAN_Charges" in sources
    }
    assert not doubles, (
        "Même mois et même montant présents à la fois en flux bancaire et en charge manuelle — "
        f"double comptage probable : {doubles}")


def test_une_charge_donne_exactement_un_flux():
    """Règle structurante : une charge = UNE charge économique. Aucune duplication par affectation
    analytique, ménage ou réserve."""
    compte: dict[str, int] = defaultdict(int)
    for f in _charges_du_flux():
        if str(f.get("source_table")) == "MASTER_FACT_MAN_Charges":
            compte[str(f.get("source_pk"))] += 1
    doublons = {k: v for k, v in compte.items() if v > 1}
    assert not doublons, f"Charges dupliquées dans le flux : {doublons}"


def test_flux_id_uniques():
    ids = [str(f.get("flux_id")) for f in _feuille(FLUX, "MASTER")]
    assert len(ids) == len(set(ids))


# ── Scénarios métier ─────────────────────────────────────────────────────────

def _charges_master() -> dict[str, dict]:
    return {str(c.get("charge_id")): c for c in _feuille(MASTER_CHARGES, "MASTER")}


def test_scenario_a_charge_conciergerie_impacte_reel_et_comptable():
    c = _charges_master().get("CHG_A_LOGICIEL")
    if c is None:
        pytest.skip("Scénario A absent du jeu de recette de cette racine.")
    assert str(c["code_impact"]) == "IC"
    assert str(c["impact_resultat_reel"]) == "OUI"
    assert str(c["impact_resultat_comptable"]) == "OUI"
    assert str(c["refacturable"]) == "NON"


def test_scenario_b_refacturable_est_rattachee_a_un_logement():
    c = _charges_master().get("CHG_B_REFACT")
    if c is None:
        pytest.skip("Scénario B absent.")
    assert str(c["refacturable"]) == "OUI"
    assert c["logement_id"] and c["proprietaire_id"], "Une refacturation sans propriétaire est une anomalie."
    assert str(c["methode_traitement"]) == "REFACTURATION_PROPRIETAIRE"


def test_scenario_c_paiement_personnel_associe_est_hors_compta():
    c = _charges_master().get("CHG_C_PERSO")
    if c is None:
        pytest.skip("Scénario C absent.")
    assert c["associe_id"], "Un paiement personnel associé sans associé résolu est une anomalie."
    assert str(c["code_impact"]) == "HC"
    assert str(c["impact_resultat_reel"]) == "OUI"
    assert str(c["impact_resultat_comptable"]) == "NON"


def test_scenario_d_hors_comptabilite_conserve_sa_justification():
    c = _charges_master().get("CHG_D_HORSCOMPTA")
    if c is None:
        pytest.skip("Scénario D absent.")
    assert str(c["code_impact"]) == "HC"
    assert str(c["impact_resultat_comptable"]) == "NON"
    assert str(c["commentaire"]).strip(), "La justification doit être conservée."


def test_scenario_e_ventilation_multi_logements_conserve_le_total():
    charges = _charges_master()
    parts = [charges.get("CHG_E_MULTI_1"), charges.get("CHG_E_MULTI_2")]
    if any(p is None for p in parts):
        pytest.skip("Scénario E absent.")
    assert sum(float(p["montant"]) for p in parts) == pytest.approx(60.00)
    assert {str(p["logement_id"]) for p in parts} == {"LOG_A2", "LOG_B1"}


def test_scenario_f_charge_fournisseur_reste_unique():
    """Facture, règlement et rapprochement bancaire vivent dans SQLite : ils ne recréent jamais
    d'impact économique. Une charge fournisseur reste une seule ligne."""
    flux = [f for f in _charges_du_flux() if str(f.get("source_pk")) == "CHG_SEED_001"]
    if not flux:
        pytest.skip("Scénario F absent.")
    assert len(flux) == 1
    assert float(flux[0]["montant"]) == pytest.approx(120.00)


# ── Contrôles : ce qui ne doit PAS entrer ────────────────────────────────────

@pytest.mark.parametrize("charge_id", ["CHG_CTRL_NONVALID", "CHG_CTRL_EXCLUE"])
def test_charge_non_validee_n_entre_pas_dans_les_flux(charge_id):
    if charge_id not in _charges_master():
        pytest.skip(f"{charge_id} absent du jeu de recette.")
    injectees = [f for f in _charges_du_flux() if str(f.get("source_pk")) == charge_id]
    assert not injectees, f"{charge_id} ne doit jamais être injectée dans les résultats."


# ── Réconciliation des totaux ────────────────────────────────────────────────

def test_reconciliation_reel_egale_comptable_plus_hors_compta():
    """Invariant du modèle : REEL = COMPTABLE + HORS_COMPTA. C'est ce qui garantit qu'aucune charge
    n'est comptée deux fois ni oubliée entre les trois visions."""
    par_vision = {str(l["vision"]): l for l in _feuille(RESULTATS, "GLOBAL")}
    if not {"REEL", "COMPTABLE", "HORS_COMPTA"} <= set(par_vision):
        pytest.skip("Visions incomplètes sur cette racine.")
    reel = float(par_vision["REEL"]["total_charges"])
    compta = float(par_vision["COMPTABLE"]["total_charges"])
    hc = float(par_vision["HORS_COMPTA"]["total_charges"])
    assert reel == pytest.approx(compta + hc, abs=0.01)


def test_les_charges_du_flux_egalent_le_total_lot10():
    """Le total de charges de Lot10 doit être exactement la somme des flux CHARGE de Lot9 :
    aucune charge n'apparaît ni ne disparaît entre les deux lots."""
    par_vision = {str(l["vision"]): l for l in _feuille(RESULTATS, "GLOBAL")}
    if "REEL" not in par_vision:
        pytest.skip("Vision REEL absente.")
    somme = sum(float(f.get("montant") or 0) for f in _charges_du_flux()
                if str(f.get("inclure_resultat_reel")) == "OUI")
    assert somme == pytest.approx(float(par_vision["REEL"]["total_charges"]), abs=0.01)


def test_le_net_proprietaire_ne_depend_pas_des_charges_conciergerie():
    """Une charge conciergerie (non refacturable, non rattachée à un logement) diminue le résultat
    de la conciergerie mais ne doit pas toucher le net propriétaire."""
    lignes = _feuille(NET, "REGLEMENT")
    if not lignes:
        pytest.skip("Net propriétaire absent.")
    colonnes = set(lignes[0])
    assert "net_proprietaire_apres_charge_mois" in colonnes
    # Le net propriétaire se construit sur payout / ménages / commissions / forfait, jamais sur
    # les charges d'exploitation de la conciergerie : aucune colonne de charge n'y figure.
    assert not {"total_charges", "charges_conciergerie"} & colonnes
