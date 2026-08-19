"""Fabrique de datasets Lot10 en SQLite, pour les tests.

Remplace les classeurs `MASTER_CALC_Commissions`/`MASTER_CALC_NetProprietaire`/
`MASTER_CALC_Resultats` que les tests devaient fabriquer : depuis la migration 0044, les lecteurs
applicatifs lisent les tables `lot10_*` du run ACTIF, sans repli Excel.

Insertion en SQL direct (pas via le moteur pandas) : ces tests vérifient les LECTEURS et les
services applicatifs, pas le calcul Lot10 lui-même — celui-ci est couvert par la parité réelle
OLD/NEW (`02_TRAVAIL/lot10_calculer_resultats.py`, commit de migration). Passer par le moteur
imposerait pandas et un jeu de flux complet pour tester un affichage.

Les noms de colonnes sont ceux des onglets legacy (repris tels quels par 0044) : une fixture écrite
pour l'ancien classeur se convertit en changeant la destination, pas les données.
"""
from __future__ import annotations

from typing import Any, Iterable

from app.db.connection import get_db

RUN_TEST = "L10-TEST"


def _inserer(conn, table: str, run_id: str, lignes: Iterable[dict[str, Any]]) -> None:
    """Insère des dicts tels quels : seules les clés fournies sont écrites, les colonnes absentes
    gardent leur défaut SQLite. Une clé inconnue de la table lèverait — c'est voulu, une fixture
    qui nomme une colonne inexistante teste autre chose que ce qu'elle croit."""
    for l in lignes:
        cols = ["run_id"] + list(l.keys())
        valeurs = [run_id] + [l[c] for c in l]
        conn.execute(
            f"INSERT INTO {table} ({', '.join(cols)}) VALUES ({', '.join(['?'] * len(cols))})",
            valeurs)


def seeder(db_path, *, commissions=(), net_vue_mois=(), net_reglement=(), net_exploitation=(),
           resultats=(), commissions_a_controler=(), run_id: str = RUN_TEST,
           statut: str = "SUCCES", actif: bool = True) -> str:
    """Crée un run Lot10 et son dataset. `actif=False` sert à vérifier qu'un run non actif n'est
    jamais servi aux lecteurs (contrepartie de l'écriture atomique côté moteur)."""
    conn = get_db(db_path)
    try:
        conn.execute("UPDATE lot10_runs SET actif = 0 WHERE actif = 1")
        conn.execute(
            "INSERT INTO lot10_runs (run_id, statut, actif, nb_commissions, nb_resultats, "
            "nb_reglements) VALUES (?,?,?,?,?,?)",
            (run_id, statut, 1 if actif else 0, len(list(commissions)), len(list(resultats)),
             len(list(net_reglement))))
        _inserer(conn, "lot10_commissions", run_id, commissions)
        _inserer(conn, "lot10_commissions_a_controler", run_id, commissions_a_controler)
        _inserer(conn, "lot10_net_vue_mois", run_id, net_vue_mois)
        _inserer(conn, "lot10_net_reglement", run_id, net_reglement)
        _inserer(conn, "lot10_net_exploitation", run_id, net_exploitation)
        _inserer(conn, "lot10_resultats", run_id, resultats)
        conn.commit()
    finally:
        conn.close()
    return run_id


def reprendre_master_reel(db_path, chemin_master, *, run_id: str = RUN_TEST) -> str | None:
    """Charge le VRAI `MASTER_CALC_NetProprietaire.xlsx` (lecture seule) dans `lot10_net_reglement`/
    `lot10_net_vue_mois`.

    Réservé aux tests qui portent sur des DONNÉES RÉELLES (12 propriétaires, PROP_0001, mois
    2025-07…) : leur valeur vient précisément de ce qu'ils vérifient des chiffres réels, pas des
    chiffres inventés. Même statut que les reprises de parité Lot9 — outil de test, jamais un
    chemin runtime : l'application, elle, ne lit plus ce classeur.

    Rend None si le classeur est absent (environnement sans données réelles) — l'appelant skippe.
    """
    from pathlib import Path

    import openpyxl

    chemin = Path(chemin_master)
    if not chemin.exists():
        return None

    wb = openpyxl.load_workbook(str(chemin), read_only=True, data_only=True)
    try:
        def _onglet(nom: str) -> list[dict[str, Any]]:
            if nom not in wb.sheetnames:
                return []
            lignes = list(wb[nom].iter_rows(values_only=True))
            if len(lignes) <= 1:
                return []
            entetes = [str(c) if c is not None else "" for c in lignes[0]]
            return [dict(zip(entetes, r)) for r in lignes[1:]
                    if any(v is not None for v in r)]

        reglement = _onglet("REGLEMENT")
        vue_mois = _onglet("VUE_MOIS")
    finally:
        wb.close()

    # Seules les colonnes réellement présentes dans la table cible sont reprises : le classeur peut
    # porter des colonnes de travail que 0044 n'a pas retenues.
    conn = get_db(db_path)
    try:
        cols_regl = {r[1] for r in conn.execute("PRAGMA table_info(lot10_net_reglement)")}
        cols_vue = {r[1] for r in conn.execute("PRAGMA table_info(lot10_net_vue_mois)")}
    finally:
        conn.close()

    return seeder(
        db_path, run_id=run_id,
        net_reglement=[{k: v for k, v in r.items() if k in cols_regl and k != "run_id"}
                       for r in reglement],
        net_vue_mois=[{k: v for k, v in r.items() if k in cols_vue and k != "run_id"}
                      for r in vue_mois])


def vider(db_path) -> None:
    """Supprime tout dataset Lot10 : l'état « aucun run actif » est un cas de test à part entière
    (le lecteur doit l'afficher, jamais lever)."""
    conn = get_db(db_path)
    try:
        for table in ("lot10_commissions", "lot10_commissions_a_controler", "lot10_net_vue_mois",
                      "lot10_net_reglement", "lot10_net_exploitation", "lot10_resultats",
                      "lot10_runs"):
            conn.execute(f"DELETE FROM {table}")
        conn.commit()
    finally:
        conn.close()
