"""Tests APP-3b-0 — Préparation REF_Assoc_Mode sur copie.

Toutes les opérations s'effectuent sur des copies dans tmp_path (external).
Le fichier réel REF_Setup.xlsm n'est jamais modifié.
Basetemp externe obligatoire (conftest.py rejette les basetemp sous APP_ROOT).
"""
from __future__ import annotations

import hashlib
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

import app.config as cfg
from app.services.ref_assoc_mode_prepare_service import (
    CONFIRMATION_EXECUTION,
    EXPECTED_HEADERS,
    INITIAL_ROWS,
    SHEET_NAME,
    TABLE_NAME,
    _sha256_file,
    _vba_snapshot,
    diagnostiquer,
    executer_migration_reelle,
    preparer_sur_copie,
)


# ── Fixtures ───────────────────────────────────────────────────────────────────

def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(str(path), "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


@pytest.fixture()
def ref_source_copy(tmp_path: Path) -> Path:
    """Copie de REF_Setup.xlsm dans tmp_path — source pour les tests."""
    dest = tmp_path / "REF_Setup_test.xlsm"
    shutil.copy2(str(cfg.REF_SETUP), str(dest))
    return dest


@pytest.fixture()
def ref_source_with_bad_sheet(tmp_path: Path) -> Path:
    """Copie de REF_Setup.xlsm avec une REF_Assoc_Mode incoherente."""
    dest = tmp_path / "REF_Setup_bad.xlsm"
    shutil.copy2(str(cfg.REF_SETUP), str(dest))
    wb = openpyxl.load_workbook(str(dest), data_only=False, keep_vba=True)
    ws = wb.create_sheet(SHEET_NAME)
    ws.append(["mauvaise_colonne", "autre_colonne"])
    ws.append(["X", "Y"])
    wb.save(str(dest))
    wb.close()
    return dest


# ── 1. Diagnostic — lecture seule ─────────────────────────────────────────────

def test_diagnostiquer_source_hash_inchange(ref_source_copy: Path, tmp_path: Path):
    hash_avant = _sha256(ref_source_copy)
    diagnostiquer(ref_source_copy)
    hash_apres = _sha256(ref_source_copy)
    assert hash_avant == hash_apres


def test_diagnostiquer_ref_assoc_mode_absent(ref_source_copy: Path):
    diag = diagnostiquer(ref_source_copy)
    assert diag["feuille_presente"] is False
    assert diag["migration_necessaire"] is True


def test_diagnostiquer_status_ok_sur_source_valide(ref_source_copy: Path):
    diag = diagnostiquer(ref_source_copy)
    assert diag["status"] == "OK"
    assert diag["sha256"] is not None


# ── 2. Préparation sur copie — hash source inchangé ───────��───────────────────

def test_preparer_source_hash_inchange(ref_source_copy: Path, tmp_path: Path):
    hash_avant = _sha256(ref_source_copy)
    preparer_sur_copie(ref_source_copy, tmp_path / "out")
    hash_apres = _sha256(ref_source_copy)
    assert hash_avant == hash_apres, "Le fichier source a été modifié pendant la préparation"


def test_preparer_invariant_dans_manifest(ref_source_copy: Path, tmp_path: Path):
    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["source_inchangee"] is True
    assert manifest["source_hash_avant"] == manifest["source_hash_apres"]


# ── 3. Préparation — statut OK ────────────────────────────────────────────────

def test_preparer_status_ok(ref_source_copy: Path, tmp_path: Path):
    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["status"] == "OK", f"Erreurs : {manifest.get('errors', [])}"


def test_preparer_feuille_presente_dans_copie(ref_source_copy: Path, tmp_path: Path):
    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["status"] == "OK"
    copie = Path(manifest["paths"]["ref_setup_prepare"])
    assert copie.exists()
    wb = openpyxl.load_workbook(str(copie), read_only=True, data_only=True, keep_vba=True)
    try:
        assert SHEET_NAME in wb.sheetnames
    finally:
        wb.close()


# ── 4. Structure de la feuille ────────────────────────────────────────────────

def test_preparer_headers_corrects(ref_source_copy: Path, tmp_path: Path):
    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["status"] == "OK"
    copie = Path(manifest["paths"]["ref_setup_prepare"])
    wb = openpyxl.load_workbook(str(copie), read_only=True, data_only=True, keep_vba=True)
    try:
        ws = wb[SHEET_NAME]
        headers = [str(cell.value or "").strip() for cell in list(ws.iter_rows(values_only=False))[0]]
    finally:
        wb.close()
    assert headers == EXPECTED_HEADERS


def test_preparer_lignes_correctes(ref_source_copy: Path, tmp_path: Path):
    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["status"] == "OK"
    copie = Path(manifest["paths"]["ref_setup_prepare"])
    wb = openpyxl.load_workbook(str(copie), read_only=True, data_only=True, keep_vba=True)
    try:
        ws = wb[SHEET_NAME]
        rows_all = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    data = [tuple(str(c).strip() if c is not None else "" for c in r) for r in rows_all[1:]]
    assert len(data) == len(INITIAL_ROWS)
    for actual, expected in zip(data, INITIAL_ROWS):
        assert actual == expected, f"Ligne incorrecte : {actual} != {expected}"


def test_preparer_table_presente(ref_source_copy: Path, tmp_path: Path):
    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["status"] == "OK"
    copie = Path(manifest["paths"]["ref_setup_prepare"])
    wb = openpyxl.load_workbook(str(copie), read_only=False, data_only=False, keep_vba=True)
    try:
        ws = wb[SHEET_NAME]
        table_names = list(ws.tables.keys())
    finally:
        wb.close()
    assert TABLE_NAME in table_names


# ── 5. Unicité ────────────────────────────────────────────────────────────────

def test_preparer_unicite_assoc_mode_id():
    ids = [r[0] for r in INITIAL_ROWS]
    assert len(ids) == len(set(ids)), f"Doublons assoc_mode_id : {ids}"


def test_preparer_unicite_mode_associe():
    pairs = [(r[1], r[2]) for r in INITIAL_ROWS]
    assert len(pairs) == len(set(pairs)), f"Doublons (mode, associe) : {pairs}"


def test_preparer_unicite_assoc_mode():
    abbrevs = [r[3] for r in INITIAL_ROWS]
    assert len(abbrevs) == len(set(abbrevs)), f"Doublons assoc_mode : {abbrevs}"


# ── 6. VBA et package sensible ────────────────────────────────────────────────

def test_preparer_vba_preserve(ref_source_copy: Path, tmp_path: Path):
    snap_avant = _vba_snapshot(ref_source_copy)
    if not snap_avant.get("vba_project_present"):
        pytest.skip("vbaProject.bin absent de la source de test")

    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["status"] == "OK", f"Erreurs préparation : {manifest.get('errors', [])}"
    assert manifest["vba_violations"] == [], f"Violations VBA : {manifest['vba_violations']}"

    copie = Path(manifest["paths"]["ref_setup_prepare"])
    snap_apres = _vba_snapshot(copie)
    assert snap_apres["vba_project_present"], "vbaProject.bin absent dans la copie"
    assert snap_avant["vba_project_sha256"] == snap_apres["vba_project_sha256"], (
        f"vbaProject.bin modifié : "
        f"{snap_avant['vba_project_sha256'][:16]}... → {snap_apres['vba_project_sha256'][:16]}..."
    )


def test_preparer_package_sensible_preserve(ref_source_copy: Path, tmp_path: Path):
    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["status"] == "OK"
    assert manifest["sensitive_violations"] == [], (
        f"Violations package sensible : {manifest['sensitive_violations']}"
    )


# ── 7. Feuilles originales préservées ─────────────────────────────────────────

def test_preparer_feuilles_originales_preservees(ref_source_copy: Path, tmp_path: Path):
    wb_src = openpyxl.load_workbook(str(ref_source_copy), read_only=True, data_only=True, keep_vba=True)
    sheets_avant = list(wb_src.sheetnames)
    wb_src.close()

    manifest = preparer_sur_copie(ref_source_copy, tmp_path / "out")
    assert manifest["status"] == "OK"

    copie = Path(manifest["paths"]["ref_setup_prepare"])
    wb_apres = openpyxl.load_workbook(str(copie), read_only=True, data_only=True, keep_vba=True)
    sheets_apres = list(wb_apres.sheetnames)
    wb_apres.close()

    for s in sheets_avant:
        assert s in sheets_apres, f"Feuille supprimée : {s}"
    assert SHEET_NAME in sheets_apres


# ── 8. Idempotence ────────────────────────────────────────────────────────────

def test_preparer_idempotent(ref_source_copy: Path, tmp_path: Path):
    out1 = tmp_path / "pass1"
    manifest1 = preparer_sur_copie(ref_source_copy, out1)
    assert manifest1["status"] == "OK"
    assert not manifest1.get("feuille_deja_presente")

    # Deuxième appel sur la copie déjà préparée
    copie_preparee = Path(manifest1["paths"]["ref_setup_prepare"])
    out2 = tmp_path / "pass2"
    manifest2 = preparer_sur_copie(copie_preparee, out2)
    assert manifest2["status"] == "OK"
    assert manifest2.get("feuille_deja_presente") is True


# ── 9. Refus feuille incoherente ──────────────────────────────────────────────

def test_preparer_refuse_feuille_incoherente(ref_source_with_bad_sheet: Path, tmp_path: Path):
    manifest = preparer_sur_copie(ref_source_with_bad_sheet, tmp_path / "out")
    assert manifest["status"] == "FEUILLE_INCOHERENTE_REFUS_ECRASEMENT", (
        f"Devrait refuser, got : {manifest['status']}"
    )
    assert manifest.get("coherence_violations"), "Devrait lister les violations"


# ── 10. Exécution réelle — flags et confirmation ──────────────────────────────

def test_flag_ref_assoc_mode_est_false():
    assert cfg.REF_ASSOC_MODE_REAL_WRITE_ENABLED is False


def test_execute_sans_flag_retourne_non_active():
    result = executer_migration_reelle(confirmation=CONFIRMATION_EXECUTION)
    assert result["status"] == "MIGRATION_REELLE_NON_ACTIVE"


def test_execute_mauvaise_confirmation_retourne_non_active_ou_refuse():
    result = executer_migration_reelle(confirmation="MAUVAIS_MOT_DE_PASSE")
    # Avec flag False, retourne NON_ACTIVE avant même d'évaluer la confirmation
    assert result["status"] in ("MIGRATION_REELLE_NON_ACTIVE", "REFUSE")


def test_execute_source_ref_inchangee_apres_refus_flag():
    hash_avant = _sha256_file(cfg.REF_SETUP)
    executer_migration_reelle(confirmation=CONFIRMATION_EXECUTION)
    hash_apres = _sha256_file(cfg.REF_SETUP)
    assert hash_avant == hash_apres, "REF_Setup.xlsm modifié alors que flag=False"
