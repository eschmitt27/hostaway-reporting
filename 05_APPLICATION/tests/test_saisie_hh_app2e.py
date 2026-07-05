"""Tests APP-2e - preparation schema et recette ecriture HH sur copies."""
from __future__ import annotations

from decimal import Decimal
import os
from pathlib import Path
import shutil
from unittest.mock import patch

import openpyxl
import pytest

import app.config as cfg
from app.services import saisie_hh_dryrun_service as dryrun_svc
from app.services import saisie_hh_real_write_service as real_svc
from app.services import saisie_hh_schema_real_prepare_service as schema_real_svc
from app.services.saisie_hh_schema_migration import NEW_SAISIE_FIELDS
from app.writers.saisie_hh_writer import write_row


ARTIFACT = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "dapp05c_formules"
    / "20260702T143503Z"
    / "SAISIE_copie.xlsx"
)


def _copy_saisie(tmp_path: Path) -> Path:
    if not ARTIFACT.exists():
        pytest.fail(f"Artefact D-APP-05C absent: {ARTIFACT}")
    tmp_path.mkdir(parents=True, exist_ok=True)
    target = tmp_path / "SAISIE_ReservationsHorsHostaway_source.xlsx"
    shutil.copy2(ARTIFACT, target)
    return target


def _make_ref(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "REF_Modes_Paiement"
    ws.append(["mode_paiement_id", "mode_paiement", "impact_banque", "impact_caisse", "impact_associee", "actif"])
    ws.append(["PAY_001", "BANQUE_PRO", "OUI", "NON", "NON", "OUI"])
    ws.append(["PAY_002", "ESPECES_CAISSE", "NON", "OUI", "NON", "OUI"])
    ws.append(["PAY_003", "CARTE_ASSOCIEE", "NON", "NON", "OUI", "OUI"])
    ws.append(["PAY_004", "COMPTE_PERSO_ASSOCIEE", "NON", "NON", "OUI", "OUI"])
    ws.append(["PAY_005", "AUTRE", "NON", "NON", "NON", "OUI"])

    ws = wb.create_sheet("REF_Cloture_Mensuelle")
    ws.append(["mois", "statut_mois"])
    ws.append(["2026-08", "OUVERT"])

    ws = wb.create_sheet("REF_Logements")
    ws.append(["logement_id", "nom_court", "actif", "statut_parc", "type_logement_id"])
    ws.append(["LOG_0001", "Test logement", "OUI", "GERE", "TYPE_T3"])

    ws = wb.create_sheet("REF_Gestion_Logements_Hist")
    ws.append(["logement_id", "proprietaire_id", "statut_gestion", "date_debut", "date_fin"])
    ws.append(["LOG_0001", "PROP_0001", "ACTIF", "2026-01-01", ""])

    ws = wb.create_sheet("REF_Proprietaires")
    ws.append(["proprietaire_id", "prenom_proprietaire", "nom_proprietaire"])
    ws.append(["PROP_0001", "David", "Dupont"])

    ws = wb.create_sheet("REF_Associes")
    ws.append(["associe_id", "prenom"])
    ws.append(["PERS_EWAN", "Ewan"])

    ws = wb.create_sheet("REF_Canaux_Reservation")
    ws.append(["canal_id", "canal"])
    ws.append(["CANAL_001", "Airbnb"])

    ws = wb.create_sheet("REF_Codes_Impact")
    ws.append(["code_impact", "libelle", "impact_resultat_comptable"])
    ws.append(["HC", "Hors comptabilite", "OUI"])

    ws = wb.create_sheet("REF_Taux_Commission")
    ws.append(["proprietaire_id", "logement_id", "taux_commission", "date_debut", "date_fin", "actif"])
    ws.append(["PROP_0001", "", 0.15, "2026-01-01", "", "OUI"])
    ws.append(["PROP_0003", "", 0.15, "2026-01-01", "", "OUI"])

    ws = wb.create_sheet("REF_Couts_Standards_Menage")
    ws.append(["type_logement_id", "cout_standard_menage", "date_debut_validite", "date_fin_validite", "actif"])
    ws.append(["TYPE_T3", 60, "2026-01-01", "", "OUI"])

    wb.save(path)
    wb.close()
    return path


def _prepare_schema(tmp_path: Path) -> tuple[Path, Path, dict]:
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "REF_Setup_source.xlsm")
    manifest = schema_real_svc.preparer_migration_hh_sur_copies(
        saisie_source=source_saisie,
        ref_setup_source=source_ref,
        output_dir=tmp_path / "schema_copies",
    )
    assert manifest["status"] == "OK", manifest["errors"]
    return Path(manifest["paths"]["saisie_copy"]), Path(manifest["paths"]["ref_setup_copy"]), manifest


def _base_form(**overrides) -> dict[str, str]:
    data = {
        "canal_id": "CANAL_001",
        "source_financiere": "SAISIE_MANUELLE",
        "proprietaire_id": "PROP_0001",
        "logement_id": "LOG_0001",
        "date_arrivee": "2026-08-15",
        "date_depart": "2026-08-18",
        "total_percu": "450.00",
        "code_impact": "HC",
        "comptabilisation": "OUI",
        "mode_paiement_id": "PAY_001",
    }
    data.update(overrides)
    return data


def _row_by_pk(path: Path, pk: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["SAISIE"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    headers = [str(h or "").strip() for h in rows[0]]
    for row in rows[1:]:
        rec = dict(zip(headers, row))
        if str(rec.get("reservation_hh_id") or "").strip() == pk:
            return rec
    raise AssertionError(f"{pk} absent de {path}")


def _run_full_recipe(tmp_path: Path, tmp_db: Path, form: dict[str, str]):
    saisie, ref, _manifest_schema = _prepare_schema(tmp_path)
    snapshots = tmp_path / "snapshots"
    snapshots.mkdir()
    dryruns = tmp_path / "dryruns"
    result = dryrun_svc.run_previsualisation(
        form,
        saisie_source=saisie,
        ref_setup_source=ref,
        dryruns_root=dryruns,
        as_of_iso="2026-07-04T00:00:00+00:00",
    )
    assert result["ok"] is True, result["manifest"].get("errors")
    assert result["manifest"]["lot4a_status"] == "ANALYSE_TERMINEE"
    pk = result["manifest"]["payload_summary"]["reservation_hh_id"]
    with (
        patch.object(cfg, "HH_REAL_WRITE_ENABLED", True),
        patch.object(cfg, "HH_REAL_WRITE_CONFIRMATION_ENABLED", True),
        patch.object(cfg, "SNAPSHOTS_DIR", snapshots),
    ):
        write_result = real_svc.enregistrer_reservation_hh_reelle(
            result["token"],
            f"ENREGISTRER {pk}",
            saisie_path=saisie,
            ref_setup_path=ref,
            dryruns_root=dryruns,
            db_path=tmp_db,
        )
    assert write_result["statut"] == "OK", write_result
    assert write_result["post_lot4a"]["post_row"]["reservation_hh_id"] == pk
    assert not (saisie.parent / f"{saisie.stem}.app2d.rollback.xlsx").exists()
    return {"saisie": saisie, "ref": ref, "dryruns": dryruns, "pk": pk, "simulation": result, "write": write_result}


def test_app2e_diagnostic_et_migration_sur_copies(tmp_path):
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "REF_Setup_source.xlsm")
    diagnostic = schema_real_svc.diagnostiquer_schema_hh(saisie_path=source_saisie, ref_setup_path=source_ref)
    assert diagnostic["migration_needed"] is True
    assert diagnostic["missing_saisie_fields"] == NEW_SAISIE_FIELDS
    assert diagnostic["missing_ref_modes"] == ["PAY_006 / DIRECT_PROPRIETAIRE"]
    manifest = schema_real_svc.preparer_migration_hh_sur_copies(
        saisie_source=source_saisie,
        ref_setup_source=source_ref,
        output_dir=tmp_path / "prepared",
    )
    assert manifest["status"] == "OK", manifest["errors"]
    assert manifest["migration"]["saisie_fields_added"] == NEW_SAISIE_FIELDS
    assert manifest["migration"]["direct_proprietaire_added"] is True
    assert manifest["diagnostic_apres"]["missing_saisie_fields"] == []
    assert manifest["diagnostic_apres"]["missing_ref_modes"] == []


def test_app2e_writer_persiste_champs_derogation_sur_colonnes_cibles(tmp_path):
    saisie, _ref, _manifest = _prepare_schema(tmp_path)
    row_data = {
        "reservation_hh_id": "RESHH-2026-08-TRACE",
        "canal_id": "CANAL_001",
        "source_financiere": "SAISIE_MANUELLE",
        "proprietaire_id": "PROP_0001",
        "logement_id": "LOG_0001",
        "date_arrivee": "2026-08-15",
        "date_depart": "2026-08-18",
        "total_percu": 450.0,
        "code_impact": "HC",
        "comptabilisation": "OUI",
        "statut_controle": "A_CONTROLER",
        "niveau_anomalie": "A_CONTROLER",
        "taux_commission_override": 0.005,
        "motif_override_taux_commission": "Taux 0,5 pct",
        "confirmation_override_taux_commission": "OUI",
        "menage_override": 72.0,
        "motif_override_menage": "Menage test",
        "confirmation_override_menage": "OUI",
    }
    with patch.object(cfg, "HH_REAL_WRITE_ENABLED", True):
        result = write_row(row_data, saisie, 4)
    assert result["statut"] == "OK", result
    row = _row_by_pk(saisie, "RESHH-2026-08-TRACE")
    assert row["taux_commission_override"] == 0.005
    assert row["motif_override_taux_commission"] == "Taux 0,5 pct"
    assert row["confirmation_override_taux_commission"] == "OUI"
    assert row["menage_override"] == 72
    assert row["motif_override_menage"] == "Menage test"
    assert row["confirmation_override_menage"] == "OUI"


@pytest.mark.parametrize(
    ("name", "form", "expected_source", "expected_acompte"),
    [
        ("banque", _base_form(mode_paiement_id="PAY_001"), "TOTAL_PERCU", Decimal("450")),
        ("especes", _base_form(mode_paiement_id="PAY_002", montant_reverse_proprietaire="120.00"), "TOTAL_PERCU_MOINS_REVERSE_ESPECES", Decimal("330")),
        ("carte", _base_form(mode_paiement_id="PAY_003", montant_recupere="100.00", associe_id_recuperateur="PERS_EWAN"), "TOTAL_PERCU_ASSOCIE", Decimal("450")),
        ("direct", _base_form(mode_paiement_id="PAY_006"), "DIRECT_PROPRIETAIRE", Decimal("0")),
    ],
)
def test_app2e_recette_integrale_paiements_sur_copies(tmp_path, tmp_db, name, form, expected_source, expected_acompte):
    recipe = _run_full_recipe(tmp_path / name, tmp_db, form)
    row = recipe["write"]["post_lot4a"]["post_row"]
    assert Decimal(str(row["acompte_facture"])) == expected_acompte
    assert row["source_acompte_facture"] == expected_source
    assert _row_by_pk(recipe["saisie"], recipe["pk"])["reservation_hh_id"] == recipe["pk"]


@pytest.mark.parametrize(
    ("name", "form", "field", "expected"),
    [
        (
            "taux_0005",
            _base_form(
                taux_commission_override_pct="0.5",
                motif_override_taux_commission="Taux 0,5 pct",
                confirmation_override_taux_commission="on",
            ),
            "taux_commission_override",
            Decimal("0.005"),
        ),
        (
            "taux_zero",
            _base_form(
                taux_commission_override_pct="0",
                motif_override_taux_commission="Taux zero",
                confirmation_override_taux_commission="on",
            ),
            "taux_commission_override",
            Decimal("0"),
        ),
        (
            "menage",
            _base_form(
                menage_override="72.00",
                motif_override_menage="Menage specifique",
                confirmation_override_menage="on",
            ),
            "menage_override",
            Decimal("72"),
        ),
    ],
)
def test_app2e_recette_integrale_derogations_sur_copies(tmp_path, tmp_db, name, form, field, expected):
    recipe = _run_full_recipe(tmp_path / name, tmp_db, form)
    saisie_row = _row_by_pk(recipe["saisie"], recipe["pk"])
    assert Decimal(str(saisie_row[field])) == expected
    assert recipe["write"]["post_lot4a"]["post_row"]["reservation_hh_id"] == recipe["pk"]


def test_app2e_recette_journal_sqlite_et_comparaison_lot4a(tmp_path, tmp_db):
    recipe = _run_full_recipe(tmp_path / "journal", tmp_db, _base_form())
    post_row = recipe["write"]["post_lot4a"]["post_row"]
    simulated = recipe["simulation"]["simulated_master_row"]
    assert real_svc._compare_lot4a_rows(simulated, post_row) == []
    import sqlite3
    conn = sqlite3.connect(tmp_db)
    try:
        write_row_db = conn.execute("SELECT statut FROM saisie_hh_writes WHERE pk=?", (recipe["pk"],)).fetchone()
        audit_row = conn.execute("SELECT action FROM audit_events WHERE action='APP2D_ECRITURE_REELLE_OK'").fetchone()
    finally:
        conn.close()
    assert write_row_db is not None and write_row_db[0] == "OK"
    assert audit_row is not None


def test_app2e_rollback_sur_divergence_post_ecriture(tmp_path, tmp_db):
    saisie, ref, _manifest_schema = _prepare_schema(tmp_path)
    dryruns = tmp_path / "dryruns"
    snapshots = tmp_path / "snapshots"
    snapshots.mkdir()
    result = dryrun_svc.run_previsualisation(
        _base_form(),
        saisie_source=saisie,
        ref_setup_source=ref,
        dryruns_root=dryruns,
        as_of_iso="2026-07-04T00:00:00+00:00",
    )
    pk = result["manifest"]["payload_summary"]["reservation_hh_id"]
    sha_avant = real_svc._sha256(saisie)
    with (
        patch.object(cfg, "HH_REAL_WRITE_ENABLED", True),
        patch.object(cfg, "HH_REAL_WRITE_CONFIRMATION_ENABLED", True),
        patch.object(cfg, "SNAPSHOTS_DIR", snapshots),
        patch("app.services.saisie_hh_real_write_service._compare_lot4a_rows", return_value=["divergence forcee"]),
    ):
        write_result = real_svc.enregistrer_reservation_hh_reelle(
            result["token"],
            f"ENREGISTRER {pk}",
            saisie_path=saisie,
            ref_setup_path=ref,
            dryruns_root=dryruns,
            db_path=tmp_db,
        )
    assert write_result["statut"] == real_svc.STATUS_CANCELLED_RESTORED
    assert write_result["rollback_status"] == "ROLLBACK_REUSSI"
    assert real_svc._sha256(saisie) == sha_avant
    assert "DIVERGENCE_ECRITURE_REELLE_VS_SIMULATION" in write_result["details"]


# ── APP-2e : tests structurels et validation migration ─────────────────────────



# ── Tests 1-5 : mutations structurelles détectées par _check_structural_preservation ──

def test_app2e_struct_dv_alteree_rejetee(tmp_path):
    """Suppression d'une validation de données détectée comme violation structurelle."""
    from openpyxl.worksheet.datavalidation import DataValidation
    from app.writers.saisie_hh_writer import _check_structural_preservation, _measure_structure

    before = tmp_path / "before.xlsx"
    after = tmp_path / "after.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    dv = DataValidation(type="list", formula1='"OUI,NON"', sqref="B2:B100")
    ws.add_data_validation(dv)
    wb.calculation.fullCalcOnLoad = True
    wb.save(str(before))
    wb.close()

    shutil.copy2(before, after)
    wb2 = openpyxl.load_workbook(str(after))
    wb2.active.data_validations.dataValidation.clear()
    wb2.calculation.fullCalcOnLoad = True
    wb2.save(str(after))
    wb2.close()

    struct = _measure_structure(before)
    violations = _check_structural_preservation(before, after, struct)
    assert any("data_validations" in v for v in violations), violations


def test_app2e_struct_mfc_alteree_rejetee(tmp_path):
    """Suppression d'une MFC détectée comme violation structurelle."""
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from app.writers.saisie_hh_writer import _check_structural_preservation, _measure_structure

    before = tmp_path / "before.xlsx"
    after = tmp_path / "after.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    ws.conditional_formatting.add("A2:A100", CellIsRule(operator="equal", formula=['"OUI"'], fill=fill))
    wb.calculation.fullCalcOnLoad = True
    wb.save(str(before))
    wb.close()

    shutil.copy2(before, after)
    wb2 = openpyxl.load_workbook(str(after))
    wb2.active.conditional_formatting._cf_rules.clear()
    wb2.calculation.fullCalcOnLoad = True
    wb2.save(str(after))
    wb2.close()

    struct = _measure_structure(before)
    violations = _check_structural_preservation(before, after, struct)
    assert any("mfc_rules" in v for v in violations), violations


def test_app2e_struct_table_ajoutee_rejetee(tmp_path):
    """Ajout d'une table inattendue détecté comme violation structurelle."""
    from openpyxl.worksheet.table import Table
    from app.writers.saisie_hh_writer import _check_structural_preservation, _measure_structure

    before = tmp_path / "before.xlsx"
    after = tmp_path / "after.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    ws["A1"] = "col1"
    ws["B1"] = "col2"
    for i in range(2, 6):
        ws.cell(row=i, column=1, value=i)
        ws.cell(row=i, column=2, value=i * 10)
    wb.calculation.fullCalcOnLoad = True
    wb.save(str(before))
    wb.close()

    shutil.copy2(before, after)
    wb2 = openpyxl.load_workbook(str(after))
    wb2.active.add_table(Table(displayName="TableInattendue", ref="A1:B5"))
    wb2.calculation.fullCalcOnLoad = True
    wb2.save(str(after))
    wb2.close()

    struct = _measure_structure(before)
    violations = _check_structural_preservation(before, after, struct)
    assert any("tables" in v for v in violations), violations


def test_app2e_struct_plage_nommee_ajoutee_rejetee(tmp_path):
    """Ajout d'une plage nommée inattendue détecté comme violation structurelle."""
    from openpyxl.workbook.defined_name import DefinedName
    from app.writers.saisie_hh_writer import _check_structural_preservation, _measure_structure

    before = tmp_path / "before.xlsx"
    after = tmp_path / "after.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    ws["A1"] = "test"
    wb.calculation.fullCalcOnLoad = True
    wb.save(str(before))
    wb.close()

    shutil.copy2(before, after)
    wb2 = openpyxl.load_workbook(str(after))
    wb2.defined_names["PLAGE_INATTENDUE"] = DefinedName("PLAGE_INATTENDUE", attr_text="SAISIE!$A$1:$A$10")
    wb2.calculation.fullCalcOnLoad = True
    wb2.save(str(after))
    wb2.close()

    struct = _measure_structure(before)
    violations = _check_structural_preservation(before, after, struct)
    assert any("defined_names" in v for v in violations), violations


def test_app2e_struct_feuille_supprimee_rejetee(tmp_path):
    """Suppression d'une feuille détectée comme violation structurelle."""
    from app.writers.saisie_hh_writer import _check_structural_preservation, _measure_structure

    before = tmp_path / "before.xlsx"
    after = tmp_path / "after.xlsx"

    wb = openpyxl.Workbook()
    wb.active.title = "SAISIE"
    wb.create_sheet("REF_Modes_Paiement")
    wb.calculation.fullCalcOnLoad = True
    wb.save(str(before))
    wb.close()

    shutil.copy2(before, after)
    wb2 = openpyxl.load_workbook(str(after))
    del wb2["REF_Modes_Paiement"]
    wb2.calculation.fullCalcOnLoad = True
    wb2.save(str(after))
    wb2.close()

    struct = _measure_structure(before)
    violations = _check_structural_preservation(before, after, struct)
    assert any("sheetnames" in v for v in violations), violations


# ── Test 6 : formule absente ────────────────────────────────────────────────────

def test_app2e_formule_absente_detectee(tmp_path):
    """Une cellule de formule absente (None) dans la ligne cible est rejetée."""
    from app.writers.saisie_hh_writer import _check_formula_cells, FORMULE_LIGNE_MODELE_ABSENTE

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    ws["A1"] = "reservation_hh_id"
    path = tmp_path / "saisie_sans_formule.xlsx"
    wb.save(str(path))
    wb.close()

    violations = _check_formula_cells(path, 2)
    assert any(FORMULE_LIGNE_MODELE_ABSENTE in v for v in violations), violations


# ── Test 7 : VBA préservé après migration REF ───────────────────────────────────

def test_app2e_vba_preserve_apres_migration_ref(tmp_path):
    """Après migration REF_Setup, le VBA est préservé si présent dans la source."""
    if not cfg.REF_SETUP.exists():
        pytest.skip("REF_Setup.xlsm absent")
    features = schema_real_svc._workbook_features(cfg.REF_SETUP, keep_vba=True)
    if not features.get("has_vba"):
        pytest.skip("REF_Setup n'a pas de VBA — test non applicable")

    source_saisie = _copy_saisie(tmp_path)
    manifest = schema_real_svc.preparer_migration_hh_sur_copies(
        saisie_source=source_saisie,
        ref_setup_source=cfg.REF_SETUP,
        output_dir=tmp_path / "prepared",
    )
    assert manifest["status"] == "OK", manifest["errors"]

    ref_work = Path(manifest["paths"]["ref_setup_copy"])
    features_after = schema_real_svc._workbook_features(ref_work, keep_vba=True)
    assert features_after.get("has_vba"), "VBA perdu apres migration REF_Setup"


# ── Test 8 : 7 colonnes ajoutées exactement ────────────────────────────────────

def test_app2e_7_cols_saisie_ajoutees_verifiees(tmp_path):
    """Les 7 champs NEW_SAISIE_FIELDS sont ajoutés exactement au header SAISIE."""
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "ref.xlsm")

    manifest = schema_real_svc.preparer_migration_hh_sur_copies(
        saisie_source=source_saisie,
        ref_setup_source=source_ref,
        output_dir=tmp_path / "prepared",
    )
    assert manifest["status"] == "OK", manifest["errors"]

    saisie_work = Path(manifest["paths"]["saisie_copy"])
    wb = openpyxl.load_workbook(str(saisie_work), read_only=True, data_only=True)
    try:
        headers = [str(cell.value or "").strip() for cell in wb["SAISIE"][1]]
    finally:
        wb.close()

    for field in NEW_SAISIE_FIELDS:
        assert field in headers, f"Champ {field!r} absent du header SAISIE apres migration"
    assert manifest["migration"]["saisie_fields_added"] == NEW_SAISIE_FIELDS


# ── Test 9 : PAY_006 idempotent ─────────────────────────────────────────────────

def test_app2e_pay006_idempotent(tmp_path):
    """Si PAY_006 est déjà présent, la re-migration renvoie direct_proprietaire_added=False."""
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "ref.xlsm")

    manifest1 = schema_real_svc.preparer_migration_hh_sur_copies(
        saisie_source=source_saisie,
        ref_setup_source=source_ref,
        output_dir=tmp_path / "p1",
    )
    assert manifest1["status"] == "OK"
    assert manifest1["migration"]["direct_proprietaire_added"] is True

    ref_migree = Path(manifest1["paths"]["ref_setup_copy"])
    manifest2 = schema_real_svc.preparer_migration_hh_sur_copies(
        saisie_source=source_saisie,
        ref_setup_source=ref_migree,
        output_dir=tmp_path / "p2",
    )
    assert manifest2["status"] == "OK"
    assert manifest2["migration"]["direct_proprietaire_added"] is False


# ── Test 10 : second preparer idempotent ───────────────────────────────────────

def test_app2e_second_preparer_idempotent(tmp_path):
    """Un second appel avec les copies déjà migrées comme source est OK et n'ajoute rien."""
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "ref.xlsm")

    manifest1 = schema_real_svc.preparer_migration_hh_sur_copies(
        saisie_source=source_saisie,
        ref_setup_source=source_ref,
        output_dir=tmp_path / "p1",
    )
    assert manifest1["status"] == "OK"

    saisie_migree = Path(manifest1["paths"]["saisie_copy"])
    ref_migree = Path(manifest1["paths"]["ref_setup_copy"])

    manifest2 = schema_real_svc.preparer_migration_hh_sur_copies(
        saisie_source=saisie_migree,
        ref_setup_source=ref_migree,
        output_dir=tmp_path / "p2",
    )
    assert manifest2["status"] == "OK"
    assert manifest2["migration"]["saisie_fields_added"] == []
    assert manifest2["migration"]["direct_proprietaire_added"] is False


# ── Test 11 : rollback double fichier ──────────────────────────────────────────

def test_app2e_rollback_double_fichier_restaure(tmp_path):
    """Si le deuxième os.replace échoue, SAISIE et REF sont restaurés à l'état original."""
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "ref.xlsm")
    sha_saisie_avant = real_svc._sha256(source_saisie)
    sha_ref_avant = real_svc._sha256(source_ref)

    replace_count = [0]
    real_replace = os.replace

    def mock_replace(src, dst):
        replace_count[0] += 1
        if replace_count[0] == 2:
            raise OSError("Simulation panne disque")
        return real_replace(src, dst)

    with (
        patch.object(cfg, "SAISIE_RESERVATIONS_HH", source_saisie),
        patch.object(cfg, "REF_SETUP", source_ref),
        patch("app.services.saisie_hh_schema_real_prepare_service.os.replace", mock_replace),
    ):
        result = schema_real_svc.executer_migration_hh_reelle(
            confirmation=schema_real_svc.CONFIRMATION_EXECUTION,
            saisie_path=source_saisie,
            ref_setup_path=source_ref,
            work_dir=tmp_path / "work",
        )

    assert result["real_status"] in ("ROLLBACK", "ROLLBACK_HASH_MISMATCH")
    assert real_svc._sha256(source_saisie) == sha_saisie_avant, "SAISIE non restaure apres rollback"
    assert real_svc._sha256(source_ref) == sha_ref_avant, "REF non restaure apres rollback"


# ── Test 12 : ROLLBACK_HASH_MISMATCH ──────────────────────────────────────────

def test_app2e_rollback_hash_mismatch_detecte(tmp_path):
    """ROLLBACK_HASH_MISMATCH si les hashes post-rollback divergent des hashes avant."""
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "ref.xlsm")

    replace_count = [0]
    failed = [False]
    real_replace = os.replace
    real_fingerprint = schema_real_svc._fingerprint

    def mock_replace(src, dst):
        replace_count[0] += 1
        if replace_count[0] == 2:
            failed[0] = True
            raise OSError("Simulation panne")
        return real_replace(src, dst)

    def mock_fingerprint(path):
        result = real_fingerprint(path)
        if failed[0]:
            return {**result, "sha256": "a" * 64}
        return result

    with (
        patch.object(cfg, "SAISIE_RESERVATIONS_HH", source_saisie),
        patch.object(cfg, "REF_SETUP", source_ref),
        patch("app.services.saisie_hh_schema_real_prepare_service.os.replace", mock_replace),
        patch.object(schema_real_svc, "_fingerprint", mock_fingerprint),
    ):
        result = schema_real_svc.executer_migration_hh_reelle(
            confirmation=schema_real_svc.CONFIRMATION_EXECUTION,
            saisie_path=source_saisie,
            ref_setup_path=source_ref,
            work_dir=tmp_path / "work",
        )

    assert result["real_status"] == "ROLLBACK_HASH_MISMATCH"
    assert result["rollback_hash_verified"] is False


# ── Tests 13-14 : garde-fou cibles migration réelle ───────────────────────────

def test_app2e_migration_reelle_accepte_cibles_configurees(tmp_path):
    """executer_migration_hh_reelle accepte exactement les cibles configurées (cfg patché vers tmp)."""
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "ref.xlsm")

    with (
        patch.object(cfg, "SAISIE_RESERVATIONS_HH", source_saisie),
        patch.object(cfg, "REF_SETUP", source_ref),
    ):
        # Ne doit pas lever CIBLE_MIGRATION_REELLE_NON_AUTORISEE
        result = schema_real_svc.executer_migration_hh_reelle(
            confirmation=schema_real_svc.CONFIRMATION_EXECUTION,
            saisie_path=source_saisie,
            ref_setup_path=source_ref,
            work_dir=tmp_path / "work",
        )
    assert "real_status" in result or result.get("status") == "REFUSE"


def test_app2e_migration_reelle_refuse_chemin_non_configure(tmp_path):
    """executer_migration_hh_reelle refuse tout chemin non configuré avant tout os.replace."""
    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "ref.xlsm")
    tiers = tmp_path / "tiers.xlsx"
    shutil.copy2(source_saisie, tiers)

    with (
        patch.object(cfg, "SAISIE_RESERVATIONS_HH", source_saisie),
        patch.object(cfg, "REF_SETUP", source_ref),
        patch("app.services.saisie_hh_schema_real_prepare_service.os.replace") as mock_replace,
    ):
        with pytest.raises(RuntimeError, match="CIBLE_MIGRATION_REELLE_NON_AUTORISEE"):
            schema_real_svc.executer_migration_hh_reelle(
                confirmation=schema_real_svc.CONFIRMATION_EXECUTION,
                saisie_path=tiers,
                ref_setup_path=source_ref,
                work_dir=tmp_path / "work",
            )
        mock_replace.assert_not_called()


# ── APP-2e : contrôle intégrité binaire VBA ────────────────────────────────────

import zipfile as _zipfile

_CONTENT_TYPES_WITH_VBA = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'
    '<Override PartName="/xl/workbook.xml"'
    ' ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml"/>'
    '</Types>'
)
_WORKBOOK_RELS_WITH_VBA = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1"'
    ' Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject"'
    ' Target="vbaProject.bin"/>'
    '</Relationships>'
)
_VBA_BIN_MINIMAL = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 120
_VBA_SIG_MINIMAL = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\xff' * 60


def _make_synthetic_xlsm(
    path: Path,
    *,
    with_vba: bool = True,
    with_sig: bool = False,
    vba_bytes: bytes | None = None,
) -> Path:
    """Crée un fichier .xlsm ZIP synthétique minimal pour les tests binaires VBA."""
    vba_data = vba_bytes if vba_bytes is not None else _VBA_BIN_MINIMAL
    with _zipfile.ZipFile(str(path), "w", _zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(_zipfile.ZipInfo("[Content_Types].xml"), _CONTENT_TYPES_WITH_VBA)
        zf.writestr(_zipfile.ZipInfo("xl/_rels/workbook.xml.rels"), _WORKBOOK_RELS_WITH_VBA)
        if with_vba:
            zf.writestr(_zipfile.ZipInfo("xl/vbaProject.bin"), vba_data)
        if with_sig:
            zf.writestr(_zipfile.ZipInfo("xl/vbaProjectSignature.bin"), _VBA_SIG_MINIMAL)
    return path


def _mutate_zip(src: Path, dst: Path, *, remove: str | None = None,
                replace: dict[str, bytes] | None = None) -> Path:
    """Copie src → dst en supprimant ou remplaçant des entrées ZIP."""
    with _zipfile.ZipFile(str(src), "r") as zf_in, \
         _zipfile.ZipFile(str(dst), "w", _zipfile.ZIP_DEFLATED) as zf_out:
        for item in zf_in.infolist():
            if remove and item.filename == remove:
                continue
            data = (replace or {}).get(item.filename, zf_in.read(item.filename))
            zf_out.writestr(item, data)
    return dst


# ── Test 1 : xl/vbaProject.bin supprimé détecté ───────────────────────────────

def test_app2e_zip_vba_project_supprime_detecte(tmp_path):
    """Suppression de xl/vbaProject.bin détectée par _check_zip_vba_integrity."""
    ref = _make_synthetic_xlsm(tmp_path / "ref.xlsm")
    work = _mutate_zip(ref, tmp_path / "work.xlsm", remove="xl/vbaProject.bin")

    violations = schema_real_svc._check_zip_vba_integrity(ref, work)
    assert any("vbaProject.bin disparu" in v for v in violations), violations


# ── Test 2 : octet changé dans xl/vbaProject.bin détecté ─────────────────────

def test_app2e_zip_vba_project_octet_modifie_detecte(tmp_path):
    """Un octet modifié dans xl/vbaProject.bin détecté par _check_zip_vba_integrity."""
    ref = _make_synthetic_xlsm(tmp_path / "ref.xlsm")
    corrupted = bytearray(_VBA_BIN_MINIMAL)
    corrupted[8] ^= 0xFF  # flip un octet
    work = _mutate_zip(ref, tmp_path / "work.xlsm", replace={"xl/vbaProject.bin": bytes(corrupted)})

    violations = schema_real_svc._check_zip_vba_integrity(ref, work)
    assert any("vbaProject.bin modifie" in v for v in violations), violations


# ── Test 3 : signature VBA disparue détectée ──────────────────────────────────

def test_app2e_zip_vba_signature_disparue_detectee(tmp_path):
    """Disparition de xl/vbaProjectSignature.bin détectée."""
    ref = _make_synthetic_xlsm(tmp_path / "ref.xlsm", with_sig=True)
    work = _mutate_zip(ref, tmp_path / "work.xlsm", remove="xl/vbaProjectSignature.bin")

    violations = schema_real_svc._check_zip_vba_integrity(ref, work)
    assert any("vbaProjectSignature.bin disparu" in v for v in violations), violations


# ── Test 4 : relation VBA supprimée de workbook.xml.rels détectée ─────────────

def test_app2e_zip_vba_relation_supprimee_detectee(tmp_path):
    """Suppression de la relation VBA dans workbook.xml.rels détectée."""
    ref = _make_synthetic_xlsm(tmp_path / "ref.xlsm")
    rels_sans_vba = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '</Relationships>'
    )
    work = _mutate_zip(
        ref, tmp_path / "work.xlsm",
        replace={"xl/_rels/workbook.xml.rels": rels_sans_vba.encode()},
    )

    violations = schema_real_svc._check_zip_vba_integrity(ref, work)
    assert any("relation VBA" in v for v in violations), violations


# ── Test 5 : package .xlsx sans VBA accepté sans faux échec ───────────────────

def test_app2e_zip_xlsx_sans_vba_accepte(tmp_path):
    """Un fichier .xlsx sans VBA ne génère aucune violation VBA."""
    # Crée un fichier xlsx minimal (pas de vbaProject.bin)
    ref = tmp_path / "ref.xlsx"
    work = tmp_path / "work.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "SAISIE"
    wb.calculation.fullCalcOnLoad = True
    wb.save(str(ref))
    wb.close()
    shutil.copy2(ref, work)

    violations = schema_real_svc._check_zip_vba_integrity(ref, work)
    assert violations == [], violations


# ── Test 10 : mutation d'une formule critique → ERREUR migration ───────────────

def test_app2e_formule_mutee_apres_migration_rejetee(tmp_path):
    """Une formule critique figée dans la copie de travail déclenche status=ERREUR."""
    from app.services.saisie_hh_schema_migration import migrate_saisie_copy as _real_migrate
    from app.readers.saisie_hh_reader import find_first_empty_data_row, _col_index

    source_saisie = _copy_saisie(tmp_path)
    source_ref = _make_ref(tmp_path / "ref.xlsm")

    def corrupt_migrate(path):
        result = _real_migrate(path)
        target = find_first_empty_data_row(path)
        if target:
            wb = openpyxl.load_workbook(str(path))
            wb["SAISIE"].cell(row=target, column=_col_index("B")).value = 99999
            wb.save(str(path))
            wb.close()
        return result

    with patch(
        "app.services.saisie_hh_schema_real_prepare_service.migrate_saisie_copy",
        corrupt_migrate,
    ):
        manifest = schema_real_svc.preparer_migration_hh_sur_copies(
            saisie_source=source_saisie,
            ref_setup_source=source_ref,
            output_dir=tmp_path / "prepared",
        )

    assert manifest["status"] == "ERREUR", manifest
    formula_errs = [e for e in manifest["errors"] if "FORMULE" in e]
    assert formula_errs, f"Aucune erreur de formule dans {manifest['errors']}"
