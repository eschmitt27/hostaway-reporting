"""Lot10 — test bloquant : aucun lecteur/service applicatif n'ouvre plus les 4 classeurs legacy.

MASTER_CALC_Flux.xlsx / MASTER_CALC_Commissions.xlsx / MASTER_CALC_NetProprietaire.xlsx /
MASTER_CALC_Resultats.xlsx sont des sorties du moteur legacy (`02_TRAVAIL/lot9_construire_flux.py`,
`lot10_calculer_resultats.py`) — depuis les migrations 0043/0044, les données vivent en SQLite
(`flux_unifies`, `lot10_*`) et les lecteurs applicatifs ne lisent plus que ces tables. Preuve active
par interception `openpyxl.load_workbook`, même principe que `test_menages_sans_excel.py` et
`test_lot9_sans_master_calc_flux.py`.
"""
from __future__ import annotations

from pathlib import Path

import pytest

INTERDITS = {
    "MASTER_CALC_Flux.xlsx",
    "MASTER_CALC_Commissions.xlsx",
    "MASTER_CALC_NetProprietaire.xlsx",
    "MASTER_CALC_Resultats.xlsx",
}


@pytest.fixture(autouse=True)
def _interdire_masters_lot10(monkeypatch):
    import openpyxl

    original = openpyxl.load_workbook

    def garde(chemin, *a, **kw):
        nom = Path(chemin).name
        if nom in INTERDITS:
            raise AssertionError(f"Master Lot9/Lot10 interdit rouvert par le NEW : {chemin}")
        return original(chemin, *a, **kw)

    monkeypatch.setattr(openpyxl, "load_workbook", garde)


@pytest.fixture
def seeded(tmp_db):
    import fixtures_lot10 as fx

    par_logement = [
        {"mois": "2026-06", "logement_id": "LOG_A1", "proprietaire_id": "PROP_A",
         "total_produits": 1000.0, "total_charges": 300.0, "resultat": 700.0, "nb_flux": 5,
         "vision": "REEL", "commentaire": ""},
        {"mois": "2026-06", "logement_id": "LOG_A1", "proprietaire_id": "PROP_A",
         "total_produits": 900.0, "total_charges": 250.0, "resultat": 650.0, "nb_flux": 4,
         "vision": "COMPTABLE", "commentaire": ""},
    ]
    vue_mois = [
        {"mois": "2026-06", "proprietaire_id": "PROP_A", "total_payout_mois": 2000,
         "total_menage_mois": 150, "total_commission_mois": 300, "charge_fixe_mensuelle": 0,
         "montant_du_conciergerie": 300, "reste_a_payer_conciergerie": 0,
         "net_proprietaire_avant_charge_mois": 1700, "net_proprietaire_apres_charge_mois": 1700,
         "nb_reservations": 4},
    ]
    commissions_a_controler = [
        {"reservation_id": "60001", "source": "vrbo", "channel_type": "VRBO",
         "statut_calcul_payout": "A_CONTROLER", "source_payout": "AUCUN_PAYOUT"},
    ]
    fx.seeder(tmp_db, resultats=par_logement, net_vue_mois=vue_mois,
              commissions_a_controler=commissions_a_controler)
    return tmp_db


def test_flux_unifie_service_sans_masters(tmp_db):
    from app.services import flux_unifie_service as svc

    resultat = svc.construire(db_path=tmp_db)
    assert resultat["ok"], resultat
    assert svc.lire(db_path=tmp_db) == []


def test_proprietaires_reglements_reader_sans_masters(seeded):
    from app.readers import proprietaires_reglements_reader as reader

    reader.vider_cache()
    reader.resultats(db_path=seeded)  # exercer la lecture sans crash
    global_ = reader.resultats_global(db_path=seeded)
    assert global_ is not None
    par_logement = reader.resultats_par_logement(db_path=seeded)
    assert any(l.get("logement_id") == "LOG_A1" for l in par_logement.lignes)
    reader.vider_cache()


def test_proprietaires_reader_sans_masters(seeded):
    from app.readers import proprietaires_reader as reader

    assert reader.calc_available() in (True, False)


def test_controles_detail_reader_commissions_a_controler_sans_masters(seeded):
    from app.readers import controles_detail_reader as detail

    detail.vider_cache()
    lignes = detail.commissions_a_controler()
    assert any(l.get("reservation_id") == "60001" for l in lignes)
    detail.vider_cache()


def test_ventes_lot12_adapter_frontiere_lot11_lot12_sans_masters(tmp_db):
    """Frontière Lot11/12 : l'adaptateur lit VUE_MOIS depuis Lot10 SQLite, jamais un master."""
    import fixtures_lot10 as fx
    from app.services import ventes_lot12_adapter_service as adapter

    vue_mois = [
        {"mois": "2026-06", "proprietaire_id": "PROP_A", "total_payout_mois": 2000,
         "total_menage_mois": 150, "total_commission_mois": 300, "charge_fixe_mensuelle": 0,
         "montant_du_conciergerie": 300, "reste_a_payer_conciergerie": 0,
         "net_proprietaire_avant_charge_mois": 1700, "net_proprietaire_apres_charge_mois": 1700,
         "nb_reservations": 4},
    ]
    fx.seeder(tmp_db, net_vue_mois=vue_mois)
    lignes = adapter.lignes_du_mois("2026-06")
    assert any(l["proprietaire_id"] == "PROP_A" for l in lignes)


def test_resultats_routes_sans_masters(client, seeded):
    r = client.get("/resultats?mois=2026-06&vision=REEL")
    assert r.status_code == 200
    r = client.get("/resultats/mensuel?mois=2026-06&vision=REEL")
    assert r.status_code == 200
    r = client.get("/resultats/proprietaires?mois=2026-06&vision=REEL")
    assert r.status_code == 200
