"""Couche HTTP du module Logements : fiche = écran central du cycle de vie.

Vérifie que création, modification, changement de propriétaire, changement de taux, archivage et
réactivation sont pilotables depuis le navigateur (routes + templates), sans jamais manipuler les
fichiers Excel directement. Complète les tests de service (`test_logements_creation.py`,
`test_logements_gestion.py`) par des tests de bout en bout HTTP.
"""
from __future__ import annotations

import openpyxl
import pytest

import app.config as cfg
from app.readers import ref_setup_reader
from test_logements import construire_referentiel


def _ref(tmp_path):
    p = tmp_path / "REF_Setup.xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("REF_Logements")
    ws.append(["logement_id", "hostaway_listing_id", "nom_logement_officiel", "nom_court", "adresse",
               "ville", "type_logement_id", "sur_hostaway", "actif", "statut_parc", "commentaire",
               "forfait_logiciel_consommables_mensuel"])
    ws.append(["LOG_A1", 900001, "Fictif A1", "A1", "1 rue", "RECETTE", "TYPE_001", "OUI", "OUI",
               "GERE", None, 0])
    ws = wb.create_sheet("REF_Gestion_Logements_Hist")
    ws.append(["gestion_id", "logement_id", "proprietaire_id", "date_debut", "date_fin",
               "statut_gestion", "source", "commentaire"])
    ws.append(["GST_1", "LOG_A1", "PROP_A", "2026-01-01", None, "ACTIF", "FICTIF", None])
    ws = wb.create_sheet("REF_Proprietaires")
    ws.append(["proprietaire_id", "nom_proprietaire", "actif"])
    for pid in ("PROP_A", "PROP_B"):
        ws.append([pid, f"Nom {pid}", "OUI"])
    ws = wb.create_sheet("REF_Types_Logements")
    ws.append(["type_logement_id", "libelle"]); ws.append(["TYPE_001", "STUDIO"])
    ws = wb.create_sheet("REF_Taux_Commission")
    ws.append(["taux_commission_id", "proprietaire_id", "logement_id", "taux_commission",
               "date_debut", "date_fin", "actif", "justification", "commentaire"])
    ws.append(["TX_A1", "PROP_A", "LOG_A1", 0.19, "2026-01-01", None, "OUI", "FICTIF", ""])
    wb.save(p); wb.close()
    return p


@pytest.fixture
def ref(tmp_path, monkeypatch):
    """Classeur pour les ÉCRITURES, référentiel SQLite pour les LECTURES.

    Les actions d'administration (créer, changer de propriétaire, changer de taux) écrivent encore
    dans REF_Setup.xlsm — c'est le périmètre de la mission « administration du référentiel ». Les
    lectures d'écran, elles, passent désormais par la base. Les deux fixtures décrivent le même
    logement, ce qui permet de vérifier que l'écran reste cohérent pendant la transition.
    """
    p = _ref(tmp_path)
    monkeypatch.setattr(cfg, "REF_SETUP", p)
    monkeypatch.setattr(ref_setup_reader, "REF_SETUP", p)

    db = construire_referentiel(
        tmp_path,
        logements=[{"logement_id": "LOG_A1", "nom_logement_officiel": "Fictif A1",
                    "nom_court": "A1", "adresse": "1 rue", "ville": "RECETTE",
                    "type_logement_id": "TYPE_001", "sur_hostaway": "OUI", "actif": "OUI",
                    "statut_parc": "GERE",
                    "forfait_logiciel_consommables_mensuel": "0"}],
        gestion=[{"gestion_id": "GST_1", "logement_id": "LOG_A1", "proprietaire_id": "PROP_A",
                  "date_debut": "2026-01-01", "date_fin": "", "statut_gestion": "ACTIF",
                  "source": "FICTIF"}],
        types=[{"type_logement_id": "TYPE_001", "type_logement": "STUDIO"}],
        taux=[{"taux_commission_id": "TX_A1", "proprietaire_id": "PROP_A",
               "logement_id": "LOG_A1", "taux_commission": "0.19", "date_debut": "2026-01-01",
               "date_fin": "", "actif": "OUI"}],
        proprietaires=[{"proprietaire_id": "PROP_A", "nom_proprietaire": "Nom PROP_A",
                        "actif": "OUI"},
                       {"proprietaire_id": "PROP_B", "nom_proprietaire": "Nom PROP_B",
                        "actif": "OUI"}])
    monkeypatch.setattr(cfg, "DB_PATH", db)
    monkeypatch.setattr(cfg, "RECETTE_MODE", True)
    monkeypatch.setattr(cfg, "RECETTE_ROOT", tmp_path.resolve())
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", True)
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_CONFIRMATION_ENABLED", True)
    return p


# ── Fiche logement = écran central ───────────────────────────────────────────

def test_fiche_affiche_etat_actuel_historique_et_actions(client, ref):
    html = client.get("/logements/LOG_A1").text
    assert "État actuel (temps réel)" in html
    assert "Modifier les informations" in html
    assert "Changer de propriétaire" in html
    assert "Changer le taux de commission" in html
    assert "Archiver ce logement" in html            # actif -> proposé
    assert "Réactiver ce logement" not in html
    assert "Historique — rattachements de gestion" in html
    assert "Historique — taux de commission" in html
    assert 'action="/logements/LOG_A1/modifier"' in html
    assert 'action="/logements/LOG_A1/changer-proprietaire"' in html
    assert 'action="/logements/LOG_A1/changer-taux-commission"' in html
    assert 'action="/logements/LOG_A1/archiver"' in html


def test_fiche_logement_inconnu_404(client, ref):
    r = client.get("/logements/LOG_ZZZ")
    assert r.status_code == 404


def test_fiche_affiche_boutons_desactives_si_flags_off(client, ref, monkeypatch):
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", False)
    html = client.get("/logements/LOG_A1").text
    assert "ces actions sont indisponibles" in html
    assert "disabled" in html


# ── Création ─────────────────────────────────────────────────────────────────

def test_page_creation_affiche_le_formulaire(client, ref):
    html = client.get("/logements/nouveau").text
    assert 'action="/logements"' in html
    assert "PROP_A" in html and "PROP_B" in html


def test_creer_puis_consulter_la_fiche(client, ref):
    r = client.post("/logements", data={
        "logement_id": "LOG_NEW", "nom_logement_officiel": "Nouveau", "nom_court": "NEW",
        "proprietaire_id": "PROP_B", "date_debut": "2026-06-01",
    }, follow_redirects=False)
    assert r.status_code == 303
    assert r.headers["location"].startswith("/logements/LOG_NEW")

    r2 = client.get(r.headers["location"])
    assert r2.status_code == 200
    assert "Logement créé" in r2.text
    # PBI (Lot13) n'a pas encore tourné : fiche minimale construite depuis REF_Setup en direct,
    # jamais un 404 pour un logement qui existe réellement.
    assert "en attente de l'export PBI" in r2.text
    assert "PROP_B" in r2.text


def test_creer_refuse_si_flags_off_reste_sur_le_formulaire(client, ref, monkeypatch):
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", False)
    r = client.post("/logements", data={
        "logement_id": "LOG_NEW", "nom_court": "NEW",
        "proprietaire_id": "PROP_B", "date_debut": "2026-06-01",
    }, follow_redirects=False)
    assert r.status_code == 303
    assert r.headers["location"].startswith("/logements/nouveau?erreur=")


# ── Modification ─────────────────────────────────────────────────────────────

def test_modifier_depuis_la_fiche(client, ref):
    r = client.post("/logements/LOG_A1/modifier",
                    data={"nom_court": "A1-bis"}, follow_redirects=False)
    assert r.status_code == 303
    assert "message=" in r.headers["location"]

    html = client.get(r.headers["location"]).text
    assert "Logement modifié" in html


# ── Changement de propriétaire (historisation) ───────────────────────────────

def test_changer_proprietaire_historise_et_message_recalcul(client, ref):
    r = client.post("/logements/LOG_A1/changer-proprietaire",
                    data={"proprietaire_id": "PROP_B", "date_debut": "2026-07-01"},
                    follow_redirects=False)
    assert r.status_code == 303
    html = client.get(r.headers["location"]).text
    assert "Propriétaire changé" in html
    assert "relancer les calculs" in html

    fiche = client.get("/logements/LOG_A1").text
    assert "Historique — rattachements de gestion" in fiche
    assert "PROP_B" in fiche and "PROP_A" in fiche       # ancienne + nouvelle ligne visibles


def test_changer_proprietaire_inconnu_affiche_erreur(client, ref):
    r = client.post("/logements/LOG_A1/changer-proprietaire",
                    data={"proprietaire_id": "PROP_ZZ", "date_debut": "2026-07-01"},
                    follow_redirects=False)
    html = client.get(r.headers["location"]).text
    assert "existe pas dans le référentiel" in html


# ── Changement de taux (historisation) ───────────────────────────────────────

def test_changer_taux_historise_et_message_recalcul(client, ref):
    r = client.post("/logements/LOG_A1/changer-taux-commission",
                    data={"taux_commission": "0.12", "date_debut": "2026-07-01"},
                    follow_redirects=False)
    assert r.status_code == 303
    html = client.get(r.headers["location"]).text
    assert "Taux de commission changé" in html
    assert "relancer Lot10" in html

    fiche = client.get("/logements/LOG_A1").text
    assert "12.0 %" in fiche or "12 %" in fiche
    assert "19.0 %" in fiche or "19 %" in fiche           # ancien taux toujours visible, jamais effacé


def test_changer_taux_invalide_affiche_erreur(client, ref):
    r = client.post("/logements/LOG_A1/changer-taux-commission",
                    data={"taux_commission": "3", "date_debut": "2026-07-01"},
                    follow_redirects=False)
    html = client.get(r.headers["location"]).text
    assert "entre 0 et 1" in html


# ── Archivage / réactivation ─────────────────────────────────────────────────

def test_archiver_puis_reactiver(client, ref):
    r = client.post("/logements/LOG_A1/archiver", data={"date_fin": "2026-06-30"},
                    follow_redirects=False)
    assert r.status_code == 303
    html = client.get(r.headers["location"]).text
    assert "archivé" in html

    fiche = client.get("/logements/LOG_A1").text
    assert "Réactiver ce logement" in fiche
    assert "Archiver ce logement" not in fiche

    r2 = client.post("/logements/LOG_A1/reactiver",
                     data={"date_debut": "2026-07-01", "proprietaire_id": "PROP_A"},
                     follow_redirects=False)
    assert r2.status_code == 303
    html2 = client.get(r2.headers["location"]).text
    assert "réactivé" in html2

    fiche2 = client.get("/logements/LOG_A1").text
    assert "Archiver ce logement" in fiche2
    assert "Réactiver ce logement" not in fiche2


def test_toutes_les_ecritures_refusent_si_flags_off(client, ref, monkeypatch):
    monkeypatch.setattr(cfg, "CHARGES_REAL_WRITE_ENABLED", False)
    routes = [
        ("/logements/LOG_A1/modifier", {"nom_court": "X"}),
        ("/logements/LOG_A1/archiver", {"date_fin": "2026-06-30"}),
        ("/logements/LOG_A1/changer-proprietaire", {"proprietaire_id": "PROP_B", "date_debut": "2026-07-01"}),
        ("/logements/LOG_A1/changer-taux-commission", {"taux_commission": "0.10", "date_debut": "2026-07-01"}),
    ]
    for url, data in routes:
        r = client.post(url, data=data, follow_redirects=False)
        html = client.get(r.headers["location"]).text
        assert "Écriture désactivée" in html
