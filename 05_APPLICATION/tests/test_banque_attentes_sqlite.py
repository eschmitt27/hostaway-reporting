"""Files d'attente de rapprochement bancaire en SQLite (Lot 8c).

Ce service ne crée aucun produit économique et ne confirme aucun rapprochement : il constitue des
files d'attente. La garantie la plus importante est négative — **aucune jointure Banque →
réservation**, jamais.

Fixtures synthétiques ; un test lit les données réelles pour la parité avec lot8c.
"""
from __future__ import annotations

from pathlib import Path

import pytest

import app.config as cfg
from app.db.connection import apply_migrations, get_db
from app.services import banque_attentes_service as att
from app.services import banque_classification_service as cls
from app.services import banque_mouvements_service as bq

COMPTE = "COMPTE_DEMO"

_COLS_REGLE = ("regle_id", "priorite", "actif", "compte_id", "type_match", "champ_cible", "motif",
               "tiers_detecte", "categorie", "type_flux_id", "code_impact", "source_economique",
               "rapprochement_requis", "validation_automatique", "niveau_risque",
               "statut_controle_defaut", "statut_classification_defaut", "date_debut_validite",
               "date_fin_validite", "commentaire")


def _regle(regle_id, priorite, motif, tiers, *, type_match=cls.MATCH_CONTIENT):
    d = {c: "" for c in _COLS_REGLE}
    d.update({"regle_id": regle_id, "priorite": str(priorite), "actif": "OUI", "compte_id": "*",
              "type_match": type_match, "champ_cible": "libelle", "motif": motif,
              "tiers_detecte": tiers, "categorie": "CAT_DEMO", "niveau_risque": "FAIBLE",
              "statut_controle_defaut": cls.ST_VALIDE,
              "statut_classification_defaut": cls.CLASS_RAPPROCHEMENT_REQUIS})
    return d


@pytest.fixture
def db(tmp_path) -> Path:
    p = tmp_path / "banque.db"
    apply_migrations(p)
    conn = get_db(p)
    try:
        trous = ", ".join(["?"] * (len(_COLS_REGLE) + 1))
        regles = [_regle("R_PLAT", 10, "VIR PLATEFORME", att.TIERS_PLATEFORME),
                  _regle("R_PROP", 20, "PROPRIO DEMO", "PROP_9001"),
                  _regle("R_AUTRE", 30, "AUTRE", "FOURNISSEUR_DEMO")]
        conn.executemany(
            f"INSERT INTO ref_banque_regles ({', '.join(_COLS_REGLE)}, import_id) "
            f"VALUES ({trous})",
            [tuple(r[c] for c in _COLS_REGLE) + ("IMP-TEST",) for r in regles])
        conn.execute("INSERT INTO ref_proprietaires (proprietaire_id, nom_proprietaire, "
                     "prenom_proprietaire, email, telephone, adresse_facturation, "
                     "mode_facturation, actif, commentaire, import_id) "
                     "VALUES ('PROP_9001','DEMO','','','','','PAR_LOGEMENT','OUI','','IMP-TEST')")
        conn.commit()
    finally:
        conn.close()
    return p


def _mvt(libelle, montant, date_op="2026-01-05"):
    return {"external_transaction_id": "", "date_operation": date_op, "date_valeur": "",
            "sens": bq.SENS_CREDIT, "montant": montant, "devise": "EUR",
            "libelle_brut": libelle, "contrepartie_brute": ""}


def _preparer(db, mouvements):
    bq.importer(mouvements, bank_account_id=COMPTE, source_type=bq.SOURCE_HISTORIQUE, db_path=db)
    cls.classer(bank_account_id=COMPTE, db_path=db)


# ── La garantie négative ────────────────────────────────────────────────────────────────────────

def test_aucune_jointure_vers_une_reservation(db):
    """BANQUE ≠ RÉSERVATION. Un versement de plateforme n'est jamais rattaché à une réservation."""
    _preparer(db, [_mvt("VIR PLATEFORME G12345678", 1200.0)])
    att.construire(bank_account_id=COMPTE, db_path=db)

    conn = get_db(db)
    try:
        n = conn.execute("SELECT COUNT(*) FROM banque_rapprochements "
                         "WHERE type_objet = 'RESERVATION'").fetchone()[0]
    finally:
        conn.close()
    assert n == 0


def test_aucun_produit_economique_cree(db):
    _preparer(db, [_mvt("VIR PLATEFORME G1", 500.0), _mvt("PROPRIO DEMO", 800.0)])
    r = att.construire(bank_account_id=COMPTE, db_path=db)
    assert r["produit_economique_cree"] == 0


def test_les_attentes_restent_de_simples_propositions(db):
    """Aucune confirmation automatique : seule une décision humaine peut valider."""
    _preparer(db, [_mvt("VIR PLATEFORME G1", 500.0)])
    att.construire(bank_account_id=COMPTE, db_path=db)
    for a in att.attentes(db_path=db):
        assert a["statut"] == att.ST_PROPOSE
        assert a["source"] == att.SOURCE_AUTO
        assert a["objet_id"] is None, "rien à rattacher — c'est tout l'objet de l'attente"


# ── Classement par motif ────────────────────────────────────────────────────────────────────────

def test_plateforme_en_attente_d_export(db):
    _preparer(db, [_mvt("VIR PLATEFORME G12345678", 1200.0)])
    r = att.construire(bank_account_id=COMPTE, db_path=db)
    assert r["plateforme"] == {"nb": 1, "total": 1200.0}
    a = att.attentes(motif=att.ATTENTE_EXPORT_PLATEFORME, db_path=db)[0]
    assert a["type_objet"] == att.TYPE_PAYOUT_PLATEFORME
    assert a["criteres"]["reference"] == "G12345678", "la référence de versement est conservée"
    assert "réservation" in a["commentaire"].lower()


def test_montant_d_ajustement_traite_a_part(db):
    """Seuil repris de lot8c : un versement de 4,97 € est un ajustement, pas un payout."""
    _preparer(db, [_mvt("VIR PLATEFORME G1", att.MONTANT_AJUSTEMENT)])
    r = att.construire(bank_account_id=COMPTE, db_path=db)
    assert r["ajustement"]["nb"] == 1 and r["plateforme"]["nb"] == 0


def test_proprietaire_en_attente_de_saisie(db):
    _preparer(db, [_mvt("PROPRIO DEMO VIREMENT", 900.0)])
    r = att.construire(bank_account_id=COMPTE, db_path=db)
    assert r["proprietaires"] == {"nb": 1, "total": 900.0}
    a = att.attentes(motif=att.ATTENTE_SAISIE_ACOMPTE, db_path=db)[0]
    assert a["type_objet"] == att.TYPE_REVERSEMENT_PROPRIETAIRE


def test_tiers_a_identifier_traite_comme_proprietaire(db):
    """Un libellé de regroupement historique n'est pas un propriétaire, mais reste à traiter."""
    conn = get_db(db)
    try:
        conn.execute("UPDATE ref_banque_regles SET tiers_detecte = 'FAMILLE_A_CONTROLER' "
                     "WHERE regle_id = 'R_PROP'")
        conn.commit()
    finally:
        conn.close()
    _preparer(db, [_mvt("PROPRIO DEMO", 100.0)])
    r = att.construire(bank_account_id=COMPTE, db_path=db)
    assert r["proprietaires"]["nb"] == 1


def test_autre_tiers_ignore(db):
    """Un fournisseur ne relève d'aucune de ces deux files."""
    _preparer(db, [_mvt("AUTRE CHOSE", 42.0)])
    r = att.construire(bank_account_id=COMPTE, db_path=db)
    assert r["nb_attentes"] == 0


# ── Recalcul et décisions humaines ──────────────────────────────────────────────────────────────

def test_recalcul_remplace_les_propositions(db):
    _preparer(db, [_mvt("VIR PLATEFORME G1", 500.0)])
    att.construire(bank_account_id=COMPTE, db_path=db)
    att.construire(bank_account_id=COMPTE, db_path=db)
    assert len(att.attentes(db_path=db)) == 1, "un recalcul remplace, il n'empile pas"


def test_une_decision_humaine_survit_au_recalcul(db):
    """Le service produit des propositions ; il ne défait jamais une décision."""
    _preparer(db, [_mvt("VIR PLATEFORME G1", 500.0)])
    att.construire(bank_account_id=COMPTE, db_path=db)
    opaque = att.attentes(db_path=db)[0]["rapprochement_id_opaque"]

    conn = get_db(db)
    try:
        conn.execute("UPDATE banque_rapprochements SET statut='CONFIRME', source='MANUEL' "
                     "WHERE rapprochement_id_opaque=?", (opaque,))
        conn.commit()
    finally:
        conn.close()

    att.construire(bank_account_id=COMPTE, db_path=db)
    conn = get_db(db)
    try:
        statut = conn.execute("SELECT statut FROM banque_rapprochements "
                              "WHERE rapprochement_id_opaque=?", (opaque,)).fetchone()[0]
    finally:
        conn.close()
    assert statut == "CONFIRME"


def test_sans_classification_le_service_refuse(db):
    bq.importer([_mvt("VIR PLATEFORME G1", 500.0)], bank_account_id=COMPTE,
                source_type=bq.SOURCE_HISTORIQUE, db_path=db)
    r = att.construire(bank_account_id=COMPTE, db_path=db)
    assert r["ok"] is False and "classification" in r["message"].lower()


def test_proprietaires_lus_du_referentiel(db):
    """lot8c portait une liste fermée codée en dur — elle se périmait à chaque ajout."""
    assert "PROP_9001" in att.proprietaires_connus(db_path=db)


# ── §11 Parité avec lot8c sur données réelles ───────────────────────────────────────────────────

RELEVE = Path(cfg.SOURCES_BRUTES) / "Banque" / \
    "BANQUE_ACTUELLE_HISTORIQUE_2025-11-03_2026-08-01.xlsx"
MASTER = Path(cfg.MASTER_BANQUE)
REF = Path(cfg.REF_SETUP)

reel_requis = pytest.mark.skipif(
    not (RELEVE.exists() and MASTER.exists() and REF.exists()),
    reason="relevé, master ou référentiel absent de cet environnement")


@reel_requis
def test_parite_avec_les_onglets_d_attente_de_lot8c(tmp_path, monkeypatch):
    """OLD (onglets RAPPROCH_*_ATTENTE) contre NEW (SQLite) : mêmes lignes, mêmes totaux.

    Les nombres sont LUS dans le master, pas écrits en dur : le volume bancaire évoluera.
    """
    import openpyxl

    from app.services import banque_adaptateurs as ad
    from app.services import ref_setup_import_service as refimp

    p = tmp_path / "reel.db"
    apply_migrations(p)
    monkeypatch.setattr(cfg, "DB_PATH", p)
    assert refimp.importer(db_path=p)["ok"]
    bq.importer(ad.depuis_xlsx_releve_consolide(RELEVE),
                bank_account_id="CM_02211_00021321603",
                source_type=bq.SOURCE_HISTORIQUE, db_path=p)
    cls.classer(db_path=p)
    r = att.construire(db_path=p)
    assert r["ok"], r

    wb = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
    try:
        def _lire(onglet):
            """Lignes de mouvement seules.

            Ces onglets se terminent par une ligne blanche et un pied « TOTAL INFORMATIF ». Les
            compter donnerait 167 là où lot8c annonce 166 : on ne retient que les lignes portant
            un identifiant de mouvement.
            """
            if onglet not in wb.sheetnames:
                return 0, 0.0
            ws = wb[onglet]
            it = ws.iter_rows(values_only=True)
            h = [str(c) for c in next(it)]
            i_mvt = h.index("mouvement_id")
            i_mt = h.index("montant_banque") if "montant_banque" in h else None
            n, total = 0, 0.0
            for row in it:
                if not str(row[i_mvt] or "").startswith("MVT"):
                    continue
                n += 1
                if i_mt is not None:
                    try:
                        total += float(row[i_mt] or 0)
                    except (TypeError, ValueError):
                        pass
            return n, round(total, 2)

        n_airbnb, t_airbnb = _lire("RAPPROCH_AIRBNB_ATTENTE")
        n_prop, t_prop = _lire("RAPPROCH_PROPRIETAIRES_ATTENTE")
    finally:
        wb.close()

    # L'onglet Airbnb mélange attente et ajustement ; on compare la somme des deux motifs.
    nb_new = r["plateforme"]["nb"] + r["ajustement"]["nb"]
    total_new = round(r["plateforme"]["total"] + r["ajustement"]["total"], 2)
    assert nb_new == n_airbnb, f"plateforme : NEW {nb_new} contre OLD {n_airbnb}"
    assert total_new == t_airbnb, f"plateforme : NEW {total_new} contre OLD {t_airbnb}"

    assert r["proprietaires"]["nb"] == n_prop
    assert r["proprietaires"]["total"] == t_prop

    conn = get_db(p)
    try:
        assert conn.execute("SELECT COUNT(*) FROM banque_rapprochements "
                            "WHERE type_objet='RESERVATION'").fetchone()[0] == 0
    finally:
        conn.close()
